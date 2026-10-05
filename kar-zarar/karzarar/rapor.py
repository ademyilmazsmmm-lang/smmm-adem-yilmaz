# -*- coding: utf-8 -*-
"""Kar/zarar sonucunun Excel raporu (cikti/kar-zarar.xlsx): firma basina bir satir, kardan zarara siralı."""

from datetime import datetime

# Defter Beyan Mali Hesap Ozeti'nin ayri kalemleri (Luca satirlarinda bos): (baslik, ayrinti anahtari)
DB_KALEMLERI = {"hasilat": "Hasılat", "diger_gelir": "Diğer Gelir", "emtia_basi": "Dönem Başı Emtia",
                "emtia_sonu": "Dönem Sonu Emtia", "giderler": "Giderler", "amortisman": "Amortisman"}
BASLIKLAR = ["Firma", "VKN / TC", "Kaynak", "Defter", "Hasılat", "Diğer Gelir", "Toplam Gelir (Satış)",
             "Dönem Başı Emtia", "Mal Alışı", "Dönem Sonu Emtia", "Giderler", "Amortisman", "Toplam Gider",
             "Kâr / Zarar", "Durum", "Not", "Dönem"]
GENISLIK = [28, 14, 14, 16, 15, 15, 17, 16, 15, 16, 15, 15, 16, 16, 10, 60, 24]
DURUM_SUTUNU = 14  # BASLIKLAR icinde "Durum"un sirasi (0'dan)
PARA_SUTUNLARI = range(4, 14)


def _siralama(s):
    # kar (vergi cikabilecek) en ustte; zarar, sonra hatalilar
    return (s.get("kar") is None, -(s.get("kar") or 0))


def toplam_gider(s):
    """Toplam gider: Defter Beyan'da donem basi emtia + mal alisi + giderler + amortisman; Luca'da mal alisi + gider.
    Eski json'larda 'toplam_gider' yoksa ayni toplam bilesenlerden hesaplanir."""
    if s.get("kar") is None:
        return None
    if s.get("toplam_gider") is not None:
        return s["toplam_gider"]
    ayr = s.get("ayrinti") or {}
    if s.get("kaynak") == "Defter Beyan":
        return round(sum(ayr.get(k) or 0.0 for k in ("emtia_basi", "mal_alis", "giderler", "amortisman")), 2)
    return round((s.get("mal_alis") or 0.0) + (s.get("gider") or 0.0), 2)


def satirlar(sonuclar):
    """BASLIKLAR sirasinda satirlar, kardan zarara. Defter Beyan kalemleri yalniz sayfada varsa dolu olur."""
    cikti = []
    for s in sorted(sonuclar.values(), key=_siralama):
        kar = s.get("kar")
        durum = (("ATLANDI" if (s.get("hata") or "").startswith("ATLANDI") else "HATA") if kar is None
                 else ("KÂR" if kar >= 0 else "ZARAR"))
        ayr = s.get("ayrinti") or {} if s.get("kaynak") == "Defter Beyan" else {}
        k = lambda ad: ayr.get(ad)  # yoksa / bos hucre: None (Excel'de bos)
        cikti.append((s["firma"], s.get("vkn") or "", s.get("kaynak") or "", s.get("defter") or "",
                      k("hasilat"), k("diger_gelir"), s.get("satis"), k("emtia_basi"), s.get("mal_alis"),
                      k("emtia_sonu"), k("giderler"), k("amortisman"), toplam_gider(s), kar, durum,
                      " ".join(x for x in (s.get("hata"),
                                           f"(Son sorguda hata: {s['son_hata']})" if s.get("son_hata") else "") if x),
                      s.get("donem") or ""))
    return cikti


def excel_yaz(yol, bas, bit, sonuclar):
    """Raporu yazar; dosya Excel'de aciksa (PermissionError) 'kar-zarar (tarih saat).xlsx' olarak yazar. Yolu dondurur."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Kâr-Zarar"
    ws.append([f"Kâr / Zarar Tahmini  {bas:%d/%m/%Y} - {bit:%d/%m/%Y}   (alınma: {datetime.now():%d/%m/%Y %H:%M})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(BASLIKLAR)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    renk = {"KÂR": "E2EFDA", "ZARAR": "FCE4D6", "HATA": "EDEDED", "ATLANDI": "FFF2CC"}
    veri = satirlar(sonuclar)
    for sat in veri:
        ws.append(list(sat))
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=renk[sat[DURUM_SUTUNU]])
        for i in PARA_SUTUNLARI:
            ws.cell(ws.max_row, i + 1).number_format = '#,##0.00;[Red]-#,##0.00'
    for i, g in enumerate(GENISLIK, 1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = "A4"
    if veri:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(BASLIKLAR))}{ws.max_row}"
    kar = sum(1 for s in veri if s[DURUM_SUTUNU] == "KÂR")
    zarar = sum(1 for s in veri if s[DURUM_SUTUNU] == "ZARAR")
    hata = sum(1 for s in veri if s[DURUM_SUTUNU] == "HATA")
    atlandi = sum(1 for s in veri if s[DURUM_SUTUNU] == "ATLANDI")
    ws.append([])
    ws.append([f"Toplam {len(veri)} firma: {kar} kâr, {zarar} zarar, {atlandi} atlandı (dönem sonu vb.), {hata} hata/okunamadı."
               " Bu bir tahmindir;"
               " dönem sonu emtia, amortisman ve gelir/gider kayıtları işlenmemiş olabilir."])
    try:
        wb.save(yol)
        return yol
    except PermissionError:
        yedek = yol.with_name(f"{yol.stem} ({datetime.now():%d.%m %H.%M}){yol.suffix}")
        wb.save(yedek)
        return yedek
