#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Luca Bot masaustu arayuzu (SMMM Adem Yilmaz).

luca-arayuz.bat ile acilir. Komut satirina gerek kalmadan:
    - Luca giris bilgileri, firma listesi, tarih araligi ve ekranlar secilir
      (ayarlar.json'a kaydedilir),
    - "Calistir" botu (luca_bot.py) arka planda gece modunda baslatir; ciktisi
      pencerede canli akar, "Durdur" o ana kadarki sonuclari kaydederek durdurur,
    - alttaki kutular (Alis Tevkifat KDV, Alis SMM, Interaktif - e-Arsiv farki,
      KDV odemesi cikabilir) rapor.json'dan hesaplanir; tiklaninca firmalari
      listeler (lucabot.gostergeler).
"""

import codecs
import json
import os
import queue
import re
import signal
import subprocess
import sys
import threading
import time
import webbrowser
from datetime import date, timedelta
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

KOK = Path(__file__).resolve().parent
sys.path.insert(0, str(KOK))

from kar_zarar_sekmesi import KarZararSekmesi  # noqa: E402
from kurulum_sihirbazi import KurulumSihirbazi  # noqa: E402
from lucabot import SURUM, SURUM_TARIHI, beyanname, firma_tablosu, gostergeler, musteri_listesi, rapor  # noqa: E402
from tablo_gorunumu import SiralaFiltreTablosu  # noqa: E402
from lucabot import kurulum as kurulum_mantigi  # noqa: E402
from lucabot.firma_listesi import devreden_kdvleri  # noqa: E402
from lucabot.ortak import (AYAR_DOSYASI, DURDUR_DOSYASI, ORNEK_AYAR, TARIH_BICIMI,  # noqa: E402
                           hedef_ay_araligi, indirme_koku, sadelestir, tarih_cozumle)
from lucabot.sabitler import BELGE_TIPLERI, EKRAN_SUTUNLARI, TUM_BELGELER  # noqa: E402

# --- gorunum: acik (varsayilan) ve koyu tema (smmmyilmaz.com: lacivert + altin) ----------
TEMALAR = {
    "koyu": dict(ZEMIN="#0B1426", PANEL="#111D35", KUTU="#0E1830", KENAR="#2A3A5C", CIZGI="#1F2C48",
                 YAZI="#E8ECF4", BASLIK_FG="#F2F4F8", SOLUK="#9AA6BD", ETIKET="#B8C2D6",
                 ALTIN_FG="#D4B263", TURUNCU="#F0B45A", KIRMIZI="#F2918A", YESIL_FG="#8CE59A",
                 UYARI_FG="#F2D98A", LOG_SOLUK="#7F8BA3", LOG_BILGI="#9AA6BD", LOG_FIRMA="#E2C47A",
                 KONSOL="#070D1A", KONSOL_FG="#CBD2DE", KONSOL_KENAR="#1B2944", SECIM="#2A3A5C",
                 HOVER="#172443", DEVRE="#5B6782", TROUGH="#1B2944"),
    "acik": dict(ZEMIN="#F2F4F8", PANEL="#FFFFFF", KUTU="#FFFFFF", KENAR="#C3CCDC", CIZGI="#DCE2EC",
                 YAZI="#1C2538", BASLIK_FG="#12203A", SOLUK="#66718A", ETIKET="#4A5670",
                 ALTIN_FG="#8A6410", TURUNCU="#B4690E", KIRMIZI="#C23B30", YESIL_FG="#1B7F3B",
                 UYARI_FG="#8A6A00", LOG_SOLUK="#7A859D", LOG_BILGI="#5B6680", LOG_FIRMA="#8A6410",
                 KONSOL="#FFFFFF", KONSOL_FG="#26324D", KONSOL_KENAR="#C3CCDC", SECIM="#D6E2F7",
                 HOVER="#E7EDF7", DEVRE="#A3ACBF", TROUGH="#DCE2EC"),
}
# her iki temada ayni: ust bant (lacivert), altin dolgu/dugme, durum renkleri
BASLIK_ZEMIN = "#091122"
BASLIK_YAZI = "#F2F4F8"
BASLIK_ETIKET = "#B8C2D6"
ALTIN = "#D4B263"
ALTIN_ACIK = "#E2C47A"
ALTIN_YAZI = "#1A1405"
YESIL = "#13804F"
TURUNCU_ZEMIN = "#F0B45A"
HATA_ZEMIN = "#B3443A"
TEMA = "acik"


def tema_uygula(ad):
    """Renk sabitlerini secilen temaya cevirir (arayuz kurulurken okunur); gecersiz ad acik temaya duser."""
    global TEMA
    TEMA = ad if ad in TEMALAR else "acik"
    globals().update(TEMALAR[TEMA])


tema_uygula("acik")

GOVDE = ("Segoe UI", 10)
GOVDE_KALIN = ("Segoe UI", 10, "bold")
KUCUK = ("Segoe UI", 9)
BOLUM = ("Segoe UI", 8, "bold")
SERIF = ("Georgia", 16, "bold")
KONSOL_YAZI = ("Consolas", 10)

EKRAN_ADLARI = {tip: ad for ad, tip in EKRAN_SUTUNLARI.items()}
BELGE_ADLARI = dict(BELGE_TIPLERI)
UZUN_BEKLEME = 30  # "Şu an" satiri bu kadar saniye degismezse turuncu, 3 katinda kirmizi
ILERLEME = re.compile(r"^\[(\d+)/(\d+)\]\s+(.+?)(?:\s+\|\s+tahmini kalan:\s*(.+))?$")
AZAMI_SATIR = 4000  # log penceresinde tutulan satir (uzun gecelerde pencere sismesin)


# --- ayarlar ---------------------------------------------------------------

def ayar_yolu():
    ozel = os.environ.get("LUCA_BOT_AYAR")
    return Path(ozel) if ozel else AYAR_DOSYASI


def ayarlari_yukle():
    """ayarlar.json (yoksa ornek). Bozuksa ValueError: uzerine yazip kaybetmeyelim."""
    for yol in (ayar_yolu(), ORNEK_AYAR):
        if yol.exists():
            try:
                return json.loads(yol.read_text(encoding="utf-8-sig"))
            except json.JSONDecodeError as e:
                raise ValueError(f"{yol.name} okunamadi (satir {e.lineno}): {e.msg}")
    return {}


def ayarlari_kaydet(ayarlar):
    yol = ayar_yolu()
    gecici = yol.with_suffix(".json.tmp")
    gecici.write_text(json.dumps(ayarlar, ensure_ascii=False, indent=2), encoding="utf-8")
    gecici.replace(yol)


def gecen_ay():
    bu_ay = date.today().replace(day=1)
    bit = bu_ay - timedelta(days=1)
    return bit.replace(day=1).strftime(TARIH_BICIMI), bit.strftime(TARIH_BICIMI)


def dosya_ac(yol):
    yol = str(yol)
    if sys.platform.startswith("win"):
        os.startfile(yol)  # noqa: S606 - kullanicinin kendi dosyasi
    elif sys.platform == "darwin":
        subprocess.Popen(["open", yol])
    else:
        subprocess.Popen(["xdg-open", yol])


def konsolsuz_yeniden_baslat():
    """Windows'ta arayuz konsollu python.exe ile acildiysa (arkada siyah pencere kalir)
    pythonw.exe ile yeniden baslatir. True: konsolsuz surec basladi, bu surec cikmali.

    .bat'in pythonw'yi bulamadigi kurulumlarda da (orn. AppData altindaki yeni tip Python)
    siyah pencere kalmasin diye arayuzun kendisi yapar. Yeni surec 1,5 sn icinde
    kapanirsa (calismadi) False doner ve arayuz bu surecte acilir.
    """
    if not sys.platform.startswith("win"):
        return False
    try:
        import ctypes
        if not ctypes.windll.kernel32.GetConsoleWindow():
            return False  # zaten konsolsuz
        exe = Path(sys.executable)
        pyw = exe.with_name("pythonw.exe")
        if exe.name.lower() == "pythonw.exe" or not pyw.exists():
            return False
        DETACHED_PROCESS, CREATE_NEW_PROCESS_GROUP = 0x00000008, 0x00000200
        yeni = subprocess.Popen([str(pyw), str(Path(__file__).resolve()), *sys.argv[1:]],
                                cwd=str(KOK), close_fds=True, stdin=subprocess.DEVNULL,
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                creationflags=DETACHED_PROCESS | CREATE_NEW_PROCESS_GROUP)
        time.sleep(1.5)
        return yeni.poll() is None
    except Exception:
        return False


def python_komutu():
    """Botu calistiracak python; arayuz pythonw ile acildiysa konsollu python.exe."""
    exe = Path(sys.executable)
    if exe.name.lower() == "pythonw.exe" and (exe.parent / "python.exe").exists():
        return str(exe.parent / "python.exe")
    return str(exe)


TL_DESENI = re.compile(r"^-?[\d.]+,\d{2} TL( \*)?$")


def liste_excel_yaz(yol, baslik, donem, sutunlar, satirlar, toplam, notlar):
    """Ozet kutusunun firma listesini .xlsx olarak yazar; TL tutarlari sayi olarak."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from lucabot.fatura_analiz import tutar_cozumle

    def hucre(deger):
        if isinstance(deger, str) and TL_DESENI.match(deger):
            return tutar_cozumle(deger.replace(" *", ""))
        return deger

    wb = Workbook()
    ws = wb.active
    ws.title = "Liste"
    ws.append([baslik])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Dönem: {donem.replace('-', ' – ')}" if donem else ""])
    ws.append([])
    ws.append(list(sutunlar))
    for h in ws[4]:
        h.font = Font(bold=True)
    for s in satirlar:
        ws.append([hucre(d) for d in s])
    if satirlar:
        ws.append([hucre(d) for d in toplam])
        for h in ws[ws.max_row]:
            h.font = Font(bold=True)
    for satir in ws.iter_rows(min_row=5):
        for h in satir:
            if isinstance(h.value, float):
                h.number_format = '#,##0.00 "TL"'
    ws.append([])
    for n in notlar:
        ws.append([n])
    ws.column_dimensions["A"].width = 38
    for harf in "BCDEF":
        ws.column_dimensions[harf].width = 18
    wb.save(yol)


# --- kucuk parcalar ----------------------------------------------------------

def bolum_basligi(ebeveyn, metin):
    return tk.Label(ebeveyn, text=metin.upper(), font=BOLUM, fg=ALTIN_FG, bg=ZEMIN, anchor="w")


def giris_kutusu(ebeveyn, degisken, gizli=False, genislik=20):
    return tk.Entry(ebeveyn, textvariable=degisken, show="•" if gizli else "", width=genislik,
                    font=GOVDE, bg=KUTU, fg=YAZI, insertbackground=YAZI, relief="flat",
                    highlightthickness=1, highlightbackground=KENAR, highlightcolor=ALTIN)


def dugme(ebeveyn, metin, komut, ana=False, **kw):
    if ana:
        renk = dict(bg=ALTIN, fg=ALTIN_YAZI, activebackground=ALTIN_ACIK, activeforeground=ALTIN_YAZI,
                    font=("Segoe UI", 11, "bold"))
    else:
        renk = dict(bg=PANEL, fg=YAZI, activebackground=HOVER, activeforeground=YAZI, font=GOVDE)
    d = tk.Button(ebeveyn, text=metin, command=komut, relief="flat", cursor="hand2",
                  bd=0, padx=14, pady=7, disabledforeground=DEVRE,
                  highlightthickness=1, highlightbackground=KENAR, **renk, **kw)
    return d


VARLIKLAR = KOK / "varliklar"
def telif():
    """Alt cubuktaki telif satiri; yil kendiliginden guncellenir."""
    return f"© {date.today().year} Adem Yılmaz — Serbest Muhasebeci Mali Müşavir · Tüm Hakları Saklıdır."


def logo_resmi(boyut, master=None):
    """varliklar/logo-<boyut>.png (Robot Stajyer) ya da None; PhotoImage'i cagiran tutmalidir."""
    try:
        return tk.PhotoImage(master=master, file=str(VARLIKLAR / f"logo-{boyut}.png"))
    except tk.TclError:
        return None


def logo(ebeveyn):
    """Robot Stajyer logosu (PNG); dosya yoksa eski altin 'AY' kutusu."""
    resim = logo_resmi(48, ebeveyn)
    if resim is not None:
        et = tk.Label(ebeveyn, image=resim, bg=BASLIK_ZEMIN, bd=0, highlightthickness=0)
        et.resim = resim  # Tk, PhotoImage'i tutmaz
        return et
    c = tk.Canvas(ebeveyn, width=46, height=46, bg=BASLIK_ZEMIN, highlightthickness=0)
    ust, alt = (0xEA, 0xD2, 0x93), (0xC4, 0x9A, 0x45)
    for y in range(46):
        o = y / 45
        renk = "#%02x%02x%02x" % tuple(int(u + (a - u) * o) for u, a in zip(ust, alt))
        c.create_line(0, y, 46, y, fill=renk)
    c.create_text(23, 24, text="AY", font=("Georgia", 15, "bold"), fill="#121A2E")
    return c


class OzetKutusu(tk.Frame):
    """Alttaki tiklanabilir gosterge kutusu."""

    def __init__(self, ebeveyn, baslik, komut, renk=None):
        renk = renk or YAZI  # tema calisma aninda belli olur (varsayilan arguman import'ta sabitlenirdi)
        super().__init__(ebeveyn, bg=PANEL, highlightthickness=1, highlightbackground=CIZGI,
                         cursor="hand2", padx=10, pady=9)
        self.komut = komut
        self.baslik = tk.Label(self, text=baslik.upper(), font=BOLUM, fg=ETIKET if renk == YAZI else renk,
                               bg=PANEL, anchor="w")
        self.deger = tk.Label(self, text="—", font=("Segoe UI", 15, "bold"), fg=renk, bg=PANEL, anchor="w")
        self.alt = tk.Label(self, text="Gör ›", font=KUCUK, fg=ALTIN_FG, bg=PANEL, anchor="w")
        for w in (self.baslik, self.deger, self.alt):
            w.pack(fill="x")
        for w in (self, self.baslik, self.deger, self.alt):
            w.bind("<Button-1>", lambda _e: self.komut())
            w.bind("<Enter>", lambda _e: self.configure(highlightbackground=ALTIN))
            w.bind("<Leave>", lambda _e: self.configure(highlightbackground=CIZGI))

    def ayarla(self, deger, alt="Gör ›"):
        self.deger.configure(text=deger)
        self.alt.configure(text=alt)


class FirmaEkranPenceresi:
    """firmalar.xlsx'i tablo olarak gosterir: her firmada hangi ekran sorgulansin (✓/X), Devreden KDV.

    Firma eklenip silinebilir; Kaydet dosyaya yazar (firma_tablosu.tabloyu_yaz).
    """

    GENISLIK = (240, 100) + (80,) * len(TUM_BELGELER) + (34,)

    def __init__(self, arayuz, yol, firmalar):
        self.arayuz, self.yol, self.firmalar = arayuz, yol, firmalar
        w = self.w = tk.Toplevel(arayuz.kok, bg=ZEMIN, padx=18, pady=14)
        w.title(f"KDV Devri ve Ekran Seçimi — {yol.name}")
        w.transient(arayuz.kok)
        w.geometry(f"{sum(self.GENISLIK) + 70}x660")

        ust = tk.Frame(w, bg=ZEMIN)
        ust.pack(fill="x")
        tk.Label(ust, text="KDV Devri ve Ekran Seçimi", font=("Georgia", 14, "bold"), fg=BASLIK_FG,
                 bg=ZEMIN).pack(side="left")
        self.v_ara = tk.StringVar()
        self.v_ara.trace_add("write", lambda *_: self._suz())
        giris_kutusu(ust, self.v_ara, genislik=24).pack(side="right", ipady=3)
        tk.Label(ust, text="Firma ara:", font=KUCUK, fg=ETIKET, bg=ZEMIN).pack(side="right", padx=6)
        tk.Label(w, text="Her firmanın Devreden KDV'sini yazın ve sorgulanacak ekranları işaretleyin."
                         " Sütun başlığına tıklayınca o ekran tüm firmalarda açılır/kapanır, firma adına"
                         " tıklayınca o firmanın tüm ekranları. ✕ firmayı listeden çıkarır.",
                 font=KUCUK, fg=SOLUK, bg=ZEMIN, anchor="w", justify="left",
                 wraplength=sum(self.GENISLIK)).pack(fill="x", pady=(4, 8))

        alt = tk.Frame(w, bg=ZEMIN)
        alt.pack(side="bottom", fill="x", pady=(10, 0))
        self.v_yeni = tk.StringVar()
        yeni = giris_kutusu(alt, self.v_yeni, genislik=26)
        yeni.pack(side="left", ipady=4)
        yeni.bind("<Return>", lambda _e: self.firma_ekle())
        dugme(alt, "Firma Ekle", self.firma_ekle).pack(side="left", padx=6)
        self.bilgi = tk.Label(alt, text="", font=KUCUK, fg=SOLUK, bg=ZEMIN)
        self.bilgi.pack(side="left", padx=10)
        dugme(alt, "Kaydet", self.kaydet, ana=True).pack(side="right")
        dugme(alt, "Vazgeç", w.destroy).pack(side="right", padx=8)
        dugme(alt, "Beyannameden Devir Al…", self.beyannameden_al).pack(side="right", padx=(0, 8))
        dugme(alt, "Luca'dan Devir Çek…", lambda: arayuz.beyanname_cek(self)).pack(side="right", padx=(0, 8))

        baslik = tk.Frame(w, bg=KUTU)
        baslik.pack(fill="x")
        basliklar = ["Firma", "Devreden KDV"] + [EKRAN_ADLARI.get(t, t).replace(" ", "\n", 1)
                                                 for t in TUM_BELGELER] + [""]
        for i, ad in enumerate(basliklar):
            baslik.grid_columnconfigure(i, minsize=self.GENISLIK[i])
            et = tk.Label(baslik, text=ad, font=BOLUM, fg=ALTIN_FG, bg=KUTU,
                          anchor="w" if i < 2 else "center", justify="center", pady=6)
            et.grid(row=0, column=i, sticky="ew")
            if 2 <= i < 2 + len(TUM_BELGELER):
                tip = TUM_BELGELER[i - 2]
                et.configure(cursor="hand2")
                et.bind("<Button-1>", lambda _e, t=tip: self._sutunu_cevir(t))

        govde = tk.Frame(w, bg=PANEL)
        govde.pack(fill="both", expand=True)
        self.tuval = tk.Canvas(govde, bg=PANEL, highlightthickness=0)
        kaydir = tk.Scrollbar(govde, command=self.tuval.yview)
        self.tuval.configure(yscrollcommand=kaydir.set)
        kaydir.pack(side="right", fill="y")
        self.tuval.pack(side="left", fill="both", expand=True)
        self.ic = tk.Frame(self.tuval, bg=PANEL)
        self.tuval.create_window((0, 0), window=self.ic, anchor="nw")
        self.ic.bind("<Configure>", lambda _e: self.tuval.configure(scrollregion=self.tuval.bbox("all")))
        for i, g in enumerate(self.GENISLIK):
            self.ic.grid_columnconfigure(i, minsize=g)

        self.satirlar, self.silinenler = [], set()
        for f in firmalar:
            self._satir_olustur(f)
        self._suz()
        w.bind_all("<MouseWheel>", self._tekerlek)
        w.bind_all("<Button-4>", lambda _e: self.tuval.yview_scroll(-3, "units"))
        w.bind_all("<Button-5>", lambda _e: self.tuval.yview_scroll(3, "units"))
        w.bind("<Destroy>", self._kapandi)

    def _satir_olustur(self, f):
        satir = {"ad": f["ad"]}
        ad = tk.Label(self.ic, text=f["ad"], font=GOVDE, fg=YAZI, bg=PANEL, anchor="w", cursor="hand2")
        ad.bind("<Button-1>", lambda _e: self._satiri_cevir(satir))
        satir["dev"] = tk.StringVar(value=f["devreden"])
        kutu = giris_kutusu(self.ic, satir["dev"], genislik=10)
        satir["secim"] = {t: tk.BooleanVar(value=t in f["ekranlar"]) for t in TUM_BELGELER}
        isaretler = [tk.Checkbutton(self.ic, variable=satir["secim"][t], bg=PANEL, activebackground=PANEL,
                                    fg=YAZI, activeforeground=YAZI, selectcolor=KUTU,
                                    highlightthickness=0, bd=0)
                     for t in TUM_BELGELER]
        sil = tk.Label(self.ic, text="✕", font=GOVDE, fg=KIRMIZI, bg=PANEL, cursor="hand2")
        sil.bind("<Button-1>", lambda _e: self.firma_sil(satir))
        satir["widgetlar"] = [ad, kutu] + isaretler + [sil]
        self.satirlar.append(satir)
        return satir

    def firma_ekle(self):
        ad = " ".join(self.v_yeni.get().split())
        if not ad:
            return
        if any(sadelestir(s["ad"]) == sadelestir(ad) for s in self.satirlar):
            messagebox.showinfo("Zaten var", f"{ad} listede zaten var.", parent=self.w)
            return
        self.silinenler.discard(ad)
        self._satir_olustur({"ad": ad, "devreden": "", "ekranlar": set(TUM_BELGELER)})
        self.v_yeni.set("")
        self.v_ara.set("")
        self._suz()
        self.tuval.update_idletasks()
        self.tuval.yview_moveto(1)

    def firma_sil(self, satir):
        if not messagebox.askyesno("Firmayı çıkar", f"{satir['ad']} listeden çıkarılsın mı?\n"
                                   "(Kaydet'e basınca dosyadan silinir; bu firma işlenmez.)", parent=self.w):
            return
        for wdg in satir["widgetlar"]:
            wdg.destroy()
        self.satirlar.remove(satir)
        self.silinenler.add(satir["ad"])
        self._suz()

    def _tekerlek(self, e):
        self.tuval.yview_scroll(int(-e.delta / 120) * 3, "units")

    def _kapandi(self, e):
        if e.widget is self.w:
            for olay in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
                self.w.unbind_all(olay)

    def gorunen(self):
        ara = sadelestir(self.v_ara.get())
        return [s for s in self.satirlar if ara in sadelestir(s["ad"])]

    def _suz(self):
        for s in self.satirlar:
            for wdg in s["widgetlar"]:
                wdg.grid_forget()
        for r, s in enumerate(self.gorunen()):
            for c, wdg in enumerate(s["widgetlar"]):
                wdg.grid(row=r, column=c, sticky="w" if c < 2 else "", padx=(8 if c == 0 else 2, 2), pady=2)
        self.tuval.yview_moveto(0)
        self.bilgi.configure(text=f"{len(self.satirlar)} firma")

    def _sutunu_cevir(self, tip):
        gorunen = self.gorunen()
        yeni = not all(s["secim"][tip].get() for s in gorunen)
        for s in gorunen:
            s["secim"][tip].set(yeni)

    def _satiri_cevir(self, satir):
        secim = satir["secim"]
        yeni = not all(v.get() for v in secim.values())
        for v in secim.values():
            v.set(yeni)

    def secimler(self):
        return [{"ad": s["ad"], "devreden": s["dev"].get().strip(),
                 "ekranlar": {t for t, v in s["secim"].items() if v.get()}} for s in self.satirlar]

    def beyannameden_al(self, yollar=None):
        """KDV1 beyannamesi PDF'lerinden Devreden KDV'yi okuyup tabloya yazar (Kaydet'e kadar dosyaya yazilmaz).

        Kontrol edilen ayin kendi beyannamesinden "101 - Önceki Dönemden Devreden",
        bir onceki ayinkinden "Sonraki Döneme Devreden" alinir.
        """
        donem = self.arayuz.secili_donem()
        if not donem:
            messagebox.showwarning("Tarih aralığı", "Ana penceredeki tarih aralığını kontrol edin;"
                                   " devir hangi dönem için alınacak buradan anlaşılıyor.", parent=self.w)
            return
        hedef = tarih_cozumle(donem.split("-")[0])
        onceki = (hedef - timedelta(days=1)).replace(day=1)
        if yollar is None:
            yollar = filedialog.askopenfilenames(
                parent=self.w, filetypes=[("Beyanname (PDF)", "*.pdf"), ("Tüm dosyalar", "*.*")],
                title=f"KDV1 beyannameleri: {onceki:%m/%Y} ya da {hedef:%m/%Y} (birden çok seçilebilir)")
        if not yollar:
            return
        okunan, okunamayan = [], []
        self.w.configure(cursor="watch")
        self.w.update_idletasks()
        try:
            for yol in yollar:
                try:
                    okunan.append(beyanname.oku(yol))
                except beyanname.BeyannameDegil as e:
                    okunamayan.append((Path(yol).name, str(e)))
        except RuntimeError as e:  # pypdf kurulu degil
            messagebox.showerror("Beyanname okunamadı", str(e), parent=self.w)
            return
        finally:
            self.w.configure(cursor="")
        eslesen, eslesmeyen, donem_disi = beyanname.devirleri_bul(
            okunan, [s["ad"] for s in self.satirlar], hedef)
        satirlar = {s["ad"]: s for s in self.satirlar}
        tl = gostergeler.tl
        sonuc = []
        yazilan = 0
        for e in sorted(eslesen, key=lambda x: x["ad"]):
            if e["tutar"] is None:
                sonuc.append((e["ad"], e["b"].unvan, "Tutar okunamadı — elle yazın", e["kaynak"],
                              e["b"].donem, ""))
                continue
            satirlar[e["ad"]]["dev"].set(tl(e["tutar"]).replace(" TL", ""))
            sonuc.append((e["ad"], e["b"].unvan, "Yazıldı", e["kaynak"], e["b"].donem, tl(e["tutar"])))
            yazilan += 1
        for b in eslesmeyen:
            tutar = b.sonraki_devreden if b.bas == onceki else b.onceki_devreden
            sonuc.append(("— (elle yazın)", b.unvan or b.dosya, "Firma bulunamadı", b.dosya, b.donem,
                          tl(tutar) if tutar is not None else ""))
        for b in donem_disi:
            sonuc.append(("—", b.unvan or b.dosya, "Dönem tutmuyor", b.dosya, b.donem, ""))
        for ad, neden in okunamayan:
            sonuc.append(("—", ad, "Okunamadı", neden, "", ""))
        self.arayuz._detay_penceresi(
            f"Beyannameden devir — {yazilan} firmaya yazıldı",
            ("Tablodaki firma", "Beyannamedeki ad", "Durum", "Kaynak", "Dönem", "Devreden KDV"),
            sonuc, (f"{yazilan} / {len(yollar)} dosya yazıldı", "", "", "", "", ""),
            [f"Kontrol edilen dönem {hedef:%m/%Y}: {hedef:%m/%Y} beyannamesinden \"101 - Önceki Dönemden"
             f" Devreden\", {onceki:%m/%Y} beyannamesinden \"Sonraki Döneme Devreden\" alındı"
             " (ikisi aynı tutardır).",
             "Tutarlar tabloya yazıldı; kontrol edip Kaydet'e basın. Eşleşmeyen firmaların devrini"
             " elle yazın."], yazi_sutunu=4)
        return sonuc

    def kaydet(self):
        try:
            firma_tablosu.tabloyu_yaz(self.yol, self.secimler(), self.silinenler)
        except PermissionError:
            messagebox.showerror("Kaydedilemedi", f"{self.yol.name} Excel'de açık; kapatıp tekrar deneyin.",
                                 parent=self.w)
            return
        except Exception as e:
            messagebox.showerror("Kaydedilemedi", f"{type(e).__name__}: {e}", parent=self.w)
            return
        self.arayuz.gostergeleri_yenile()
        self.w.destroy()


class LucaListesiPenceresi:
    """Luca'dan cekilen firma listesinin firmalar.xlsx'ten farki; "Tabloya Uygula" ile dosyaya yazilir."""

    def __init__(self, arayuz, yol, yil, luca_sayisi, plan):
        self.arayuz, self.yol, self.plan = arayuz, yol, plan
        w = self.w = tk.Toplevel(arayuz.kok, bg=ZEMIN, padx=20, pady=16)
        w.title(f"Luca'dan Firma Listesi — {yil}")
        w.transient(arayuz.kok)
        w.geometry("860x560")
        tk.Label(w, text=f"Luca'da {yil} yılında {luca_sayisi} firma var", font=("Georgia", 14, "bold"),
                 fg=BASLIK_FG, bg=ZEMIN, anchor="w").pack(fill="x")
        tk.Label(w, text=f"{yol.name} ile karşılaştırma: {len(plan['yeni'])} yeni firma,"
                         f" {len(plan['guncellenecek'])} firmada açılış/kapanış güncellenecek,"
                         f" {plan['ayni']} firma aynı, {len(plan['luca_da_yok'])} firma Luca'nın {yil}"
                         " listesinde yok (listeden çıkarılabilir).", font=KUCUK, fg=SOLUK, bg=ZEMIN, anchor="w", justify="left",
                 wraplength=800).pack(fill="x", pady=(2, 10))

        alt = tk.Frame(w, bg=ZEMIN)
        alt.pack(side="bottom", fill="x", pady=(10, 0))
        self.v_yeni = tk.BooleanVar(value=True)
        self.v_sil = tk.BooleanVar(value=True)
        secenekler = tk.Frame(alt, bg=ZEMIN)
        secenekler.pack(side="left")
        tk.Checkbutton(secenekler, text="Yeni firmaları ekle (tüm ekranlar işaretli)", variable=self.v_yeni,
                       font=GOVDE, **Arayuz._kutu_renk()).pack(anchor="w")
        tk.Checkbutton(secenekler, text=f"Luca'nın {yil} listesinde olmayan firmaları listeden çıkar",
                       variable=self.v_sil, font=GOVDE, **Arayuz._kutu_renk()).pack(anchor="w")
        dugme(alt, "Vazgeç", w.destroy).pack(side="right")
        dugme(alt, "Tabloya Uygula", self.uygula, ana=True).pack(side="right", padx=8)
        tk.Label(w, text="Listede kalan firmaların ekran seçimleri ve Devreden KDV'leri değişmez;"
                         " yazmadan önce dosyanın yedeği alınır. Luca'da kapanışı boş olan firmanın tablodaki"
                         " kapanışı korunur; kapanış tarihi dönem sonuysa (31/12) firma açık sayılır.",
                 font=KUCUK, fg=SOLUK, bg=ZEMIN, anchor="w", justify="left", wraplength=800
                 ).pack(side="bottom", fill="x", pady=(4, 0))

        sutunlar = ("Durum", "Firma", "Açılış", "Kapanış")
        agac = self.agac = ttk.Treeview(w, columns=sutunlar, show="headings", style="Liste.Treeview")
        for s, g in zip(sutunlar, (250, 300, 100, 170)):
            agac.heading(s, text=s, anchor="w")
            agac.column(s, anchor="w", width=g, stretch=s == "Firma")
        for r in plan["yeni"]:
            agac.insert("", "end", values=("Yeni — eklenecek", r["ad"], r.get("acilis", ""),
                                           r.get("kapanis", "") or "—"))
        for g in plan["guncellenecek"]:
            eski, yeni = g["kapanis"]
            agac.insert("", "end", values=("Güncellenecek", g["ad"], g["acilis"][1] or "—",
                                           f"{eski or '—'} → {yeni or '—'}" if eski != yeni else yeni or "—"))
        for ad in plan["luca_da_yok"]:
            agac.insert("", "end", values=("Çıkarılacak (Luca'da yok)", ad, "", ""))
        agac.pack(fill="both", expand=True)

    def uygula(self):
        try:
            yedek = firma_tablosu.luca_plani_uygula(self.yol, self.plan, yeni_ekle=self.v_yeni.get(),
                                                     eksikleri_sil=self.v_sil.get())
        except PermissionError:
            messagebox.showerror("Kaydedilemedi", f"{self.yol.name} Excel'de açık; kapatıp tekrar deneyin.",
                                 parent=self.w)
            return
        except Exception as e:
            messagebox.showerror("Kaydedilemedi", f"{type(e).__name__}: {e}", parent=self.w)
            return
        self.arayuz.gostergeleri_yenile()
        self.arayuz._log_ekle(f"Luca firma listesi {self.yol.name} dosyasına uygulandı (yedek: {yedek.name})\n",
                              "ok")
        self.w.destroy()


# --- ana pencere ---------------------------------------------------------------

class Arayuz:
    BOT = KOK / "luca_bot.py"

    def __init__(self, kok):
        self.kok = kok
        self.surec = None
        self.okuyucu = None
        self.kuyruk = queue.Queue()
        self.yarim_satir = ""
        self.durdurma_istendi = False
        self.mod = "calisma"  # "firma_listesi": surec Luca'dan musteri listesi cekiyor
        self.liste_yili = None
        self.beyanname_penceresi = None  # "Luca'dan Devir Çek"i baslatan KDV Devri penceresi
        self.luca_penceresi = None  # son acilan "Luca'dan firma listesi" onizleme penceresi
        self.son_islem = ("", 0.0)  # (son log satiri, geldigi an): "su an ne yapiyor"
        self.gostergeler = {"tevkifat": [], "smm": [], "fark": [], "kdv": []}
        try:
            self.ayarlar = ayarlari_yukle()
        except ValueError as e:
            messagebox.showerror("Ayar dosyası bozuk", f"{e}\n\nDosyayı Not Defteri ile düzeltin.")
            raise SystemExit(1)

        kok.title("Dijital Stajyer")
        try:  # pencere/gorev cubugu simgesi
            self._simge = logo_resmi(256, kok) or logo_resmi(96, kok)
            if self._simge is not None:
                kok.iconphoto(True, self._simge)
        except tk.TclError:
            pass
        kok.geometry("1140x780")
        kok.minsize(1020, 720)
        kok.protocol("WM_DELETE_WINDOW", self.kapat)
        self._kur()
        if kurulum_mantigi.sihirbaz_gerekli(self.ayarlar):  # yeni kullanici: Luca bilgileri hic girilmemis
            kok.after(400, self.sihirbazi_ac)

    def _kur(self):
        """Tum arayuzu secili temayla (yeniden) kurar; tema degisince de cagrilir."""
        kok = self.kok
        tema_uygula(self.ayarlar.get("tema") or "acik")
        kok.configure(bg=ZEMIN)
        self._stiller()
        self._degiskenler()
        self._baslik()
        self._sekme_cubugu()
        self._alt_cubuk()  # govdeden once paketlenir ki her zaman en altta kalsin
        govde = tk.Frame(kok, bg=ZEMIN)
        self.fatura_govde = govde
        self._sol_panel(govde)
        tk.Frame(govde, bg=CIZGI, width=1).pack(side="left", fill="y")
        self._sag_panel(govde)
        self.kz = KarZararSekmesi(self, kok, sys.modules[__name__])
        self.muavin_govde = self._muavin_taslagi(kok)
        self.sekme_sec("fatura")
        self._giris_ozetini_yaz()
        self.alt_sekme_sec("surec")
        self.gostergeleri_yenile()
        self.dongu_id = kok.after(100, self._dongu)

    def sihirbazi_ac(self):
        """Yeni kullanici kurulum sihirbazi (Luca/Defter Beyan bilgileri, tercihler, firma listesi)."""
        if getattr(self, "sihirbaz", None) is not None:
            try:
                if self.sihirbaz.w.winfo_exists():
                    self.sihirbaz.w.lift()
                    return
            except tk.TclError:
                pass
        self.sihirbaz = KurulumSihirbazi(self, sys.modules[__name__])

    def tema_degistir(self):
        """Acik <-> koyu: ayar kaydedilir, arayuz ayni pencerede yeniden kurulur (calisma sirasinda yapilmaz)."""
        if self.surec or self.kz.calisiyor():
            messagebox.showinfo("Çalışma sürüyor", "Tema, çalışma bitince değiştirilebilir.", parent=self.kok)
            return
        a = self.ayarlar
        a["tema"] = "koyu" if TEMA == "acik" else "acik"
        a["baslangic_tarihi"], a["bitis_tarihi"] = self.v_bas.get().strip(), self.v_bit.get().strip()
        a["firma_listesi"] = self.v_liste.get()
        a["arayuz_ekranlar"] = [t for t in TUM_BELGELER if self.v_ekran[t].get()]
        try:
            ayarlari_kaydet(a)
        except OSError:
            pass
        self.kok.after_cancel(self.dongu_id)
        for w in self.kok.winfo_children():
            w.destroy()
        self._kur()

    # -- kurulum ---------------------------------------------------------------

    def _stiller(self):
        s = ttk.Style(self.kok)
        try:
            s.theme_use("clam")
        except tk.TclError:
            pass
        s.configure("Altin.Horizontal.TProgressbar", troughcolor=TROUGH, background=ALTIN,
                    bordercolor=TROUGH, lightcolor=ALTIN, darkcolor=ALTIN, thickness=8)
        s.configure("Liste.Treeview", background=PANEL, fieldbackground=PANEL, foreground=YAZI,
                    rowheight=26, font=GOVDE, bordercolor=CIZGI)
        s.configure("Liste.Treeview.Heading", background=KUTU, foreground=ALTIN_FG, font=BOLUM,
                    relief="flat")
        s.map("Liste.Treeview", background=[("selected", SECIM)], foreground=[("selected", YAZI)])

    def _degiskenler(self):
        a = self.ayarlar
        vars_bas, vars_bit = gecen_ay()
        self.v_bas = tk.StringVar(value=a.get("baslangic_tarihi") or vars_bas)
        self.v_bit = tk.StringVar(value=a.get("bitis_tarihi") or vars_bit)
        self.v_liste = tk.StringVar(value=a.get("firma_listesi") or "")
        self.v_firma = tk.StringVar()
        self.v_devam = tk.BooleanVar(value=True)
        secili = set(a.get("arayuz_ekranlar") or TUM_BELGELER)
        self.v_ekran = {t: tk.BooleanVar(value=t in secili) for t in TUM_BELGELER}
        self.v_hepsi = tk.BooleanVar(value=all(v.get() for v in self.v_ekran.values()))

    def _baslik(self):
        b = tk.Frame(self.kok, bg=BASLIK_ZEMIN, height=74)
        b.pack(fill="x")
        b.pack_propagate(False)
        sol = tk.Frame(b, bg=BASLIK_ZEMIN)
        sol.pack(side="left", padx=22)
        logo(sol).pack(side="left", pady=14)
        yazi = tk.Frame(sol, bg=BASLIK_ZEMIN)
        yazi.pack(side="left", padx=12)
        tk.Label(yazi, text="Dijital Stajyer", font=SERIF, fg=BASLIK_YAZI, bg=BASLIK_ZEMIN).pack(anchor="w")
        tk.Label(yazi, text="SMMM OFİSİ - DİJİTAL ASİSTAN", font=("Segoe UI", 8, "bold"),
                 fg=ALTIN, bg=BASLIK_ZEMIN).pack(anchor="w")
        sag = tk.Frame(b, bg=BASLIK_ZEMIN)
        sag.pack(side="right", padx=22)
        site = tk.Label(sag, text="smmmyilmaz.com", font=KUCUK, fg=ALTIN, bg=BASLIK_ZEMIN, cursor="hand2")
        site.pack(side="right", padx=(12, 0))
        site.bind("<Button-1>", lambda _e: webbrowser.open("https://smmmyilmaz.com"))
        tema = tk.Label(sag, text="🌙 Koyu tema" if TEMA == "acik" else "☀ Açık tema", font=KUCUK,
                        fg=BASLIK_YAZI, bg="#16213A", padx=8, pady=3, cursor="hand2")
        tema.pack(side="right", padx=(12, 0))
        tema.bind("<Button-1>", lambda _e: self.tema_degistir())
        self.durum_etiketi = tk.Label(sag, text="  Hazır  ", font=("Segoe UI", 9, "bold"),
                                      fg="#FFFFFF", bg=YESIL, padx=6, pady=2)
        self.durum_etiketi.pack(side="right", padx=(12, 0))
        tk.Label(sag, text="Luca · e-Fatura / e-Arşiv · Kâr / Zarar", font=KUCUK,
                 fg=BASLIK_ETIKET, bg=BASLIK_ZEMIN).pack(side="right")
        tk.Frame(self.kok, bg=CIZGI, height=1).pack(fill="x")

    def _sekme_cubugu(self):
        """Ust sekmeler: buyuk, kalin yazi; secili sekmenin altinda altin cizgi."""
        cubuk = tk.Frame(self.kok, bg=PANEL)
        cubuk.pack(fill="x")
        self.sekme_dugmeleri, self.sekme_cizgileri = {}, {}
        for anahtar, ad in (("fatura", "Fatura İndirme"), ("kz", "Kâr / Zarar"), ("muavin", "Muavin")):
            hucre = tk.Frame(cubuk, bg=PANEL, cursor="hand2")
            hucre.pack(side="left")
            d = tk.Label(hucre, text=ad, font=("Segoe UI", 12, "bold"), padx=28, pady=11, cursor="hand2", bg=PANEL)
            d.pack()
            cizgi = tk.Frame(hucre, bg=PANEL, height=4)
            cizgi.pack(fill="x")
            for w in (hucre, d, cizgi):
                w.bind("<Button-1>", lambda _e, k=anahtar: self.sekme_sec(k))
            d.bind("<Enter>", lambda _e, k=anahtar: self.aktif_sekme != k and self.sekme_dugmeleri[k].configure(fg=YAZI))
            d.bind("<Leave>", lambda _e, k=anahtar: self.aktif_sekme != k and self.sekme_dugmeleri[k].configure(fg=ETIKET))
            self.sekme_dugmeleri[anahtar], self.sekme_cizgileri[anahtar] = d, cizgi
        self.aktif_sekme = None
        tk.Frame(self.kok, bg=KENAR, height=1).pack(fill="x")

    def _alt_cubuk(self):
        """Pencerenin en alti: telif satiri ve surum (tiklayinca Hakkinda)."""
        c = tk.Frame(self.kok, bg=BASLIK_ZEMIN)  # her iki temada koyu (ust bantla ayni)
        c.pack(side="bottom", fill="x")
        tk.Label(c, text=telif(), font=KUCUK, fg=BASLIK_ETIKET, bg=BASLIK_ZEMIN, anchor="w"
                 ).pack(side="left", padx=18, pady=6)
        s = tk.Label(c, text=f"Sürüm {SURUM}", font=KUCUK, fg=ALTIN, bg=BASLIK_ZEMIN, cursor="hand2")
        s.pack(side="right", padx=18)
        s.bind("<Button-1>", lambda _e: self.hakkinda())

    def hakkinda(self):
        messagebox.showinfo(
            "Dijital Stajyer — Hakkında",
            f"Dijital Stajyer\nSürüm {SURUM}  ({SURUM_TARIHI})\n\n"
            "Luca e-Fatura / e-Arşiv indirme, kâr/zarar tahmini ve rapor aracı.\n\n"
            f"Lisans sahibi: Adem Yılmaz, Serbest Muhasebeci Mali Müşavir\n{telif()}", parent=self.kok)

    def _muavin_taslagi(self, kok):
        """Henuz yapilmadi: Luca muavin dokumu ile gelen/giden faturalari karsilastirma ekrani icin yer tutucu."""
        g = tk.Frame(kok, bg=ZEMIN)
        orta = tk.Frame(g, bg=ZEMIN)
        orta.place(relx=0.5, rely=0.42, anchor="center")
        tk.Label(orta, text="🛠", font=("Segoe UI Emoji", 40), fg=ALTIN_FG, bg=ZEMIN).pack()
        tk.Label(orta, text="Çalışma var", font=SERIF, fg=BASLIK_FG, bg=ZEMIN).pack(pady=(8, 4))
        tk.Label(orta, text="Luca mükerrer / eksik fatura tespiti ekranı yapım aşamasında.\n"
                 "Gelen/giden faturalar Luca muavin dökümüyle karşılaştırılacak.", font=GOVDE, fg=SOLUK,
                 bg=ZEMIN, justify="center").pack()
        return g

    def sekme_sec(self, anahtar):
        self.fatura_govde.pack_forget()
        self.kz.govde.pack_forget()
        self.muavin_govde.pack_forget()
        {"fatura": self.fatura_govde, "kz": self.kz.govde, "muavin": self.muavin_govde}[anahtar].pack(
            fill="both", expand=True)
        for k, d in self.sekme_dugmeleri.items():
            d.configure(fg=ALTIN_FG if k == anahtar else ETIKET)
            self.sekme_cizgileri[k].configure(bg=ALTIN if k == anahtar else PANEL)
        self.aktif_sekme = anahtar
        if anahtar == "kz":
            self.kz.sonucu_goster()  # fatura indirme sekmesinde yeni ay indirilmis olabilir

    def _sol_panel(self, govde):
        p = tk.Frame(govde, bg=ZEMIN, width=330, padx=22, pady=16)
        p.pack(side="left", fill="y")
        p.pack_propagate(False)

        bolum_basligi(p, "Luca Girişi").pack(fill="x")
        satir = tk.Frame(p, bg=ZEMIN)
        satir.pack(fill="x", pady=(6, 14))
        self.giris_ozeti = tk.Label(satir, text="", font=GOVDE, fg=YAZI, bg=ZEMIN, anchor="w",
                                    justify="left")
        self.giris_ozeti.pack(side="left", fill="x", expand=True)
        dugme(satir, "Değiştir", self.giris_penceresi).pack(side="right")
        dugme(p, "Kurulum Sihirbazı…", self.sihirbazi_ac).pack(fill="x", pady=(0, 12))

        bolum_basligi(p, "Firma Listesi").pack(fill="x")
        self.luca_liste_dugmesi = dugme(p, "Luca'dan Firma Listesini Çek…", self.firma_listesi_cek)
        self.luca_liste_dugmesi.pack(fill="x", pady=(6, 0))
        dugme(p, "KDV Devri ve Ekran Seçimi…", self.firma_ekran_penceresi).pack(fill="x", pady=(6, 0))
        tk.Label(p, text="Sadece bu firma (boş = listedeki hepsi)", font=KUCUK, fg=ETIKET,
                 bg=ZEMIN, anchor="w").pack(fill="x", pady=(10, 3))
        giris_kutusu(p, self.v_firma).pack(fill="x", ipady=4)

        bolum_basligi(p, "Tarih Aralığı").pack(fill="x", pady=(16, 0))
        t = tk.Frame(p, bg=ZEMIN)
        t.pack(fill="x", pady=(6, 0))
        for i, (ad, v) in enumerate((("Başlangıç", self.v_bas), ("Bitiş", self.v_bit))):
            k = tk.Frame(t, bg=ZEMIN)
            k.grid(row=0, column=i, sticky="ew", padx=(0, 8) if i == 0 else 0)
            tk.Label(k, text=ad, font=KUCUK, fg=ETIKET, bg=ZEMIN, anchor="w").pack(fill="x")
            e = giris_kutusu(k, v, genislik=12)
            e.pack(fill="x", ipady=4)
            e.bind("<FocusOut>", lambda _e: self.gostergeleri_yenile())
        t.columnconfigure(0, weight=1)
        t.columnconfigure(1, weight=1)

        bolum_basligi(p, "Ekranlar").pack(fill="x", pady=(16, 4))
        tk.Checkbutton(p, text="Hepsini Seç", variable=self.v_hepsi, command=self._hepsi_degisti,
                       font=GOVDE_KALIN, **self._kutu_renk()).pack(anchor="w")
        tk.Frame(p, bg=CIZGI, height=1).pack(fill="x", pady=4)
        g = tk.Frame(p, bg=ZEMIN)
        g.pack(fill="x")
        for i, tip in enumerate(TUM_BELGELER):
            tk.Checkbutton(g, text=EKRAN_ADLARI.get(tip, tip), variable=self.v_ekran[tip],
                           command=self._ekran_degisti, font=GOVDE, **self._kutu_renk()
                           ).grid(row=i // 2, column=i % 2, sticky="w")
        tk.Checkbutton(p, text="Bugün bitenleri atla (kaldığı yerden)", variable=self.v_devam,
                       font=KUCUK, **self._kutu_renk()).pack(anchor="w", pady=(12, 0))

        alt = tk.Frame(p, bg=ZEMIN)
        alt.pack(side="bottom", fill="x")
        self.calistir_dugmesi = dugme(alt, "Çalıştır", self.calistir, ana=True)
        self.calistir_dugmesi.pack(side="left", fill="x", expand=True, padx=(0, 6))
        self.durdur_dugmesi = dugme(alt, "Durdur", self.durdur, state="disabled")
        self.durdur_dugmesi.pack(side="left", fill="x", expand=True)

    @staticmethod
    def _kutu_renk():
        return dict(bg=ZEMIN, fg=YAZI, selectcolor=KUTU, activebackground=ZEMIN,
                    activeforeground=YAZI, highlightthickness=0, bd=0, anchor="w")

    def _alt_sekmeler(self, ebeveyn):
        """Sag panelin 'Süreç / Firma Durumu / Hatalı' alt sekmeleri."""
        cubuk = tk.Frame(ebeveyn, bg=ZEMIN)
        cubuk.pack(fill="x", padx=22, pady=(10, 0))
        self.alt_dugmeler = {}
        for anahtar, ad in (("surec", "Süreç"), ("durum", "Firma Durumu"), ("hata", "Hatalı / İnmeyen")):
            d = tk.Label(cubuk, text=ad, font=GOVDE_KALIN, padx=14, pady=6, cursor="hand2", bg=PANEL)
            d.pack(side="left", padx=(0, 4))
            d.bind("<Button-1>", lambda _e, k=anahtar: self.alt_sekme_sec(k))
            self.alt_dugmeler[anahtar] = d

    def alt_sekme_sec(self, anahtar):
        for k, c in self.alt_cerceveler.items():
            c.pack_forget()
        self.alt_cerceveler[anahtar].pack(fill="both", expand=True)
        for k, d in self.alt_dugmeler.items():
            d.configure(fg=ALTIN_FG if k == anahtar else SOLUK, bg=ZEMIN if k == anahtar else CIZGI)
        if anahtar != "surec":
            self.gostergeleri_yenile()

    def _sag_panel(self, govde):
        dis = tk.Frame(govde, bg=ZEMIN)
        dis.pack(side="left", fill="both", expand=True)
        self._alt_sekmeler(dis)
        p = tk.Frame(dis, bg=ZEMIN, padx=22, pady=16)
        self.alt_cerceveler = {"surec": p,
                               "durum": tk.Frame(dis, bg=ZEMIN, padx=22, pady=12),
                               "hata": tk.Frame(dis, bg=ZEMIN, padx=22, pady=12)}
        self._durum_sekmesi(self.alt_cerceveler["durum"])
        self._hata_sekmesi(self.alt_cerceveler["hata"])

        ust = tk.Frame(p, bg=ZEMIN)
        ust.pack(fill="x")
        self.ilerleme_etiketi = tk.Label(ust, text="Hazır. Ayarları kontrol edip Çalıştır'a basın.",
                                         font=GOVDE_KALIN, fg=YAZI, bg=ZEMIN, anchor="w")
        self.ilerleme_etiketi.pack(side="left")
        self.kalan_etiketi = tk.Label(ust, text="", font=KUCUK, fg=SOLUK, bg=ZEMIN)
        self.kalan_etiketi.pack(side="right")
        self.ilerleme = ttk.Progressbar(p, style="Altin.Horizontal.TProgressbar", maximum=100)
        self.ilerleme.pack(fill="x", pady=(8, 2))
        # bot ne yapiyor ve o adimda ne zamandir bekliyor: takilan adim gozle gorulsun
        self.son_islem_etiketi = tk.Label(p, text="", font=KUCUK, fg=SOLUK, bg=ZEMIN, anchor="w")
        self.son_islem_etiketi.pack(fill="x", pady=(0, 8))

        # alt kisim once yerlestirilir ki log alani kalan yeri doldursun
        alt = tk.Frame(p, bg=ZEMIN)
        alt.pack(side="bottom", fill="x", pady=(10, 0))
        dugme(alt, "Rapor Dosyasını Aç", self.raporu_ac).pack(side="left")
        dugme(alt, "İndirilenler Klasörü", self.klasoru_ac).pack(side="left", padx=8)
        self.durum_yazisi = tk.Label(alt, text="", font=KUCUK, fg=SOLUK, bg=ZEMIN)
        self.durum_yazisi.pack(side="right")

        kutular = tk.Frame(p, bg=ZEMIN)
        kutular.pack(side="bottom", fill="x", pady=(10, 0))
        self.donem_etiketi = tk.Label(kutular, text="", font=KUCUK, fg=SOLUK, bg=ZEMIN, anchor="w")
        self.donem_etiketi.pack(fill="x", pady=(0, 6))
        sira = tk.Frame(kutular, bg=ZEMIN)
        sira.pack(fill="x")
        self.kutu = {
            "tevkifat": OzetKutusu(sira, "Alış Tevkifat KDV", lambda: self.detay("tevkifat")),
            "smm": OzetKutusu(sira, "Alış SMM", lambda: self.detay("smm")),
            "fark": OzetKutusu(sira, "İnteraktif Farkı", lambda: self.detay("fark"), TURUNCU),
            "kdv": OzetKutusu(sira, "KDV Ödemesi", lambda: self.detay("kdv"), KIRMIZI),
            "hata": OzetKutusu(sira, "Hata / İnmeyen", lambda: self.detay("hata"), TURUNCU),
        }
        for i, k in enumerate(self.kutu.values()):
            k.grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else 8, 0))
            sira.columnconfigure(i, weight=1, uniform="kutu")

        cerceve = tk.Frame(p, bg=KONSOL, highlightthickness=1, highlightbackground=KONSOL_KENAR)
        cerceve.pack(fill="both", expand=True)
        self.log = tk.Text(cerceve, bg=KONSOL, fg=KONSOL_FG, font=KONSOL_YAZI, relief="flat",
                           wrap="none", padx=12, pady=10, insertbackground=YAZI, state="disabled",
                           highlightthickness=0, bd=0)
        kaydir = tk.Scrollbar(cerceve, command=self.log.yview, bg=KONSOL, troughcolor=KONSOL,
                              activebackground=KENAR, relief="flat", bd=0)
        self.log.configure(yscrollcommand=kaydir.set)
        kaydir.pack(side="right", fill="y")
        self.log.pack(side="left", fill="both", expand=True)
        for etiket, renk in (("ok", YESIL_FG), ("uyari", UYARI_FG), ("soluk", LOG_SOLUK),
                             ("hata", KIRMIZI), ("firma", LOG_FIRMA), ("bilgi", LOG_BILGI)):
            self.log.tag_configure(etiket, foreground=renk)

    # -- Firma Durumu / Hatali sekmeleri ---------------------------------------------------

    @staticmethod
    def durum_renkleri():
        return {"tamam": YESIL_FG, "uyari": TURUNCU, "hata": KIRMIZI, "soluk": SOLUK}

    def _durum_sekmesi(self, p):
        """rapor.xlsx'teki 'firma durumu' tablosu: arama, Dönem/Durum filtresi, basliga tiklayinca siralama."""
        tk.Label(p, text="Firma durumu — rapor.xlsx ile aynı veri; başlığa tıklayınca sıralanır, kutulardan süzülür.",
                 font=KUCUK, fg=SOLUK, bg=ZEMIN, anchor="w").pack(fill="x", pady=(0, 6))
        alt = tk.Frame(p, bg=ZEMIN)
        alt.pack(side="bottom", fill="x", pady=(8, 0))
        dugme(alt, "Rapor Dosyasını Aç", self.raporu_ac).pack(side="left")
        dugme(alt, "İndirilenler Klasörü", self.klasoru_ac).pack(side="left", padx=8)
        yazi = {0, 1, 2, 3, rapor.BASLIKLAR.index("Eksik/Fazla Faturalar"), rapor.BASLIKLAR.index("Not"),
                rapor.BASLIKLAR.index("Son İşlem")}
        genislik = [{0: 190, 1: 175, 2: 150, 3: 300}.get(i, 90) for i in range(len(rapor.BASLIKLAR))]
        for ad in ("Eksik/Fazla Faturalar", "Not", "Son İşlem"):
            genislik[rapor.BASLIKLAR.index(ad)] = 200
        self.durum_tablosu = SiralaFiltreTablosu(
            p, sys.modules[__name__], rapor.BASLIKLAR, genislik, yazi=yazi, filtreler=("Dönem", "Durum"),
            varsayilan_filtre={"Dönem": self.secili_donem() or "Tümü"}, yatay=True,
            etiketler=self.durum_renkleri())
        self.durum_tablosu.pack(fill="both", expand=True)
        self._durum_donemi = None

    def _hata_sekmesi(self, p):
        """Sorgulamada hata alan / faturasi inmeyen ekranlar; secilenler (ya da gorunenlerin hepsi) tekrar sorgulanir."""
        self.hata_bilgisi = tk.Label(p, text="", font=KUCUK, fg=SOLUK, bg=ZEMIN, anchor="w", justify="left",
                                     wraplength=760)
        self.hata_bilgisi.pack(fill="x", pady=(0, 6))
        alt = tk.Frame(p, bg=ZEMIN)
        alt.pack(side="bottom", fill="x", pady=(8, 0))
        dugme(alt, "Seçilenleri Tekrar Sorgula", lambda: self._hatalilari_sorgula(False), ana=True
              ).pack(side="left")
        dugme(alt, "Görünenlerin Hepsini Tekrar Sorgula", lambda: self._hatalilari_sorgula(True)
              ).pack(side="left", padx=8)
        dugme(alt, "Görünenleri Seç", lambda: self.hata_tablosu.hepsini_sec()).pack(side="left")
        self.hata_tablosu = SiralaFiltreTablosu(
            p, sys.modules[__name__], ("Firma", "Ekran", "Durum", "İnmeyen", "Not"), (200, 200, 170, 80, 300),
            yazi={0, 1, 2, 4}, filtreler=("Ekran", "Durum"), secim="extended", etiketler=self.durum_renkleri())
        self.hata_tablosu.pack(fill="both", expand=True)
        self.hata_listesi = []

    def _hatalilari_sorgula(self, gorunenlerin_hepsi):
        t = self.hata_tablosu
        sira = t.gorunen_indeksler() if gorunenlerin_hepsi else t.secili_indeksler()
        if not sira:
            messagebox.showinfo("Satır seçin", "Tekrar sorgulanacak satırları seçin (Ctrl / Shift ile birden çok)"
                                " ya da 'Görünenlerin Hepsini Tekrar Sorgula'yı kullanın.", parent=self.kok)
            return
        self.tekrar_sorgula([self.hata_listesi[i] for i in sira])

    def _tablolari_yenile(self, kayitlar, donem):
        """rapor.json'dan Firma Durumu ve Hatali tablolarini doldurur (her firmadan sonra cagrilir)."""
        satirlar, etiketler = [], []
        for firma in sorted(kayitlar):
            k = rapor._tamamla(dict(kayitlar[firma]))
            k.setdefault("firma", firma)
            satir = rapor._satir(k)
            durum = rapor._genel_durum(k["durumlar"])
            if durum.startswith(rapor.SORUNLU_DURUMLAR):
                etiket = "hata"
            elif satir[3]:
                etiket = "uyari"
            elif not durum or durum.startswith("bekliyor"):
                etiket = "soluk"
            else:
                etiket = "tamam"
            satirlar.append(tuple(satir))
            etiketler.append(etiket)
        self.durum_tablosu.doldur(satirlar, etiketler)
        if donem != self._durum_donemi:  # ana penceredeki tarih araligi degisti: Donem filtresi ona gecer
            self._durum_donemi = donem
            ad = donem if any(str(s[1]) == donem for s in satirlar) else "Tümü"
            self.durum_tablosu.filtre_ayarla("Dönem", ad)

        self.hata_listesi = list(self.gostergeler.get("hata", []))
        hata_satirlari = [(x["firma"], BELGE_ADLARI.get(x["ekran"], x["ekran"]), self._durum_metni(x),
                           x["inmeyen"] or "", x["not"]) for x in self.hata_listesi]
        self.hata_tablosu.doldur(hata_satirlari, ["hata" if str(x["durum"]).startswith("hata") else "uyari"
                                                  for x in self.hata_listesi])
        self.hata_bilgisi.configure(
            text=(f"Dönem: {donem.replace('-', ' – ')} (ana penceredeki tarih aralığı). " if donem else "")
                 + f"{len(self.hata_listesi)} ekranda sorun var. Satırları seçip tekrar sorgulatın;"
                   " yalnızca seçilen firma ve ekranlar çalışır, sonuçlar rapora işlenir.")

    # -- ayarlar ---------------------------------------------------------------

    def _giris_ozetini_yaz(self):
        a = self.ayarlar
        if a.get("uye_no") and a.get("kullanici_adi") and a.get("parola"):
            self.giris_ozeti.configure(text=f"{a['kullanici_adi']}  (üye {a['uye_no']})", fg=YAZI)
        else:
            self.giris_ozeti.configure(text="Giriş bilgisi girilmedi", fg=TURUNCU)

    def giris_penceresi(self):
        w = tk.Toplevel(self.kok, bg=ZEMIN, padx=24, pady=20)
        w.title("Luca Giriş Bilgileri")
        w.transient(self.kok)
        w.resizable(False, False)
        alanlar = [("Üye No", "uye_no", False), ("Kullanıcı Adı", "kullanici_adi", False),
                   ("Parola", "parola", True),
                   ("Doğrulama Anahtarı (iki aşamalı giriş, isteğe bağlı)", "dogrulama_anahtari", True)]
        degerler = {}
        for ad, anahtar, gizli in alanlar:
            tk.Label(w, text=ad, font=KUCUK, fg=ETIKET, bg=ZEMIN, anchor="w").pack(fill="x", pady=(8, 3))
            degerler[anahtar] = tk.StringVar(value=str(self.ayarlar.get(anahtar) or ""))
            giris_kutusu(w, degerler[anahtar], gizli=gizli, genislik=38).pack(fill="x", ipady=4)
        tk.Label(w, text="Bilgiler bu bilgisayardaki ayarlar.json dosyasına kaydedilir.",
                 font=KUCUK, fg=SOLUK, bg=ZEMIN).pack(anchor="w", pady=(12, 0))

        def kaydet():
            for anahtar, v in degerler.items():
                self.ayarlar[anahtar] = v.get().strip() if anahtar != "parola" else v.get()
            ayarlari_kaydet(self.ayarlar)
            self._giris_ozetini_yaz()
            w.destroy()

        alt = tk.Frame(w, bg=ZEMIN)
        alt.pack(fill="x", pady=(16, 0))
        dugme(alt, "Kaydet", kaydet, ana=True).pack(side="right")
        dugme(alt, "Vazgeç", w.destroy).pack(side="right", padx=8)
        w.grab_set()

    def _liste_tam_yolu(self):
        yol = self.v_liste.get().strip()
        if not yol:
            return None
        p = Path(yol)
        return p if p.is_absolute() else KOK / p

    def _listeyi_ayarla(self, yol):
        p = Path(yol)
        try:
            yol = str(p.relative_to(KOK))
        except ValueError:
            yol = str(p)
        self.v_liste.set(yol)
        self.ayarlar["firma_listesi"] = yol
        ayarlari_kaydet(self.ayarlar)
        self.gostergeleri_yenile()

    def firma_ekran_penceresi(self):
        yol = self._liste_tam_yolu()
        if not yol or not yol.exists():
            yol = KOK / "firmalar.xlsx"
            if not yol.exists():
                try:
                    firmalar = sorted(json.loads(self.rapor_yolu().read_text(encoding="utf-8")))
                except (OSError, ValueError):
                    firmalar = []
                if not messagebox.askyesno(
                        "Firma listesi yok",
                        "Henüz firma listesi seçilmedi. Yeni bir firmalar.xlsx oluşturulsun mu?\n"
                        + (f"Daha önce işlenen {len(firmalar)} firma hazır eklenir; "
                           if firmalar else "Firmaları açılan pencereden ekleyebilirsiniz; ")
                        + "başka firma ekleyip çıkarabilirsiniz.", parent=self.kok):
                    return
                firma_tablosu.sablon_olustur(yol, firmalar, ornek=False)
            self._listeyi_ayarla(yol)
        try:
            firmalar = firma_tablosu.tabloyu_oku(yol)
        except Exception as e:
            messagebox.showerror("Okunamadı", f"{yol.name} okunamadı: {e}", parent=self.kok)
            return
        return FirmaEkranPenceresi(self, yol, firmalar)

    def _hepsi_degisti(self):
        for v in self.v_ekran.values():
            v.set(self.v_hepsi.get())

    def _ekran_degisti(self):
        self.v_hepsi.set(all(v.get() for v in self.v_ekran.values()))

    # -- calistirma --------------------------------------------------------------

    def komut(self):
        """luca_bot.py komutu; hata varsa ValueError (kullaniciya gosterilir)."""
        bas, bit = self.v_bas.get().strip(), self.v_bit.get().strip()
        try:
            if tarih_cozumle(bit) < tarih_cozumle(bas):
                raise ValueError("Bitiş tarihi başlangıçtan önce olamaz.")
        except ValueError as e:
            if "olamaz" in str(e):
                raise
            raise ValueError("Tarihleri GG/AA/YYYY biçiminde yazın (örnek 01/09/2026).")
        secili = [t for t in TUM_BELGELER if self.v_ekran[t].get()]
        if not secili:
            raise ValueError("En az bir ekran seçin.")
        komut = [python_komutu(), "-u", str(self.BOT), "--bitince-kapat",
                 "--baslangic", bas, "--bitis", bit]
        if len(secili) == len(TUM_BELGELER):
            komut.append("--hepsi")
        else:
            for t in secili:
                komut += ["--belge-tipi", t]
        if self.v_firma.get().strip():
            komut += ["--firma", self.v_firma.get().strip()]
        if not self.v_devam.get():
            komut.append("--bastan")
        return komut

    def calistir(self):
        if self.surec:
            return
        a = self.ayarlar
        if not self._giris_tamam_mi():
            return
        try:
            komut = self.komut()
        except ValueError as e:
            messagebox.showwarning("Kontrol edin", str(e), parent=self.kok)
            return
        a["baslangic_tarihi"], a["bitis_tarihi"] = self.v_bas.get().strip(), self.v_bit.get().strip()
        a["firma_listesi"] = self.v_liste.get()
        a["arayuz_ekranlar"] = [t for t in TUM_BELGELER if self.v_ekran[t].get()]
        ayarlari_kaydet(a)

        self._baslat(komut, "Luca'ya giriş yapılıyor…")

    def firma_listesi_cek(self):
        """Luca'nin Yönetici > Müşteri Listesi ekranindan secili yilin firmalarini cekip tabloyla karsilastirir."""
        if self.surec or not self._giris_tamam_mi():
            return
        try:
            yil = tarih_cozumle(self.v_bas.get().strip()).year
        except ValueError:
            yil = date.today().year
        if not messagebox.askyesno(
                "Firmaları Luca'dan çek",
                f"Luca'nın Yönetici › Müşteri İşlemleri › Müşteri Listesi ekranından {yil} yılında"
                " açık olan firmalar (açılış/kapanış tarihleriyle) okunacak.\n\n"
                "Tarayıcı açılıp kendiliğinden kapanır; bitince firma listenizle farkı gösterilir,"
                " siz onaylamadan dosyaya bir şey yazılmaz. Devam edilsin mi?", parent=self.kok):
            return
        self.liste_yili = yil
        komut = [python_komutu(), "-u", str(self.BOT), "--bitince-kapat",
                 "--firma-listesi-cek", "--yil", str(yil)]
        self._baslat(komut, f"Luca'dan {yil} firma listesi çekiliyor…", mod="firma_listesi")

    def beyanname_cek(self, pencere):
        """Luca'nin GIB Beyanname Takip ekranindan KDV1 PDF'lerini indirir; bitince devirler pencereye yazilir."""
        if self.surec:
            messagebox.showinfo("Çalışıyor", "Önce çalışan işlem bitsin ya da Durdur'a basın.", parent=pencere.w)
            return
        if not self._giris_tamam_mi():
            return
        donem = self.secili_donem()
        if not donem:
            messagebox.showwarning("Tarih aralığı", "Ana penceredeki tarih aralığını kontrol edin;"
                                   " hangi dönemin devri alınacağı buradan anlaşılıyor.", parent=pencere.w)
            return
        hedef = tarih_cozumle(donem.split("-")[0])
        onceki = (hedef - timedelta(days=1)).replace(day=1)
        if not messagebox.askyesno(
                "Devri Luca'dan çek",
                f"Kontrol edilen dönem {hedef:%m/%Y}: Luca'nın Muhasebe › Beyannameler › GİB Beyanname Takip"
                f" ekranında {onceki:%m/%Y} dönemi, KDV1, Onaylanmış olarak listelenip tüm beyannameler tek"
                " seferde indirilecek ve devreden KDV'ler okunacak (birkaç dakika sürebilir).\n\n"
                "Devreden KDV'ler tabloya yazılır, Kaydet'e kadar dosyaya geçmez. Devam edilsin mi?",
                parent=pencere.w):
            return
        self.beyanname_penceresi = pencere
        komut = [python_komutu(), "-u", str(self.BOT), "--bitince-kapat", "--beyanname-cek",
                 "--baslangic", hedef.strftime(TARIH_BICIMI), "--bitis", donem.split("-")[1]]
        self._baslat(komut, "Luca'dan beyannameler alınıyor…", mod="beyanname")

    def _beyannameler_alindi(self):
        dosya = indirme_koku(self.ayarlar) / "luca-beyannameler.json"
        pencere, self.beyanname_penceresi = self.beyanname_penceresi, None
        try:
            veri = json.loads(dosya.read_text(encoding="utf-8"))
            yollar = [y for y in veri["dosyalar"] if Path(y).exists()]
        except (OSError, ValueError, KeyError):
            messagebox.showerror("Beyannameler okunamadı", f"{dosya.name} okunamadı.", parent=self.kok)
            return
        if pencere is None or not pencere.w.winfo_exists():
            messagebox.showinfo("Beyannameler alındı", f"{len(yollar)} beyanname PDF'i alındı; devirleri"
                                " yazmak için KDV Devri penceresinde \"Beyannameden Devir Al\"ı kullanın.",
                                parent=self.kok)
            return
        pencere.w.lift()
        pencere.beyannameden_al(yollar)

    def luca_listesini_goster(self):
        """Cekilen Luca listesini firmalar.xlsx ile karsilastirip onizleme penceresini acar."""
        dosya = indirme_koku(self.ayarlar) / musteri_listesi.MUSTERI_LISTESI_DOSYASI
        try:
            yil, kayitlar = musteri_listesi.oku(dosya)
        except ValueError as e:
            messagebox.showerror("Liste okunamadı", str(e), parent=self.kok)
            return None
        yol = self._liste_tam_yolu()
        if not yol or not yol.exists():
            yol = KOK / "firmalar.xlsx"
            if not yol.exists():
                firma_tablosu.sablon_olustur(yol, [], ornek=False)
            self._listeyi_ayarla(yol)
        try:
            plan = firma_tablosu.luca_plani(yol, yil, kayitlar)
        except Exception as e:
            messagebox.showerror("Karşılaştırılamadı", f"{yol.name} okunamadı: {e}", parent=self.kok)
            return None
        if not (plan["yeni"] or plan["guncellenecek"] or plan["luca_da_yok"]):
            messagebox.showinfo("Firma listesi", f"Luca'da {yil} yılında {len(kayitlar)} firma var;"
                                f" {yol.name} ile aynı, değişiklik yok.", parent=self.kok)
            return None
        self.luca_penceresi = LucaListesiPenceresi(self, yol, yil, len(kayitlar), plan)
        return self.luca_penceresi

    def _giris_tamam_mi(self):
        a = self.ayarlar
        if a.get("uye_no") and a.get("kullanici_adi") and a.get("parola"):
            return True
        messagebox.showwarning("Giriş bilgisi eksik",
                               "Program Luca'ya kendisi girer; önce Üye No, Kullanıcı Adı ve"
                               " Parola'yı girin.", parent=self.kok)
        self.giris_penceresi()
        return False

    def _baslat(self, komut, ilk_etiket, mod="calisma"):
        if self.kz.calisiyor():
            messagebox.showinfo("Başka çalışma sürüyor", "Kâr / Zarar çalışıyor; bitince başlatın "
                                "(ikisi aynı anda Luca'ya girmemeli).", parent=self.kok)
            return
        ortam = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUNBUFFERED="1")
        bayrak = 0
        if sys.platform.startswith("win"):
            # Bot pencere (konsol) acmadan calisir; arayuz de konsolsuz (pythonw) acilabilir.
            # "Durdur" bu yuzden sinyal degil durdur dosyasiyla calisir (bkz. durdur()).
            bayrak = subprocess.CREATE_NEW_PROCESS_GROUP | getattr(subprocess, "CREATE_NO_WINDOW", 0x08000000)
        try:
            DURDUR_DOSYASI.unlink()  # onceki calismadan kalmis istek yeni calismayi durdurmasin
        except OSError:
            pass
        self._log_temizle()
        self._log_ekle("Başlatılıyor…\n", "bilgi")
        try:
            self.surec = subprocess.Popen(komut, cwd=str(KOK), env=ortam, stdin=subprocess.DEVNULL,
                                          stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                          bufsize=0, creationflags=bayrak)
        except OSError as e:
            messagebox.showerror("Başlatılamadı", f"Program başlatılamadı: {e}", parent=self.kok)
            return
        self.durdurma_istendi = False
        self.yarim_satir = ""
        self.son_islem = ("", 0.0)
        self.okuyucu = threading.Thread(target=self._oku, args=(self.surec,), daemon=True)
        self.okuyucu.start()
        self.ilerleme.configure(value=0)
        self.ilerleme_etiketi.configure(text=ilk_etiket)
        self.kalan_etiketi.configure(text="")
        self._durum("Çalışıyor", ALTIN, ALTIN_YAZI)
        self.mod = mod
        self.calistir_dugmesi.configure(state="disabled")
        self.luca_liste_dugmesi.configure(state="disabled")
        self.durdur_dugmesi.configure(state="normal", text="Durdur")

    def _oku(self, surec):
        cozucu = codecs.getincrementaldecoder("utf-8")(errors="replace")
        while True:
            parca = surec.stdout.read(4096)
            if not parca:
                break
            self.kuyruk.put(cozucu.decode(parca))
        self.kuyruk.put(cozucu.decode(b"", final=True))

    def durdur(self):
        if not self.surec:
            return
        if self.durdurma_istendi:
            if messagebox.askyesno("Zorla kapat", "Program hâlâ kapanmadı. Zorla kapatılsın mı?\n"
                                   "(Son firmanın sonucu kaydedilmemiş olabilir.)", parent=self.kok):
                self.surec.kill()
            return
        self.durdurma_istendi = True
        # Asil yol: bot bekleme adimlarinda bu dosyayi gorup sonuclari kaydederek durur
        # (konsolsuz calismada da gecerli). Sinyal ek olarak gonderilir; ulasmazsa sorun degil.
        try:
            DURDUR_DOSYASI.write_text("durdur", encoding="utf-8")
        except OSError:
            pass
        try:
            self.surec.send_signal(signal.CTRL_BREAK_EVENT if sys.platform.startswith("win")
                                   else signal.SIGINT)
        except (OSError, ValueError):
            pass
        self._log_ekle("\nDurduruluyor… o ana kadarki sonuçlar kaydediliyor.\n", "uyari")
        self.durdur_dugmesi.configure(text="Zorla Kapat")

    def _dongu(self):
        metin = ""
        try:
            while True:
                metin += self.kuyruk.get_nowait()
        except queue.Empty:
            pass
        if metin:
            self._metni_isle(metin)
        # surec bittiginde borudaki son satirlar da okunmus olmali
        if (self.surec and self.surec.poll() is not None and not self.okuyucu.is_alive()
                and self.kuyruk.empty()):
            self._bitti(self.surec.returncode)
        self._son_islemi_goster()
        self.kz.tikla()
        self.dongu_id = self.kok.after(150, self._dongu)

    def _metni_isle(self, metin):
        satirlar = (self.yarim_satir + metin.replace("\r\n", "\n")).split("\n")
        self.yarim_satir = satirlar.pop()
        for satir in satirlar:
            self._satir(satir)

    def _satir(self, satir):
        sade = satir.strip()
        m = ILERLEME.match(sade)
        etiket = None
        if sade and set(sade) - set("=-"):
            self.son_islem = (sade, time.monotonic())
        if m:
            sira, toplam = int(m.group(1)), int(m.group(2))
            self.ilerleme_etiketi.configure(text=f"İşleniyor: {m.group(3)}  ({sira} / {toplam} firma)")
            self.kalan_etiketi.configure(text=f"Tahmini kalan: {m.group(4)}" if m.group(4) else "")
            self.ilerleme.configure(value=(sira - 1) * 100 / max(toplam, 1))
            self.gostergeleri_yenile()  # bir onceki firma rapora yazildi
            etiket = "firma"
        elif "[OK]" in sade:
            etiket = "ok"
        elif "[!!]" in sade or sade.startswith(("DIKKAT", "UYARI")) or "UYARI:" in sade:
            etiket = "uyari"
        elif "[--]" in sade or set(sade) <= set("=-") and sade:
            etiket = "soluk"
        elif "HATA" in sade:
            etiket = "hata"
        elif sade.startswith("Luca'da") or "LUCA GIRIS" in sade:
            self.ilerleme_etiketi.configure(text="Firma listesi okunuyor…")
        self._log_ekle(satir + "\n", etiket)

    def _bitti(self, kod):
        if self.yarim_satir:
            self._satir(self.yarim_satir)
            self.yarim_satir = ""
        self.surec.stdout.close()
        self.surec = None
        self.calistir_dugmesi.configure(state="normal")
        self.luca_liste_dugmesi.configure(state="normal")
        self.durdur_dugmesi.configure(state="disabled", text="Durdur")
        self.gostergeleri_yenile()
        mod, self.mod = self.mod, "calisma"
        if mod == "firma_listesi":
            if kod == 0 and not self.durdurma_istendi:
                self.ilerleme.configure(value=100)
                self.ilerleme_etiketi.configure(text="Luca'dan firma listesi alındı.")
                self._durum("Tamamlandı", YESIL, "#FFFFFF")
                self.luca_listesini_goster()
            elif self.durdurma_istendi or kod in (130, -2):
                self.ilerleme_etiketi.configure(text="Firma listesi çekme durduruldu.")
                self._durum("Durduruldu", TURUNCU_ZEMIN, ALTIN_YAZI)
            else:
                self.ilerleme_etiketi.configure(
                    text="Firma listesi alınamadı — log'a bakın (indirilenler\\…\\tani klasöründe ekran görüntüsü var).")
                self._durum("Hata", HATA_ZEMIN, "#FFFFFF")
        elif mod == "beyanname":
            if kod == 0 and not self.durdurma_istendi:
                self.ilerleme.configure(value=100)
                self.ilerleme_etiketi.configure(text="Luca'dan beyannameler alındı.")
                self._durum("Tamamlandı", YESIL, "#FFFFFF")
                self._beyannameler_alindi()
            elif self.durdurma_istendi or kod in (130, -2):
                self.ilerleme_etiketi.configure(text="Beyanname alma durduruldu.")
                self._durum("Durduruldu", TURUNCU_ZEMIN, ALTIN_YAZI)
            else:
                self.ilerleme_etiketi.configure(
                    text="Beyannameler alınamadı — log'a bakın (günlük klasörün tani klasöründe ekran görüntüsü var).")
                self._durum("Hata", HATA_ZEMIN, "#FFFFFF")
        elif kod == 0 and not self.durdurma_istendi:
            self.ilerleme.configure(value=100)
            self.ilerleme_etiketi.configure(text="Tamamlandı. Rapor güncellendi.")
            self._durum("Tamamlandı", YESIL, "#FFFFFF")
        elif self.durdurma_istendi or kod in (130, -2):
            self.ilerleme_etiketi.configure(text="Durduruldu. Yeniden Çalıştır'a basınca kaldığı yerden sürer.")
            self._durum("Durduruldu", TURUNCU_ZEMIN, ALTIN_YAZI)
        else:
            self.ilerleme_etiketi.configure(text="Hata ile bitti — log'un sonuna bakın.")
            self._durum("Hata", HATA_ZEMIN, "#FFFFFF")

    def _son_islemi_goster(self):
        """'Şu an: <son log satırı> — 14 sn' (çalışırken); uzun beklemede renk değişir."""
        metin, an = self.son_islem
        if not (self.surec and metin):
            yazi, renk = "", SOLUK
        else:
            gecen = int(time.monotonic() - an)
            yazi = f"Şu an: {metin[:110]}  —  {gecen} sn"
            renk = KIRMIZI if gecen >= UZUN_BEKLEME * 3 else TURUNCU if gecen >= UZUN_BEKLEME else SOLUK
        if self.son_islem_etiketi.cget("text") != yazi:
            self.son_islem_etiketi.configure(text=yazi, fg=renk)

    def _durum(self, metin, zemin, yazi):
        self.durum_etiketi.configure(text=f"  {metin}  ", bg=zemin, fg=yazi)

    # -- log ---------------------------------------------------------------------

    def _log_ekle(self, metin, etiket=None):
        alttaydi = self.log.yview()[1] >= 0.999
        self.log.configure(state="normal")
        self.log.insert("end", metin, (etiket,) if etiket else ())
        fazla = int(self.log.index("end-1c").split(".")[0]) - AZAMI_SATIR
        if fazla > 0:
            self.log.delete("1.0", f"{fazla + 1}.0")
        self.log.configure(state="disabled")
        if alttaydi:
            self.log.see("end")

    def _log_temizle(self):
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")

    def log_metni(self):
        return self.log.get("1.0", "end-1c")

    # -- ozet kutulari ------------------------------------------------------------

    def secili_donem(self):
        try:
            bas, bit = hedef_ay_araligi(tarih_cozumle(self.v_bas.get().strip()),
                                        tarih_cozumle(self.v_bit.get().strip()))
        except ValueError:
            return None
        return f"{bas:{TARIH_BICIMI}}-{bit:{TARIH_BICIMI}}"

    def rapor_yolu(self):
        return indirme_koku(self.ayarlar) / "rapor.json"

    def gostergeleri_yenile(self):
        donem = self.secili_donem()
        try:
            kayitlar = json.loads(self.rapor_yolu().read_text(encoding="utf-8"))
        except (OSError, ValueError):
            kayitlar = {}
        try:
            devreden = devreden_kdvleri(self.v_liste.get())
        except Exception:
            devreden = {}
        self.gostergeler = gostergeler.hesapla(kayitlar, devreden, donem)
        self._tablolari_yenile(kayitlar, donem)
        t = gostergeler.toplamlar(self.gostergeler)
        tl = gostergeler.tl
        self.kutu["tevkifat"].ayarla(tl(t["tevkifat"]), f"{t['tevkifat_adet']} fatura ›")
        self.kutu["smm"].ayarla(tl(t["smm"]), f"{t['smm_adet']} makbuz ›")
        self.kutu["fark"].ayarla(f"{t['fark']} fatura")
        self.kutu["kdv"].ayarla(f"{t['kdv_firma']} firma")
        self.kutu["hata"].ayarla(f"{t['hata']} ekran", "Gör / tekrarla ›")
        if donem:
            bas, bit = donem.split("-")
            self.donem_etiketi.configure(text=f"Dönem özeti: {bas} – {bit}   (iptal/itiraz edilen"
                                              " faturalar tutarlara dahil değildir)")

    @staticmethod
    def _durum_metni(x):
        durum = x["durum"]
        if durum == "tamam" or durum.startswith("tamam"):
            if x["inmeyen"]:
                return "Bazı faturalar inmedi"
            return {"tamam (excel eksik)": "Excel inmedi", "tamam (iptal eksik)": "İptal/itiraz eksik"}.get(
                durum, durum)
        return {"kaynaktan inmedi": "Kaynaktan inmedi", "dosya inmedi": "Belge paketi inmedi",
                "excel inmedi": "Excel inmedi", "ekran acilmadi": "Ekran açılmadı"}.get(durum, durum)

    def tekrar_sorgula(self, liste, pencere=None):
        """Hata alan firma/ekranlari (yalniz bunlari) ana penceredeki tarih aralığı için yeniden calistirir."""
        if self.surec:
            messagebox.showinfo("Çalışıyor", "Önce çalışan işlem bitsin ya da Durdur'a basın.", parent=self.kok)
            return
        if not self._giris_tamam_mi():
            return
        bas, bit = self.v_bas.get().strip(), self.v_bit.get().strip()
        try:
            if tarih_cozumle(bit) < tarih_cozumle(bas):
                raise ValueError
        except ValueError:
            messagebox.showwarning("Tarih aralığı", "Ana penceredeki tarih aralığını kontrol edin"
                                   " (GG/AA/YYYY).", parent=self.kok)
            return
        istek = {}
        for x in liste:
            istek.setdefault(x["firma"], []).append(x["ekran"])
        if not messagebox.askyesno(
                "Tekrar sorgula",
                f"{len(istek)} firmada {len(liste)} ekran yeniden sorgulanacak ({bas} – {bit}).\n"
                "Yalnızca bu ekranlar çalışır; sonuçlar rapora işlenir. Devam edilsin mi?", parent=self.kok):
            return
        yol = indirme_koku(self.ayarlar) / "tekrar-listesi.json"
        yol.write_text(json.dumps(istek, ensure_ascii=False, indent=1), encoding="utf-8")
        if pencere is not None:
            pencere.destroy()
        komut = [python_komutu(), "-u", str(self.BOT), "--bitince-kapat", "--baslangic", bas, "--bitis", bit,
                 "--tekrar-listesi", str(yol)]
        self._baslat(komut, "Hatalı ekranlar yeniden sorgulanıyor…")

    def detay(self, tur):
        tl = gostergeler.tl
        g = self.gostergeler[tur]
        if tur == "tevkifat":
            baslik = "Alış Tevkifat KDV — firmalar"
            sutunlar = ("Firma", "Fatura", "Tevkifat KDV")
            satirlar = [(x["firma"], x["adet"], tl(x["tutar"]) + (" *" if x["tahmini"] else ""))
                        for x in g]
            toplam = ("Toplam", sum(x["adet"] for x in g), tl(sum(x["tutar"] for x in g)))
            notlar = []
            if any(x["tahmini"] for x in g):
                notlar.append("* Excel'de tevkifat tutarı sütunu bulunamadı; faturanın KDV'si"
                              " gösterildi (tevkif edilen kısım orana göre bundan azdır).")
            notlar.append("KDV2 için: iptal/itiraz edilen tevkifatlı faturalar da sayıya dahildir,"
                          " tutara dahil değildir.")
        elif tur == "smm":
            baslik = "Alış SMM — firmalar"
            sutunlar = ("Firma", "Makbuz", "Tutar")
            satirlar = [(x["firma"], x["adet"], tl(x["tutar"])) for x in g]
            toplam = ("Toplam", sum(x["adet"] for x in g), tl(sum(x["tutar"] for x in g)))
            notlar = []
        elif tur == "fark":
            baslik = "İnteraktif V.D. − e-Arşiv Alış farkı"
            sutunlar = ("Firma / Eksik-fazla fatura", "İnteraktif", "e-Arşiv Alış", "Fark")
            satirlar = []
            for x in g:
                satirlar.append((x["firma"], x["interaktif"], x["earsiv"],
                                 f"{x['fark']:+d}  ({'e-Arşiv’de eksik' if x['fark'] > 0 else 'e-Arşiv’de fazla'})"))
                # fatura satiri: ismin ilk kelimesi, fatura no sonu, tutar
                for yazi, etiket in ((x["eksik"], "e-Arşiv’de eksik"), (x["fazla"], "e-Arşiv’de fazla")):
                    for unvan, no, *kalan in yazi:
                        tutar = kalan[0] if kalan else 0
                        satirlar.append((f"      ↳ {unvan or '?'}  {gostergeler.kisa_fatura_no(no)}"
                                         + (f"  {tl(tutar)}" if tutar else ""), "", "", etiket))
                if x["eslesmedi"]:
                    satirlar.append(("      ↳ iki listenin fatura numaraları tutmuyor; tek tek gösterilemiyor",
                                     "", "", ""))
                elif not x["eksik"] and not x["fazla"]:
                    satirlar.append(("      ↳ sayı farkı var ama hangi fatura olduğu belirlenemedi", "", "", ""))
            toplam = ("Toplam", sum(x["interaktif"] for x in g), sum(x["earsiv"] for x in g),
                      f"{sum(abs(x['fark']) for x in g)} fatura")
            notlar = ["Fatura satırı: ismin ilk kelimesi, fatura numarası (ilk 3 karakter..son 3 hane, örn. GIB..756) ve tutar."
                      " Tüm liste rapor.xlsx › İndirilen Faturalar sayfasında da yazar."]
            return self._detay_penceresi(baslik, sutunlar, satirlar, toplam, notlar, yazi_sutunu=1,
                                         genislik=[430, 90, 110, 220])
        elif tur == "hata":
            baslik = "Hata alınan / inmeyen ekranlar"
            sutunlar = ("Firma", "Ekran", "Durum", "İnmeyen", "Not")
            satirlar = [(x["firma"], BELGE_ADLARI.get(x["ekran"], x["ekran"]), self._durum_metni(x),
                         x["inmeyen"] or "", x["not"]) for x in g]
            toplam = ("Toplam", f"{len(g)} ekran", "", sum(x["inmeyen"] for x in g), "")
            notlar = ["İnmeyen: GİB'de olup Luca/kaynak sunucudan inmeyen fatura sayısı. \"Tekrar Sorgula\" yalnızca"
                      " bu firma ve ekranları, ana penceredeki tarih aralığı için yeniden çalıştırır."]
            return self._detay_penceresi(
                baslik, sutunlar, satirlar, toplam, notlar, yazi_sutunu={0, 1, 2, 4},
                genislik=[190, 190, 200, 80, 330],
                ek_dugme=("Bunları Tekrar Sorgula", lambda w, g=g: self.tekrar_sorgula(g, w)) if g else None)
        else:
            baslik = "KDV ödemesi çıkabilecek firmalar"
            sutunlar = ("Firma", "Satış KDV", "Alış KDV", "Devreden", "Tahmini Ödeme")
            satirlar = [(x["firma"], tl(x["satis"]), tl(x["alis"]),
                         tl(x["devreden"]) if x["devreden"] is not None else "girilmedi",
                         tl(x["odeme"])) for x in g]
            toplam = ("Toplam", "", "", "", tl(sum(x["odeme"] for x in g)))
            notlar = ["Tahmini: Satış KDV − Alış KDV − Devreden KDV (iptaller hariç). Diğer beyan"
                      " kalemleri hesaba girmez; kesin tutar değil, uyarıdır.",
                      "Devreden KDV, firma listesindeki \"Devreden KDV\" sütunundan okunur;"
                      " boşsa 0 sayılır."]
        self._detay_penceresi(baslik, sutunlar, satirlar, toplam, notlar)

    def _detay_penceresi(self, baslik, sutunlar, satirlar, toplam, notlar, yazi_sutunu=1,
                         genislik=None, ek_dugme=None):
        """yazi_sutunu: bastaki bu kadar sutun yazidir (sola yaslanir), kalanlar tutar (saga);
        sutun numaralari kumesi de verilebilir. genislik: sutun basina piksel (yoksa varsayilan).
        ek_dugme: (yazi, komut(pencere)) altta ayrica gosterilir."""
        yazilar = set(range(yazi_sutunu)) if isinstance(yazi_sutunu, int) else set(yazi_sutunu)
        if genislik is None:
            genislik = [260 if i == 0 else 170 if i in yazilar else 130 for i in range(len(sutunlar))]
        w = tk.Toplevel(self.kok, bg=ZEMIN, padx=20, pady=16)
        w.title(baslik)
        w.transient(self.kok)
        w.geometry(f"{min(sum(genislik) + 70, 1500)}x520")
        tk.Label(w, text=baslik, font=("Georgia", 14, "bold"), fg="#F2F4F8", bg=ZEMIN,
                 anchor="w").pack(fill="x")
        donem = self.secili_donem()
        if donem:
            tk.Label(w, text=f"Dönem: {donem.replace('-', ' – ')}", font=KUCUK, fg=SOLUK, bg=ZEMIN,
                     anchor="w").pack(fill="x", pady=(2, 10))
        alt = tk.Frame(w, bg=ZEMIN)
        alt.pack(side="bottom", fill="x", pady=(10, 0))
        for n in notlar:
            tk.Label(w, text=n, font=KUCUK, fg=SOLUK, bg=ZEMIN, anchor="w", justify="left",
                     wraplength=max(sum(genislik), 600)).pack(side="bottom", fill="x", pady=(4, 0))
        cerceve = tk.Frame(w, bg=ZEMIN)
        cerceve.pack(fill="both", expand=True)
        agac = ttk.Treeview(cerceve, columns=sutunlar, show="headings", style="Liste.Treeview")
        kaydir = ttk.Scrollbar(cerceve, orient="vertical", command=agac.yview)
        agac.configure(yscrollcommand=kaydir.set)
        for i, s in enumerate(sutunlar):
            yon = "w" if i in yazilar else "e"
            agac.heading(s, text=s, anchor=yon)
            # ilk sutun artan yeri alir; digerleri sabit ve sigacak genislikte
            agac.column(s, anchor=yon, width=genislik[i], minwidth=genislik[i], stretch=(i == 0))
        agac.tag_configure("toplam", foreground=ALTIN_FG, font=GOVDE_KALIN)
        agac.tag_configure("alt", foreground=ETIKET)
        if not satirlar:
            agac.insert("", "end", values=("Bu dönemde kayıt yok",) + ("",) * (len(sutunlar) - 1))
        for s in satirlar:
            agac.insert("", "end", values=s, tags=("alt",) if str(s[0]).startswith("      ") else ())
        if satirlar:
            agac.insert("", "end", values=toplam, tags=("toplam",))
        kaydir.pack(side="right", fill="y")
        agac.pack(side="left", fill="both", expand=True)

        def excele_indir():
            ad = re.sub(r"[^\w-]+", "-", sadelestir(baslik.split("—")[0]).lower()).strip("-")
            yol = filedialog.asksaveasfilename(
                title="Excel olarak kaydet", parent=w, defaultextension=".xlsx",
                initialdir=str(indirme_koku(self.ayarlar)),
                initialfile=f"{ad}-{(donem or '').replace('/', '.')}.xlsx",
                filetypes=[("Excel", "*.xlsx")])
            if not yol:
                return
            try:
                liste_excel_yaz(yol, baslik, donem, sutunlar, satirlar, toplam, notlar)
            except PermissionError:
                messagebox.showerror("Kaydedilemedi", "Dosya Excel'de açık; kapatıp tekrar deneyin.",
                                     parent=w)
                return
            dosya_ac(yol)

        dugme(alt, "Kapat", w.destroy).pack(side="right")
        dugme(alt, "Excel olarak indir", excele_indir, ana=not ek_dugme).pack(side="right", padx=8)
        if ek_dugme:
            dugme(alt, ek_dugme[0], lambda: ek_dugme[1](w), ana=True).pack(side="right")
        return w

    # -- dosyalar / kapanis -------------------------------------------------------

    def raporu_ac(self):
        yol = indirme_koku(self.ayarlar) / "rapor.xlsx"
        if not yol.exists():
            messagebox.showinfo("Rapor yok", "Henüz rapor oluşmadı; önce bir çalıştırma yapın.",
                                parent=self.kok)
            return
        dosya_ac(yol)

    def klasoru_ac(self):
        dosya_ac(indirme_koku(self.ayarlar))

    def kapat(self):
        if self.kz.calisiyor():
            if not messagebox.askyesno("Çalışma sürüyor", "Kâr / Zarar çalışıyor. Durdurup çıkılsın mı?",
                                       parent=self.kok):
                return
            self.kz.kapat()
        if self.surec:
            if not messagebox.askyesno("Çalışma sürüyor", "Program çalışıyor. Durdurup çıkılsın mı?",
                                       parent=self.kok):
                return
            self.durdur()
            try:
                self.surec.wait(timeout=20)
            except subprocess.TimeoutExpired:
                self.surec.kill()
        self.kok.after_cancel(self.dongu_id)
        self.kok.destroy()


def main():
    if konsolsuz_yeniden_baslat():
        return
    if sys.platform.startswith("win"):
        try:  # yuksek cozunurluklu ekranlarda bulanik yazi olmasin
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    kok = tk.Tk()
    try:
        Arayuz(kok)
        kok.mainloop()
    except SystemExit:
        raise
    except Exception:
        import traceback
        hata = traceback.format_exc()
        (KOK / "arayuz-hata.log").write_text(hata, encoding="utf-8")
        messagebox.showerror("Beklenmeyen hata", f"{hata[-1500:]}\n\nAyrıntı: arayuz-hata.log")
        raise


if __name__ == "__main__":
    main()
