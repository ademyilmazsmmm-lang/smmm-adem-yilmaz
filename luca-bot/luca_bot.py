#!/usr/bin/env python3
"""Luca portalindan tum firmalar icin e-fatura / e-arsiv belgelerini toplu ceker ve indirir."""

import argparse
import csv
import json
import os
import re
import shutil
import sys
import zipfile
import threading
import time
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path

from playwright.sync_api import sync_playwright


import rapor  # gunluk toplu rapor (rapor.xlsx / rapor.csv)
import eposta  # calisma bitince ozet e-postasi

KOK = Path(__file__).resolve().parent
AYAR_DOSYASI = KOK / "ayarlar.json"
ORNEK_AYAR = KOK / "ayarlar.ornek.json"

GIRIS_URL = "https://www.luca.com.tr"  # uygulama adresine dogrudan gidilince "LUCA HATA" veriyor
GIRIS_SAYFASI = "https://agiris.luca.com.tr/LUCASSO/giris.erp"  # ortak giris ekrani
UYGULAMA_PARCASI = "/Luca/"
SISTEM_GIRIS = ["Sistem Giriş", "Sistem Girisi", "Sistem Giris"]
GIRIS_DUGMESI = ["GİRİŞ", "Giriş", "GIRIS", "Giris"]
URUN_ADAYLARI = ["LUCA MALİ MÜŞAVİR PAKETİ", "Mali Müşavir Paketi",
                 "MALİ MÜŞAVİR PAKETİ", "Luca Mali Müşavir", "MALİ MÜŞAVİR"]
# Iki asamali dogrulama ekranini taniyan yazilar
DOGRULAMA_ISARETLERI = ["İki Aşamalı Doğrulama", "Doğrulama Kodu", "Güvenlik Kodu",
                        "SMS ile gönderilen", "Tek Kullanımlık Şifre"]
# Iki asamali dogrulama KAPALIYKEN Luca captcha soruyor; bot bunu gecemez
CAPTCHA_ISARETLERI = ["Captcha doğrulaması", "Captcha doğrulama", "güvenlik kontrolünü"]
DOGRULAMA_ONAY = ["Tamam", "Doğrula", "Onayla", "GİRİŞ", "Giriş"]
UST_MENU = "Akıllı Entegrasyon Noktası"
MODUL_ADAYLARI = ["İşletme Defteri", "Ser.Mes.Defteri", "Serbest Meslek Defteri",
                  "Basit Usül", "Basit Usul", "Genel Muhasebe", "Bilanço Defteri",
                  "Muhasebe", "Defter"]

BELGE_TIPLERI = {
    "e-arsiv-alis": "e-Arşiv Alış Faturaları",
    "e-arsiv-satis": "e-Arşiv Satış Faturaları",
    "e-fatura-alis": "e-Fatura Alış Faturaları",
    "e-fatura-satis": "e-Fatura Satış Faturaları",
    "gib-5000": "GİB 5000/30000",
    "turmob-alis": "TÜRMOB Ent. Alış Faturaları",
    "turmob-satis": "TÜRMOB Ent. Satış Faturaları",
    "esmm-alis": "GİB e-SMM Alış",
    "esmm-satis": "GİB e-SMM Satış",
    # Akilli Entegrasyon Noktasi altinda degil, modul menusunun kendisinde:
    "e-arsiv-interaktif": "E-Arşiv Faturaları Sorgulama",
}

# menusu iki kademeli olan (Akilli Entegrasyon Noktasi araciligi olmayan) ekranlar
IKI_KADEMELI = {"e-arsiv-interaktif"}

# "GİB'den İptal/İtiraz Sorgula" butonu yalnizca bu ekranlarda var; digerlerinde
# butonu aramak her ekranda yarim dakika bosa gidiyordu
IPTAL_EKRANLARI = {"e-arsiv-alis", "e-arsiv-interaktif"}

# Belge (XML/zip) indirilmeyen ekranlar. Excel her ekranda iniyor: tevkifatli
# faturalar (KDV2) ancak Excel'deki sutunlardan guvenilir sekilde gorulebiliyor.
SADECE_EXCEL = {"e-arsiv-satis", "e-fatura-alis", "e-fatura-satis",
                "turmob-alis", "turmob-satis", "esmm-alis", "esmm-satis"}

# "--hepsi" ile calistirilacak sira: alis/satis ekranlari, en sonda karsilastirma
TUM_BELGELER = ["e-arsiv-alis", "e-arsiv-satis", "e-fatura-alis", "e-fatura-satis",
                "gib-5000", "turmob-alis", "turmob-satis", "esmm-alis", "esmm-satis",
                "e-arsiv-interaktif"]

# Interaktif Vergi Dairesi ekrani
INTERAKTIF_SORGU = "İnteraktif V.D'sinden E-Arşiv Faturalarını Sorgula"
INTERAKTIF_CAPASI = "Luca Proxy ile Sorgula"  # secenek penceresinin kendi yazisi
INTERAKTIF_SERVIS = "GİB Servis ile Sorgula"
INTERAKTIF_LISTELE = "Mevcut E-Arşiv Faturalarını Listele"
INTERAKTIF_TEKRAR = 2  # Luca ilk sorguda hep getirmiyor, iki kez calistiriliyor
# Ekranin altinda "1 / 1 (Toplam Kayit Sayisi: 7)" yazar; sorgunun bitip
# listeyi doldurdugunu anlamanin en guvenilir yolu bu
KAYIT_SAYISI_DESENI = re.compile(r"Toplam\s*Kay[ıi]t\s*Say[ıi]s[ıi]\s*[:=]?\s*(\d+)", re.I)

KAPAT_METINLERI = ["Bir daha gösterme"]  # sayfadaki "Tamam"/"Kapat" baska islevlere ait olabiliyor
DIYALOG_ONAY = ["Belgeleri Getir", "Sorgula", "Onayla", "Uygula"]
INDIRME_ONAY = ["Seçilenleri İndir", "Belgeleri İndir", "Dosyaları İndir", "İndir", "Onayla"]
ISLEM_BITTI = "sona erdi"
INDIRILEMEDI = "indirilemedi"
ISLEM_ISARETLERI = ["İşlem Takip", "sorgulandı", "belge kaydı bulundu", "Otomatik aşağı kaydır"]
# GIB'e ulasilamadiginda Luca bu uyariyi verip bekliyor; bosuna beklememek icin
GIB_HATA_ISARETLERI = ["VERILER GETIRILIRKEN HATA", "GIB INTERNET SITESINDEN",
                       "GIB INTERNET E-ARSIV", "ERISILEMEDI", "BAGLANTI KURULAMADI"]
# Firmanin o servise abonesi/yetkisi yoksa Luca bu SOAP hatasini yaziyor ve
# pencere hic kapanmiyor; ayni ekranin kalan tarih araliklarini denemek bos
YETKI_ISARETLERI = ["IZNINIZ BULUNMAMAKTADIR", "YETKINIZ BULUNMAMAKTADIR",
                    "SOAP FAULT", "YETKISIZ ISLEM"]
YETKI_YOK = -2  # islem_takibini_bekle bu ekranin atlanmasi gerektigini boyle soyler


def ekran_tamamlanmis_mi(durum):
    """rapor.json'daki bir ekran durumu, o ekranin o gun bittigi anlamina mi geliyor.

    Bilgisayar/tarayici yarida kapanip calisma yeniden baslatildiginda, zaten
    tamamlanmis ekranlar tekrar acilmasin diye kullanilir. "hata: ..." ve
    "kaynaktan inmedi" gercek basarisizlik oldugu icin tamamlanmis sayilmaz;
    calisma yeniden baslayinca o ekranlar tekrar denenir.
    """
    return bool(durum) and (durum.startswith("tamam") or durum.startswith("atlandi")
                             or durum in ("fatura yok", "donem disi"))


def gib_hatasi(metin):
    duz = sadelestir(metin or "")
    return any(isaret in duz for isaret in GIB_HATA_ISARETLERI)


def yetki_hatasi(metin):
    duz = sadelestir(metin or "")
    return any(isaret in duz for isaret in YETKI_ISARETLERI)

KISAYOLLAR = {  # butonlarin kendi ipuclarinda yazan kisayollar (tiklama engellenirse kullanilir)
    "GİB'den Getir": "Alt+g",
    "Seçilenleri İndir": "Alt+z",
    "Yenile": "Alt+l",
    "Excel": "Alt+e",
    "Belge Seç": "Alt+b",
}

AYAR = {"azami_saniye": 900, "durgunluk_saniye": 180, "indirme_saniye": 30, "iptal_itiraz": True, "donem_degistir": True, "chrome_gunlugu": False, "tarayici": None, "profil_yerel": False, "indirmeyi_yakala": True, "azami_fatura": 500}  # ayarlar.json ile degistirilebilir

TARIH_BICIMI = "%d/%m/%Y"
AZAMI_GUN = 7  # GIB sorgusu tek seferde en fazla 7 gun kabul ediyor (eskiden 30)
# Aylik sorguyu kabul eden ekranlar 7 gunluk parcalamaya gerek duymuyor
AYLIK_AZAMI_GUN = 30
AYLIK_SORGU = {"e-arsiv-interaktif", "gib-5000"}

# Ayni faturalar iki ekranda da goruldugu icin toplam sayilirken bir kez
# sayilmali; diger ekranlar (satis, e-fatura, e-SMM ...) ayri belgeler
ORTUSEN_EKRANLAR = {"e-arsiv-alis", "e-arsiv-interaktif"}
COKME_DENEMESI = 3  # tarayici indirme sirasinda cokerse firma kac kez tekrar denensin

# GIB'den iptal/itiraz sorgulama (fatura listesi indikten sonra calisir)
IPTAL_DUGME_ADAYLARI = ["GİB'den İptal/İtiraz Sorgula", "GİB'den iptal/itiraz Sorgula",
                        "iptal/itiraz Sorgula", "İptal/İtiraz Sorgula"]
IPTAL_ONAY = ["İptal/İtiraz Sorgula", "İptal/itiraz Sorgula", "İptal/İtiraz sorgula"]
IPTAL_CAPASI = "Raporlanma Tarihi"  # diyalogun kendi aciklama yazisi
IPTAL_DESENI = re.compile(r"IPTAL|ITIRAZ")
TEVKIFAT_DESENI = re.compile(r"TEVKIFAT")


def yaz(mesaj, log_dosyasi=None):
    print(mesaj, flush=True)
    if log_dosyasi:
        with open(log_dosyasi, "a", encoding="utf-8") as f:
            f.write(f"{time.strftime('%H:%M:%S')} {mesaj}\n")


def sadelestir(metin):
    """Turkce karakter ve buyuk/kucuk harf farkini yok sayarak karsilastirma icin."""
    metin = metin.replace("ı", "i").replace("İ", "i")
    metin = unicodedata.normalize("NFKD", metin)
    metin = "".join(c for c in metin if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", metin).strip().upper()


def karsilastir(metin):
    """Arama icin: Turkce/buyuk-kucuk farki ve bosluklar yok sayilir ('yakups' -> 'YAKUP SÖĞÜ')."""
    return sadelestir(metin).replace(" ", "")


def dosya_adi_yap(metin):
    metin = metin.replace("ı", "i").replace("İ", "I")
    metin = unicodedata.normalize("NFKD", metin)
    metin = "".join(c for c in metin if not unicodedata.combining(c))
    metin = re.sub(r"[^A-Za-z0-9._ -]", "_", metin).strip()
    return re.sub(r"\s+", " ", metin) or "isimsiz"


def tarih_cozumle(metin):
    return datetime.strptime(metin.strip(), TARIH_BICIMI).date()


def hedef_ay_araligi(baslangic, bitis):
    """Indirme/listeleme araligi: baslangic tarihinin ayi.

    Agustos faturalari eylulde de duzenlenebildigi icin kullanici
    01/08/2026-15/09/2026 gibi bir aralik veriyor. Eylul tarihleri yalnizca
    gec duzenlenen agustos faturalarini GIB'den cekmek icin sorgulanir;
    indirme ve listeleme agustosun tamami uzerinden yapilir.

    Sorgu araligi ayin sonundan once bitse bile (orn. 01/08-02/08 gibi kisa bir
    deneme) listeleme yine 01/08-31/08 olur: aksi halde daha once GIB'den
    cekilmis agustos faturalari "Belge Ara" suzgecine takilip listede
    gorunmuyordu.
    """
    ay_sonu = (baslangic.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    return baslangic, ay_sonu


def tarih_araliklari(baslangic, bitis, gun=AZAMI_GUN):
    """AZAMI_GUN gunluk parcalara boler; her parca bir oncekinin bitis tarihinden basliyor."""
    if bitis < baslangic:
        raise ValueError("Bitis tarihi baslangictan once olamaz")
    araliklar = []
    su_an = baslangic
    while su_an < bitis:
        son = min(su_an + timedelta(days=gun), bitis)
        araliklar.append((su_an.strftime(TARIH_BICIMI), son.strftime(TARIH_BICIMI)))
        su_an = son
    if not araliklar:
        araliklar.append((baslangic.strftime(TARIH_BICIMI), bitis.strftime(TARIH_BICIMI)))
    return araliklar


def icinde_bulunulan_ay():
    bugun = date.today()
    bas = bugun.replace(day=1)
    bit = (bas + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    return bas, bit


def ayarlari_oku():
    kaynak = AYAR_DOSYASI if AYAR_DOSYASI.exists() else ORNEK_AYAR
    with open(kaynak, encoding="utf-8") as f:
        return json.load(f)


def cerceveler(page):
    """Luca eski bir Struts uygulamasi; ekranlar farkli frame'lere dagilmis olabiliyor."""
    return list(page.frames)


def bul(page, kurucu, sure=15000, gorunur=True):
    bitis = time.time() + sure / 1000
    son_hata = None
    while time.time() < bitis:
        for fr in cerceveler(page):
            try:
                loc = kurucu(fr)
                if loc.count() == 0:
                    continue
                ilk = loc.first
                if not gorunur or ilk.is_visible():
                    return fr, ilk
            except Exception as e:  # frame gezinme sirasinda kopabiliyor
                son_hata = e
        page.wait_for_timeout(300)
    raise LookupError(f"Ogeye ulasilamadi (son hata: {son_hata})")


def metinle_bul(page, metin, sure=15000):
    """Once butonun kendisi, sonra tam metin, en son kapsayici aranir.

    'Tamam' ve 'Vazgec' ayni kapsayicida oldugu icin kapsayiciya tiklamak
    yanlislikla Vazgec'e denk gelip secimi iptal ediyordu.
    """
    # ilk iki yontem kisa tutulur; bulunamayan her metin icin 3 kez tam sure
    # beklemek firma+menu adimlarinda saniyeler kaybettiriyordu
    pay = max(300, sure // 4)
    son_hata = None
    for kurucu, kurucu_sure in (
        (lambda f: f.get_by_role("button", name=metin, exact=True), pay),
        (lambda f: f.get_by_text(metin, exact=True), pay),
        (lambda f: f.get_by_text(metin, exact=False), sure),
    ):
        try:
            return bul(page, kurucu, sure=kurucu_sure)
        except LookupError as e:
            son_hata = e
    raise son_hata


def varsa_tikla(page, metinler, sure=1500):
    for metin in metinler:
        try:
            _, loc = metinle_bul(page, metin, sure=sure)
            loc.click()
            page.wait_for_timeout(400)
            return metin
        except Exception:
            continue
    return None


MENU_KELIMELERI = ("ISLEMLERI", "ISLEMLER", "BEYANNAME", "RAPOR", "LISTESI", "HESAP PLANI",
                   "MAKBUZ", "DEFTER", "HIZLI ERISIM", "SORGULAMA", "FATURALARI")


def menu_listesi_mi(secenekler):
    isabet = sum(1 for s in secenekler if any(k in sadelestir(s) for k in MENU_KELIMELERI))
    return isabet >= 3


def firma_adaylari(page):
    adaylar = []
    for fr in cerceveler(page):
        try:
            secimler = fr.locator("select")
            for i in range(secimler.count()):
                sec = secimler.nth(i)
                if not sec.is_visible():
                    continue
                temiz = [s.strip() for s in sec.locator("option").all_inner_texts() if s.strip()]
                if len(temiz) < 5:
                    continue
                adaylar.append((fr, sec, temiz))
        except Exception:
            continue
    return adaylar


def secili_metin(sec):
    try:
        return sec.evaluate("el => el.selectedIndex >= 0 ? el.options[el.selectedIndex].text : ''") or ""
    except Exception:
        return ""


def firma_secici(page):
    """Firma listesi, secili degeri sayfa basliginda gecen liste (baslik ornegi: 'AKIN COBAN [ 2026 ]')."""
    adaylar = firma_adaylari(page)
    if not adaylar:
        raise LookupError("Firma listesi (select) bulunamadi")

    try:
        baslik = sadelestir(page.title())
    except Exception:
        baslik = ""
    if baslik:
        for aday in adaylar:
            secili = sadelestir(secili_metin(aday[1]))
            if len(secili) >= 3 and secili in baslik:
                return aday

    firma_gibi = [a for a in adaylar if not menu_listesi_mi(a[2])]
    if firma_gibi:
        return max(firma_gibi, key=lambda a: len(a[2]))
    raise LookupError("Firma listesi ayirt edilemedi")


def acik_pencere(page):
    """Luca diyaloglari kapanmazsa arkadaki butonlar tiklanamiyor (pointer events engelleniyor)."""
    for fr in cerceveler(page):
        try:
            loc = fr.locator(".luca-open-window")
            if loc.count() and loc.first.is_visible():
                return fr, loc.first
        except Exception:
            continue
    return None, None


def diyalogda_tikla(page, metinler, sure=4000):
    _, pencere = indirme_diyalogu(page)
    return pencerede_tikla(page, pencere, metinler, sure)


DIYALOG_CAPASI = "Tüm faturaları seçmek için"
FATURA_YOK_CAPASI = "fatura bulunamadı"


def metinli_diyalog(page, capa):
    """Diyalogu sinif adina degil, icindeki yaziya gore bulur."""
    for fr in cerceveler(page):
        for secici in ("div", "table", "form"):
            try:
                loc = fr.locator(secici).filter(has_text=capa)
                if loc.count() and loc.last.is_visible():
                    return fr, loc.last
            except Exception:
                continue
    return None, None


def pencerede_tikla(page, pencere, metinler, sure=4000):
    """Butona pencerenin icinden basar (yan yana duran Tamam/Iptal karismasin diye)."""
    if pencere is None:
        return None
    for metin in metinler:
        for kurucu in (lambda: pencere.get_by_role("button", name=metin, exact=True),
                       lambda: pencere.get_by_text(metin, exact=True),
                       lambda: pencere.get_by_text(metin, exact=False)):
            try:
                loc = kurucu()
                if loc.count() and loc.last.is_visible():
                    loc.last.click(timeout=sure)
                    page.wait_for_timeout(600)
                    return metin
            except Exception:
                continue
    return None


# Sorgu bitince Luca bazen Islem Takip yerine bilgi penceresi gosteriyor
BILGI_CAPALARI = ("adet fatura bulundu", "fatura bulundu", "işlem tamamlandı",
                  "sorgulama tamamlandı", "kayıt bulunamadı")


def bilgi_penceresini_kapat(page):
    """Sorgu sonucu bildiren pencereyi kapatir ve yazisini dondurur.

    Bu pencere ciktiginda sorgu bitmistir; Islem Takip penceresini beklemeye
    devam etmek bosuna zaman kaybi oluyordu.
    """
    for capa in BILGI_CAPALARI:
        fr, pencere = metinli_diyalog(page, capa)
        if pencere is None:
            continue
        try:
            metin = " ".join((pencere.inner_text() or "").split())[:80]
        except Exception:
            metin = capa
        if not pencerede_tikla(page, pencere, ["Tamam", "Kapat"]):
            varsa_tikla(page, ["Tamam"], sure=2000)
        page.wait_for_timeout(500)
        return metin
    return ""


def fatura_yok_penceresini_kapat(page):
    """'Her hangi bir fatura bulunamadi' penceresi Tamam beklerken akisi kilitliyor."""
    _, pencere = metinli_diyalog(page, FATURA_YOK_CAPASI)
    if pencere is None:
        return False
    if not pencerede_tikla(page, pencere, ["Tamam"]):
        varsa_tikla(page, ["Tamam"], sure=2000)
    page.wait_for_timeout(500)
    return True


def indirme_diyalogu(page):
    fr, pencere = metinli_diyalog(page, DIYALOG_CAPASI)
    if pencere is not None:
        return fr, pencere
    _, pencere = acik_pencere(page)
    return (None, pencere) if pencere is not None else (None, None)


def diyalogda_tumunu_sec(page):
    """Indirme penceresindeki 'Tum faturalari secmek icin buraya' baglantisina tiklar.

    Pencerede iki 'buraya' var; ilki tum faturalar, ikincisi yalnizca
    onaylanmis faturalar icin.
    """
    _, pencere = indirme_diyalogu(page)
    if pencere is None:
        return False
    for kurucu in (
        lambda: pencere.get_by_text(DIYALOG_CAPASI, exact=False).first.get_by_text("buraya", exact=False),
        lambda: pencere.get_by_role("link", name="buraya", exact=True),
        lambda: pencere.get_by_text("buraya", exact=True),
        lambda: pencere.get_by_text("buraya", exact=False),
    ):
        try:
            loc = kurucu()
            if loc.count() and loc.first.is_visible():
                loc.first.click(timeout=4000)
                page.wait_for_timeout(1000)
                return True
        except Exception:
            continue
    return False


def acik_pencereleri_kapat(page, log=None):
    for deneme in range(4):
        fr, pencere = acik_pencere(page)
        if pencere is None:
            return True
        if deneme == 0 and log:
            yaz("    Acik Luca penceresi kapatiliyor", log)
        if varsa_tikla(page, ["Kapat"], sure=1500):
            page.wait_for_timeout(600)
            continue
        kapandi = False
        for secici in ("[class*='close']", "[class*='kapat']", "[onclick*='close']", "[onclick*='Kapat']"):
            try:
                dugme = fr.locator(f".luca-open-window {secici}").first
                if dugme.count() and dugme.is_visible():
                    dugme.click(timeout=3000)
                    kapandi = True
                    break
            except Exception:
                continue
        if not kapandi:
            try:
                page.keyboard.press("Escape")
            except Exception:
                pass
        page.wait_for_timeout(700)
    return acik_pencere(page)[1] is None


def firma_dogrula(page, firma_adi, sure=8000):
    """Secim gerceklesti mi: sekme basligi secili firmanin adini tasiyor (orn. 'DENTAL [ 2026 ]')."""
    hedef = sadelestir(firma_adi)
    bitis = time.time() + sure / 1000
    while time.time() < bitis:
        try:
            if hedef and hedef in sadelestir(page.title()):
                return True
        except Exception:
            pass
        page.wait_for_timeout(500)
    return False


DONEM_DESENI = re.compile(r"(\d{2}[./]\d{2}[./]\d{4})\s*-\s*(\d{2}[./]\d{2}[./]\d{4})")


def calisma_donemi(page):
    """Firma adinin altindaki donem kutusu (orn. '08/04/2022 - 31/12/2022')."""
    for fr in cerceveler(page):
        try:
            kutular = fr.locator("select")
            for i in range(min(kutular.count(), 12)):
                eslesme = DONEM_DESENI.search(secili_metin(kutular.nth(i)) or "")
                if eslesme:
                    return (tarih_cozumle(eslesme.group(1).replace(".", "/")),
                            tarih_cozumle(eslesme.group(2).replace(".", "/")))
        except Exception:
            continue
    return None, None


def donem_secici(page):
    """Calisma donemi listesi: secenekleri '01/01/2026 - 31/12/2026' bicimindeki select."""
    for fr in cerceveler(page):
        try:
            kutular = fr.locator("select")
            for i in range(min(kutular.count(), 12)):
                sec = kutular.nth(i)
                secenekler = [m.strip() for m in sec.locator("option").all_inner_texts()]
                if any(DONEM_DESENI.search(m) for m in secenekler):
                    return sec, secenekler
        except Exception:
            continue
    return None, []


def donem_araligi(metin):
    eslesme = DONEM_DESENI.search(metin or "")
    if not eslesme:
        return None, None
    return (tarih_cozumle(eslesme.group(1).replace(".", "/")),
            tarih_cozumle(eslesme.group(2).replace(".", "/")))


def donem_ortusuyor(bas, bit, istenen_bas, istenen_bit):
    """Donem okunamadiysa (None) engel cikarilmaz, sorgu denenir."""
    if bit and bit < istenen_bas:
        return False
    if bas and bas > istenen_bit:
        return False
    return True


def uygun_donem(secenekler, istenen_bas, istenen_bit):
    """Istenen yila ait donemi secer.

    Ayni yil farkli firmalarda farkli basliyor: yeni kurulanda
    '15/04/2026 - 31/12/2026', eskide '01/01/2026 - 31/12/2026'. Bu yuzden
    once yil tutturulur, gun/ay onemli degil.
    """
    adaylar = []
    for metin in secenekler:
        bas, bit = donem_araligi(metin)
        if bit and donem_ortusuyor(bas, bit, istenen_bas, istenen_bit):
            adaylar.append((metin, bit))
    if not adaylar:
        return None
    for yil in (istenen_bit.year, istenen_bas.year):
        for metin, bit in adaylar:
            if bit.year == yil:
                return metin
    return adaylar[0][0]


def donem_ayarla(page, firma_adi, istenen_bas, istenen_bit, log=None):
    """Firma eski donemde acilmissa uygun donemi secer; yoksa False doner.

    Luca her firmada en son kullanilan donemi hatirliyor. Acik bir firma 2025
    doneminde kalmis olabiliyor; bunlari atlamak yerine donemi degistirmek
    gerekiyor. Gercekten kapanmis firmalarda uygun donem secenegi bulunmaz.
    """
    bas, bit = calisma_donemi(page)
    if donem_ortusuyor(bas, bit, istenen_bas, istenen_bit):
        return True

    if not AYAR["donem_degistir"]:
        yaz(f"    Firma donemi {bas:%d/%m/%Y}-{bit:%d/%m/%Y}, donem degistirme kapali", log)
        return False

    sec, secenekler = donem_secici(page)
    if sec is None:
        yaz("    Donem listesi bulunamadi, sorgu yine de denenecek", log)
        return True

    uygun = uygun_donem(secenekler, istenen_bas, istenen_bit)
    if uygun is None:
        yaz(f"    Firma donemi {bas:%d/%m/%Y}-{bit:%d/%m/%Y}, uygun donem secenegi yok", log)
        return False

    yaz(f"    Donem degistiriliyor: {uygun}", log)
    for _ in range(2):
        try:
            sec.select_option(label=uygun, timeout=10000)
        except Exception:
            pass
        page.wait_for_timeout(600)
        varsa_tikla(page, ["Tamam"], sure=3000)
        page.wait_for_timeout(2000)
        if donem_ortusuyor(*calisma_donemi(page), istenen_bas, istenen_bit):
            if firma_dogrula(page, firma_adi, sure=6000):
                return True
            yaz("    UYARI: donem degisti ama firma dogrulanamadi, atlandi", log)
            return False
        acik_pencereleri_kapat(page)
        sec, _ = donem_secici(page)
        if sec is None:
            break
    yaz("    UYARI: donem degistirilemedi", log)
    return False


def firma_sec(page, firma_adi, log=None):
    """Firmanin bulundugu listeyi adiyla secer; secim 'Tamam' ile onaylanip dogrulanir."""
    # giris sonrasi acik kalan bilgi penceresi Tamam'a basilmasini engelliyordu
    varsa_tikla(page, KAPAT_METINLERI, sure=600)
    acik_pencereleri_kapat(page)

    hedef = None
    for deneme in range(3):  # onceki firmadan sonra liste gec yuklenebiliyor
        for _, sec, secenekler in firma_adaylari(page):
            if firma_adi in secenekler:
                hedef = sec
                break
        if hedef is not None:
            break
        acik_pencereleri_kapat(page)
        page.wait_for_timeout(3000)
    if hedef is None:
        raise LookupError(f"'{firma_adi}' acik listelerin hicbirinde bulunamadi")

    for _ in range(3):
        try:
            hedef.select_option(label=firma_adi, timeout=15000)
        except Exception:
            pass
        page.wait_for_timeout(600)
        varsa_tikla(page, ["Tamam"], sure=3000)
        page.wait_for_timeout(1500)
        if firma_dogrula(page, firma_adi):
            break
        varsa_tikla(page, KAPAT_METINLERI, sure=1200)
        sayfayi_toparla(page)  # gorunmez diyalog Tamam'i engelliyor olabilir
        page.wait_for_timeout(1000)
    else:
        # dogrulanmadan devam edilirse baska firmanin faturalari cekilir; bu firmayi atla
        raise LookupError(f"'{firma_adi}' secimi onaylanamadi (Tamam gecmedi), firma atlandi")

    try:
        page.wait_for_load_state("networkidle", timeout=5000)
    except Exception:
        pass


def gorunur_mu(page, metin, sure=1500):
    """Sadece varlik kontrolu; tiklama icin kullanilmadigindan tek (hizli) arama yeter."""
    try:
        bul(page, lambda f: f.get_by_text(metin, exact=False), sure=sure)
        return True
    except LookupError:
        return False


_SON_MODUL = None  # ilk firmada calisan modul adi; sonraki firmalarda once bu denenir


def modul_menusunu_ac(page, dogrula):
    """Modul menusunu (Isletme Defteri vb.) acar.

    Dokuz aday sirayla denenince her ekranda saniyeler gidiyordu; bir kez
    calisan ad hatirlanip sonraki cagrilarda ilk sirada deneniyor.
    """
    global _SON_MODUL
    adaylar = list(MODUL_ADAYLARI)
    one_al = [m for m in (_SON_MODUL, ekrandaki_modul(page)) if m in adaylar]
    for m in reversed(one_al):  # ekranda gorunen en one, sonra son calisan
        adaylar.remove(m)
        adaylar.insert(0, m)
    for modul in adaylar:
        if menu_ogesini_ac(page, modul, sure=1200, dogrula=dogrula):
            _SON_MODUL = modul
            return True
    return False


def menu_ogesini_ac(page, metin, sure=4000, dogrula=None):
    """Eski Luca menuleri kimi yerde hover, kimi yerde tiklama ile aciliyor."""
    try:
        _, oge = metinle_bul(page, metin, sure=sure)
    except LookupError:
        return False
    for eylem in ("hover", "click"):
        try:
            getattr(oge, eylem)()
        except Exception:
            continue
        page.wait_for_timeout(700)
        if dogrula is None or dogrula():
            return True
    return bool(dogrula()) if dogrula else True


def dugmeye_bas(page, metin, sure=8000):
    """Once butona tiklar; tiklama engellenirse butonun kendi klavye kisayolunu dener."""
    try:
        _, dugme = metinle_bul(page, metin, sure=sure)
        dugme.click(timeout=8000)
        return True
    except Exception:
        pass
    kisayol = KISAYOLLAR.get(metin)
    if kisayol:
        try:
            page.keyboard.press(kisayol)
            page.wait_for_timeout(600)
            return True
        except Exception:
            pass
    return False


# Gorunur menu yazilari tek JS cagrisiyla alinir; oge basina is_visible +
# inner_text cagirmak cerceve basina yuzlerce gidis donus demekti
MENU_METNI_JS = """els => els.map(e => ({
  t: ((e.innerText || '').trim()).slice(0, 60),
  g: !!(e.offsetParent || e.getClientRects().length)
}))"""


def menu_metinleri(page, sinir=40):
    bulunan = []
    for fr in cerceveler(page):
        try:
            ogeler = fr.locator("a, td, span")
            bilgiler = ogeler.evaluate_all(MENU_METNI_JS)
        except Exception:
            continue
        for b in bilgiler:
            metin = (b.get("t") or "").strip()
            if b.get("g") and 3 <= len(metin) <= 40 and metin not in bulunan:
                bulunan.append(metin)
                if len(bulunan) >= sinir:
                    return bulunan
    return bulunan


def ekrandaki_modul(page):
    """Ust cubukta gercekten duran modul adi; yoksa None.

    Dokuz adayi tek tek tiklamayi denemek yerine once ekranda hangisinin
    yazdigina bakilir: bir JS cagrisi, dokuz bosa arama yerine.
    """
    try:
        gorunen = {karsilastir(m) for m in menu_metinleri(page, sinir=80)}
    except Exception:
        return None
    for modul in MODUL_ADAYLARI:
        k = karsilastir(modul)
        if k and any(g.startswith(k) for g in gorunen):
            return modul
    return None


def _cerceve_imzasi(page):
    """Acik cercevelerin adresleri: ekran degisti mi anlamak icin."""
    try:
        return tuple(sorted(f.url for f in page.frames))
    except Exception:
        return ()


def ekran_hazir_bekle(page, isaret, onceki_imza=None, azami=6000):
    """Menu tiklandiktan sonra ekranin kendi butonu gorunene kadar bekler.

    Eskiden burada sabit bekleme + networkidle vardi; ekran 1 saniyede acilsa
    bile her menude 6-8 saniye harcaniyordu. Artik ekran hazir olur olmaz
    devam edilir, bulunamazsa eski sureye kadar beklenir.
    """
    basla = time.time()
    while True:
        page.wait_for_timeout(200)
        gecen = (time.time() - basla) * 1000
        if gecen >= azami:
            return False
        # "GİB'den Getir" her ekranda var; yeni ekran acilmadan onceki ekranin
        # butonunu gorup devam edersek yanlis ekrani sorgulariz. Once cerceve
        # adreslerinin degismesi beklenir; 2 saniye sonra bu kontrol birakilir
        # (kimi ekran ayni adrese yukleniyor).
        if onceki_imza and gecen < 2000 and _cerceve_imzasi(page) == onceki_imza:
            continue
        if gorunur_mu(page, isaret, sure=400):
            return True


def modul_menusunden_git(page, hedef, isaret=None):
    """Modul menusu (orn. Isletme Defteri) -> madde. Ara menu yok."""
    hedef_gorunur = lambda: gorunur_mu(page, hedef, sure=1200)
    for deneme in range(3):
        if not hedef_gorunur():
            modul_menusunu_ac(page, hedef_gorunur)
        try:
            _, madde = metinle_bul(page, hedef, sure=8000)
        except LookupError:
            page.wait_for_timeout(1000)
            continue
        onceki_imza = _cerceve_imzasi(page)
        madde.click(timeout=8000)
        if not ekran_hazir_bekle(page, isaret or hedef, onceki_imza, azami=8000):
            try:
                page.wait_for_load_state("networkidle", timeout=3000)
            except Exception:
                pass
        return
    raise LookupError(f"'{hedef}' menu maddesi bulunamadi. Gorunen menuler: {menu_metinleri(page)}")


def menuye_git(page, belge_tipi):
    hedef = BELGE_TIPLERI[belge_tipi]
    # Ekran acildiginda mutlaka gorunen buton: bekleme bunu gorunce biter
    isaret = INTERAKTIF_LISTELE if belge_tipi in IKI_KADEMELI else "GİB'den Getir"
    if belge_tipi in IKI_KADEMELI:
        return modul_menusunden_git(page, hedef, isaret)
    ust_gorunur = lambda: gorunur_mu(page, UST_MENU, sure=1200)
    hedef_gorunur = lambda: gorunur_mu(page, hedef, sure=1200)

    alt = None
    for deneme in range(3):  # menu kimi zaman hover'da acilip hemen kapaniyor
        if not ust_gorunur():
            if not modul_menusunu_ac(page, ust_gorunur):
                if deneme == 2:
                    raise LookupError(
                        f"'{UST_MENU}' menusu acilamadi. Sayfada gorunen menuler: {menu_metinleri(page)}"
                    )
                page.wait_for_timeout(1000)
                continue

        # ust menu her seferinde acilir: ekran basligi da hedef metni icerebildigi
        # icin "zaten gorunuyor" kontrolune guvenip menuyu atlamak yanlis ekrana
        # tiklamaya yol aciyor
        menu_ogesini_ac(page, UST_MENU, sure=6000, dogrula=hedef_gorunur)
        try:
            _, alt = metinle_bul(page, hedef, sure=4000)
            break
        except LookupError:
            page.wait_for_timeout(1000)

    if alt is None:
        raise LookupError(f"'{hedef}' menu maddesi bulunamadi. Gorunen menuler: {menu_metinleri(page)}")
    onceki_imza = _cerceve_imzasi(page)
    alt.click()
    if not ekran_hazir_bekle(page, isaret, onceki_imza, azami=6000):
        try:
            page.wait_for_load_state("networkidle", timeout=3000)
        except Exception:
            pass


# Tum kutularin degeri/nitelikleri tek seferde okunur: her kutu icin ayri ayri
# is_visible/input_value/get_attribute cagirmak ekran basina 10-25 sn suruyordu.
KUTU_BILGISI = """els => els.map(el => ({
  d: el.value || '',
  n: ((el.name || '') + ' ' + (el.id || '') + ' ' + (el.className || '')).toLowerCase(),
  g: !!(el.offsetParent || el.getClientRects().length)
}))"""

TARIH_NITELIGI = re.compile(r"tarih|date")


_SON_KUTU_HATASI = ""  # tani icin: tek JS cagrisi neden tutmadi


def _kutu_bilgileri(kutular):
    """Kutularin degeri/nitelikleri: once tek JS cagrisi, olmazsa tek tek.

    Luca cerceveleri sik sik yeniden yuklendigi icin JS calistirma
    "execution context destroyed" ile dusebiliyor; o zaman yavas ama
    saglam olan eski yola donulur.
    """
    global _SON_KUTU_HATASI
    try:
        bilgiler = kutular.evaluate_all(KUTU_BILGISI)
        if bilgiler:
            return bilgiler
    except Exception as e:
        _SON_KUTU_HATASI = type(e).__name__
    try:
        sayi = min(kutular.count(), 80)
    except Exception as e:
        _SON_KUTU_HATASI = type(e).__name__
        return []
    bilgiler = []
    for i in range(sayi):
        try:
            kutu = kutular.nth(i)
            nitelik = " ".join(x for x in (kutu.get_attribute("name"), kutu.get_attribute("id"),
                                           kutu.get_attribute("class")) if x)
            bilgiler.append({"d": kutu.input_value() or "", "n": nitelik.lower(),
                             "g": kutu.is_visible()})
        except Exception:
            bilgiler.append({"d": "", "n": "", "g": False})
    return bilgiler


def _tarih_kutulari(kapsayici):
    try:
        kutular = kapsayici.locator("input[type=text], input:not([type])")
    except Exception:
        return []
    return [kutular.nth(i) for i, b in enumerate(_kutu_bilgileri(kutular))
            if b.get("g") and (TARIH_DESENI.search(b.get("d") or "")
                               or TARIH_NITELIGI.search(b.get("n") or ""))]


def tarih_kutulari(page, kapsam=None, sure=6000):
    """Tarih kutulari; kapsam (acik pencere) verilirse once orada aranir.

    Tum sayfa taranirsa ekranin arkasindaki suzgec kutulari one geciyor ve
    tarih acik pencereye degil listeye yaziliyordu. Pencere hazir olmadan
    bakilirsa kutu bulunamadigi icin kisa araliklarla tekrar denenir.
    """
    bitis = time.time() + sure / 1000
    while True:
        adaylar = _tarih_kutulari(kapsam) if kapsam is not None else []
        if len(adaylar) < 2:
            hepsi = []
            for fr in cerceveler(page):
                hepsi.extend(_tarih_kutulari(fr))
            if len(hepsi) > len(adaylar):
                adaylar = hepsi
        if len(adaylar) >= 2 or time.time() >= bitis:
            return adaylar
        page.wait_for_timeout(400)


def kutuya_yaz(kutu, deger):
    try:
        kutu.fill(deger)
        if (kutu.input_value() or "").strip() == deger:
            return True
    except Exception:
        pass
    try:
        kutu.click()
        kutu.press("Control+a")
        kutu.type(deger, delay=40)
        if (kutu.input_value() or "").strip() == deger:
            return True
    except Exception:
        pass
    try:  # salt okunur / datepicker bagli alanlar icin son care
        kutu.evaluate(
            "(el, v) => { el.removeAttribute('readonly'); el.value = v;"
            " el.dispatchEvent(new Event('input', {bubbles:true}));"
            " el.dispatchEvent(new Event('change', {bubbles:true})); }",
            deger,
        )
        return (kutu.input_value() or "").strip() == deger
    except Exception:
        return False


def metin_iceren_sayfa(page, metin, sure=800):
    """Islem Takip penceresi ayri bir pencerede acilabildigi icin tum sayfalara bakilir."""
    sayfalar = [page]
    try:
        sayfalar += [p for p in page.context.pages if p is not page and not p.is_closed()]
    except Exception:
        pass
    for p in sayfalar:
        try:
            metinle_bul(p, metin, sure=sure)
            return p
        except LookupError:
            continue
    return None


GIB_HATA_METINLERI = ["veriler getirilirken hata", "e-Arşiv sistemine giriş",
                      "GİB internet sitesinden"]


def ekranda_gib_hatasi(page):
    """Islem Takip penceresi acilmadan ekrana dusen GIB hata uyarisi."""
    for fr in cerceveler(page):
        for metin in GIB_HATA_METINLERI:
            try:
                loc = fr.get_by_text(metin, exact=False)
                if loc.count() and loc.first.is_visible():
                    return True
            except Exception:
                continue
    return False


def islem_gunlugu(page):
    """Islem Takip penceresinin metni; pencere kapaliysa bos doner.

    Tarih diyalogu da .luca-open-window oldugu icin yalnizca islem gunlugu
    isaretlerini tasiyan pencere kabul edilir.
    """
    isaretler = [i.lower() for i in ISLEM_ISARETLERI] + [ISLEM_BITTI]
    sayfalar = [page]
    try:
        sayfalar += [p for p in page.context.pages if p is not page and not p.is_closed()]
    except Exception:
        pass
    for p in sayfalar:
        for fr in cerceveler(p):
            try:
                loc = fr.locator(".luca-open-window")
                for i in range(loc.count()):
                    pencere = loc.nth(i)
                    if not pencere.is_visible():
                        continue
                    metin = pencere.inner_text() or ""
                    if any(isaret in metin.lower() for isaret in isaretler):
                        return metin
            except Exception:
                continue
    return ""


def indirilemeyen_sayisi(sayfa):
    """Islem gunlugundeki \"url'li fatura indirilemedi\" satirlarini sayar."""
    gunluk = islem_gunlugu(sayfa)
    if gunluk:
        return gunluk.lower().count(INDIRILEMEDI)
    for fr in cerceveler(sayfa):
        try:
            if fr.get_by_text(ISLEM_BITTI, exact=False).count() == 0:
                continue
            return fr.locator("body").inner_text().lower().count(INDIRILEMEDI)
        except Exception:
            continue
    return 0


def islem_takibini_bekle(page, log, azami_saniye=900, durgunluk_saniye=180,
                         pencere_bekleme=12, en_az_saniye=3):
    """Sorgu bitene kadar bekler; indirilemeyen fatura sayisini dondurur (-1: tamamlanmadi).

    Bitis uc sekilde anlasilir: gunlukte "sona erdi" yazmasi, pencerenin
    kendiliginden kapanmasi (hizli biten sorgularda boyle oluyor) veya
    yazinin durgunluk_saniye boyunca hic degismemesi.
    """
    basla = time.time()
    pencere_goruldu = False
    son_gunluk = ""
    son_degisim = time.time()
    son_bildirim = 0

    while True:
        gecen = time.time() - basla

        if gecen >= en_az_saniye and fatura_yok_penceresini_kapat(page):
            yaz(f"    Luca: fatura bulunamadi ({int(gecen)} sn)", log)
            return 0

        if gecen >= en_az_saniye:  # "... adet fatura bulundu" da sorgunun bittigini soyler
            bilgi = bilgi_penceresini_kapat(page)
            if bilgi:
                yaz(f"    Luca: {bilgi} ({int(gecen)} sn)", log)
                return 0

        gunluk = islem_gunlugu(page)

        if gunluk:
            if not pencere_goruldu:
                pencere_goruldu = True
                yaz("    İşlem Takip penceresi acildi, sorgu suruyor...", log)
            if gunluk != son_gunluk:
                son_gunluk = gunluk
                son_degisim = time.time()

            # onceki sorgunun "sona erdi" yazisi ekranda kalmis olabilir
            if ISLEM_BITTI in gunluk.lower() and gecen >= en_az_saniye:
                basarisiz = gunluk.lower().count(INDIRILEMEDI)
                yaz(f"    GİB sorgusu tamamlandi ({int(gecen)} sn)"
                    + (f", {basarisiz} fatura indirilemedi" if basarisiz else ""), log)
                varsa_tikla(page, ["Kapat"], sure=4000)
                page.wait_for_timeout(1000)
                return basarisiz

            if yetki_hatasi(gunluk):
                yaz(f"    Bu firmanin bu servise yetkisi yok ({int(gecen)} sn),"
                    " ekran atlaniyor", log)
                varsa_tikla(page, ["Kapat", "Tamam"], sure=3000)
                acik_pencereleri_kapat(page)
                fatura_yok_penceresini_kapat(page)
                return YETKI_YOK

            if gib_hatasi(gunluk):
                yaz(f"    GİB'e ulasilamadi ({int(gecen)} sn), bu sorgu atlaniyor", log)
                varsa_tikla(page, ["Kapat", "Tamam"], sure=3000)
                acik_pencereleri_kapat(page)
                fatura_yok_penceresini_kapat(page)
                return 0

            if time.time() - son_degisim > durgunluk_saniye:
                sure_metni = (f"{int(durgunluk_saniye // 60)} dk" if durgunluk_saniye >= 60
                              else f"{int(durgunluk_saniye)} sn")
                yaz(f"    Sorgu {sure_metni} boyunca ilerlemedi, takildi sayiliyor", log)
                varsa_tikla(page, ["Kapat"], sure=3000)
                acik_pencereleri_kapat(page)
                fatura_yok_penceresini_kapat(page)
                return son_gunluk.lower().count(INDIRILEMEDI)

        elif pencere_goruldu:
            # pencere kendiliginden kapandi: sorgu bitmis demektir
            basarisiz = son_gunluk.lower().count(INDIRILEMEDI)
            yaz(f"    GİB sorgusu tamamlandi ({int(gecen)} sn, pencere kapandi)"
                + (f", {basarisiz} fatura indirilemedi" if basarisiz else ""), log)
            return basarisiz

        elif ekranda_gib_hatasi(page):
            yaz(f"    GİB'e ulasilamadi ({int(gecen)} sn), bu sorgu atlaniyor", log)
            varsa_tikla(page, ["Kapat", "Tamam"], sure=3000)
            acik_pencereleri_kapat(page)
            return 0

        elif gecen > pencere_bekleme:
            yaz(f"    İşlem Takip penceresi {int(gecen)} sn icinde gorunmedi, devam ediliyor", log)
            return 0

        if gecen > azami_saniye:
            yaz(f"    UYARI: GİB sorgusu {int(gecen)} sn sonra zaman asimina ugradi", log)
            varsa_tikla(page, ["Kapat"], sure=3000)
            return -1

        if gecen - son_bildirim >= 15:
            son_bildirim = gecen
            yaz(f"    ... bekleniyor ({int(gecen)} sn)", log)
        page.wait_for_timeout(1500)


def gibden_getir(page, baslangic, bitis, log):
    yaz(f"    GİB'den Getir aciliyor ({baslangic} - {bitis})", log)
    acik_pencereleri_kapat(page, log)  # onceki sorgudan kalan pencere tiklamayi engelliyor
    fatura_yok_penceresini_kapat(page)  # onceki sorgunun bildirimi yeni sorguya karismasin
    if not dugmeye_bas(page, "GİB'den Getir"):
        raise LookupError("'GİB'den Getir' butonuna basilamadi")
    page.wait_for_timeout(2500)

    kutular = tarih_kutulari(page)
    if len(kutular) < 2:
        yaz(f"    UYARI: tarih kutulari bulunamadi ({len(kutular)} adet"
            + (f", son hata: {_SON_KUTU_HATASI}" if _SON_KUTU_HATASI else "")
            + "), Luca varsayilani kullanilacak", log)
    else:
        kutuya_yaz(kutular[0], baslangic)
        kutuya_yaz(kutular[1], bitis)
        try:
            yaz(f"    Tarih araligi girildi: {kutular[0].input_value()} - {kutular[1].input_value()}", log)
        except Exception:
            pass

    tiklanan = varsa_tikla(page, DIYALOG_ONAY, sure=5000)
    if not tiklanan:
        raise LookupError(f"'Belgeleri Getir' butonu bulunamadi. Gorunen ogeler: {menu_metinleri(page, 20)}")
    yaz(f"    '{tiklanan}' tiklandi, sorgu basladi", log)

    basarisiz = islem_takibini_bekle(page, log, azami_saniye=AYAR["azami_saniye"],
                                     durgunluk_saniye=AYAR["durgunluk_saniye"])
    acik_pencereleri_kapat(page, log)
    return basarisiz


def listeyi_yenile(page, log, ek=""):
    """Sorgu sonrasi liste kendiliginden tazelenmiyor; tum sorgular bitince bir kez."""
    yaz("    Liste yenileniyor" + (f" ({ek})" if ek else ""), log)
    dugmeye_bas(page, "Yenile", sure=5000)
    page.wait_for_timeout(3000)
    try:
        page.wait_for_load_state("networkidle", timeout=60000)
    except Exception:
        pass


# Pencerenin onay dugmesi de "Belge Ara" yaziyor; arac cubugundaki ayni adli
# butona tekrar basmamak icin her zaman pencerenin icinden tiklanir
BELGE_ARA_ONAY = ["Belge Ara", "Ara", "Sorgula", "Tamam", "Uygula"]
BELGE_ARA_CAPALARI = ("Muhasebeleşmiş", "Belge Numarası", "Tarih Aralığı")


def pencere_acik_mi(pencere):
    if pencere is None:
        return False
    try:
        return pencere.is_visible()
    except Exception:
        return False


def belge_ara(page, bas, bit, log):
    """Listeyi istenen tarih araligina getirir (indirme ay geneli olsun diye).

    Sorgular 7 gunluk parcalarla yapildigi icin ekranda son parcanin filtresi
    kaliyordu; indirmeden once tum donem yeniden aranir.
    """
    acik_pencereleri_kapat(page, log)
    if not dugmeye_bas(page, "Belge Ara", sure=5000):
        yaz("    UYARI: 'Belge Ara' butonu bulunamadi, liste oldugu gibi kullanilacak", log)
        return False
    page.wait_for_timeout(1500)
    pencere = None
    for capa in BELGE_ARA_CAPALARI:
        _, pencere = metinli_diyalog(page, capa)
        if pencere is not None:
            break
    kutular = tarih_kutulari(page, pencere)
    if len(kutular) < 2:
        yaz("    UYARI: 'Belge Ara' tarih kutulari bulunamadi", log)
        acik_pencereleri_kapat(page, log)
        return False
    kutuya_yaz(kutular[0], bas)
    kutuya_yaz(kutular[1], bit)
    onay = pencerede_tikla(page, pencere, BELGE_ARA_ONAY, sure=5000) if pencere is not None else None
    if not onay:  # pencere taninmadi: tarih kutusunda Enter de aramayi baslatiyor
        try:
            kutular[1].press("Enter")
            onay = "Enter"
        except Exception:
            pass
    page.wait_for_timeout(1500)
    if onay and onay != "Enter" and pencere_acik_mi(pencere):
        # tiklanan oge baslik olabilir; pencere hala duruyorsa Enter ile aranir
        try:
            kutular[1].press("Enter")
            onay = f"{onay}+Enter"
        except Exception:
            pass
    yaz(f"    Belge Ara: {bas} - {bit}"
        + (f" ('{onay}' tiklandi)" if onay else " (UYARI: onay butonu bulunamadi)"), log)
    page.wait_for_timeout(3000)
    try:
        page.wait_for_load_state("networkidle", timeout=30000)
    except Exception:
        pass
    return bool(onay)


# GIB ekranlari tarihi 12/08/2026, 12.08.2026, 12-08-2026 ya da 2026-08-12 yazabiliyor
TARIH_DESENI = re.compile(r"\d{2}[./-]\d{2}[./-]\d{4}|\d{4}-\d{2}-\d{2}")
# ETTN'siz belge numarasi: 3 harf + 13 rakam (orn. GIB2026000000011)
BELGE_NO_DESENI = re.compile(r"[A-Za-z]{3}\d{13}")


def fatura_satiri_mi(hucreler, en_az=3):
    """Satir fatura satiri mi: yeterli hucre + tarih ya da belge numarasi.

    Interaktif V.D. ekraninda tarih bicimi degisebildigi icin belge numarasi
    da olcut alinir; yoksa dolu listeler bos gorunuyordu.
    """
    if len(hucreler) < en_az:
        return False
    return any(TARIH_DESENI.search(h) or BELGE_NO_DESENI.search(h) for h in hucreler)


def cerceveden_satirlar(fr):
    try:
        satirlar = fr.locator("tr")
        adet = min(satirlar.count(), 500)
    except Exception:
        return []
    veriler = []
    for i in range(adet):
        try:
            hucreler = [h.strip() for h in satirlar.nth(i).locator("td").all_inner_texts() if h.strip()]
        except Exception:
            continue
        if fatura_satiri_mi(hucreler, 3):
            veriler.append(hucreler)
    return veriler


# Satir hucreleri: once gercek hucre ogeleri, olmazsa dogrudan cocuklar,
# en son satir metni (izgara <td> kullanmadiginda metin tek parca geliyordu)
HUCRE_CIKAR = """el => {
  const s = el.closest('tr, [role=row], li')
    || (el.parentElement && el.parentElement.parentElement);
  if (!s) return [];
  const metin = e => (e.innerText || e.textContent || '').trim();
  let h = [...s.querySelectorAll('td, th, [role=gridcell], [role=cell]')].map(metin);
  if (h.filter(Boolean).length < 2) h = [...s.children].map(metin);
  if (h.filter(Boolean).length < 2) h = metin(s).split(/\\t|\\n|\\s{2,}/);
  return h.map(t => t.trim()).filter(Boolean);
}"""


def kutulardan_satirlar(kutular, sayi):
    """Satirlari isaret kutularindan cikarir.

    Fatura izgarasi her zaman <tr>/<td> degil; kutunun en yakin satir
    atasinin metni alinirsa yapidan bagimsiz calisir ve satir sayisi
    isaretlenen kayit sayisiyla dogal olarak ayni olur.
    """
    if kutular is None:
        return []
    satirlar = []
    for i in range(sayi):
        try:
            hucreler = kutular.nth(i).evaluate(HUCRE_CIKAR) or []
        except Exception:
            continue
        hucreler = [h.strip() for h in hucreler if h and h.strip()]
        if fatura_satiri_mi(hucreler, 3):
            satirlar.append(hucreler)
    return satirlar


def tabloyu_oku(page, tercih=None):
    """Fatura tablosu, isaret kutulariyla ayni cerceveden okunur.

    Tum cerceveler taranirsa baska ekranlardan kalan tablolar da fatura
    sanilip olmayan satirlar listeleniyordu.
    """
    if tercih is not None:
        # Bu cercevede satir yoksa liste gercekten bostur; genel taramaya
        # dusulurse baska ekranlarda kalan tablolar fatura sanilıyor.
        return tercih, cerceveden_satirlar(tercih)

    en_iyi = (None, [])
    for fr in cerceveler(page):
        veriler = cerceveden_satirlar(fr)
        if len(veriler) > len(en_iyi[1]):
            en_iyi = (fr, veriler)
    return en_iyi


SECIM_SECICILERI = ("input[type=checkbox]", "[role=checkbox]",
                    "img[src*='check']", "img[src*='tick']", "[class*='checkbox']")


def veri_kutusu_sayisi(kutular, adet, sinir=14):
    """Kutulardan kaci gercek bir fatura satirinda duruyor.

    Ekranda sutun secimi ve suzgec satirlarinin da isaret kutusu var; bunlar
    fatura satiri sanilinca liste bos okunuyordu.
    """
    bulunan = 0
    for i in range(min(adet, sinir)):
        try:
            hucreler = [h for h in (kutular.nth(i).evaluate(HUCRE_CIKAR) or []) if h]
        except Exception:
            continue
        if fatura_satiri_mi(hucreler, 3):
            bulunan += 1
    return bulunan


def secim_kutulari(page):
    """Fatura satirlarindaki isaret kutularini bulur.

    Yalnizca "en cok kutu" olcut alinirsa baslik/suzgec satirinin kutulari
    secilebiliyor. Once kutunun durdugu satirin fatura satiri olup olmadigina
    bakilir; hicbir cercevede veri satiri yoksa eski davranisa (en cok kutu)
    dusulur.
    """
    en_iyi = (None, 0, None)
    en_iyi_veri = 0
    yedek = (None, 0, None)
    for fr in cerceveler(page):
        for secici in SECIM_SECICILERI:
            try:
                loc = fr.locator(secici)
                adet = loc.count()
                if not adet or not loc.first.is_visible():
                    continue
            except Exception:
                continue
            if adet > yedek[1]:
                yedek = (loc, adet, fr)
            veri = veri_kutusu_sayisi(loc, adet)
            if veri > en_iyi_veri or (veri and veri == en_iyi_veri and adet > en_iyi[1]):
                en_iyi, en_iyi_veri = (loc, adet, fr), veri
            if veri:  # bu cercevede fatura satiri bulundu, digerlerini deneme
                break
    return en_iyi if en_iyi[0] is not None else yedek


def isaretli_sayisi(kutular, sayi):
    if kutular is None:
        return 0
    try:
        return sum(1 for i in range(sayi) if kutular.nth(i).is_checked())
    except Exception:
        return 0


def veri_satir_indisleri(fr):
    try:
        satirlar = fr.locator("tr")
        adet = min(satirlar.count(), 500)
    except Exception:
        return None, []
    indisler = []
    for i in range(adet):
        try:
            hucreler = [h.strip() for h in satirlar.nth(i).locator("td").all_inner_texts() if h.strip()]
        except Exception:
            continue
        if fatura_satiri_mi(hucreler, 3):
            indisler.append(i)
    return satirlar, indisler


def belge_sec_diyalogu(page, satir_sayisi=0):
    """Luca'nin kendi secim yolu: Belge Seç -> Tümünü Seç -> Tamam.

    Satirlari tek tek isaretlemek yerine "Belge Seçiniz" penceresi kullanilir;
    pencere listedeki butun belgeleri (ekranda gorunmeyenler dahil) secer.
    """
    if not dugmeye_bas(page, "Belge Seç", sure=4000):
        return 0
    page.wait_for_timeout(900)
    try:  # pencere acildiysa icindeki "Tümünü Seç" gorunur olur
        _, dugme = metinle_bul(page, "Tümünü Seç", sure=5000)
        dugme.click(timeout=5000)
    except Exception:
        varsa_tikla(page, ["Kapat"])
        return 0
    page.wait_for_timeout(600)
    try:
        _, tamam = metinle_bul(page, "Tamam", sure=5000)
        tamam.click(timeout=5000)
    except Exception:
        varsa_tikla(page, ["Kapat"])
        return 0
    page.wait_for_timeout(1000)
    kutular, sayi, _ = secim_kutulari(page)
    return isaretli_sayisi(kutular, sayi) or satir_sayisi


def hepsini_sec(page, fr=None, satir_sayisi=0):
    secilen = belge_sec_diyalogu(page, satir_sayisi)  # Luca'nin kendi yolu
    if secilen:
        return secilen

    kutular, sayi, kutu_cercevesi = secim_kutulari(page)
    if fr is None:
        fr = kutu_cercevesi

    if sayi:
        try:  # baslik satirindaki kutu genelde hepsini isaretler
            kutular.first.check(timeout=5000)
            page.wait_for_timeout(800)
            isaretli = isaretli_sayisi(kutular, sayi)
            if isaretli > 1:
                return isaretli
        except Exception:
            pass
        secilen = 0
        for i in range(sayi):
            try:
                kutu = kutular.nth(i)
                if kutu.is_visible() and not kutu.is_checked():
                    kutu.check(timeout=2000)
                    secilen += 1
            except Exception:
                continue
        if secilen:
            page.wait_for_timeout(500)
            return secilen

    # Luca'nin kendi ipucu: tabloda bosluk tusu satir seciyor
    satirlar, indisler = veri_satir_indisleri(fr)
    if satirlar is not None and indisler:
        try:
            satirlar.nth(indisler[0]).click(timeout=5000)
            page.wait_for_timeout(400)
            for _ in indisler:
                page.keyboard.press("Space")
                page.wait_for_timeout(150)
                page.keyboard.press("ArrowDown")
                page.wait_for_timeout(150)
            return isaretli_sayisi(kutular, sayi) or len(indisler)
        except Exception:
            pass
    return 0


def uyari_metni(page):
    """Luca uyarisi (orn. 'Lutfen indirilecek faturalari seciniz') varsa metnini dondurur."""
    for fr in cerceveler(page):
        try:
            loc = fr.get_by_text("Lütfen", exact=False)
            if loc.count() and loc.first.is_visible():
                return " ".join((loc.first.inner_text() or "").split())[:90]
        except Exception:
            continue
    return None


CD_DESENI = re.compile(r'filename\*?=(?:UTF-8'')?"?([^";]+)"?', re.I)
INDIRME_TURLERI = ("application/zip", "application/octet-stream", "application/x-zip",
                   "application/vnd.ms-excel", "application/vnd.openxmlformats",
                   "application/force-download", "application/download")


def yanit_dosya_mi(basliklar):
    """Sunucu yaniti indirilecek bir dosya mi (ekran yerine dosya)."""
    cd = (basliklar.get("content-disposition") or "").lower()
    ct = (basliklar.get("content-type") or "").lower()
    return "attachment" in cd or any(t in ct for t in INDIRME_TURLERI)


UZANTILAR = [("zip", ".zip"), ("spreadsheetml", ".xlsx"), ("ms-excel", ".xls"),
             ("csv", ".csv"), ("pdf", ".pdf"), ("xml", ".xml")]


def yanit_dosya_adi(basliklar, yedek_ad):
    """Dosya adi once Content-Disposition'dan, yoksa icerik turunden uretilir."""
    eslesme = CD_DESENI.search(basliklar.get("content-disposition") or "")
    if eslesme and eslesme.group(1).strip():
        return dosya_adi_yap(eslesme.group(1).strip())
    ct = (basliklar.get("content-type") or "").lower()
    uzanti = next((u for anahtar, u in UZANTILAR if anahtar in ct), ".dat")
    return f"{yedek_ad}{uzanti}"


def indir_yakalayarak(page, dugme_metni, hedef_klasor, on_ek, log, azami_saniye=30,
                      pencere_acilir=True):
    """Dosyayi tarayiciya indirtmeden, istegi yakalayip kendimiz kaydederiz.

    Chrome indirmeyi kaydettigi anda cokuyor (AddKeepAlive kDownloadInProgress).
    Bu yolda istek route ile yakalanir, govde Python tarafinda alinir ve
    tarayicinin indirme mekanizmasi hic devreye girmez.
    """
    ctx = page.context
    alinan = {}
    inenler = []  # route ile yakalanamazsa tarayici indirmesi yedek kalsin
    dinleyici = lambda d: inenler.append(d)
    page.on("download", dinleyici)

    def yonlendir(route):
        try:
            yanit = route.fetch()
        except Exception:
            try:
                route.continue_()
            except Exception:
                pass
            return
        try:
            basliklar = {k.lower(): v for k, v in (yanit.headers or {}).items()}
            if not alinan and yanit_dosya_mi(basliklar):
                alinan["ad"] = yanit_dosya_adi(basliklar, on_ek)
                alinan["govde"] = yanit.body()
                # abort edilirse Luca'nin cercevesi "sayfa kullanilamiyor"
                # hatasina dusup ekrani bozuyordu; 204 ile tarayici bulundugu
                # sayfada kalir, indirme de baslamaz
                try:
                    route.fulfill(status=204, body="")
                except Exception:
                    route.abort()
                return
            route.fulfill(response=yanit)
        except Exception:
            try:
                route.continue_()
            except Exception:
                pass

    ctx.route("**/*", yonlendir)
    try:
        # onceki adimdan kalan bildirim yeni indirmeye karismasin
        fatura_yok_penceresini_kapat(page)
        if not dugmeye_bas(page, dugme_metni, sure=8000):
            yaz(f"    '{dugme_metni}' butonuna basilamadi, atlandi", log)
            return None
        page.wait_for_timeout(2000)

        pencere = indirme_diyalogu(page)[1] if pencere_acilir else None
        onay = None
        if pencere is not None:
            if diyalogda_tumunu_sec(page):
                yaz("    Onay penceresinde 'tum faturalar' secildi", log)
            onay = diyalogda_tikla(page, INDIRME_ONAY)
            if onay:
                yaz(f"    Onay penceresinde '{onay}' tiklandi", log)
            else:
                yaz("    UYARI: onay penceresi tiklanamadi", log)

        sure = 5 if (pencere is not None and not onay) else azami_saniye
        uyari = None
        bitis = time.time() + sure
        basla = time.time()
        while time.time() < bitis and "govde" not in alinan and not inenler:
            if time.time() - basla >= 2 and fatura_yok_penceresini_kapat(page):
                uyari = "Luca: fatura bulunamadi"
                break
            uyari = uyari_metni(page)
            if uyari:
                break
            page.wait_for_timeout(500)

        # tiklama gectigi halde dosya gelmediyse butonun kendi kisayolu denenir
        kisayol = KISAYOLLAR.get(dugme_metni)
        if kisayol and not uyari and "govde" not in alinan and not inenler:
            yaz(f"    Dosya gelmedi, '{dugme_metni}' kisayolu deneniyor ({kisayol})", log)
            try:
                page.keyboard.press(kisayol)
            except Exception:
                pass
            bitis = time.time() + min(azami_saniye, 15)
            while time.time() < bitis and "govde" not in alinan and not inenler:
                page.wait_for_timeout(500)

        if "govde" in alinan:
            yol = hedef_klasor / f"{on_ek}_{alinan['ad']}"
            yol.write_bytes(alinan["govde"])
            yaz(f"    indirildi (yakalanarak): {yol.name}", log)
            return yol

        if inenler:  # istek yakalanamadi ama tarayici indirdi
            dosya = inenler[0]
            yol = hedef_klasor / f"{on_ek}_{dosya.suggested_filename}"
            dosya.save_as(str(yol))
            yaz(f"    indirildi: {yol.name}", log)
            return yol

        if uyari:
            yaz(f"    Luca uyarisi: {uyari}", log)
        else:
            yaz(f"    '{dugme_metni}' icin {int(sure)} sn icinde dosya gelmedi", log)
        return None
    except Exception as e:
        yaz(f"    '{dugme_metni}' indirilemedi ({type(e).__name__}: {e})", log)
        return None
    finally:
        try:
            ctx.unroute("**/*", yonlendir)
        except Exception:
            pass
        try:
            page.remove_listener("download", dinleyici)
        except Exception:
            pass
        if sayfa_canli(page):
            acik_pencereleri_kapat(page)


def indir(page, dugme_metni, hedef_klasor, on_ek, log, azami_saniye=30, pencere_acilir=True):
    """Indirme akisi: arac cubugu butonu -> pencerede 'tum faturalar' -> pencerede indir.

    expect_download yerine olay dinleyip beklenir; boylece Luca "faturalari
    seciniz" uyarisi verdiginde bos yere zaman asimi beklenmez.
    """
    indirilenler = []
    dinleyici = lambda d: indirilenler.append(d)
    page.on("download", dinleyici)
    try:
        fatura_yok_penceresini_kapat(page)  # onceki adimdan kalan bildirim
        if not dugmeye_bas(page, dugme_metni, sure=8000):
            yaz(f"    '{dugme_metni}' butonuna basilamadi, atlandi", log)
            return None
        page.wait_for_timeout(2000)

        # Excel pencere acmaz; onceki islemden kalan pencere onay sanilmasin
        pencere = indirme_diyalogu(page)[1] if pencere_acilir else None
        diyalog_var = pencere is not None
        onay = None
        if diyalog_var:
            if diyalogda_tumunu_sec(page):
                yaz("    Onay penceresinde 'tum faturalar' secildi", log)
            onay = diyalogda_tikla(page, INDIRME_ONAY)
            if onay:
                yaz(f"    Onay penceresinde '{onay}' tiklandi", log)
            else:
                yaz("    UYARI: onay penceresi tiklanamadi", log)

        # onay penceresi acilip tiklanamadiysa dosya gelmeyecek; Excel gibi
        # pencere acmayan butonlarda ise normal sure beklenir
        sure = 5 if (diyalog_var and not onay) else azami_saniye
        uyari = None
        bitis = time.time() + sure
        basla = time.time()
        while time.time() < bitis and not indirilenler:
            if time.time() - basla >= 2 and fatura_yok_penceresini_kapat(page):
                uyari = "Luca: fatura bulunamadi"
                break
            uyari = uyari_metni(page)
            if uyari:
                break
            page.wait_for_timeout(500)

        if indirilenler:
            dosya = indirilenler[0]
            ad = f"{on_ek}_{dosya.suggested_filename}"
            yol = hedef_klasor / ad
            dosya.save_as(str(yol))
            yaz(f"    indirildi: {ad}", log)
            return yol

        if uyari:
            yaz(f"    Luca uyarisi: {uyari}", log)
        else:
            yaz(f"    '{dugme_metni}' icin {int(sure)} sn icinde dosya gelmedi", log)
        return None
    except Exception as e:
        yaz(f"    '{dugme_metni}' indirilemedi ({type(e).__name__})", log)
        return None
    finally:
        try:
            page.remove_listener("download", dinleyici)
        except Exception:
            pass
        acik_pencereleri_kapat(page)  # pencere acik kalirsa sonraki adimlar kilitleniyor


def interaktif_kayit_sayisi(page):
    """Ekran altindaki 'Toplam Kayit Sayisi' degeri; bulunamazsa None."""
    for fr in cerceveler(page):
        try:
            loc = fr.get_by_text("Kayıt Sayısı", exact=False)
            for i in range(min(loc.count(), 3)):
                eslesme = KAYIT_SAYISI_DESENI.search(loc.nth(i).inner_text() or "")
                if eslesme:
                    return int(eslesme.group(1))
        except Exception:
            continue
    return None


def listeyi_bekle(page, log, azami_saniye=120):
    """Sorgu sonrasi listenin dolmasini bekler (kayit sayisi > 0)."""
    basla = time.time()
    son_bildirim = 0
    while time.time() - basla < azami_saniye:
        sayi = interaktif_kayit_sayisi(page)
        if sayi:
            yaz(f"    Liste doldu: {sayi} kayit ({int(time.time() - basla)} sn)", log)
            return sayi
        if fatura_yok_penceresini_kapat(page):
            yaz("    Luca: fatura bulunamadi", log)
            return 0
        if sayi == 0 and time.time() - basla > 10:
            yaz("    Liste bos (Toplam Kayit Sayisi: 0)", log)
            return 0
        gecen = time.time() - basla
        if gecen - son_bildirim >= 15:
            son_bildirim = gecen
            yaz(f"    ... liste bekleniyor ({int(gecen)} sn)", log)
        page.wait_for_timeout(1500)
    yaz(f"    Liste {azami_saniye} sn icinde dolmadi", log)
    return None


def radyo_sec(page, pencere, metin):
    """Secenek penceresindeki radyo dugmesini yazisina gore isaretler."""
    hedef = karsilastir(metin)
    kapsam = pencere if pencere is not None else page
    try:
        radyolar = kapsam.locator("input[type=radio]")
        for i in range(min(radyolar.count(), 10)):
            r = radyolar.nth(i)
            try:
                cevre = r.evaluate(
                    "el => { const s = el.closest('tr, li, label, div')"
                    " || el.parentElement; return s ? s.innerText : ''; }") or ""
            except Exception:
                cevre = ""
            if hedef in karsilastir(cevre):
                try:
                    r.check(timeout=4000)
                except Exception:
                    r.click(timeout=4000)
                page.wait_for_timeout(400)
                return True
    except Exception:
        pass
    # radyo bulunamadiysa yazisina tiklamak da secimi yapar
    return bool(pencerede_tikla(page, pencere, [metin], sure=3000))


def interaktif_sorgula(page, araliklar, log, listeleme_araligi=None):
    """Interaktif Vergi Dairesi ekraninda e-arsiv faturalarini GIB servisinden ceker.

    Akis: tarih araligi -> "İnteraktif V.D'sinden E-Arşiv Faturalarını Sorgula"
    -> acilan pencerede "GİB Servis ile Sorgula" -> ayni isimli onay butonu.
    Luca ilk sorguda listeyi her zaman doldurmadigi icin iki kez calistirilir.
    """
    calisan = 0
    for sira, (bas, bit) in enumerate(araliklar):
        # Luca ilk sorguda 9000/9001 hatasi verip bos donuyor; yalnizca ilk
        # aralikta iki kez sorulur, sonrakilerde tek sorgu yetiyor
        tekrar = INTERAKTIF_TEKRAR if sira == 0 else 1
        for tur in range(1, tekrar + 1):
            acik_pencereleri_kapat(page, log)
            fatura_yok_penceresini_kapat(page)

            kutular = tarih_kutulari(page)
            if len(kutular) >= 2:
                kutuya_yaz(kutular[0], bas)
                kutuya_yaz(kutular[1], bit)
            else:
                yaz("    UYARI: tarih kutulari bulunamadi, Luca varsayilani kullanilacak", log)
            yaz(f"    Interaktif V.D. sorgusu ({bas} - {bit}) {tur}/{tekrar}", log)

            if not dugmeye_bas(page, INTERAKTIF_SORGU, sure=8000):
                yaz(f"    '{INTERAKTIF_SORGU}' butonuna basilamadi", log)
                return calisan
            page.wait_for_timeout(2000)

            _, pencere = metinli_diyalog(page, INTERAKTIF_CAPASI)
            if pencere is None:
                # pencere yoksa onay butonu da yok; arac cubugu butonuna tekrar
                # basmamak icin burada durulur
                yaz("    UYARI: 'GİB Servis ile Sorgula' penceresi acilmadi", log)
                return calisan
            if radyo_sec(page, pencere, INTERAKTIF_SERVIS):
                yaz(f"    '{INTERAKTIF_SERVIS}' secildi", log)
            else:
                yaz(f"    UYARI: '{INTERAKTIF_SERVIS}' secenegi isaretlenemedi", log)
            if not pencerede_tikla(page, pencere, [INTERAKTIF_SORGU], sure=5000):
                yaz("    UYARI: pencerede sorgu butonu tiklanamadi", log)
                acik_pencereleri_kapat(page, log)
                return calisan

            # bu ekranda Islem Takip penceresi acilmiyor; listenin dolmasi beklenir
            listeyi_bekle(page, log, azami_saniye=min(AYAR["azami_saniye"], 120))
            acik_pencereleri_kapat(page, log)
            calisan += 1
            page.wait_for_timeout(1500)

    # son sorgunun tarih filtresi ekranda kaliyor; listeleme tum donem icin yapilir
    if listeleme_araligi:
        acik_pencereleri_kapat(page, log)
        fatura_yok_penceresini_kapat(page)
        kutular = tarih_kutulari(page)
        if len(kutular) >= 2:
            kutuya_yaz(kutular[0], listeleme_araligi[0])
            kutuya_yaz(kutular[1], listeleme_araligi[1])
            yaz(f"    Listeleme araligi: {listeleme_araligi[0]} - {listeleme_araligi[1]}", log)

    # bu ekranda sorgu listeyi kendiliginden doldurmuyor
    if dugmeye_bas(page, INTERAKTIF_LISTELE, sure=5000):
        yaz(f"    '{INTERAKTIF_LISTELE}' tiklandi", log)
        page.wait_for_timeout(3000)
    else:
        yaz(f"    UYARI: '{INTERAKTIF_LISTELE}' butonu bulunamadi", log)
    return calisan


def iptal_itiraz_sorgula(page, araliklar, log, interaktif=False):
    """Listedeki faturalar icin GIB'den iptal/itiraz durumunu sorgular.

    Luca akisi: arac cubugundan "GİB'den İptal/İtiraz Sorgula" -> acilan
    pencerede tarih araligi -> "İptal/İtiraz Sorgula". Sorgu tum listeye
    uygulandigi icin faturalari onceden isaretlemek gerekmiyor.
    Buradaki tarih, faturanin GIB'e raporlanma tarihidir.
    """
    calisan = 0
    for bas, bit in araliklar:
        acik_pencereleri_kapat(page, log)
        fatura_yok_penceresini_kapat(page)
        if not varsa_tikla(page, IPTAL_DUGME_ADAYLARI, sure=4000):
            yaz("    'GİB'den İptal/İtiraz Sorgula' butonu bulunamadi, atlandi", log)
            return calisan
        page.wait_for_timeout(2000)

        _, pencere = metinli_diyalog(page, IPTAL_CAPASI)
        kutular = tarih_kutulari(page, pencere)
        if len(kutular) >= 2:
            kutuya_yaz(kutular[0], bas)
            kutuya_yaz(kutular[1], bit)
            try:  # gercekten pencereye yazildi mi
                girilen = f"{kutular[0].input_value()} - {kutular[1].input_value()}"
            except Exception:
                girilen = f"{bas} - {bit}"
            yaz(f"    Iptal/itiraz sorgusu ({girilen})", log)
        else:
            yaz("    UYARI: iptal/itiraz tarih kutulari bulunamadi, Luca varsayilani kullanilacak", log)

        onay = pencerede_tikla(page, pencere, IPTAL_ONAY) if pencere is not None else None
        if not onay:
            # pencere taninmadiysa tam metinle ara; arac cubugu butonu farkli yazildigi
            # icin tam eslesme yanlislikla ona denk gelmez
            onay = varsa_tikla(page, IPTAL_ONAY, sure=4000)
        if not onay:
            yaz("    UYARI: iptal/itiraz sorgu butonu tiklanamadi", log)
            acik_pencereleri_kapat(page, log)
            return calisan

        # interaktif V.D. ekraninda Islem Takip penceresi hic acilmiyor; kisa
        # beklenir ama bilgi penceresi kontrolu (3 sn) yine de calissin
        islem_takibini_bekle(page, log, azami_saniye=AYAR["azami_saniye"],
                             durgunluk_saniye=AYAR["durgunluk_saniye"],
                             pencere_bekleme=4 if interaktif else 12)
        acik_pencereleri_kapat(page, log)
        calisan += 1

    if interaktif:
        # bu ekranda Yenile listeyi bosaltiyor; liste yeniden listelenmeli
        yaz("    Liste yeniden listeleniyor (iptal/itiraz sonrasi)", log)
        dugmeye_bas(page, INTERAKTIF_LISTELE, sure=5000)
        page.wait_for_timeout(3000)
    else:
        listeyi_yenile(page, log, "iptal/itiraz sonrasi")
    return calisan


def iptal_itiraz_satirlari(satirlar):
    """Durum sutununda iptal/itiraz gecen satirlar."""
    return [s for s in satirlar if IPTAL_DESENI.search(sadelestir(" ".join(s)))]


# Fatura numarasi 16 hane: 3 on ek + 4 haneli yil + 9 haneli sira
# (orn. ABC2026000000123). On ekte rakam da olabildigi icin harf sarti yok.
FATURA_NO_DESENI = re.compile(r"\b([A-Za-z0-9ÇĞİÖŞÜçğıöşü]{3}\d{13})\b")
# katı desen tutmazsa: 16 haneli, en az bir rakam iceren herhangi bir kod
FATURA_NO_YEDEK = re.compile(r"\b(?=[A-Za-z0-9]{16}\b)[A-Za-z0-9]*\d[A-Za-z0-9]*\b")
# tevkifat isaretleri: ekran yazisi, UBL etiketi ve KDV tevkifat vergi kodu
XML_TEVKIFAT = ("tevkifat", "withholdingtaxtotal", ">9015<", "kdvtevkifat")


def fatura_kimligi(satir):
    """Bir liste satirindan (unvan ilk kelimesi, fatura no) cikarir."""
    no = ""
    for desen in (FATURA_NO_DESENI, FATURA_NO_YEDEK):
        for hucre in satir:
            eslesme = desen.search(hucre.replace(" ", "").replace("-", ""))
            if eslesme:
                no = eslesme.group(0).upper()
                break
        if no:
            break
    if not no:
        return None
    unvan = ""
    for hucre in satir:
        metin = hucre.strip()
        if metin.upper() == no or TARIH_DESENI.search(metin):
            continue
        harf = sum(1 for c in metin if c.isalpha())
        if harf >= 3 and harf > sum(1 for c in unvan if c.isalpha()):
            unvan = metin
    ilk_kelime = (unvan.split() or [""])[0].strip(".,")
    return ilk_kelime, no


def fatura_kimlikleri(satirlar):
    return [k for k in (fatura_kimligi(s) for s in satirlar) if k]


def zipten_tevkifatlilar(zip_yolu):
    """Inen belge paketindeki XML'lerde tevkifat arar; fatura numaralarini dondurur.

    Ekranda tevkifat sutunu olmayabildigi gibi XML de bozuk inebiliyor;
    bu yuzden iki kaynak birlikte kullanilir.
    """
    bulunan = set()
    try:
        with zipfile.ZipFile(zip_yolu) as z:
            for ad in z.namelist():
                if not ad.lower().endswith((".xml", ".html", ".htm")):
                    continue
                try:
                    icerik = z.read(ad).decode("utf-8", "ignore").lower()
                except Exception:
                    continue
                if any(isaret in icerik for isaret in XML_TEVKIFAT):
                    eslesme = FATURA_NO_DESENI.search(ad.replace(" ", ""))
                    bulunan.add(eslesme.group(1).upper() if eslesme else ad)
    except Exception:
        return set()
    return bulunan


def _hucre_metni(h):
    if h is None:
        return ""
    if isinstance(h, (datetime, date)):
        return h.strftime(TARIH_BICIMI)
    return str(h).strip()


def _excel_tam_oku(ac, log=None, sadece_ilk=False):
    """read_only okuma bos donerse dosyayi normal modda okur (yavas ama saglam)."""
    basliklar, satirlar = [], []
    try:
        wb = ac(False)
    except Exception as e:
        yaz(f"    Excel tam okumada acilamadi ({type(e).__name__})", log)
        return basliklar, satirlar
    try:
        for ws in (wb.worksheets[:1] if sadece_ilk else wb.worksheets):
            for ham in ws.iter_rows(values_only=True):
                hucreler = [_hucre_metni(h) for h in ham]
                while hucreler and not hucreler[-1]:
                    hucreler.pop()
                if len([h for h in hucreler if h]) < 2:
                    continue
                if not basliklar:
                    basliklar = hucreler
                else:
                    satirlar.append(hucreler)
    except Exception as e:
        yaz(f"    Excel tam okunurken hata ({type(e).__name__})", log)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return basliklar, satirlar


def excelden_tablo(yol, log=None, sadece_ilk=False):
    """Inen Excel'i (basliklar, satirlar) olarak okur; sutunlar yerinde kalir.

    Ekrandaki tabloyu kazimak yerine inen dosyayi kaynak almak daha saglam:
    sutun basliklari belli oldugu icin tevkifat ve iptal/itiraz dogrudan
    kendi sutunlarindan okunabiliyor.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        yaz("    openpyxl kurulu degil, Excel okunamadi", log)
        return [], []
    import warnings

    def ac(read_only):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return load_workbook(str(yol), read_only=read_only, data_only=True)

    try:
        wb = ac(True)
    except Exception as e:
        yaz(f"    Excel acilamadi ({type(e).__name__})", log)
        return [], []

    basliklar, satirlar = [], []
    try:
        sayfalar = wb.worksheets[:1] if sadece_ilk else wb.worksheets
        for ws in sayfalar:
            try:  # Luca dosyalarinda boyut bilgisi eksik; olmazsa satirlar bos geliyor
                ws.reset_dimensions()
            except Exception:
                pass
            for ham in ws.iter_rows(values_only=True):
                hucreler = [_hucre_metni(h) for h in ham]
                while hucreler and not hucreler[-1]:
                    hucreler.pop()
                if len([h for h in hucreler if h]) < 2:
                    continue
                if not basliklar:
                    basliklar = hucreler
                else:
                    satirlar.append(hucreler)
        if not satirlar:  # read_only modu bos donduyse dosyayi tumuyle acip tekrar dene
            sayfalar = ", ".join(ws.title for ws in wb.worksheets) or "-"
            basliklar, satirlar = _excel_tam_oku(ac, log, sadece_ilk)
            if not satirlar:
                yaz(f"    Excel bos geldi (sayfalar: {sayfalar},"
                    f" baslik: {len(basliklar)} sutun)", log)
    except Exception as e:
        yaz(f"    Excel okunurken hata ({type(e).__name__})", log)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return basliklar, satirlar


def excelden_satirlar(yol, log=None):
    basliklar, satirlar = excelden_tablo(yol, log)
    if not satirlar:
        yaz(f"    Excel'de veri yok (baslik: {len(basliklar)} sutun)", log)
    return satirlar


def sutun_indeksi(basliklar, *anahtarlar):
    """Basligi anahtari iceren ilk sutunun sirasi; yoksa None."""
    for i, baslik in enumerate(basliklar):
        duz = sadelestir(baslik)
        if any(a in duz for a in anahtarlar):
            return i
    return None


BOS_DEGERLER = {"", "0", "0,00", "0.00", "-", "YOK", "HAYIR"}


def sutunlu_satirlar(basliklar, satirlar, *anahtarlar):
    """Belirtilen sutunu dolu olan satirlar; sutun yoksa None doner."""
    i = sutun_indeksi(basliklar, *anahtarlar)
    if i is None:
        return None
    return [s for s in satirlar
            if i < len(s) and sadelestir(s[i]) not in BOS_DEGERLER]


def tutar_cozumle(metin):
    """'1.234,56' / '1234,56' / '1234.56' (openpyxl'in dogrudan verdigi sayi) -> float.

    Bos ya da sayi olmayan metin 0.0 sayilir.
    """
    s = str(metin or "").strip()
    if not s:
        return 0.0
    s = re.sub(r"[^0-9,.\-]", "", s)  # TL isareti, bosluk vb. temizlenir
    if not s or s in ("-", ".", ","):
        return 0.0
    if "," in s:  # Turkce bicim: binlik nokta, ondalik virgul
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _tutar_sutunlari(basliklar, anahtarlar, haric=()):
    """Basligi anahtarlardan birini iceren (ve haric olanlari icermeyen) tum sutun indeksleri.

    Bazi ekranlarda KDV oran basina ayri sutunlarda gosteriliyor
    (orn. "KDV Matrahı (%20)", "KDV Matrahı (%10)"); oran onemli olmadigi
    icin ayni turden birden fazla sutun varsa hepsi toplanir.
    """
    idx = []
    for i, baslik in enumerate(basliklar):
        duz = sadelestir(baslik)
        if any(h in duz for h in haric):
            continue
        if any(a in duz for a in anahtarlar):
            idx.append(i)
    return idx


def matrah_kdv_sutunlari(basliklar):
    """(matrah sutunlari, KDV tutari sutunlari) indeksleri."""
    matrah = _tutar_sutunlari(basliklar, ("MATRAH",))
    if not matrah:  # bazi ekranlarda sutun adi "matrah" gecmiyor
        matrah = _tutar_sutunlari(basliklar, ("MAL HIZMET TOPLAM TUTARI", "MAL HIZMET TUTARI"))
    kdv = _tutar_sutunlari(basliklar, ("KDV",), haric=("ORAN",))
    return matrah, kdv


def _satirlarin_tutari(satirlar, sutunlar):
    return sum(tutar_cozumle(s[i]) for s in satirlar for i in sutunlar if i < len(s))


def satir_toplam_tutari(basliklar, satir):
    """Bir faturanin genel toplam tutari (eksik/fazla fatura listesinde gosterilir)."""
    for anahtar in ("GENEL TOPLAM", "VERGILER DAHIL TOPLAM TUTAR", "ODENECEK TUTAR", "TOPLAM TUTAR"):
        i = sutun_indeksi(basliklar, anahtar)
        if i is not None and i < len(satir):
            return tutar_cozumle(satir[i])
    return 0.0


def fatura_kimlikleri_tutarli(basliklar, satirlar):
    """(unvan, no, tutar) uclusu: eksik/fazla fatura listesinde tutar da gorunsun."""
    sonuc = []
    for s in satirlar:
        k = fatura_kimligi(s)
        if not k:
            continue
        unvan, no = k
        sonuc.append((unvan, no, satir_toplam_tutari(basliklar, s)))
    return sonuc


def _iki_kaynaktan(sutundan, metinden, sutun_adi):
    """Sutun ve satir metni bulgularini birlestirir.

    Luca durumu bazen kendi sutununda degil "Onay Durumu" gibi bir sutunda
    yaziyor; yalnizca sutuna bakilinca bos sutun bulgusu satir metnindeki
    kaydi orten bir sonuc veriyordu.
    """
    metinden = metinden or []
    if sutundan is None:
        return metinden, "satir metni"
    birlesik = list(sutundan)
    bilinen = {id(s) for s in sutundan}
    birlesik += [s for s in metinden if id(s) not in bilinen]
    if len(birlesik) == len(sutundan):
        return birlesik, sutun_adi
    return birlesik, f"{sutun_adi} + satir metni" if sutundan else "satir metni"


def excelden_sonuca_isle(sonuc, yol, klasor, log):
    """Inen Excel'i asil kaynak alir: satirlar, tevkifat ve iptal/itiraz.

    Ekran kazimaya gore guvenilir, cunku sutun basliklari belli.
    """
    basliklar, satirlar = excelden_tablo(yol, log)
    if not satirlar:
        return []
    yaz(f"    Excel'den {len(satirlar)} satir okundu", log)
    sonuc["fatura_sayisi"] = len(satirlar)
    sonuc["faturalar"] = [list(k) for k in fatura_kimlikleri_tutarli(basliklar, satirlar)]
    with open(klasor / "liste.csv", "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows([basliklar] + satirlar)

    if sonuc.get("belge_tipi") in rapor.TEVKIFAT_EKRANLARI:
        tevkifatlilar, kaynak = _iki_kaynaktan(
            sutunlu_satirlar(basliklar, satirlar, "TEVKIFAT"),
            tevkifatli_satirlar(satirlar), "Tevkifat sutunu")
    else:  # satis ekrani: tevkifat KDV2 beyanina girmiyor
        tevkifatlilar, kaynak = [], ""
    if len(tevkifatlilar) > sonuc.get("tevkifat", 0):
        sonuc["tevkifat"] = len(tevkifatlilar)
    if tevkifatlilar:
        with open(klasor / "tevkifatli.csv", "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows([basliklar] + tevkifatlilar)
        yaz(f"    DIKKAT: {len(tevkifatlilar)} tevkifatli alis faturasi ({kaynak}, KDV2)", log)

    iptaller, _ = _iki_kaynaktan(
        sutunlu_satirlar(basliklar, satirlar, "IPTAL", "ITIRAZ"),
        iptal_itiraz_satirlari(satirlar), "Iptal sutunu")
    if len(iptaller) > sonuc.get("iptal_itiraz", 0):
        sonuc["iptal_itiraz"] = len(iptaller)
    if iptaller:
        with open(klasor / "iptal-itiraz.csv", "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows([basliklar] + iptaller)
        yaz(f"    DIKKAT: {len(iptaller)} faturada iptal/itiraz var", log)

    # KDV2 kontrolu icin toplam matrah/KDV: iptal/itiraz olan faturalar
    # zaten hicbir zaman gerceklesmedigi icin toplama dahil edilmez.
    iptal_id = {id(s) for s in iptaller}
    sayilan_satirlar = [s for s in satirlar if id(s) not in iptal_id]
    matrah_sutunlari, kdv_sutunlari = matrah_kdv_sutunlari(basliklar)
    if sayilan_satirlar and not (matrah_sutunlari or kdv_sutunlari):
        yaz("    UYARI: Matrah/KDV sutunu bulunamadi, tutar toplanamadi"
            f" (basliklar: {', '.join(b for b in basliklar if b)})", log)
    else:
        sonuc["matrah"] = _satirlarin_tutari(sayilan_satirlar, matrah_sutunlari)
        sonuc["kdv"] = _satirlarin_tutari(sayilan_satirlar, kdv_sutunlari)
    return satirlar


def tevkifatli_satirlar(satirlar, belge_tipi=None):
    """Ekranda 'tevkifat' yazan satirlar (KDV2 icin isaret).

    Satis ekranlarinda aranmaz: KDV2 beyani alis faturalarindaki tevkifat icin
    verildiginden satis tarafindaki tevkifat uyari uretmemeli.
    """
    if belge_tipi is not None and belge_tipi not in rapor.TEVKIFAT_EKRANLARI:
        return []
    return [s for s in satirlar if TEVKIFAT_DESENI.search(sadelestir(" ".join(s)))]


def indirme_islevi():
    return indir_yakalayarak if AYAR.get("indirmeyi_yakala") else indir


def firma_isle(page, firma, belge_tipi, araliklar, cikti_kok, log, azami_deneme=3,
               firma_secili=False):
    sonuc = {"firma": firma, "belge_tipi": belge_tipi, "fatura_sayisi": 0,
             "durum": "", "dosyalar": [], "indirilemeyen": 0, "iptal_itiraz": 0,
             "tevkifat": 0, "donem": "", "not": "", "faturalar": [],
             "matrah": 0, "kdv": 0}
    klasor = cikti_kok / dosya_adi_yap(firma) / belge_tipi
    klasor.mkdir(parents=True, exist_ok=True)

    acik_pencereleri_kapat(page, log)
    if not firma_secili:  # ikinci belge tipinde firma zaten secili
        yaz("    Firma seciliyor", log)
        firma_sec(page, firma, log)
    istenen_bas = tarih_cozumle(araliklar[0][0])
    istenen_bit = tarih_cozumle(araliklar[-1][1])
    indirme_bas, indirme_bit = hedef_ay_araligi(istenen_bas, istenen_bit)
    indirme_araligi = (indirme_bas.strftime(TARIH_BICIMI), indirme_bit.strftime(TARIH_BICIMI))
    if not donem_ayarla(page, firma, istenen_bas, istenen_bit, log):
        sonuc["durum"] = "donem disi"
        return sonuc
    donem_bas, donem_bit = calisma_donemi(page)
    if donem_bas and donem_bit:
        sonuc["donem"] = f"{indirme_bas:%d/%m/%Y}-{indirme_bit:%d/%m/%Y}"
    if indirme_bit != istenen_bit and not firma_secili:
        yaz(f"    Hedef donem: {indirme_araligi[0]} - {indirme_araligi[1]}"
            f" (GIB sorgusu {istenen_bas:%d/%m/%Y} - {istenen_bit:%d/%m/%Y})", log)

    yaz("    Menuye gidiliyor", log)
    menu_basla = time.time()
    menuye_git(page, belge_tipi)
    gecen_menu = time.time() - menu_basla
    if gecen_menu > 5:  # nerede beklendigi gunlukten anlasilsin
        yaz(f"    Menu {int(gecen_menu)} sn'de acildi", log)
    # Ekranda hic fatura yoksa Luca acilista "Her hangi bir fatura bulunamadi"
    # penceresi gosteriyor; Tamam denmeden ekranla hicbir sey yapilamiyor.
    # Tek deneme yeter: pencere gec cikarsa sorgu adimlari yine kapatiyor,
    # burada beklemek her ekranda bos yere saniyeler harciyordu.
    if fatura_yok_penceresini_kapat(page):
        yaz("    'Fatura bulunamadi' penceresi kapatildi", log)

    interaktif = belge_tipi in IKI_KADEMELI
    sadece_excel = belge_tipi in SADECE_EXCEL  # bu ekranlarda XML inmez, Excel iner
    # Interaktif V.D. ekraninda fatura, duzenlendigi tarihle degil ait oldugu
    # donemle listelendigi icin hedef ayin disini sorgulamaya gerek yok.
    # GIB 5000/30000 de artik aylik sorgu kabul ediyor; orada tarih araligi
    # daralmaz, yalnizca parca boyu 7 gunden 30 gune cikar.
    if interaktif:
        sorgu_araliklari = tarih_araliklari(indirme_bas, indirme_bit, AYLIK_AZAMI_GUN)
    elif belge_tipi in AYLIK_SORGU:
        sorgu_araliklari = tarih_araliklari(istenen_bas, istenen_bit, AYLIK_AZAMI_GUN)
    else:
        sorgu_araliklari = araliklar
    kalan_hata = 0
    if interaktif:
        interaktif_sorgula(page, sorgu_araliklari, log, indirme_araligi)
    else:
        yetkisiz = False
        for bas, bit in sorgu_araliklari:
            if yetkisiz:  # servise yetki yok: kalan tarih araliklari denenmez
                break
            onceki_hata = None
            for deneme in range(1, azami_deneme + 1):
                try:
                    basarisiz = gibden_getir(page, bas, bit, log)
                except LookupError as e:
                    # bu ekranda GIB sorgusu yoksa liste yine de okunur
                    yaz(f"    UYARI: sorgu yapilamadi ({e}); ekrandaki liste kullanilacak", log)
                    sonuc["not"] = "GIB sorgusu bu ekranda yok"
                    basarisiz = 0
                    onceki_hata = None
                    break
                if basarisiz == YETKI_YOK:
                    sonuc["not"] = "bu firmanin bu servise yetkisi yok"
                    yetkisiz = True
                    basarisiz = 0
                    break
                if basarisiz <= 0:
                    break
                # sayi azalmiyorsa karsi sunucu yanit vermiyor demektir; tekrar denemek bos
                if onceki_hata is not None and basarisiz >= onceki_hata:
                    yaz(f"    {basarisiz} fatura tekrarda da inmedi (kaynak sunucu yanit vermiyor)", log)
                    kalan_hata += basarisiz
                    break
                if deneme == azami_deneme:
                    yaz(f"    {basarisiz} fatura {azami_deneme} denemede de indirilemedi", log)
                    kalan_hata += basarisiz
                    break
                onceki_hata = basarisiz
                yaz(f"    Tekrar sorgulaniyor ({deneme + 1}/{azami_deneme})", log)
                page.wait_for_timeout(5000)
        # liste her 7 gunluk parcada degil, tum sorgular bitince bir kez tazelenir
        listeyi_yenile(page, log)
        # ekranda son parcanin filtresi kalmasin; indirme tum donem uzerinden
        belge_ara(page, indirme_araligi[0], indirme_araligi[1], log)
    sonuc["indirilemeyen"] = kalan_hata

    yaz("    Tablo okunuyor", log)
    sayi = None
    kutular, kutu_sayisi, kutu_cercevesi = secim_kutulari(page)
    satirlar = kutulardan_satirlar(kutular, kutu_sayisi)
    fr = kutu_cercevesi
    if not satirlar and interaktif:
        # sorgu listeyi kendiliginden doldurmadiysa kayitli faturalari listele
        sayi = interaktif_kayit_sayisi(page)
        yaz(f"    Ekrandaki kayit sayisi: {sayi if sayi is not None else 'okunamadi'}", log)
        if not sayi:
            dugmeye_bas(page, INTERAKTIF_LISTELE, sure=5000)
            sayi = listeyi_bekle(page, log, azami_saniye=30)
        kutular, kutu_sayisi, kutu_cercevesi = secim_kutulari(page)
        satirlar = kutulardan_satirlar(kutular, kutu_sayisi)
        fr = kutu_cercevesi
    if not satirlar:
        fr, satirlar = tabloyu_oku(page, kutu_cercevesi)
    # liste ekrandan okunamazsa sorun degil: asil kaynak inen Excel
    sonuc["fatura_sayisi"] = len(satirlar) or (sayi or 0)
    yaz(f"    {len(satirlar)} satir listelendi"
        + (f" (ekranda {sayi} kayit)" if sayi and not satirlar else ""), log)
    tevkifatlilar, ekran_tevkifat = [], set()
    excel_alindi = False

    if not satirlar and interaktif and not sayi:
        # Ekran kayit sayisini da veremiyorsa liste hakkinda hicbir sey bilmiyoruz;
        # bu durumda Excel hemen alinir. Sayi biliniyorsa Excel iptal sorgusundan
        # sonra indirilir, boyle iptal/itiraz durumu da dosyaya yansir.
        # e-Arsiv ekraninda liste okunamayabiliyor; Excel'i indirip oradan okuruz
        # (bu ekranda Excel tum listeyi indirdigi icin secim yapilmaz)
        yol = indirme_islevi()(page, "Excel", klasor, "liste", log,
                               azami_saniye=AYAR["indirme_saniye"], pencere_acilir=False)
        if yol:
            excel_alindi = True
            sonuc["dosyalar"].append(yol.name)
            satirlar = excelden_sonuca_isle(sonuc, yol, klasor, log)
            if not satirlar:
                yaz("    Excel'de de satir bulunamadi", log)
                if sayi:  # ekran sayiyi biliyor, en azindan o rapora gecsin
                    sonuc["fatura_sayisi"] = sayi
                    sonuc["not"] = f"ekranda {sayi} kayit var, liste okunamadi"

    if satirlar:
        with open(klasor / "liste.csv", "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(satirlar)
        sonuc["faturalar"] = [list(k) for k in fatura_kimlikleri(satirlar)]
        tevkifatlilar = tevkifatli_satirlar(satirlar, belge_tipi)
        ekran_tevkifat = {no for _, no in fatura_kimlikleri(tevkifatlilar)}
        sonuc["tevkifat"] = len(tevkifatlilar)
        if tevkifatlilar:
            with open(klasor / "tevkifatli.csv", "w", encoding="utf-8-sig", newline="") as f:
                csv.writer(f).writerows(tevkifatlilar)
            yaz(f"    DIKKAT: {len(tevkifatlilar)} tevkifatli alis faturasi (KDV2)", log)

    # satir listesi cikarilamasa bile ekranda kayit varsa islemler yapilmali
    # (iptal/itiraz bu yuzden atlaniyordu)
    satir_sayisi = len(satirlar) or (sayi or 0)

    # Cok faturali firmalar (e-fatura portalinden elle indirilenler) atlanir:
    # binlerce belgeyi indirmek gece calismasinin tamamini tiketiyor.
    azami = AYAR.get("azami_fatura") or 0
    ekran_sayisi = interaktif_kayit_sayisi(page) or 0  # sayfalamada satir sayisi yaniltir
    gercek_sayi = max(satir_sayisi, ekran_sayisi)
    if azami and gercek_sayi > azami:
        yaz(f"    ATLANDI: {gercek_sayi} fatura (sinir {azami}); portalden elle indirilecek", log)
        sonuc["fatura_sayisi"] = gercek_sayi
        sonuc["durum"] = "atlandi (cok fatura)"
        sonuc["not"] = f"{gercek_sayi} fatura, sinir {azami}: portalden elle indirin"
        return sonuc
    if satir_sayisi:
        # Belge indir (XML): interaktif V.D. ekraninda bu buton yok, secim de gerekmiyor
        if not interaktif and not sadece_excel:
            secilen = hepsini_sec(page, fr, satir_sayisi)
            if secilen:
                yaz(f"    {secilen} kayit isaretlendi, indirme basliyor", log)
            else:
                yaz("    UYARI: hicbir kayit isaretlenemedi, indirme yine de denenecek", log)
            yol = indirme_islevi()(page, "Seçilenleri İndir", klasor, "belgeler", log,
                                   azami_saniye=AYAR["indirme_saniye"])
            if yol:
                sonuc["dosyalar"].append(yol.name)
                if yol.suffix.lower() == ".zip":
                    # XML bozuk inebildigi gibi ekranda da sutun olmayabiliyor;
                    # iki kaynagin birlesimi alinir
                    xml_tevkifat = zipten_tevkifatlilar(yol)
                    if xml_tevkifat:
                        sonuc["tevkifat"] = max(len(ekran_tevkifat | xml_tevkifat),
                                                len(tevkifatlilar), len(xml_tevkifat))
                        yaz(f"    XML'de {len(xml_tevkifat)} tevkifatli fatura bulundu"
                            f" (toplam {sonuc['tevkifat']})", log)
            if not sayfa_canli(page):
                if yol:
                    # belge paketi elimizde; firmayi tekrar sorgulamaya gerek yok
                    yaz("    Belgeler indi ama tarayici kapandi;"
                        " Excel'in indirilme sorgusu bu firmada atlandi", log)
                    sonuc["durum"] = "tamam (iptal eksik)"
                    sonuc["not"] = "tarayici belge indirmeden sonra kapandi"
                    return sonuc
                # dosya yok: firmayi 'tamam' sayma, ana dongu bastan denesin
                raise RuntimeError("tarayici indirme sirasinda kapandi")


        # İptal/itiraz sorgusu Excel'den ONCE yapılır (interaktif V.D. için flow: sorgu -> iptal -> excel)
        if AYAR["iptal_itiraz"] and belge_tipi in IPTAL_EKRANLARI:
            try:
                # Sorgu tum listeye uygulanir; fatura isaretlemeye gerek yok.
                # Sorgu ile ayni araliklar kullanilir: GIB alis ekraninda 7 gunluk
                # parcalar (tek seferde sorulunca iptaller cikmiyor), interaktif
                # V.D. ekraninda donemin tamami icin tek sorgu.
                basarili = iptal_itiraz_sorgula(page, sorgu_araliklari, log, interaktif)

                if basarili:
                    # sorgu durum sutununu degistirir; liste yeniden okunur
                    kutular, kutu_sayisi, yeni_fr = secim_kutulari(page)
                    yeni_satirlar = kutulardan_satirlar(kutular, kutu_sayisi)
                    if yeni_satirlar:
                        satirlar, fr = yeni_satirlar, yeni_fr
                        sonuc["fatura_sayisi"] = len(satirlar)
                        with open(klasor / "liste.csv", "w", encoding="utf-8-sig", newline="") as f:
                            csv.writer(f).writerows(satirlar)
                        sonuc["tevkifat"] = len(tevkifatli_satirlar(satirlar, belge_tipi))
                    iptaller = iptal_itiraz_satirlari(satirlar)
                    sonuc["iptal_itiraz"] = len(iptaller)
                    if iptaller:
                        with open(klasor / "iptal-itiraz.csv", "w", encoding="utf-8-sig", newline="") as f:
                            csv.writer(f).writerows(iptaller)
                        yaz(f"    DIKKAT: {len(iptaller)} faturada iptal/itiraz var", log)
                    else:
                        yaz("    Iptal/itiraz kaydi yok", log)
                elif interaktif:
                    yaz("    UYARI: İnteraktif V.D.'de iptal sorgusu başarısız, Excel al devam edecek", log)
            except Exception as e:
                yaz(f"    Iptal/itiraz sorgusu yapilamadi ({type(e).__name__}: {e})", log)
                if interaktif:
                    yaz("    UYARI: İnteraktif V.D.'de iptal sorgusu hata, Excel al devam edecek", log)
                acik_pencereleri_kapat(page, log)

        # Excel al: İptal sorgusu sonrasında (interaktif V.D. için) veya normal flow'ta
        if not excel_alindi:
            acik_pencereleri_kapat(page, log)
            fatura_yok_penceresini_kapat(page)
            page.wait_for_timeout(2000)  # İptal sorgusu sonrası sayfa stabilize olması için
            if not interaktif:
                # iptal sonrasi Yenile suzgeci sifirliyor; Excel hedef donemi kapsasin
                belge_ara(page, indirme_araligi[0], indirme_araligi[1], log)
            # İptal sorgusu sonrası satır sayısı değişmiş olabilir, güncelle
            satir_sayisi_guncel = len(satirlar) or satir_sayisi
            kutular, kutu_sayisi, guncel_fr = secim_kutulari(page)
            # Interaktif V.D. ekraninda Excel tum listeyi indiriyor; secim gerekmiyor.
            # GIB alis ekrani ise "once faturalari seciniz" uyarisi veriyor.
            if guncel_fr is not None and not interaktif:
                fr = guncel_fr
                secilen = hepsini_sec(page, fr, satir_sayisi_guncel) or 0
                if secilen == 0:
                    yaz("    UYARI: Faturalar secilemedi", log)
                page.wait_for_timeout(3000)
            yol = indirme_islevi()(page, "Excel", klasor, "liste", log,
                                   azami_saniye=max(AYAR["indirme_saniye"], 60), pencere_acilir=False)
            if yol:
                sonuc["dosyalar"].append(yol.name)
                # Excel iptal/itiraz ve tevkifat sutunlarini icerdigi icin her iki
                # ekranda da asil kaynak odur; ekran kazima yalnizca yedek
                excel_satirlari = excelden_sonuca_isle(sonuc, yol, klasor, log)
                if excel_satirlari:
                    satirlar = excel_satirlari


        # belge paketi inmediyse firma tamamlanmis sayilmaz; ozette goze carpsin
        sonuc["durum"] = ("tamam" if (interaktif or sadece_excel or sonuc["dosyalar"])
                          else "dosya inmedi")
    else:
        # GIB'de fatura vardi ama kaynak sunucudan inmedi: "fatura yok" demek yaniltici
        sonuc["durum"] = "kaynaktan inmedi" if kalan_hata else "fatura yok"
    return sonuc


_duraklama_islenen = set()


def duraklamalari_engelle(ctx, page):
    """Luca sayfalarindaki 'debugger' duraklamalari sekmeyi dondurdugu icin atlanir."""
    if page is None or page in _duraklama_islenen:
        return
    _duraklama_islenen.add(page)
    try:
        cdp = ctx.new_cdp_session(page)
    except Exception:
        return

    def devam_et(_=None):
        try:
            cdp.send("Debugger.resume")
        except Exception:
            pass

    for komut, parametre in (
        ("Debugger.enable", None),
        ("Debugger.setSkipAllPauses", {"skip": True}),
        ("Debugger.resume", None),  # sayfa zaten duraklamissa serbest birak
    ):
        try:
            cdp.send(komut, parametre) if parametre else cdp.send(komut)
        except Exception:
            continue
    try:
        cdp.on("Debugger.paused", devam_et)
    except Exception:
        pass


def giris_bilgileri(ayarlar):
    """ayarlar.json'daki giris bilgileri; eksikse None."""
    uye = str(ayarlar.get("uye_no") or "").strip()
    kullanici = str(ayarlar.get("kullanici_adi") or "").strip()
    parola = str(ayarlar.get("parola") or "")
    return (uye, kullanici, parola) if uye and kullanici and parola else None


def dogrulama_ekrani_mi(page):
    """Girisden sonra iki asamali dogrulama ekrani cikmis mi."""
    for isaret in DOGRULAMA_ISARETLERI:
        if gorunur_mu(page, isaret, sure=800):
            return isaret
    return ""


def captcha_ekrani_mi(page):
    """Captcha ekrani: iki asamali dogrulama kapaliyken Luca bunu soruyor."""
    for isaret in CAPTCHA_ISARETLERI:
        if gorunur_mu(page, isaret, sure=800):
            return isaret
    return ""


def dogrulama_kodu(ayarlar, log=None):
    """ayarlar.json'daki gizli anahtardan o anki dogrulama kodunu uretir.

    Luca'da iki asamali dogrulamayi "kimlik dogrulayici uygulama" ile
    acarsaniz kurulum ekranindaki gizli anahtari dogrulama_anahtari alanina
    yazin; kodu bot kendisi hesaplar ve giris tumuyle otomatik olur.
    Uretilemezse sebebi yazilir: sessizce elle girise dusmek kullaniciya
    neyin eksik oldugunu soylemiyordu.
    """
    anahtar = str(ayarlar.get("dogrulama_anahtari") or "").replace(" ", "")
    if not anahtar:
        yaz("    ayarlar.json'da dogrulama_anahtari yok, kod elle girilmeli", log)
        return ""
    try:
        import pyotp
    except ImportError:
        yaz("    pyotp kurulu degil (kurulum.bat calistirin), kod elle girilmeli", log)
        return ""
    try:
        return pyotp.TOTP(anahtar).now()
    except Exception as e:
        yaz(f"    Dogrulama kodu uretilemedi ({type(e).__name__}), anahtari kontrol edin", log)
        return ""


def dogrulama_kodunu_gir(page, kod):
    """Dogrulama kodunu ekrandaki kutuya yazip onaylar."""
    for secici in ("input[type=text]:visible", "input[type=tel]:visible",
                   "input[type=number]:visible", "input[type=password]:visible"):
        try:
            loc = page.locator(secici)
            if not loc.count():
                continue
            kutu = loc.first
            kutuya_yaz(kutu, kod)
            if not varsa_tikla(page, DOGRULAMA_ONAY, sure=4000):
                kutu.press("Enter")
            page.wait_for_timeout(4000)
            return True
        except Exception:
            continue
    return False


def urun_sec(page, log=None, sure=30000):
    """Giris sonrasi cikan urun secim ekranindan Mali Musavir paketini secer.

    Kutular giristen birkac saniye sonra beliriyor; tek seferlik tiklama
    denemesi erken kaldigi icin bot bu ekranda bekliyordu.
    """
    bitis = time.time() + sure / 1000
    while True:
        try:
            if UYGULAMA_PARCASI in page.url:  # uygulama zaten acildi
                return True
        except Exception:
            return False
        tiklanan = varsa_tikla(page, URUN_ADAYLARI, sure=1200)
        if tiklanan:
            yaz(f"    Urun secildi: {tiklanan}", log)
            page.wait_for_timeout(3000)
            return True
        if time.time() >= bitis:
            return False
        page.wait_for_timeout(1000)


def otomatik_giris(page, ayarlar, log):
    """Luca girisini ayarlar.json'daki bilgilerle kendisi yapar.

    Bilgiler eksikse ya da form bulunamazsa False doner; bu durumda
    kullanici elle giris yapar, akis degismez.
    """
    bilgiler = giris_bilgileri(ayarlar)
    if bilgiler is None:
        return False
    uye, kullanici, parola = bilgiler
    yaz("Otomatik giris yapiliyor...", log)
    try:
        if UYGULAMA_PARCASI in page.url:  # oturum zaten acik
            return True
        if "giris.erp" not in page.url.lower():
            varsa_tikla(page, SISTEM_GIRIS, sure=5000)
            page.wait_for_timeout(2500)
        if "giris.erp" not in page.url.lower():
            page.goto(GIRIS_SAYFASI)
        page.wait_for_timeout(2500)

        parolalar = page.locator("input[type=password]:visible")
        metinler = page.locator("input[type=text]:visible, input:not([type]):visible")
        if not parolalar.count() or metinler.count() < 2:
            yaz("UYARI: giris formu bulunamadi, elle giris yapin", log)
            return False
        kutuya_yaz(metinler.nth(0), uye)
        kutuya_yaz(metinler.nth(1), kullanici)
        kutuya_yaz(parolalar.first, parola)
        if not varsa_tikla(page, GIRIS_DUGMESI, sure=5000):
            parolalar.first.press("Enter")
        page.wait_for_timeout(6000)

        captcha = captcha_ekrani_mi(page)
        if captcha:
            yaz("    Captcha ekrani cikti; bot bunu gecemez.", log)
            yaz("    Luca'da iki asamali dogrulamayi acarsaniz captcha kalkar"
                " (Luca bunu giris ekraninda kendisi yaziyor).", log)
            return False

        isaret = dogrulama_ekrani_mi(page)
        if isaret:
            kod = dogrulama_kodu(ayarlar, log)
            if kod and dogrulama_kodunu_gir(page, kod):
                yaz("    Dogrulama kodu girildi", log)
                if not dogrulama_ekrani_mi(page):
                    urun_sec(page, log)
                    return True
                yaz("    UYARI: dogrulama kodu kabul edilmedi", log)
            # kod uretilemiyorsa (SMS vb.) elle girilmesi beklenir
            dakika = max(0, int(ayarlar.get("dogrulama_bekleme_dakika", 5)))
            yaz(f"    Iki asamali dogrulama ekrani ('{isaret}')."
                f" Kodu elle girin, {dakika} dk bekleniyor...", log)
            bitis = time.time() + dakika * 60
            while time.time() < bitis and dogrulama_ekrani_mi(page):
                page.wait_for_timeout(3000)
            if dogrulama_ekrani_mi(page):
                yaz("    UYARI: dogrulama tamamlanmadi, giris yapilamadi", log)
                return False
            yaz("    Dogrulama tamamlandi", log)

        urun_sec(page, log)  # urun secim ekrani cikarsa
        return True
    except Exception as e:
        yaz(f"UYARI: otomatik giris yapilamadi ({type(e).__name__}), elle giris yapin", log)
        return False


def kullanici_bekle(ctx, mesaj):
    """Duz input() Playwright olaylarini dondurur; beklerken sayfa olaylari islenmeye devam etmeli."""
    hazir = threading.Event()

    def oku():
        try:
            input(mesaj)
        finally:
            hazir.set()

    threading.Thread(target=oku, daemon=True).start()
    while not hazir.is_set():
        acik = [p for p in ctx.pages if not p.is_closed()]
        for p in acik:
            duraklamalari_engelle(ctx, p)
        if acik:
            try:
                acik[0].wait_for_timeout(250)
                continue
            except Exception:
                pass
        time.sleep(0.25)


def kullanici_metni_al(ctx, mesaj):
    """kullanici_bekle ile ayni, ama ENTER'a basilmadan once yazilan metni dondurur."""
    hazir = threading.Event()
    kutu = {"metin": ""}

    def oku():
        try:
            kutu["metin"] = input(mesaj)
        finally:
            hazir.set()

    threading.Thread(target=oku, daemon=True).start()
    while not hazir.is_set():
        acik = [p for p in ctx.pages if not p.is_closed()]
        for p in acik:
            duraklamalari_engelle(ctx, p)
        if acik:
            try:
                acik[0].wait_for_timeout(250)
                continue
            except Exception:
                pass
        time.sleep(0.25)
    return kutu["metin"]


def profil_klasoru(log=None):
    """Tarayici profilinin yeri.

    Varsayilan, calisan surumdeki gibi program klasorudur. OneDrive disina
    almak icin --profil-yerel kullanilir (o kuram dogrulanmadi, secenek olarak
    duruyor).
    """
    eski = KOK / ".tarayici-profili"
    if not AYAR.get("profil_yerel"):
        return eski
    yerel = os.environ.get("LOCALAPPDATA") or os.environ.get("XDG_CACHE_HOME")
    if not yerel:
        return eski
    yeni = Path(yerel) / "luca-bot" / "tarayici-profili"
    if yeni.exists():
        return yeni
    try:
        yeni.parent.mkdir(parents=True, exist_ok=True)
        if eski.exists():
            shutil.move(str(eski), str(yeni))
            yaz(f"Tarayici profili OneDrive disina tasindi: {yeni}", log)
        else:
            yeni.mkdir(parents=True, exist_ok=True)
        return yeni
    except Exception as e:
        yaz(f"UYARI: profil tasinamadi ({type(e).__name__}), eski konum kullanilacak", log)
        return eski


# Calisan surumdeki sira: once kurulu Chrome. --tarayici ile degistirilebilir.
TARAYICILAR = [("chrome", "Google Chrome"), ("msedge", "Microsoft Edge"), (None, "Playwright Chromium")]


def kanal_profili(profil, kanal):
    """Her tarayicinin kendi profili olmali; Edge, Chrome'un profilini acamiyor."""
    if kanal == "chrome":
        return profil  # mevcut profil Chrome'a ait, oturum korunsun
    return profil.parent / f"{profil.name}-{kanal or 'chromium'}"


def tarayici_ac(pw, profil, log, gunluk=False):
    """Once bilgisayarda kurulu Chrome/Edge denenir; Playwright'in kendi tarayicisi son care.

    Chrome bu makinede indirme sirasinda cokuyorsa --tarayici edge ile
    digerine gecilebilir.
    """
    hatalar = []
    tercih = AYAR.get("tarayici")  # "chrome" / "edge" / "chromium" ya da None
    kanallar = {"chrome": "chrome", "edge": "msedge", "chromium": None}
    if tercih in kanallar:
        adaylar = [t for t in TARAYICILAR if t[0] == kanallar[tercih]]
    else:
        adaylar = TARAYICILAR
    for kanal, ad in adaylar:
        secenekler = {"channel": kanal} if kanal else {}
        kanal_yolu = kanal_profili(profil, kanal)
        try:
            kanal_yolu.mkdir(parents=True, exist_ok=True)
            ctx = pw.chromium.launch_persistent_context(
                str(kanal_yolu), headless=False, accept_downloads=True,
                args=["--start-maximized"] + (["--enable-logging", "--v=1"] if gunluk else []),
                # Playwright varsayilan olarak --disable-breakpad geciyor; tani
                # modunda kaldiriyoruz ki cokme Windows olay gunlugune dussun
                ignore_default_args=["--enable-automation"]
                + (["--disable-breakpad"] if gunluk else []),
                chromium_sandbox=True, no_viewport=True, **secenekler
            )
            yaz(f"Tarayici: {ad}", log)
            ctx.on("page", lambda p: duraklamalari_engelle(ctx, p))
            ctx.on("close", lambda _: yaz("UYARI: tarayici kapandi (Chrome cokmus olabilir)", log))
            for p in ctx.pages:
                duraklamalari_engelle(ctx, p)
            return ctx
        except Exception as e:
            hatalar.append(f"  {ad}: {str(e).splitlines()[0][:120]}")
    raise RuntimeError(
        "Hicbir tarayici acilamadi:\n" + "\n".join(hatalar)
        + "\n\nCozum: Google Chrome kurun (google.com/chrome) veya"
        " internet baglantisi duzelince tarayici-indir.bat dosyasini calistirin."
    )


def uygulama_sayfasi_bul(ctx):
    """Giris sonrasi Luca birkac pencere aciyor; firma listesini iceren sayfayi sec."""
    acik = [p for p in ctx.pages if not p.is_closed()]
    for p in acik:
        try:
            firma_secici(p)
            return p
        except Exception:
            continue
    for p in acik:
        try:
            if UYGULAMA_PARCASI in p.url and "giris" not in p.url.lower():
                return p
        except Exception:
            continue
    return None


def bekci_sekmesi_ac(ctx):
    """Kapali: ikinci sekme acmak indirme sirasindaki cokmeyi tetikliyordu."""
    return


def firma_sayisi(page):
    """Sayfadaki firma listesinde kac firma var (liste yuklenmis mi kontrolu)."""
    try:
        return len(firma_secici(page)[2])
    except Exception:
        return 0


def tarayiciyi_yeniden_baslat(pw, profil, eski_ctx, log, bekleme_saniye=90, asgari_firma=5):
    """Chrome cokerse yeniden acar; profil oturumu tasidigi icin genelde giris gerekmez.

    Luca ekranini kendiliginden bulursa (ctx, page) doner, bulamazsa (ctx, None).
    """
    try:
        if eski_ctx is not None:
            eski_ctx.close()
    except Exception:
        pass
    yaz("Tarayici yeniden aciliyor...", log)
    try:
        ctx = tarayici_ac(pw, profil, log, AYAR.get("chrome_gunlugu", False))
    except Exception as e:
        yaz(f"Tarayici yeniden acilamadi: {type(e).__name__}: {e}", log)
        return None, None
    page = ctx.pages[0] if ctx.pages else ctx.new_page()
    bekci_sekmesi_ac(ctx)
    try:
        page.bring_to_front()
        page.goto(GIRIS_URL)
    except Exception:
        pass
    bitis = time.time() + bekleme_saniye
    while time.time() < bitis:
        uygulama = uygulama_sayfasi_bul(ctx)
        # liste yarim yuklendiyse firma bulunamiyor; dolmasini bekle
        if uygulama is not None and firma_sayisi(uygulama) >= asgari_firma:
            yaz(f"Luca ekrani bulundu ({firma_sayisi(uygulama)} firma), devam ediliyor.", log)
            return ctx, uygulama
        try:
            page.wait_for_timeout(2000)
        except Exception:
            break
    yaz("Yeniden girise ihtiyac var; Luca oturumu profilden acilmadi.", log)
    return ctx, None


def sayfa_canli(page):
    try:
        return page is not None and not page.is_closed()
    except Exception:
        return False


def sayfayi_kurtar(ctx, log=None):
    """Calisilan sayfa kapanirsa tarayicida acik kalan Luca sayfasina gecer."""
    try:
        yeni = uygulama_sayfasi_bul(ctx)
    except Exception:
        yeni = None
    if yeni is None:
        return None
    yaz("    Sayfa kapanmisti, acik Luca sayfasina gecildi", log)
    try:
        duraklamalari_engelle(ctx, yeni)
    except Exception:
        pass
    return yeni


def sayfalari_ozetle(ctx):
    satirlar = []
    for p in ctx.pages:
        if p.is_closed():
            continue
        try:
            secenekler = []
            for fr in p.frames:
                try:
                    kutular = fr.locator("select")
                    for i in range(min(kutular.count(), 6)):
                        adet = kutular.nth(i).locator("option").count()
                        if adet:
                            secenekler.append(str(adet))
                except Exception:
                    continue
            satirlar.append(
                f"  - {p.url}\n"
                f"      baslik: {p.title()[:60]} | frame: {len(p.frames)}"
                f" | liste secenek sayilari: {', '.join(secenekler) or 'yok'}"
            )
        except Exception as e:
            satirlar.append(f"  - (sayfa okunamadi: {type(e).__name__})")
    return "\n".join(satirlar) or "  (acik sayfa yok)"


def oturumu_yenile(page, ctx, ayarlar, log):
    """Luca oturumu dustugunde giris sayfasina donup yeniden girer.

    Oturum dustugunde ekranda ne menu ne firma listesi kaliyor; bot bunu
    firma hatasi sanip pes ediyordu. Giris bilgileri ayarlar.json'da varsa
    gece calismasi kaldigi yerden surebilir.
    """
    if giris_bilgileri(ayarlar) is None:
        return None
    yaz("    Luca oturumu yenileniyor...", log)
    try:
        page.goto(GIRIS_URL)
        page.wait_for_timeout(2000)
    except Exception:
        return None
    if not otomatik_giris(page, ayarlar, log):
        return None
    bitis = time.time() + 60
    while time.time() < bitis:
        uygulama = uygulama_sayfasi_bul(ctx)
        if uygulama is not None:
            yaz("    Oturum yenilendi, devam ediliyor", log)
            return uygulama
        try:
            page.wait_for_timeout(2000)
        except Exception:
            return None
    return None


def sayfayi_toparla(page):
    """Hata sonrasi acik kalan diyaloglari kapatir.

    Sayfa yeniden YUKLENMEZ: Luca uygulama adresine dogrudan gidilince oturumu
    reddedip "LUCA HATA" veriyor ve sonraki tum firmalar basarisiz oluyordu.
    """
    acik_pencereleri_kapat(page)
    for _ in range(2):
        try:
            page.keyboard.press("Escape")
            page.wait_for_timeout(300)
        except Exception:
            break


# firmalar.xlsx'teki ekran sutunlari: baslik -> belge tipi
EKRAN_SUTUNLARI = {
    "e-Arşiv Alış": "e-arsiv-alis",
    "e-Arşiv Satış": "e-arsiv-satis",
    "e-Fatura Alış": "e-fatura-alis",
    "e-Fatura Satış": "e-fatura-satis",
    "GİB 5000/30000": "gib-5000",
    "TÜRMOB Alış": "turmob-alis",
    "TÜRMOB Satış": "turmob-satis",
    "e-SMM Alış": "esmm-alis",
    "e-SMM Satış": "esmm-satis",
    "İnteraktif V.D.": "e-arsiv-interaktif",
}
# ekran sutununa bunlardan biri yazilirsa o ekran o firmada acilmaz
ATLA_DEGERLERI = {"X", "HAYIR", "YOK", "ATLA", "-", "0"}


def sutun_tam_indeksi(basliklar, ad):
    """Basligi birebir eslesen sutun (e-Arsiv Alis ile Satis karismasin diye)."""
    aranan = sadelestir(ad)
    for i, baslik in enumerate(basliklar):
        if sadelestir(baslik) == aranan:
            return i
    return None


def firma_listesini_oku(yol, log=None):
    """Islenecek firmalar: {kisa ad: (kapanis tarihi, atlanacak ekranlar)}.

    Excel'de "Kısa Adı" ve "Kapanış Tarihi" sutunlari aranir. Kisa ad Luca'nin
    firma listesinde gorunen adla ayni oldugu icin eslestirme dogrudan yapilir.
    Ekran sutunlarina X yazilan belge tipleri o firmada hic acilmaz (orn.
    e-faturasi olmayan firmada e-Fatura Alis/Satis).
    """
    yol = Path(yol)
    if not yol.is_absolute():
        yol = KOK / yol
    if not yol.exists():
        yaz(f"UYARI: firma listesi bulunamadi: {yol}", log)
        return {}
    # yalnizca ilk sayfa: dosyadaki aciklama sayfasi firma sanilmasin
    basliklar, satirlar = excelden_tablo(yol, log, sadece_ilk=True)
    if not satirlar:
        yaz(f"UYARI: firma listesi okunamadi: {yol}", log)
        return {}
    ad_i = sutun_indeksi(basliklar, "KISA AD")
    kapanis_i = sutun_indeksi(basliklar, "KAPANIS")
    if ad_i is None:
        ad_i = 0
    ekran_i = {}
    for baslik, tip in EKRAN_SUTUNLARI.items():
        i = sutun_tam_indeksi(basliklar, baslik)
        if i is not None:
            ekran_i[tip] = i
    liste = {}
    for satir in satirlar:
        ad = satir[ad_i].strip() if ad_i < len(satir) else ""
        if not ad:
            continue
        kapanis = None
        if kapanis_i is not None and kapanis_i < len(satir):
            try:
                kapanis = tarih_cozumle(satir[kapanis_i])
            except Exception:
                kapanis = None
        atlanan = {tip for tip, i in ekran_i.items()
                   if i < len(satir) and sadelestir(satir[i]) in ATLA_DEGERLERI}
        liste[ad] = (kapanis, atlanan)
    return liste


def listede_bul(ad, liste):
    """Luca adi listedeki hangi kisa ada denk geliyor (kisaltilmis adlar icin).

    Once birebir eslesme aranir; yoksa bas kismi tutanlardan EN UZUN olani
    secilir. Aksi halde "ADEM", "ADEM MERGE" firmasiyla da esleserek yanlis
    firmanin ayarlarini uyguluyordu.
    """
    k = karsilastir(ad)
    if not k:
        return None
    en_iyi, en_uzun = None, -1
    for liste_adi in liste:
        a = karsilastir(liste_adi)
        if not a:
            continue
        if a == k:
            return liste_adi
        # Luca adlari kisaltarak gosterdigi icin listedeki daha uzun ad da tutar.
        # Ters yon ("ADEM" listedeki ad, Luca'da "ADEM AKÇAY") ancak yeterince
        # uzun adlarda kabul edilir; kisa adlar baska firmalara yapisiyordu.
        if (a.startswith(k) or (len(a) >= 8 and k.startswith(a))) and len(a) > en_uzun:
            en_iyi, en_uzun = liste_adi, len(a)
    return en_iyi


def ozet_yaz(ozet_yolu, kalan_yolu, sonuclar, bekleyenler, belge_tipi):
    """Ozeti her firmadan sonra yeniden yazar; islenmeyenler 'bekliyor' olarak gorunur."""
    with open(ozet_yolu, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Firma", "Belge Tipi", "Fatura Sayisi", "Indirilemeyen", "Iptal/Itiraz",
                    "Durum", "Dosyalar"])
        for s in sonuclar:
            w.writerow([s["firma"], s["belge_tipi"], s["fatura_sayisi"], s.get("indirilemeyen", 0),
                        s.get("iptal_itiraz", 0), s["durum"], "; ".join(s["dosyalar"])])
        for firma in bekleyenler:
            w.writerow([firma, belge_tipi, "", "", "", "bekliyor", ""])
    kalan_yolu.write_text(", ".join(bekleyenler), encoding="utf-8")


def hata_kaydet(page, klasor, firma):
    klasor.mkdir(parents=True, exist_ok=True)
    ad = dosya_adi_yap(firma)
    try:
        page.screenshot(path=str(klasor / f"{ad}.png"), full_page=True)
    except Exception:
        pass
    try:
        (klasor / f"{ad}.html").write_text(page.content(), encoding="utf-8")
    except Exception:
        pass


def main():
    ayarlar = ayarlari_oku()
    p = argparse.ArgumentParser(description="Luca toplu e-fatura indirme botu")
    p.add_argument("--firma", action="append",
                   help="Sadece bu firma(lar) islensin; virgulle ayirarak birden fazla yazilabilir")
    p.add_argument("--belge-tipi", action="append", choices=list(BELGE_TIPLERI),
                   help="Birden fazla kez verilebilir; her firmada sirayla islenir")
    p.add_argument("--hepsi", action="store_true",
                   help="Tum belge tiplerini sirayla isle (menudeki butun ekranlar)")
    p.add_argument("--karsilastir", action="store_true",
                   help="Iki e-arsiv ekranini da calistir (Akilli Entegrasyon + Interaktif V.D.)")
    p.add_argument("--limit", type=int, help="Ilk N firma ile sinirla")
    p.add_argument("--baslangic", help="GG/AA/YYYY (ayarlar.json'daki degeri ezer)")
    p.add_argument("--bitis", help="GG/AA/YYYY (ayarlar.json'daki degeri ezer)")
    p.add_argument("--tarayici-indirsin", action="store_true",
                   help="Dosyayi tarayici indirsin (varsayilan: istek yakalanip kaydedilir)")
    p.add_argument("--profil-yerel", action="store_true",
                   help="Tarayici profilini program klasoru yerine %LOCALAPPDATA% altinda tut")
    p.add_argument("--tarayici", choices=["chrome", "edge", "chromium"],
                   help="Hangi tarayici kullanilsin (Chrome cokuyorsa edge deneyin)")
    p.add_argument("--chrome-gunlugu", action="store_true",
                   help="Chrome cokerse sebebini yazmasi icin ayrintili gunluk tut")
    p.add_argument("--donem-degistirme", action="store_true",
                   help="Donemi degistirme; eski donemdeki firmalari atla (sorun cikarsa)")
    p.add_argument("--azami-fatura", type=int,
                   help="Bu sayidan cok faturasi olan firmalari atla (0: sinir yok)")
    p.add_argument("--iptal-itiraz-atla", action="store_true",
                   help="Faturalari indir ama GIB iptal/itiraz sorgusunu yapma")
    p.add_argument("--listele", action="store_true", help="Sadece firma listesini yazdir, islem yapma")
    p.add_argument("--bitince-kapat", action="store_true",
                   help="Is bitince ENTER beklemeden tarayiciyi kapat (gece calistirma icin)")
    args = p.parse_args()
    if args.hepsi:
        args.belge_tipi = list(TUM_BELGELER)
    elif args.karsilastir:
        args.belge_tipi = ["e-arsiv-alis", "e-arsiv-interaktif"]
    elif not args.belge_tipi:
        varsayilan = ayarlar.get("belge_tipi", "e-arsiv-alis")
        args.belge_tipi = varsayilan if isinstance(varsayilan, list) else [varsayilan]

    bas_metin = args.baslangic or ayarlar.get("baslangic_tarihi")
    bit_metin = args.bitis or ayarlar.get("bitis_tarihi")
    if bas_metin and bit_metin:
        try:
            baslangic, bitis = tarih_cozumle(bas_metin), tarih_cozumle(bit_metin)
        except ValueError:
            p.error("Tarihler GG/AA/YYYY biciminde olmali, orn: 01/08/2026")
        if bitis < baslangic:
            p.error("Bitis tarihi baslangictan once olamaz")
    else:
        baslangic, bitis = icinde_bulunulan_ay()
    araliklar = tarih_araliklari(baslangic, bitis)
    azami_deneme = max(1, int(ayarlar.get("tekrar_deneme", 3)))
    AYAR["azami_saniye"] = max(60, int(float(ayarlar.get("sorgu_azami_dakika", 30)) * 60))
    AYAR["durgunluk_saniye"] = max(30, int(float(ayarlar.get("durgunluk_dakika", 3)) * 60))
    AYAR["indirme_saniye"] = max(3, int(ayarlar.get("indirme_bekleme_saniye", 30)))
    # cok faturali firmalar atlanir (0: sinir yok)
    AYAR["azami_fatura"] = max(0, int(args.azami_fatura if args.azami_fatura is not None
                                      else ayarlar.get("azami_fatura", 500)))
    AYAR["iptal_itiraz"] = bool(ayarlar.get("iptal_itiraz_sorgula", True)) and not args.iptal_itiraz_atla
    AYAR["donem_degistir"] = bool(ayarlar.get("donem_degistir", True)) and not args.donem_degistirme
    AYAR["chrome_gunlugu"] = bool(args.chrome_gunlugu)
    AYAR["profil_yerel"] = bool(args.profil_yerel) or bool(ayarlar.get("profil_yerel", False))
    AYAR["indirmeyi_yakala"] = (bool(ayarlar.get("indirmeyi_yakala", True))
                                and not args.tarayici_indirsin)
    if AYAR["chrome_gunlugu"]:
        # Playwright'in tarayici cikis mesajlarini ekrana bassin; cokme sebebi
        # genelde burada yaziyor ("Target crashed", exit code, stderr)
        os.environ["DEBUG"] = "pw:browser"
    AYAR["tarayici"] = args.tarayici or ayarlar.get("tarayici") or None
    hata_siniri = max(1, int(ayarlar.get("ardisik_hata_siniri", 5)))

    cikti_kok = Path(ayarlar.get("indirme_klasoru") or "indirilenler").expanduser()
    if not cikti_kok.is_absolute():
        cikti_kok = KOK / cikti_kok
    calisma = cikti_kok / date.today().isoformat()
    calisma.mkdir(parents=True, exist_ok=True)
    log = calisma / "calisma.log"
    profil = profil_klasoru(log)
    with sync_playwright() as pw:
        ctx = tarayici_ac(pw, profil, log, AYAR["chrome_gunlugu"])
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.bring_to_front()
        page.goto(GIRIS_URL)

        if otomatik_giris(page, ayarlar, log):
            # muhasebe ekrani kendiliginden acilana kadar beklenir; bu sirada
            # urun secim ekrani gec belirmis olabilir, tekrar denenir
            bitis = time.time() + 90
            while time.time() < bitis and uygulama_sayfasi_bul(ctx) is None:
                urun_sec(page, log, sure=0)
                page.wait_for_timeout(2000)
        if uygulama_sayfasi_bul(ctx) is None and args.bitince_kapat:
            # gece modunda ENTER'a basacak kimse yok; bosuna beklenmez
            yaz("\nOtomatik giris yapilamadi (dogrulama kodu ya da sifre sorunu).", log)
            yaz("Gece modunda elle giris beklenmez, calisma baslatilmadi.", log)
            ctx.close()
            return
        if uygulama_sayfasi_bul(ctx) is None:
            yaz("\n>>> Tarayicida Luca'ya giris yapin.", log)
            yaz(">>> Girisden sonra MUHASEBE EKRANINI acin (sag ustte firma listesi gorunen ekran).", log)
            kullanici_bekle(ctx, ">>> O ekran acikken ENTER'a basin: ")
        else:
            yaz("Giris yapildi, muhasebe ekrani bulundu.", log)

        uygulama = None
        for deneme in range(1, 6):
            uygulama = uygulama_sayfasi_bul(ctx)
            if uygulama:
                break
            yaz(f"\nFirma listesi olan ekran bulunamadi ({deneme}/5). Acik pencereler:", log)
            yaz(sayfalari_ozetle(ctx), log)
            tani = calisma / "tani"
            tani.mkdir(parents=True, exist_ok=True)
            for sira, p in enumerate([x for x in ctx.pages if not x.is_closed()], 1):
                try:
                    p.screenshot(path=str(tani / f"deneme{deneme}-sayfa{sira}.png"))
                    (tani / f"deneme{deneme}-sayfa{sira}.html").write_text(p.content(), encoding="utf-8")
                except Exception:
                    pass
            yaz(f"Ekran goruntuleri kaydedildi: {tani}", log)
            yaz("\nLuca'da muhasebe modulunu acip firma listesinin gorundugu ekrana gelin.", log)
            kullanici_bekle(ctx, ">>> Hazir oldugunuzda ENTER'a basin (vazgecmek icin pencereyi kapatin): ")

        if uygulama is None:
            yaz("\nMuhasebe ekrani bulunamadi, islem durduruldu.", log)
            yaz("Yukaridaki pencere listesini gonderirseniz duzeltirim.", log)
            input(">>> Kapatmak icin ENTER: ")
            ctx.close()
            return

        page = uygulama
        duraklamalari_engelle(ctx, page)
        page.bring_to_front()
        page.on("dialog", lambda d: d.accept())
        uygulama_url = page.url
        yaz(f"Calisilan sayfa: {uygulama_url}", log)

        varsa_tikla(page, KAPAT_METINLERI)
        _, _, firmalar = firma_secici(page)
        tumu = list(firmalar)
        yaz(f"{len(firmalar)} firma bulundu", log)

        if args.listele:
            for f in firmalar:
                print(" -", f)
            dosya = calisma / "firmalar.txt"
            dosya.write_text("\n".join(firmalar), encoding="utf-8")
            yaz(f"\nListe dosyaya da yazildi: {dosya}", log)
            kullanici_bekle(ctx, ">>> Kapatmak icin ENTER: ")
            ctx.close()
            return

        if args.firma:
            aranan = [(p.strip(), karsilastir(p)) for deger in args.firma for p in deger.split(",") if p.strip()]
            bulunamayan = [ham for ham, a in aranan if not any(a in karsilastir(f) for f in firmalar)]
            if bulunamayan:
                yaz(f"Eslesmeyen arama: {', '.join(bulunamayan)}", log)
            firmalar = [f for f in firmalar if any(a in karsilastir(f) for _, a in aranan)]
            if not firmalar:
                yaz(f"\n'{', '.join(args.firma)}' ile eslesen firma yok.", log)
                yaz("Listedeki ilk 30 kayit:", log)
                for ad in tumu[:30]:
                    yaz(f"  - {ad}", log)
                yaz("\nNot: Luca adlari kisaltarak gosterebiliyor, adin bas kismini yazin.", log)
                kullanici_bekle(ctx, ">>> Kapatmak icin ENTER: ")
                ctx.close()
                return
            yaz(f"Eslesen firma(lar): {', '.join(firmalar)}", log)
        atlama_adlari = [a for a in ayarlar.get("atlanacak_firmalar", []) if a.strip()]
        if atlama_adlari:
            def ayni_firma(ad, atlanan):
                # Luca adlari kisaltarak gosterdigi icin bas kismi tutan ad da atlanir
                k, a = karsilastir(ad), karsilastir(atlanan)
                return bool(k) and bool(a) and (k.startswith(a) or a.startswith(k))

            atlananlar = [f for f in firmalar if any(ayni_firma(f, a) for a in atlama_adlari)]
            firmalar = [f for f in firmalar if f not in atlananlar]
            if atlananlar:
                yaz(f"Atlanan firma ({len(atlananlar)}): {', '.join(atlananlar)}", log)
            # --firma ile calisirken zaten tek firma var; eslesmeyen adlar dogal
            if not args.firma:
                eslesmeyen = [a for a in atlama_adlari
                              if not any(ayni_firma(f, a) for f in atlananlar)]
                if eslesmeyen:
                    yaz(f"UYARI: atlama listesinde eslesmeyen ad: {', '.join(eslesmeyen)}", log)
        firma_atlanan = {}  # firma -> o firmada acilmayacak belge tipleri
        liste_yolu = ayarlar.get("firma_listesi")
        if liste_yolu:
            izinli = firma_listesini_oku(liste_yolu, log)
            if izinli:
                yaz(f"Firma listesi: {liste_yolu} ({len(izinli)} firma)", log)
                kalanlar, disarida, kapanmis = [], [], []
                for f in firmalar:
                    liste_adi = listede_bul(f, izinli)
                    if liste_adi is None:
                        disarida.append(f)
                        continue
                    kapanis, atlanan_ekranlar = izinli[liste_adi]
                    # donem baslamadan kapanmis firmada aranacak fatura yok
                    if kapanis is not None and kapanis < baslangic:
                        kapanmis.append(f"{f} ({kapanis:%d/%m/%Y})")
                        continue
                    kalanlar.append(f)
                    if atlanan_ekranlar:
                        firma_atlanan[f] = atlanan_ekranlar
                firmalar = kalanlar
                if disarida:
                    yaz(f"Listede olmayan {len(disarida)} firma atlandi", log)
                if kapanmis:
                    yaz(f"Donem oncesi kapanan {len(kapanmis)} firma atlandi: "
                        + ", ".join(kapanmis[:12])
                        + (f" ... (+{len(kapanmis) - 12})" if len(kapanmis) > 12 else ""), log)
                if firma_atlanan:
                    toplam = sum(len(v) for v in firma_atlanan.values())
                    yaz(f"Listede {len(firma_atlanan)} firmada {toplam} ekran"
                        " isaretlenmis, o ekranlar acilmayacak", log)

        if args.limit:
            firmalar = firmalar[: args.limit]

        # Calisma yarida kesilip (bilgisayar kapanmasi, elektrik vb.) ayni gun
        # yeniden baslatilirsa, o gune ait rapor.json'da zaten tamamlanmis
        # gorunen ekranlar bulunur. Gece modunda (kimse cevap veremez) kaldigi
        # yerden otomatik devam edilir; elle calistirmada kullaniciya sorulur,
        # cunku bazen kasitli olarak hepsinin yeniden taranmasi istenebilir.
        bugun_tamam = {}
        onceki_rapor_yolu = calisma / "rapor.json"
        if onceki_rapor_yolu.exists():
            try:
                onceki_rapor = json.loads(onceki_rapor_yolu.read_text(encoding="utf-8"))
            except Exception:
                onceki_rapor = {}
            aday_tamam = {}
            for firma, kayit in onceki_rapor.items():
                durumlar = kayit.get("durumlar", {}) if isinstance(kayit, dict) else {}
                tamam_tipler = {tip for tip in args.belge_tipi
                               if ekran_tamamlanmis_mi(durumlar.get(tip, ""))}
                if tamam_tipler:
                    aday_tamam[firma] = tamam_tipler
            if aday_tamam:
                toplam = sum(len(v) for v in aday_tamam.values())
                if args.bitince_kapat:
                    bugun_tamam = aday_tamam
                    yaz(f"Bugun daha once {len(aday_tamam)} firmada {toplam} ekran tamamlanmis"
                        " (yarida kalan calisma), gece modunda kaldigi yerden devam ediliyor", log)
                else:
                    print(f"\nBu klasorde ({calisma.name}) bugun daha once {len(aday_tamam)} firmada"
                          f" {toplam} ekran tamamlanmis gorunuyor (yarida kalan bir calisma olabilir).")
                    cevap = kullanici_metni_al(
                        ctx, ">>> [D]evam: kaldigi yerden surer, tamamlanmis ekranlar tekrar"
                        " acilmaz  /  [B]astan: hepsi yeniden taranir (varsayilan D): ")
                    if cevap.strip().lower().startswith("b"):
                        yaz("Kullanici secimi: hepsi bastan taranacak", log)
                    else:
                        bugun_tamam = aday_tamam
                        yaz(f"Kullanici secimi: kaldigi yerden devam ({len(aday_tamam)} firmada"
                            f" {toplam} ekran atlanacak)", log)

        yaz(f"Islenecek firma sayisi: {len(firmalar)} | belge tipi: {', '.join(args.belge_tipi)}", log)
        yaz(f"Tarih araligi: {araliklar[0][0]} - {araliklar[-1][1]} ({len(araliklar)} sorgu/firma)\n", log)

        sonuclar = []
        ardisik_hata = 0
        ozet = calisma / "ozet.csv"
        kalan_dosya = calisma / "kalan-firmalar.txt"
        ozet_yaz(ozet, kalan_dosya, [], firmalar, args.belge_tipi[0])  # bastan yazilir ki yarida kalsa da dosya olsun

        def durumu_kaydet(kalanlar):
            ozet_yaz(ozet, kalan_dosya, sonuclar, kalanlar, args.belge_tipi[0])
            try:
                rapor.guncelle(calisma, sonuclar, kalanlar, args.belge_tipi[0])
            except Exception as e:  # rapor yazilamazsa calisma durmasin
                yaz(f"    UYARI: rapor guncellenemedi ({type(e).__name__}: {e})", log)

        def hatayi_yaz(firma, e):
            yaz(f"    HATA: {type(e).__name__}: {e}", log)
            if sayfa_canli(page):
                hata_kaydet(page, calisma / "hatalar", firma)
            sonuclar.append({"firma": firma, "belge_tipi": args.belge_tipi[0], "fatura_sayisi": 0,
                             "durum": f"hata: {type(e).__name__}", "dosyalar": [],
                             "indirilemeyen": 0, "iptal_itiraz": 0, "tevkifat": 0,
                             "donem": "", "not": str(e)[:120]})

        def sayfa_hazirla(page, ctx):
            """Sayfa olduyse once acik sekmeye gecer, o da yoksa tarayiciyi yeniden acar."""
            if sayfa_canli(page):
                return page, ctx
            yeni_sayfa = sayfayi_kurtar(ctx, log)
            if sayfa_canli(yeni_sayfa):
                return yeni_sayfa, ctx
            yeni_ctx, yeni_sayfa = tarayiciyi_yeniden_baslat(pw, profil, ctx, log)
            if yeni_ctx is None or not sayfa_canli(yeni_sayfa):
                return None, yeni_ctx if yeni_ctx is not None else ctx
            return yeni_sayfa, yeni_ctx

        tarayici_gitti = False
        for i, firma in enumerate(firmalar, 1):
            yaz(f"[{i}/{len(firmalar)}] {firma}", log)
            page, ctx = sayfa_hazirla(page, ctx)
            if page is None:
                tarayici_gitti = True
                break

            try:
                atlanacak = firma_atlanan.get(firma, set())
                tamamlanmis = bugun_tamam.get(firma, set())
                # ilk ekran atlanmis olabilir; firma gercekten secildi mi izlenir
                secildi = False
                for tip in args.belge_tipi:
                    if tip in atlanacak:  # firmalar.xlsx'te X isaretli ekran
                        yaz(f"  -- {BELGE_TIPLERI[tip]}: listede atlanmis", log)
                        continue
                    if tip in tamamlanmis:  # yarida kalan calismadan zaten tamamlanmis
                        yaz(f"  -- {BELGE_TIPLERI[tip]}: bugun tamamlanmis, atlaniyor", log)
                        continue
                    if len(args.belge_tipi) > 1:
                        yaz(f"  -- {BELGE_TIPLERI[tip]}", log)
                    if secildi and sayfa_canli(page):
                        sayfayi_toparla(page)  # onceki ekrandan kalan diyaloglar
                    sonuc = firma_isle(page, firma, tip, araliklar, calisma, log,
                                       azami_deneme, firma_secili=secildi)
                    secildi = True
                    sonuclar.append(sonuc)
                    if sonuc["durum"].startswith(("atlandi", "donem disi")):
                        # donemi tutmayan ya da sinir ustu firma: kalan ekranlar taranmaz
                        yaz("    Bu firmanin kalan ekranlari atlandi", log)
                        break
                ardisik_hata = 0
            except Exception as e:
                # sayfa kapandiysa hata firmanin degil tarayicinin; firmayi yakmadan
                # tarayici toparlanip bir kez daha denenir
                if not sayfa_canli(page):
                    # cokme rastgele; tarayiciyi toparlayip firmayi birkac kez dene
                    for tur in range(1, COKME_DENEMESI + 1):
                        page, ctx = sayfa_hazirla(page, ctx)
                        if page is None:
                            tarayici_gitti = True
                            break
                        yaz(f"    {firma} yeniden deneniyor ({tur}/{COKME_DENEMESI})", log)
                        try:
                            secildi = False
                            for tip in args.belge_tipi:
                                if tip in firma_atlanan.get(firma, set()):
                                    continue
                                if tip in bugun_tamam.get(firma, set()):
                                    continue
                                sonuclar.append(firma_isle(page, firma, tip, araliklar, calisma,
                                                           log, azami_deneme,
                                                           firma_secili=secildi))
                                secildi = True
                            ardisik_hata = 0
                            break
                        except Exception as e2:
                            if sayfa_canli(page):  # cokme degil, gercek hata
                                hatayi_yaz(firma, e2)
                                sayfayi_toparla(page)
                                ardisik_hata += 1
                                break
                            if tur == COKME_DENEMESI:
                                hatayi_yaz(firma, e2)
                                ardisik_hata += 1
                    if tarayici_gitti:
                        break
                else:
                    hatayi_yaz(firma, e)
                    sayfayi_toparla(page)
                    ardisik_hata += 1

            # her firmadan sonra guncellenir: gece yarida kalirsa sabah nerede kalindigi gorulur
            durumu_kaydet(firmalar[i:])

            # pes etmeden once oturumu yenilemeyi dene: ust uste hatalarin
            # sebebi genelde firma degil, dusmus Luca oturumu oluyor
            if 2 <= ardisik_hata < hata_siniri and sayfa_canli(page):
                yeni_sayfa = oturumu_yenile(page, ctx, ayarlar, log)
                if yeni_sayfa is not None:
                    page = yeni_sayfa
                    ardisik_hata = 0

            if ardisik_hata >= hata_siniri:
                yaz(f"\nUst uste {hata_siniri} firma basarisiz oldu; Luca oturumu bozulmus olabilir.", log)
                yaz("Islem durduruldu. Tarayicidan Luca'ya tekrar girip yeniden calistirin.", log)
                break

        if tarayici_gitti:
            islenen = len(sonuclar)
            durumu_kaydet(firmalar[islenen:])
            yaz("\nTarayici kapandi ve yeniden acilamadi (Luca oturumu dustu).", log)
            yaz(f"Kalan {len(firmalar) - islenen} firma 'bekliyor' olarak birakildi, hata yazilmadi.", log)
            yaz("Tarayiciyi acip Luca'ya girin ve programi yeniden calistirin.", log)

        basarili = sum(1 for s in sonuclar if s["durum"] == "tamam")
        bos = sum(1 for s in sonuclar if s["durum"] == "fatura yok")
        # her belge tipi ayri bir sonuc satiri; firma sayisi ile karistirilmasin.
        # e-Arsiv Alis ile Interaktif V.D. ayni faturalari gosterdigi icin o
        # ikisinden yalnizca yuksek olani sayilir, digerleri ayri belgelerdir
        firma_basi, ortusen_basi = {}, {}
        for s in sonuclar:
            firma = s["firma"]
            sayi = s["fatura_sayisi"] or 0
            if s.get("belge_tipi") in ORTUSEN_EKRANLAR:
                ortusen_basi[firma] = max(ortusen_basi.get(firma, 0), sayi)
            else:
                firma_basi[firma] = firma_basi.get(firma, 0) + sayi
            firma_basi.setdefault(firma, 0)
        toplam_fatura = sum(firma_basi.values()) + sum(ortusen_basi.values())
        atlanan_ekran = sum(len(firma_atlanan.get(f, ())) for f in firma_basi)
        yaz(f"\nBitti. {len(firma_basi)} firma, {len(sonuclar)} ekran:"
            f" {basarili} ekranda fatura indi, {bos} ekran bos,"
            f" {len(sonuclar) - basarili - bos} ekran sorunlu"
            + (f", {atlanan_ekran} ekran listede atlanmis" if atlanan_ekran else "")
            + f". Toplam {toplam_fatura} fatura listelendi.", log)
        yaz(f"Dosyalar: {calisma}", log)
        yaz(f"Ozet: {ozet}", log)
        yaz(f"Rapor: {calisma / 'rapor.xlsx'} (aksiyon gereken firmalar en ustte)", log)

        # ozet e-postasi: gonderilemezse calisma yine de tamamlanmis sayilir
        donem_metni = next((s["donem"] for s in sonuclar if s.get("donem")), "")
        try:
            eposta.gonder(ayarlar, calisma, sonuclar,
                          lambda mesaj: yaz(mesaj, log), donem_metni)
        except Exception as e:
            yaz(f"UYARI: e-posta adimi hata verdi ({type(e).__name__}: {e})", log)

        if not args.bitince_kapat:
            input(">>> Tarayiciyi kapatmak icin ENTER'a basin: ")
        ctx.close()


if __name__ == "__main__":
    sys.exit(main())
