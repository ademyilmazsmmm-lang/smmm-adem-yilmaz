# -*- coding: utf-8 -*-
"""Luca Bot arayuzunun "Kâr / Zarar" sekmesi.

Repodaki ayri kar-zarar/ programini (kar_zarar.py) arka planda calistirir; ciktisi canli
akar, sonuc cikti/kar-zarar.json'dan okunup firma basina bir satir gosterilir. Fatura
indirme calismasiyla ayni anda calismaz (ikisi de Luca'ya girer). Luca'ya henuz islenmemis
doneme ait indirilmis faturalar varsa "faturalar dahil" kar/zarar ayrica hesaplanir
(bkz. lucabot.kar_zarar_ozet).
"""

import codecs
import json
import os
import queue
import re
import signal
import subprocess
import sys
import threading
import time
from datetime import date, timedelta
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from lucabot import gostergeler, kar_zarar_ozet
from lucabot.ortak import TARIH_BICIMI, indirme_koku, tarih_cozumle

KOK = Path(__file__).resolve().parent
# luca-bot ile yan yana (repo koku) ya da icinde durabilir
ADAYLAR = (KOK.parent / "kar-zarar", KOK / "kar-zarar")
# Luca girisi icin luca-bot ayarlarindan her calismada kar-zarar/ayarlar.json'a aktarilan anahtarlar
ORTAK_ANAHTARLAR = ("uye_no", "kullanici_adi", "parola", "dogrulama_anahtari", "dogrulama_bekleme_dakika",
                    "tarayici", "tarayici_yolu", "tarayici_sandbox", "profil_yerel", "giris_adresi",
                    "donem_degistir")
KAYNAKLAR = (("Mizan (Excel)", "mizan"), ("Hesap Planı Listesi", "hesap-plani"))
ILERLEME = re.compile(r"^\[(\d+)/(\d+)\]\s+(.+)$")
UZUN_BEKLEME = 30


def kar_zarar_klasoru():
    for yol in ADAYLAR:
        if (yol / "kar_zarar.py").exists():
            return yol
    return None


def varsayilan_donem(bugun=None):
    """Ocak basindan, Luca'ya islenmis olmasi beklenen son ay sonuna: bugun 5 Ekim ise 01/01 - 31/08."""
    bugun = bugun or date.today()
    onceki = bugun.replace(day=1) - timedelta(days=1)  # gecen ay sonu (faturalari indirilen ay)
    bit = onceki.replace(day=1) - timedelta(days=1)    # bir onceki ay sonu (Luca'ya islenmis)
    return bit.replace(month=1, day=1).strftime(TARIH_BICIMI), bit.strftime(TARIH_BICIMI)


def ayar_hazirla(klasor, ana_ayarlar, ek):
    """kar-zarar/ayarlar.json'u olusturur/gunceller: luca-bot girisi + sekmedeki alanlar. Yolunu dondurur."""
    yol = klasor / "ayarlar.json"
    ayar = {}
    for kaynak in (yol, klasor / "ayarlar.ornek.json"):
        if kaynak.exists():
            try:
                ayar = json.loads(kaynak.read_text(encoding="utf-8-sig"))
            except ValueError as e:
                raise ValueError(f"{kaynak.name} okunamadı: {e}")
            break
    for anahtar in ORTAK_ANAHTARLAR:
        if ana_ayarlar.get(anahtar) not in (None, ""):
            ayar[anahtar] = ana_ayarlar[anahtar]
    if not ayar.get("vkn_listesi") and ana_ayarlar.get("firma_listesi"):
        ayar["vkn_listesi"] = ana_ayarlar["firma_listesi"]  # Vergi No sutunu varsa VKN'si bulunamayanlar tamamlanir
    for anahtar in ("indirme_klasoru",):
        ayar.pop(anahtar, None)
    ayar.update(ek)
    gecici = yol.with_suffix(".json.tmp")
    gecici.write_text(json.dumps(ayar, ensure_ascii=False, indent=2), encoding="utf-8")
    gecici.replace(yol)
    return yol


class KarZararSekmesi:
    def __init__(self, arayuz, ebeveyn, ui):
        """ui: luca_arayuz modulu (renkler ve kucuk bilesenler); dongusel import olmasin diye verilir."""
        self.a = arayuz
        self.ui = ui
        self.klasor = kar_zarar_klasoru()
        self.surec = None
        self.okuyucu = None
        self.kuyruk = queue.Queue()
        self.yarim_satir = ""
        self.durdurma_istendi = False
        self.son_islem = ("", 0.0)
        self.satirlar = []
        self.ham = []
        self.donem_metni = ""
        self.govde = tk.Frame(ebeveyn, bg=ui.ZEMIN)
        self._degiskenler()
        if self.klasor is None:
            tk.Label(self.govde, text="kar-zarar klasörü bulunamadı.\n\nluca-bot klasörünün yanında (aynı üst "
                     "klasörde) `kar-zarar` klasörü olmalı.", font=ui.GOVDE, fg=ui.TURUNCU, bg=ui.ZEMIN,
                     justify="left").pack(padx=30, pady=30, anchor="w")
            return
        self._sol_panel()
        tk.Frame(self.govde, bg=ui.CIZGI, width=1).pack(side="left", fill="y")
        self._sag_panel()
        self.sonucu_yukle()

    # -- kurulum ---------------------------------------------------------------

    def _kz_ayar(self):
        try:
            return json.loads((self.klasor / "ayarlar.json").read_text(encoding="utf-8-sig"))
        except (OSError, ValueError):
            return {}

    def _degiskenler(self):
        bas, bit = varsayilan_donem()
        kz = self._kz_ayar() if self.klasor else {}
        self.v_bas = tk.StringVar(value=bas)
        self.v_bit = tk.StringVar(value=bit)
        self.v_firma = tk.StringVar()
        ad = {d: a for a, d in KAYNAKLAR}
        self.v_kaynak = tk.StringVar(value=ad.get(kz.get("luca_kaynagi"), KAYNAKLAR[0][0]))
        self.v_db_kod = tk.StringVar(value=kz.get("defterbeyan_kullanici", ""))
        self.v_db_sifre = tk.StringVar(value=kz.get("defterbeyan_sifre", ""))
        self.v_fatura = tk.BooleanVar(value=False)  # "Taranan Faturaları Dahil Et" düğmesiyle açılır

    def _sol_panel(self):
        u = self.ui
        p = tk.Frame(self.govde, bg=u.ZEMIN, width=330, padx=22, pady=16)
        p.pack(side="left", fill="y")
        p.pack_propagate(False)

        u.bolum_basligi(p, "Dönem (Luca'ya işlenmiş)").pack(fill="x")
        t = tk.Frame(p, bg=u.ZEMIN)
        t.pack(fill="x", pady=(6, 0))
        for i, (ad, v) in enumerate((("Başlangıç", self.v_bas), ("Bitiş", self.v_bit))):
            k = tk.Frame(t, bg=u.ZEMIN)
            k.grid(row=0, column=i, sticky="ew", padx=(0, 8) if i == 0 else 0)
            tk.Label(k, text=ad, font=u.KUCUK, fg=u.ETIKET, bg=u.ZEMIN, anchor="w").pack(fill="x")
            u.giris_kutusu(k, v, genislik=12).pack(fill="x", ipady=4)
        t.columnconfigure(0, weight=1)
        t.columnconfigure(1, weight=1)

        tk.Label(p, text="Sadece bu firma (boş = hepsi)", font=u.KUCUK, fg=u.ETIKET, bg=u.ZEMIN,
                 anchor="w").pack(fill="x", pady=(14, 3))
        u.giris_kutusu(p, self.v_firma).pack(fill="x", ipady=4)

        u.bolum_basligi(p, "Luca firmaları (genel muhasebe)").pack(fill="x", pady=(16, 0))
        kutu = ttk.Combobox(p, textvariable=self.v_kaynak, values=[a for a, _ in KAYNAKLAR],
                            state="readonly", font=u.GOVDE)
        kutu.pack(fill="x", pady=(6, 0), ipady=3)

        u.bolum_basligi(p, "Defter Beyan (işletme / SMK)").pack(fill="x", pady=(16, 0))
        tk.Label(p, text="Kullanıcı kodu", font=u.KUCUK, fg=u.ETIKET, bg=u.ZEMIN, anchor="w").pack(fill="x", pady=(6, 3))
        u.giris_kutusu(p, self.v_db_kod).pack(fill="x", ipady=4)
        tk.Label(p, text="Şifre", font=u.KUCUK, fg=u.ETIKET, bg=u.ZEMIN, anchor="w").pack(fill="x", pady=(8, 3))
        u.giris_kutusu(p, self.v_db_sifre, gizli=True).pack(fill="x", ipady=4)
        tk.Label(p, text="Güvenlik kodunu açılan tarayıcıda siz yazarsınız.", font=u.KUCUK, fg=u.SOLUK,
                 bg=u.ZEMIN, anchor="w", justify="left", wraplength=280).pack(fill="x", pady=(6, 0))


        alt = tk.Frame(p, bg=u.ZEMIN)
        alt.pack(side="bottom", fill="x")
        self.calistir_dugmesi = u.dugme(alt, "Çalıştır", self.calistir, ana=True)
        self.calistir_dugmesi.pack(side="left", fill="x", expand=True, padx=(0, 6))
        self.durdur_dugmesi = u.dugme(alt, "Durdur", self.durdur, state="disabled")
        self.durdur_dugmesi.pack(side="left", fill="x", expand=True)

    def _sag_panel(self):
        u = self.ui
        p = tk.Frame(self.govde, bg=u.ZEMIN, padx=22, pady=16)
        p.pack(side="left", fill="both", expand=True)
        self.ilerleme_etiketi = tk.Label(p, text="Hazır. Dönemi kontrol edip Çalıştır'a basın.",
                                         font=u.GOVDE_KALIN, fg=u.YAZI, bg=u.ZEMIN, anchor="w")
        self.ilerleme_etiketi.pack(fill="x")
        self.ilerleme = ttk.Progressbar(p, style="Altin.Horizontal.TProgressbar", maximum=100)
        self.ilerleme.pack(fill="x", pady=(8, 2))
        self.son_islem_etiketi = tk.Label(p, text="", font=u.KUCUK, fg=u.SOLUK, bg=u.ZEMIN, anchor="w")
        self.son_islem_etiketi.pack(fill="x", pady=(0, 6))

        alt = tk.Frame(p, bg=u.ZEMIN)
        alt.pack(side="bottom", fill="x", pady=(10, 0))
        u.dugme(alt, "Excel olarak indir", self.excele_indir).pack(side="left")
        u.dugme(alt, "Programın Excel Raporu", self.raporu_ac).pack(side="left", padx=8)
        self.durum_yazisi = tk.Label(alt, text="", font=u.KUCUK, fg=u.SOLUK, bg=u.ZEMIN)
        self.durum_yazisi.pack(side="right")

        # log kucuk tutulur; asil yer sonuc tablosunun
        cerceve = tk.Frame(p, bg=u.KONSOL, highlightthickness=1, highlightbackground="#1B2944", height=130)
        cerceve.pack(side="bottom", fill="x", pady=(10, 0))
        cerceve.pack_propagate(False)
        self.log = tk.Text(cerceve, bg=u.KONSOL, fg="#CBD2DE", font=u.KONSOL_YAZI, relief="flat", wrap="none",
                           padx=12, pady=8, insertbackground=u.YAZI, state="disabled", highlightthickness=0, bd=0)
        kaydir = tk.Scrollbar(cerceve, command=self.log.yview, bg=u.KONSOL, troughcolor=u.KONSOL,
                              activebackground=u.KENAR, relief="flat", bd=0)
        self.log.configure(yscrollcommand=kaydir.set)
        kaydir.pack(side="right", fill="y")
        self.log.pack(side="left", fill="both", expand=True)
        for etiket, renk in (("ok", "#8CE59A"), ("uyari", "#F2D98A"), ("soluk", "#7F8BA3"),
                             ("hata", u.KIRMIZI), ("firma", u.ALTIN_ACIK), ("bilgi", "#9AA6BD")):
            self.log.tag_configure(etiket, foreground=renk)

        self.donem_etiketi = tk.Label(p, text="", font=u.KUCUK, fg=u.SOLUK, bg=u.ZEMIN, anchor="w")
        self.donem_etiketi.pack(fill="x", pady=(2, 0))
        satir = tk.Frame(p, bg=u.ZEMIN)
        satir.pack(fill="x", pady=(2, 6))
        self.fatura_dugmesi = u.dugme(satir, "Taranan Faturaları Dahil Et", self.fatura_degistir, ana=True)
        self.fatura_dugmesi.pack(side="left")
        self.fatura_etiketi = tk.Label(satir, text="", font=u.KUCUK, fg=u.ALTIN, bg=u.ZEMIN, anchor="w",
                                       justify="left", wraplength=420)
        self.fatura_etiketi.pack(side="left", padx=12, fill="x", expand=True)
        self.ozet_etiketi = tk.Label(p, text="", font=u.GOVDE_KALIN, fg=u.ALTIN, bg=u.ZEMIN, anchor="w",
                                     justify="left", wraplength=640)
        self.ozet_etiketi.pack(side="bottom", fill="x", pady=(6, 0))
        tablo = tk.Frame(p, bg=u.ZEMIN)
        tablo.pack(fill="both", expand=True)
        sutunlar = ("Firma", "Kaynak", "Dönem kârı / zararı", "Fatura farkı", "Faturalar dahil", "Not")
        genislik = (150, 90, 155, 120, 155, 80)
        self.agac = ttk.Treeview(tablo, columns=sutunlar, show="headings", style="Liste.Treeview")
        kay = ttk.Scrollbar(tablo, orient="vertical", command=self.agac.yview)
        self.agac.configure(yscrollcommand=kay.set)
        for i, (s, g) in enumerate(zip(sutunlar, genislik)):
            yon = "e" if i in (2, 3, 4) else "w"
            self.agac.heading(s, text=s, anchor=yon)
            self.agac.column(s, anchor=yon, width=g, minwidth=80, stretch=(i == 0 or i == 5))
        self.agac.tag_configure("kar", foreground="#8CE59A")
        self.agac.tag_configure("zarar", foreground=u.KIRMIZI)
        self.agac.tag_configure("hata", foreground=u.SOLUK)
        kay.pack(side="right", fill="y")
        self.agac.pack(side="left", fill="both", expand=True)
        self.agac.bind("<<TreeviewSelect>>", self._secildi)

    # -- sonuc -------------------------------------------------------------------

    def cikti_klasoru(self):
        return Path(os.environ.get("KARZARAR_CIKTI") or self.klasor / "cikti")

    def _rapor_kayitlari(self):
        try:
            return json.loads((indirme_koku(self.a.ayarlar) / "rapor.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return {}

    def sonucu_yukle(self):
        try:
            self.donem_metni, sonuclar = kar_zarar_ozet.sonuclari_oku(self.cikti_klasoru() / "kar-zarar.json")
        except ValueError:
            self.donem_metni, sonuclar = "", []
        self.ham = sonuclar
        self.sonucu_goster()

    def sonucu_goster(self):
        if self.klasor is None:
            return
        aralik = kar_zarar_ozet.donem_coz(self.donem_metni)
        kayitlar = self._rapor_kayitlari() if self.v_fatura.get() else {}
        bit = aralik[1] if aralik else date.today()
        self.satirlar = kar_zarar_ozet.satirlar(self.ham, kayitlar, bit)
        tl = gostergeler.tl
        self.agac.delete(*self.agac.get_children())
        for i, s in enumerate(self.satirlar):
            kar, dahil = s.get("kar"), s.get("kar_dahil")
            fark = (s["fatura_satis"] - s["fatura_alis"]) if s.get("fatura_satis") is not None else None
            if kar is None:
                etiket, ilk, ikinci, ucuncu = "hata", "—", "—", "—"
                not_ = s.get("hata") or ""
            else:
                etiket = "kar" if kar >= 0 else "zarar"
                ilk = tl(kar) + (" kâr" if kar >= 0 else " zarar")
                ikinci = tl(fark) if fark is not None else "—"
                ucuncu = (tl(dahil) + (" kâr" if dahil >= 0 else " zarar")) if dahil is not None else "—"
                not_ = s.get("fatura_not") or ""
            self.agac.insert("", "end", iid=str(i), tags=(etiket,),
                             values=(s["firma"], s.get("kaynak") or "", ilk, ikinci, ucuncu, not_))
        if aralik:
            self.donem_etiketi.configure(
                text=f"Sonuç dönemi: {aralik[0]:%d/%m/%Y} – {aralik[1]:%d/%m/%Y}   ({len(self.satirlar)} firma)"
                     "   Bu bir tahmindir; asıl inceleme firmada yapılır.")
        else:
            self.donem_etiketi.configure(text="Henüz sonuç yok.")
        self._ust_ozet(kayitlar, aralik)

    def fatura_degistir(self):
        """Dugme: indirilmis faturalari kar/zarara kat (ya da cikar)."""
        self.v_fatura.set(not self.v_fatura.get())
        self.sonucu_goster()

    def _ust_ozet(self, kayitlar, aralik):
        """Fatura donemi bilgisi + toplam: 'dahil edilince ne olur' tek bakista gorunsun."""
        donem, adet = kar_zarar_ozet.fatura_donemi(self._rapor_kayitlari())
        self.fatura_dugmesi.configure(text="Faturaları Hariç Tut" if self.v_fatura.get()
                                      else "Taranan Faturaları Dahil Et")
        if not self.v_fatura.get():
            self.fatura_etiketi.configure(
                text=(f"Faturalar hesaba katılmıyor. İndirilmiş: {donem.replace('-', ' – ')} ({adet} firma)"
                      if donem else "Faturalar hesaba katılmıyor; indirilmiş fatura yok."))
        elif not donem:
            self.fatura_etiketi.configure(text="İndirilmiş fatura yok; önce Fatura İndirme'den faturaları indirin.")
        else:
            self.fatura_etiketi.configure(
                text=f"İndirilmiş faturalar: {donem.replace('-', ' – ')} ({adet} firma) — "
                     "Luca'ya işlenmemiş bu dönemin satış − alışı kâra eklenir.")
        t = kar_zarar_ozet.toplamlar(self.satirlar)
        if not t["firma"]:
            self.ozet_etiketi.configure(text="")
            return
        tl = gostergeler.tl

        def kz(x):
            return tl(x) + (" kâr" if x >= 0 else " zarar")

        metin = f"Toplam ({t['firma']} firma): dönem {kz(t['kar'])}"
        if t["fatura_firma"]:
            metin += f"  →  faturalar dahil edilince {kz(t['dahil'])}  ({t['fatura_firma']} firmanın faturası eklendi)"
        self.ozet_etiketi.configure(text=metin)

    def _secildi(self, _e=None):
        secili = self.agac.selection()
        if secili:
            self.ozet_etiketi.configure(text=f"{self.satirlar[int(secili[0])]['firma']}: "
                                             + kar_zarar_ozet.ozet_cumlesi(self.satirlar[int(secili[0])]))

    def excele_indir(self):
        if not self.satirlar:
            messagebox.showinfo("Sonuç yok", "Önce çalıştırın.", parent=self.a.kok)
            return
        aralik = kar_zarar_ozet.donem_coz(self.donem_metni)
        donem = f"{aralik[0]:%d/%m/%Y}-{aralik[1]:%d/%m/%Y}" if aralik else ""
        yol = filedialog.asksaveasfilename(
            title="Excel olarak kaydet", parent=self.a.kok, defaultextension=".xlsx",
            initialfile=f"kar-zarar-{donem.replace('/', '.')}.xlsx", filetypes=[("Excel", "*.xlsx")])
        if not yol:
            return
        sutunlar = ["Firma", "Kaynak", "Satış", "Toplam Gider", "Kâr / Zarar", "Fatura Dönemi", "Fatura Satış",
                    "Fatura Alış", "Faturalar Dahil Kâr / Zarar", "Not"]
        satirlar = []
        for s in self.satirlar:
            gider = None if s.get("kar") is None else (s.get("toplam_gider") if s.get("toplam_gider") is not None
                                                       else round((s.get("mal_alis") or 0) + (s.get("gider") or 0), 2))
            satirlar.append([s["firma"], s.get("kaynak") or "", s.get("satis"), gider, s.get("kar"),
                             s.get("fatura_donem") or "", s.get("fatura_satis"), s.get("fatura_alis"),
                             s.get("kar_dahil"), s.get("hata") or s.get("fatura_not") or ""])
        notlar = ["Tahmindir: faturalardan mal alışı ile gider ayrılamaz (alış tek kalem sayılır); maaş, amortisman "
                  "gibi yevmiye kalemleri faturada yoktur.",
                  "Faturalar dahil = dönem kârı + indirilen faturaların (KDV hariç) satışı − alışı."]
        try:
            self._excel_yaz(yol, donem, sutunlar, satirlar, notlar)
        except PermissionError:
            messagebox.showerror("Kaydedilemedi", "Dosya Excel'de açık; kapatıp tekrar deneyin.", parent=self.a.kok)
            return
        self.ui.dosya_ac(yol)

    @staticmethod
    def _excel_yaz(yol, donem, sutunlar, satirlar, notlar):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Kâr-Zarar"
        ws.append(["Kâr / Zarar Tahmini"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([f"Dönem: {donem.replace('-', ' – ')}" if donem else ""])
        ws.append([])
        ws.append(sutunlar)
        for h in ws[4]:
            h.font = Font(bold=True)
        for s in satirlar:
            ws.append(s)
        for satir in ws.iter_rows(min_row=5):
            for h in satir:
                if isinstance(h.value, (int, float)):
                    h.number_format = '#,##0.00;[Red]-#,##0.00'
        ws.append([])
        for n in notlar:
            ws.append([n])
        ws.column_dimensions["A"].width = 34
        for harf in "BCDEFGHIJ":
            ws.column_dimensions[harf].width = 18
        ws.column_dimensions["J"].width = 50
        wb.save(yol)

    def raporu_ac(self):
        yol = self.cikti_klasoru() / "kar-zarar.xlsx"
        if not yol.exists():
            messagebox.showinfo("Rapor yok", "Henüz rapor oluşmadı; önce çalıştırın.", parent=self.a.kok)
            return
        self.ui.dosya_ac(yol)

    # -- calistirma ------------------------------------------------------------------

    def calisiyor(self):
        return self.surec is not None

    def komut(self):
        bas, bit = self.v_bas.get().strip(), self.v_bit.get().strip()
        try:
            if tarih_cozumle(bit) < tarih_cozumle(bas):
                raise ValueError("Bitiş tarihi başlangıçtan önce olamaz.")
        except ValueError as e:
            if "olamaz" in str(e):
                raise
            raise ValueError("Tarihleri GG/AA/YYYY biçiminde yazın (örnek 01/01/2026).")
        kaynak = dict(KAYNAKLAR).get(self.v_kaynak.get(), "mizan")
        ayar_hazirla(self.klasor, self.a.ayarlar, {
            "luca_kaynagi": kaynak,
            "defterbeyan_kullanici": self.v_db_kod.get().strip(),
            "defterbeyan_sifre": self.v_db_sifre.get()})
        komut = [self.ui.python_komutu(), "-u", str(self.klasor / "kar_zarar.py"), "--bitince-kapat",
                 "--tarih", f"{bas}-{bit}"]
        if self.v_firma.get().strip():
            komut += ["--firma", self.v_firma.get().strip()]
        return komut

    def calistir(self):
        if self.a.surec:
            messagebox.showinfo("Başka çalışma sürüyor", "Fatura indirme/liste çekme çalışıyor; bitince başlatın "
                                "(ikisi aynı anda Luca'ya girmemeli).", parent=self.a.kok)
            return
        if not self.a._giris_tamam_mi():
            return
        try:
            komut = self.komut()
        except ValueError as e:
            messagebox.showerror("Eksik bilgi", str(e), parent=self.a.kok)
            return
        self._baslat(komut)

    def _baslat(self, komut):
        ortam = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUNBUFFERED="1",
                     LUCA_BOT_AYAR=str(self.klasor / "ayarlar.json"))
        bayrak = 0
        if sys.platform.startswith("win"):
            bayrak = subprocess.CREATE_NEW_PROCESS_GROUP | getattr(subprocess, "CREATE_NO_WINDOW", 0x08000000)
        try:
            (self.klasor / "durdur.istek").unlink()
        except OSError:
            pass
        self._log_temizle()
        self._log_ekle("Başlatılıyor…\n", "bilgi")
        try:
            self.surec = subprocess.Popen(komut, cwd=str(self.klasor), env=ortam, stdin=subprocess.DEVNULL,
                                          stdout=subprocess.PIPE, stderr=subprocess.STDOUT, bufsize=0,
                                          creationflags=bayrak)
        except OSError as e:
            messagebox.showerror("Başlatılamadı", f"Program başlatılamadı: {e}", parent=self.a.kok)
            return
        self.durdurma_istendi = False
        self.yarim_satir = ""
        self.son_islem = ("", 0.0)
        self.okuyucu = threading.Thread(target=self._oku, args=(self.surec,), daemon=True)
        self.okuyucu.start()
        self.ilerleme.configure(value=0)
        self.ilerleme_etiketi.configure(text="Kâr/zarar hesaplanıyor…")
        self.a._durum("Çalışıyor", self.ui.ALTIN, self.ui.ALTIN_YAZI)
        self.calistir_dugmesi.configure(state="disabled")
        self.durdur_dugmesi.configure(state="normal", text="Durdur")
        self.a.calistir_dugmesi.configure(state="disabled")
        self.a.luca_liste_dugmesi.configure(state="disabled")

    def _oku(self, surec):
        cozucu = codecs.getincrementaldecoder("utf-8")(errors="replace")
        while True:
            parca = surec.stdout.read(4096)
            if not parca:
                break
            self.kuyruk.put(cozucu.decode(parca))
        self.kuyruk.put(cozucu.decode(b"", final=True))

    def durdur(self):
        if not self.surec:
            return
        if self.durdurma_istendi:
            if messagebox.askyesno("Zorla kapat", "Program hâlâ kapanmadı. Zorla kapatılsın mı?", parent=self.a.kok):
                self.surec.kill()
            return
        self.durdurma_istendi = True
        try:
            (self.klasor / "durdur.istek").write_text("durdur", encoding="utf-8")
        except OSError:
            pass
        try:
            self.surec.send_signal(signal.CTRL_BREAK_EVENT if sys.platform.startswith("win") else signal.SIGINT)
        except (OSError, ValueError):
            pass
        self._log_ekle("\nDurduruluyor… o ana kadarki sonuçlar kaydediliyor.\n", "uyari")
        self.durdur_dugmesi.configure(text="Zorla Kapat")

    def tikla(self):
        """Arayuz dongusu: kuyruktaki ciktiyi isler, surec bittiyse sonucu toplar."""
        metin = ""
        try:
            while True:
                metin += self.kuyruk.get_nowait()
        except queue.Empty:
            pass
        if metin:
            self._metni_isle(metin)
        if (self.surec and self.surec.poll() is not None and not self.okuyucu.is_alive()
                and self.kuyruk.empty()):
            self._bitti(self.surec.returncode)
        self._son_islemi_goster()

    def _metni_isle(self, metin):
        satirlar = (self.yarim_satir + metin.replace("\r\n", "\n")).split("\n")
        self.yarim_satir = satirlar.pop()
        for satir in satirlar:
            self._satir(satir)

    def _satir(self, satir):
        sade = satir.strip()
        etiket = None
        if sade and set(sade) - set("=-"):
            self.son_islem = (sade, time.monotonic())
        m = ILERLEME.match(sade)
        if m:
            sira, toplam = int(m.group(1)), int(m.group(2))
            self.ilerleme_etiketi.configure(text=f"İşleniyor: {m.group(3)}  ({sira} / {toplam} firma)")
            self.ilerleme.configure(value=(sira - 1) * 100 / max(toplam, 1))
            self.sonucu_yukle()
            etiket = "firma"
        elif "[OK]" in sade:
            etiket = "ok"
        elif "[!!]" in sade or sade.startswith(("DIKKAT", "UYARI")) or "UYARI:" in sade:
            etiket = "uyari"
        elif "[--]" in sade or sade and set(sade) <= set("=-"):
            etiket = "soluk"
        elif "HATA" in sade:
            etiket = "hata"
        self._log_ekle(satir + "\n", etiket)

    def _bitti(self, kod):
        if self.yarim_satir:
            self._satir(self.yarim_satir)
            self.yarim_satir = ""
        self.surec.stdout.close()
        self.surec = None
        self.calistir_dugmesi.configure(state="normal")
        self.durdur_dugmesi.configure(state="disabled", text="Durdur")
        self.a.calistir_dugmesi.configure(state="normal")
        self.a.luca_liste_dugmesi.configure(state="normal")
        self.sonucu_yukle()
        u = self.ui
        if kod == 0 and not self.durdurma_istendi:
            self.ilerleme.configure(value=100)
            self.ilerleme_etiketi.configure(text=f"Tamamlandı: {len(self.satirlar)} firma.")
            self.a._durum("Tamamlandı", u.YESIL, "#FFFFFF")
        elif self.durdurma_istendi or kod in (130, -2):
            self.ilerleme_etiketi.configure(text="Durduruldu; o ana kadarki sonuçlar tabloda.")
            self.a._durum("Durduruldu", u.TURUNCU, u.ALTIN_YAZI)
        else:
            self.ilerleme_etiketi.configure(text="Hata ile bitti — log'un sonuna bakın.")
            self.a._durum("Hata", "#B3443A", "#FFFFFF")

    def _son_islemi_goster(self):
        metin, an = self.son_islem
        if not (self.surec and metin):
            yazi, renk = "", self.ui.SOLUK
        else:
            gecen = int(time.monotonic() - an)
            yazi = f"Şu an: {metin[:110]}  —  {gecen} sn"
            renk = (self.ui.KIRMIZI if gecen >= UZUN_BEKLEME * 3 else
                    self.ui.TURUNCU if gecen >= UZUN_BEKLEME else self.ui.SOLUK)
        if self.son_islem_etiketi.cget("text") != yazi:
            self.son_islem_etiketi.configure(text=yazi, fg=renk)

    def _log_ekle(self, metin, etiket=None):
        alttaydi = self.log.yview()[1] >= 0.999
        self.log.configure(state="normal")
        self.log.insert("end", metin, (etiket,) if etiket else ())
        fazla = int(self.log.index("end-1c").split(".")[0]) - self.ui.AZAMI_SATIR
        if fazla > 0:
            self.log.delete("1.0", f"{fazla + 1}.0")
        self.log.configure(state="disabled")
        if alttaydi:
            self.log.see("end")

    def _log_temizle(self):
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")

    def log_metni(self):
        return self.log.get("1.0", "end-1c")

    def kapat(self):
        """Pencere kapanirken calisan program varsa durdurur."""
        if self.surec:
            self.durdur()
            try:
                self.surec.wait(timeout=20)
            except subprocess.TimeoutExpired:
                self.surec.kill()
