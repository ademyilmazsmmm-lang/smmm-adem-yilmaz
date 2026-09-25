# -*- coding: utf-8 -*-
"""rapor.xlsx: gunun mutabakat ve ozet raporu (OpenPyXL).

Sayfalar:
    Özet                  genel tablo: kac firma/ekran, kac fatura, toplam matrah/KDV,
                          ekran bazinda dagilim
    Firma Durumu          her firma tek satir; aksiyon gerekenler en ustte, renkli
    İndirilen Faturalar   her fatura tek satir; e-Arsiv Alis <-> Interaktif V.D.
                          karsilastirmasi (Mutabakat sutunu)
    Dosyalar              inen her dosya ve klasoru (tiklayinca klasor acilir)
    Hatalar ve Uyarılar   sorunlu/bekleyen ekranlar, sebebi, ne yapilmasi gerektigi
                          ve (varsa) hata aninin ekran goruntusu

Veri rapor.json'dan gelir, yani ayni gun yapilan butun calismalari kapsar.
"""

from datetime import datetime
from pathlib import Path

from . import rapor

BASLIK_RENGI = "1F4E78"
UYARI_RENGI = "FFF2CC"
HATA_RENGI = "F8CBAD"
IYI_RENGI = "E2EFDA"
PARA_BICIMI = '#,##0.00'

# Firma Durumu sayfasinda sayiya cevrilecek (TL) sutunlar
PARA_SUTUNLARI = {"Alış Matrah", "Alış KDV", "Satış Matrah", "Satış KDV"}

EKRAN_ADLARI = dict(rapor.SUTUNLAR)

NE_YAPMALI = [
    ("hata", "Programı tekrar çalıştırın ([D]evam seçeneği yalnızca bu ekranları dener)."
             " Tekrarlarsa ekran görüntüsünü gönderin."),
    ("dosya inmedi", "Tekrar çalıştırın; belge paketi inmedi."),
    ("kaynaktan inmedi", "GİB kaynağı yanıt vermedi; bir süre sonra tekrar sorgulayın."),
    ("tamam (iptal eksik)", "Belgeler indi, iptal/itiraz ve Excel eksik; tekrar çalıştırın."),
    ("bekliyor", "Çalışma bu firmaya gelmeden durdu; tekrar çalıştırın."),
    ("atlandi", "Çok fatura var; e-Fatura portalinden elle indirin."),
    ("donem disi", "Firmanın bu döneme ait çalışma dönemi yok (kapanmış olabilir)."),
]


def _ne_yapmali(durum, not_metni=""):
    if "yetkisi yok" in (not_metni or ""):
        return "Firmanın bu servise yetkisi yok; firmalar.xlsx'te bu ekrana X koyabilirsiniz."
    if "takildi" in (not_metni or ""):
        return "GİB sorgusu yanıt vermedi; daha sonra tekrar deneyin."
    for anahtar, oneri in NE_YAPMALI:
        if (durum or "").startswith(anahtar):
            return oneri
    return ""


def _durum_grubu(durum):
    durum = durum or ""
    if durum == "tamam (iptal eksik)":
        return "sorunlu"
    if durum.startswith("tamam"):
        return "basarili"
    if durum == "fatura yok":
        return "bos"
    if durum.startswith("atlandi") or durum == "donem disi":
        return "atlanan"
    if durum == "bekliyor":
        return "bekliyor"
    return "sorunlu"


def _sayi(metin):
    """'1.234,56' -> 1234.56; bos -> None (hucre bos kalsin)."""
    if metin in ("", None):
        return None
    if isinstance(metin, (int, float)):
        return metin
    s = str(metin).replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return metin


# --- bicimlendirme yardimcilari -------------------------------------------

def _stiller():
    from openpyxl.styles import Alignment, Font, PatternFill
    return {
        "baslik_font": Font(bold=True, color="FFFFFF"),
        "baslik_dolgu": PatternFill("solid", fgColor=BASLIK_RENGI),
        "orta": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "ust": Alignment(vertical="top", wrap_text=True),
        "uyari": PatternFill("solid", fgColor=UYARI_RENGI),
        "hata": PatternFill("solid", fgColor=HATA_RENGI),
        "iyi": PatternFill("solid", fgColor=IYI_RENGI),
        "kalin": Font(bold=True),
        "buyuk": Font(bold=True, size=14, color=BASLIK_RENGI),
        "bag": Font(color="0563C1", underline="single"),
    }


def _tablo(ws, basliklar, satirlar, genislikler, stil, ilk_satir=1, filtre=True):
    """Baslikli tablo yazar; baslik satirini sabitler ve suzgec ekler."""
    from openpyxl.utils import get_column_letter
    for j, b in enumerate(basliklar, 1):
        h = ws.cell(row=ilk_satir, column=j, value=b)
        h.font = stil["baslik_font"]
        h.fill = stil["baslik_dolgu"]
        h.alignment = stil["orta"]
    for i, satir in enumerate(satirlar, ilk_satir + 1):
        for j, deger in enumerate(satir, 1):
            h = ws.cell(row=i, column=j, value=deger)
            h.alignment = stil["ust"]
    for j, g in enumerate(genislikler, 1):
        ws.column_dimensions[get_column_letter(j)].width = g
    if filtre:
        ws.freeze_panes = ws.cell(row=ilk_satir + 1, column=1)
        son = ilk_satir + max(len(satirlar), 1)
        ws.auto_filter.ref = f"A{ilk_satir}:{get_column_letter(len(basliklar))}{son}"


def _klasor_bagi(hucre, yol, stil):
    if not yol:
        return
    try:
        hucre.hyperlink = Path(yol).resolve().as_uri()
        hucre.font = stil["bag"]
    except Exception:
        pass


# --- sayfalar ----------------------------------------------------------------

def _ozet_sayfasi(ws, kayitlar, satirlar, stil):
    ws.title = "Özet"
    donemler = sorted({k.get("donem") for k in kayitlar.values() if k.get("donem")})
    ws["A1"] = "Luca Bot - Mutabakat Özeti"
    ws["A1"].font = stil["buyuk"]
    ws["A2"] = (f"Oluşturulma: {datetime.now():%d/%m/%Y %H:%M}"
                + (f"    Dönem: {', '.join(donemler)}" if donemler else ""))

    ekran_sayim = {}
    toplam = {"basarili": 0, "bos": 0, "atlanan": 0, "sorunlu": 0, "bekliyor": 0}
    for k in kayitlar.values():
        for tip, durum in k["durumlar"].items():
            grup = _durum_grubu(durum)
            toplam[grup] += 1
            e = ekran_sayim.setdefault(tip, {"firma": 0, "basarili": 0, "bos": 0, "atlanan": 0,
                                             "sorunlu": 0, "bekliyor": 0, "fatura": 0, "dosya": 0})
            e["firma"] += 1
            e[grup] += 1
            e["fatura"] += k["sayilar"].get(tip, 0) or 0
            dosya = k.get("dosya") or {}
            e["dosya"] += dosya.get(tip, 0) if isinstance(dosya, dict) else 0

    def grup_toplami(alan, ekranlar):
        return sum(rapor._grup_toplami(k, alan, ekranlar) for k in kayitlar.values())

    def mukerrersiz_fatura(k):
        grup_basi = {}
        for tip, sayi in k["sayilar"].items():
            anahtar = rapor.ortusen_grubu(tip)
            grup_basi[anahtar] = max(grup_basi.get(anahtar, 0), sayi or 0)
        return sum(grup_basi.values())

    aksiyonlu = sum(1 for s in satirlar if s[3])
    gostergeler = [
        ("Firma sayısı", len(kayitlar)),
        ("Aksiyon gereken firma", aksiyonlu),
        ("Fatura inen ekran", toplam["basarili"]),
        ("Boş ekran (fatura yok)", toplam["bos"]),
        ("Atlanan ekran (dönem dışı / çok fatura)", toplam["atlanan"]),
        ("Sorunlu ekran (hata / inmedi)", toplam["sorunlu"]),
        ("Bekleyen ekran (henüz işlenmedi)", toplam["bekliyor"]),
        ("Listelenen fatura (aynı fatura iki ekranda sayılmaz)",
         sum(mukerrersiz_fatura(k) for k in kayitlar.values())),
        ("Tevkifatlı alış faturası (KDV2)",
         sum(rapor._en_yuksek(k["tevkifat"]) for k in kayitlar.values())),
        ("İptal/itiraz edilmiş fatura", sum(rapor._en_yuksek(k["iptal"]) for k in kayitlar.values())),
        ("Alış matrahı (TL)", grup_toplami("matrah", rapor.ALIS_EKRANLARI)),
        ("Alış KDV (TL)", grup_toplami("kdv", rapor.ALIS_EKRANLARI)),
        ("Satış matrahı (TL)", grup_toplami("matrah", rapor.SATIS_EKRANLARI)),
        ("Satış KDV (TL)", grup_toplami("kdv", rapor.SATIS_EKRANLARI)),
        ("İnen dosya", sum(rapor._dosya_sayisi(k) or 0 for k in kayitlar.values())),
    ]
    _tablo(ws, ["Gösterge", "Değer"], gostergeler, [52, 20], stil, ilk_satir=4, filtre=False)
    for i, (ad, _) in enumerate(gostergeler, 5):
        h = ws.cell(row=i, column=2)
        if "(TL)" in ad:
            h.number_format = PARA_BICIMI
        if ad.startswith("Sorunlu") and h.value:
            h.fill = stil["hata"]
        elif ad.startswith("Aksiyon") and h.value:
            h.fill = stil["uyari"]

    bas = 5 + len(gostergeler) + 2
    ws.cell(row=bas - 1, column=1, value="Ekran bazında").font = stil["kalin"]
    ekran_satirlari = []
    for tip, ad in rapor.SUTUNLAR:
        e = ekran_sayim.get(tip)
        if e:
            ekran_satirlari.append([ad, e["firma"], e["basarili"], e["bos"],
                                    e["atlanan"], e["sorunlu"], e["bekliyor"], e["fatura"], e["dosya"]])
    _tablo(ws, ["Ekran", "Firma", "Fatura inen", "Boş", "Atlanan", "Sorunlu", "Bekliyor",
                "Fatura adedi", "İnen dosya"],
           ekran_satirlari, [52, 20, 12, 8, 10, 10, 10, 13, 11], stil, ilk_satir=bas, filtre=False)
    ws.freeze_panes = None


def _firma_durumu_sayfasi(ws, satirlar, stil):
    ws.title = "Firma Durumu"
    para = {i for i, b in enumerate(rapor.BASLIKLAR) if b in PARA_SUTUNLARI}
    duzeltilmis = [[_sayi(d) if j in para else d for j, d in enumerate(s)] for s in satirlar]
    genislik = ([26, 22, 16, 46] + [13] * len(rapor.SUTUNLAR)
                + [8, 46, 12, 12, 11, 14, 14, 14, 14, 11, 24, 16])
    _tablo(ws, rapor.BASLIKLAR, duzeltilmis, genislik, stil)
    for i, satir in enumerate(duzeltilmis, 2):
        if satir[3]:
            dolgu = stil["hata"] if str(satir[3]).startswith("HATA") else stil["uyari"]
            for j in range(1, len(rapor.BASLIKLAR) + 1):
                ws.cell(row=i, column=j).fill = dolgu
        for j in para:
            ws.cell(row=i, column=j + 1).number_format = PARA_BICIMI


def _mutabakat_durumlari(kayit):
    """e-Arsiv Alis ile Interaktif V.D. listelerini fatura numarasiyla eslestirir."""
    eksik, fazla = rapor.fatura_farklari(kayit)
    if eksik is None:
        return {}
    if eksik and fazla and not rapor._ortak_fatura_var(kayit):
        return {"*": "karşılaştırılamadı (numaralar farklı biçimde)"}
    eksik_no = {rapor._no(s[1]) for s in eksik}
    fazla_no = {rapor._no(s[1]) for s in fazla}
    durumlar = {}
    for tip in ("e-arsiv-alis", "e-arsiv-interaktif"):
        for s in kayit["faturalar"].get(tip) or []:
            no = rapor._no(s[1]) if len(s) > 1 else ""
            if no in eksik_no:
                durumlar[(tip, no)] = "e-Arşiv'de YOK (İnteraktif'te var)"
            elif no in fazla_no:
                durumlar[(tip, no)] = "İnteraktif'te YOK (e-Arşiv'de var)"
            elif no:
                durumlar[(tip, no)] = "iki ekranda da var"
    return durumlar


def _fatura_sayfasi(ws, kayitlar, stil):
    ws.title = "İndirilen Faturalar"
    satirlar, renk = [], []
    for firma in sorted(kayitlar):
        k = kayitlar[firma]
        mutabakat = _mutabakat_durumlari(k)
        for tip, ad in rapor.SUTUNLAR:
            yon = ("Alış" if tip in rapor.ALIS_EKRANLARI or tip == "e-arsiv-interaktif"
                   else "Satış" if tip in rapor.SATIS_EKRANLARI else "")
            for f in k["faturalar"].get(tip) or []:
                if len(f) < 2:
                    continue
                no = rapor._no(f[1])
                durum = mutabakat.get((tip, no)) or (mutabakat.get("*", "") if tip in (
                    "e-arsiv-alis", "e-arsiv-interaktif") else "")
                satirlar.append([firma, ad, yon, f[0], f[1], rapor._tutar(f) or None, durum])
                renk.append("YOK" in durum)
    _tablo(ws, ["Firma", "Ekran", "Yön", "Karşı Taraf", "Fatura No", "Tutar (TL)", "Mutabakat"],
           satirlar, [26, 24, 8, 26, 20, 15, 34], stil)
    for i, (satir, uyari) in enumerate(zip(satirlar, renk), 2):
        ws.cell(row=i, column=6).number_format = PARA_BICIMI
        if uyari:
            for j in range(1, 8):
                ws.cell(row=i, column=j).fill = stil["uyari"]
    if not satirlar:
        ws.cell(row=2, column=1, value="Fatura ayrıntısı yok (henüz indirilen liste yok).")


def _dosya_sayfasi(ws, kayitlar, stil):
    ws.title = "Dosyalar"
    satirlar, klasorler = [], []
    for firma in sorted(kayitlar):
        k = kayitlar[firma]
        for tip, ad in rapor.SUTUNLAR:
            for dosya in k["dosya_adlari"].get(tip) or []:
                klasor = k["klasorler"].get(tip, "")
                satirlar.append([firma, ad, k["durumlar"].get(tip, ""), dosya, klasor])
                klasorler.append(klasor)
    _tablo(ws, ["Firma", "Ekran", "Durum", "Dosya", "Klasör (tıklayın)"],
           satirlar, [26, 24, 16, 40, 70], stil)
    for i, klasor in enumerate(klasorler, 2):
        _klasor_bagi(ws.cell(row=i, column=5), klasor, stil)


def _hata_sayfasi(ws, kayitlar, stil):
    ws.title = "Hatalar ve Uyarılar"
    satirlar, gorseller = [], []
    for firma in sorted(kayitlar):
        k = kayitlar[firma]
        for tip, durum in k["durumlar"].items():
            not_metni = k["notlar"].get(tip, "")
            grup = _durum_grubu(durum)
            sorunlu = grup in ("sorunlu", "bekliyor") or durum == "donem disi"
            uyari = bool(not_metni) and not sorunlu
            inmeyen = k["inmeyen"].get(tip, 0)
            if not (sorunlu or uyari or inmeyen):
                continue
            aciklama = not_metni
            if inmeyen:
                aciklama = (aciklama + " | " if aciklama else "") + f"{inmeyen} fatura kaynaktan inmedi"
            satirlar.append([firma, EKRAN_ADLARI.get(tip, tip), durum, "HATA" if grup == "sorunlu"
                             else "Uyarı", aciklama, _ne_yapmali(durum, not_metni),
                             "görüntüyü aç" if k["goruntuler"].get(tip) else ""])
            gorseller.append(k["goruntuler"].get(tip, ""))
    _tablo(ws, ["Firma", "Ekran", "Durum", "Tür", "Açıklama", "Ne yapmalı", "Ekran görüntüsü"],
           satirlar, [26, 24, 20, 8, 50, 60, 16], stil)
    for i, (satir, gorsel) in enumerate(zip(satirlar, gorseller), 2):
        dolgu = stil["hata"] if satir[3] == "HATA" else stil["uyari"]
        for j in range(1, 7):
            ws.cell(row=i, column=j).fill = dolgu
        _klasor_bagi(ws.cell(row=i, column=7), gorsel, stil)
    if not satirlar:
        h = ws.cell(row=2, column=1, value="Hata yok - bütün ekranlar sorunsuz işlendi.")
        h.fill = stil["iyi"]


def yaz(yol, kayitlar, satirlar):
    """rapor.xlsx'i bastan yazar; openpyxl yoksa None doner (rapor.csv yeterli)."""
    try:
        from openpyxl import Workbook
    except ImportError:
        return None
    for k in kayitlar.values():
        rapor._tamamla(k)
    stil = _stiller()
    wb = Workbook()
    _ozet_sayfasi(wb.active, kayitlar, satirlar, stil)
    _firma_durumu_sayfasi(wb.create_sheet(), satirlar, stil)
    _fatura_sayfasi(wb.create_sheet(), kayitlar, stil)
    _dosya_sayfasi(wb.create_sheet(), kayitlar, stil)
    _hata_sayfasi(wb.create_sheet(), kayitlar, stil)

    gecici = Path(yol).with_suffix(".tmp.xlsx")
    wb.save(str(gecici))
    try:
        gecici.replace(yol)
    except PermissionError:
        gecici.unlink(missing_ok=True)
        raise  # dosya Excel'de acik; cagiran kullaniciya soyler
    return yol
