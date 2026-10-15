# -*- coding: utf-8 -*-
"""Interaktif V.D. - e-Arsiv Alis farki icin Excel: suzulebilir duz fatura tablosu + firma ozeti.

    Faturalar     her fatura tek satir (Firma, Durum, Kaynak, Karsi taraf, Fatura no, Tutar); tutar SAYI,
                  boylece Excel'de firmaya ve tutara (orn. >= 1000) gore suzulur / siralanir
    Firma Özeti   firma basina eksik/fazla adet ve tutar, esik altinda gizlenenler
    Bilgi         donem, uygulanan tutar esigi, aciklama
"""

from datetime import datetime

from . import gostergeler as g

SIRA = {g.EKSIK: 0, g.FAZLA: 1, g.ESLESMEDI: 2, g.BELIRSIZ: 3}
PARA = "#,##0.00"


def yaz(yol, donem, fark, tum_satirlar, esik=0):
    """fark: hesapla()["fark"]; tum_satirlar: fark_satirlari(fark) (esiksiz); esik: bu TL'den kucukler yazilmaz."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    baslik_font, baslik_dolgu = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    dolgu = {g.EKSIK: PatternFill("solid", fgColor="FFF2CC"), g.FAZLA: PatternFill("solid", fgColor="FCE4D6"),
             g.ESLESMEDI: PatternFill("solid", fgColor="EDEDED"), g.BELIRSIZ: PatternFill("solid", fgColor="EDEDED")}
    satirlar = g.esik_uygula(tum_satirlar, esik)
    satirlar.sort(key=lambda r: (r["firma"], SIRA.get(r["durum"], 9), -(r["tutar"] or 0)))

    def tablo(ws, basliklar, genislik):
        for j, b in enumerate(basliklar, 1):
            h = ws.cell(row=1, column=j, value=b)
            h.font, h.fill = baslik_font, baslik_dolgu
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for j, w in enumerate(genislik, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A2"

    wb = Workbook()
    ws = wb.active
    ws.title = "Faturalar"
    tablo(ws, ["Firma", "Durum", "Kaynak", "Karşı taraf (ünvan)", "Fatura no", "Tutar (TL)"],
          [28, 30, 16, 36, 22, 16])
    for i, r in enumerate(satirlar, 2):
        for j, d in enumerate((r["firma"], r["durum"], r["kaynak"], r["unvan"], r["no"], r["tutar"]), 1):
            ws.cell(row=i, column=j, value=d)
        ws.cell(row=i, column=6).number_format = PARA
        for j in range(1, 7):
            ws.cell(row=i, column=j).fill = dolgu.get(r["durum"], PatternFill())
    ws.auto_filter.ref = f"A1:F{max(len(satirlar), 1) + 1}"
    if satirlar:
        son = len(satirlar) + 1
        ws.cell(row=son + 1, column=1, value="TOPLAM (süzülen satırlar)").font = Font(bold=True)
        h = ws.cell(row=son + 1, column=6, value=f"=SUBTOTAL(109,F2:F{son})")
        h.font, h.number_format = Font(bold=True), PARA

    # ---- firma ozeti
    oz = wb.create_sheet("Firma Özeti")
    tablo(oz, ["Firma", "İnteraktif adet", "e-Arşiv adet", "Fark", "Eksik adet", "Eksik tutar (TL)",
               "Fazla adet", "Fazla tutar (TL)", "Not"], [28, 14, 14, 10, 12, 18, 12, 18, 60])
    gorunen = {}
    for r in satirlar:
        gorunen.setdefault(r["firma"], []).append(r)
    tumu = {}
    for r in tum_satirlar:
        tumu.setdefault(r["firma"], []).append(r)
    for i, x in enumerate(sorted(fark, key=lambda x: x["firma"]), 2):
        f = x["firma"]
        eksik = [r for r in gorunen.get(f, []) if r["durum"] == g.EKSIK]
        fazla = [r for r in gorunen.get(f, []) if r["durum"] == g.FAZLA]
        gizli = [r for r in tumu.get(f, []) if r not in gorunen.get(f, [])]
        notlar = []
        if x["eslesmedi"]:
            notlar.append("İki listenin fatura numaraları tutmuyor; Faturalar sayfasında iki listenin tamamı var, "
                          "tutar/ünvanla elle karşılaştırın.")
        elif any(r["durum"] == g.BELIRSIZ for r in tumu.get(f, [])):
            notlar.append("Sayı farkı var ama hangi fatura olduğu belirlenemedi.")
        if gizli:
            notlar.append(f"Tutar eşiği ({g.tl(esik)}) altında {len(gizli)} fatura gizlendi "
                          f"({g.tl(sum(r['tutar'] or 0 for r in gizli))}).")
        for j, d in enumerate((f, x["interaktif"], x["earsiv"], x["fark"], len(eksik), sum(r["tutar"] or 0 for r in eksik),
                               len(fazla), sum(r["tutar"] or 0 for r in fazla), " ".join(notlar)), 1):
            oz.cell(row=i, column=j, value=d)
        for j in (6, 8):
            oz.cell(row=i, column=j).number_format = PARA
        oz.cell(row=i, column=9).alignment = Alignment(wrap_text=True, vertical="top")
    oz.auto_filter.ref = f"A1:I{max(len(fark), 1) + 1}"
    if fark:
        son = len(fark) + 1
        oz.cell(row=son + 1, column=1, value="TOPLAM (süzülen satırlar)").font = Font(bold=True)
        for j in (2, 3, 5, 6, 7, 8):
            harf = get_column_letter(j)
            h = oz.cell(row=son + 1, column=j, value=f"=SUBTOTAL(109,{harf}2:{harf}{son})")
            h.font = Font(bold=True)
            if j in (6, 8):
                h.number_format = PARA

    # ---- bilgi
    bi = wb.create_sheet("Bilgi")
    bi.column_dimensions["A"].width = 26
    bi.column_dimensions["B"].width = 110
    bilgi = [("Rapor", "İnteraktif V.D. − e-Arşiv Alış farkı"),
             ("Dönem", donem.replace("-", " – ") if donem else ""),
             ("Tutar eşiği", f"{g.tl(esik)} ve üstü yazıldı" if esik else "Eşik yok (tüm faturalar)"),
             ("Hazırlanma", datetime.now().strftime("%d/%m/%Y %H:%M")),
             ("Eksik", "GİB'de (İnteraktif V.D.) görünüp Luca'da (e-Arşiv Alış) olmayan fatura."),
             ("Fazla", "Luca'da (e-Arşiv Alış) görünüp GİB'de (İnteraktif V.D.) olmayan fatura."),
             ("Eşleşmedi", "İki listenin fatura numaraları hiç tutmuyorsa tek tek eşleştirilemez; iki listenin "
                           "tamamı yazılır."),
             ("Kullanım", "Faturalar sayfasında Firma sütunundaki süzgeçten firmayı, Tutar sütunundaki "
                          "'Sayı Süzgeçleri'nden tutarı (örn. ≥ 1000) seçin; toplam satırı süzülene göre değişir.")]
    for i, (a, b) in enumerate(bilgi, 1):
        bi.cell(row=i, column=1, value=a).font = Font(bold=True)
        bi.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(yol)
