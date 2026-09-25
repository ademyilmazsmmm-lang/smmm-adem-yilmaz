# -*- coding: utf-8 -*-
"""Inen dosyalarin analizi: Excel'den satirlar, tevkifat, iptal/itiraz, matrah/KDV.

Tarayiciya dokunmaz; yalnizca diskteki dosyalarla ve satir listeleriyle
calisir, bu yuzden Luca olmadan da test edilebilir (bkz. testler/).
"""

import csv
import re
import warnings
import zipfile
from datetime import date, datetime

from .ortak import TARIH_BICIMI, sadelestir, yaz
from .sabitler import (FATURA_NO_DESENI, FATURA_NO_YEDEK, IPTAL_DESENI,
                       TARIH_DESENI, TEVKIFAT_DESENI, XML_TEVKIFAT)
from . import rapor

BOS_DEGERLER = {"", "0", "0,00", "0.00", "-", "YOK", "HAYIR"}


def csv_yaz(yol, satirlar):
    """Excel'de Turkce karakterler dogru gorunsun diye BOM'lu UTF-8 yazar."""
    with open(yol, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(satirlar)


# --- satir icinden fatura bilgisi -----------------------------------------

def iptal_itiraz_satirlari(satirlar):
    """Durum sutununda iptal/itiraz gecen satirlar."""
    return [s for s in satirlar if IPTAL_DESENI.search(sadelestir(" ".join(s)))]


def tevkifatli_satirlar(satirlar, belge_tipi=None):
    """Ekranda 'tevkifat' yazan satirlar (KDV2 icin isaret).

    Satis ekranlarinda aranmaz: KDV2 beyani alis faturalarindaki tevkifat icin
    verildiginden satis tarafindaki tevkifat uyari uretmemeli.
    """
    if belge_tipi is not None and belge_tipi not in rapor.TEVKIFAT_EKRANLARI:
        return []
    return [s for s in satirlar if TEVKIFAT_DESENI.search(sadelestir(" ".join(s)))]


def fatura_kimligi(satir):
    """Bir liste satirindan (unvan ilk kelimesi, fatura no) cikarir; bulunamazsa None."""
    no = ""
    for desen in (FATURA_NO_DESENI, FATURA_NO_YEDEK):
        for hucre in satir:
            eslesme = desen.search(hucre.replace(" ", "").replace("-", ""))
            if eslesme:
                no = eslesme.group(0).upper()
                break
        if no:
            break
    if not no:
        return None
    unvan = ""
    for hucre in satir:
        metin = hucre.strip()
        if metin.upper() == no or TARIH_DESENI.search(metin):
            continue
        harf = sum(1 for c in metin if c.isalpha())
        if harf >= 3 and harf > sum(1 for c in unvan if c.isalpha()):
            unvan = metin
    ilk_kelime = (unvan.split() or [""])[0].strip(".,")
    return ilk_kelime, no


def fatura_kimlikleri(satirlar):
    return [k for k in (fatura_kimligi(s) for s in satirlar) if k]


def zipten_tevkifatlilar(zip_yolu):
    """Inen belge paketindeki XML'lerde tevkifat arar; fatura numaralarini dondurur.

    Ekranda tevkifat sutunu olmayabildigi gibi XML de bozuk inebiliyor;
    bu yuzden iki kaynak birlikte kullanilir.
    """
    bulunan = set()
    try:
        with zipfile.ZipFile(zip_yolu) as z:
            for ad in z.namelist():
                if not ad.lower().endswith((".xml", ".html", ".htm")):
                    continue
                try:
                    icerik = z.read(ad).decode("utf-8", "ignore").lower()
                except Exception:
                    continue
                if any(isaret in icerik for isaret in XML_TEVKIFAT):
                    eslesme = FATURA_NO_DESENI.search(ad.replace(" ", ""))
                    bulunan.add(eslesme.group(1).upper() if eslesme else ad)
    except Exception:
        return set()
    return bulunan


# --- Excel okuma ----------------------------------------------------------

def _hucre_metni(h):
    if h is None:
        return ""
    if isinstance(h, (datetime, date)):
        return h.strftime(TARIH_BICIMI)
    return str(h).strip()


def _sayfalari_oku(wb, sadece_ilk, satir_en_az=2):
    """Calisma kitabindan (basliklar, satirlar).

    En az iki hucresi dolu ilk satir baslik sayilir (ustteki tek hucrelik
    "FIRMA LISTESI" gibi basliklar atlanir). Veri satirlari icin en az
    satir_en_az dolu hucre aranir.
    """
    basliklar, satirlar = [], []
    for ws in (wb.worksheets[:1] if sadece_ilk else wb.worksheets):
        try:  # Luca dosyalarinda boyut bilgisi eksik; olmazsa satirlar bos geliyor
            ws.reset_dimensions()
        except Exception:
            pass
        for ham in ws.iter_rows(values_only=True):
            hucreler = [_hucre_metni(h) for h in ham]
            while hucreler and not hucreler[-1]:
                hucreler.pop()
            dolu = len([h for h in hucreler if h])
            if not basliklar:
                if dolu >= 2:
                    basliklar = hucreler
            elif dolu >= satir_en_az:
                satirlar.append(hucreler)
    return basliklar, satirlar


def excelden_tablo(yol, log=None, sadece_ilk=False, satir_en_az=2):
    """Inen Excel'i (basliklar, satirlar) olarak okur; sutunlar yerinde kalir.

    Ekrandaki tabloyu kazimak yerine inen dosyayi kaynak almak daha saglam:
    sutun basliklari belli oldugu icin tevkifat ve iptal/itiraz dogrudan
    kendi sutunlarindan okunabiliyor. Once hizli (read_only) okunur; bos
    donerse dosya tumuyle acilip tekrar denenir. Hicbir durumda hata
    firlatmaz, okunamazsa ([], []) doner.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        yaz("    openpyxl kurulu degil, Excel okunamadi", log)
        return [], []

    def ac(read_only):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return load_workbook(str(yol), read_only=read_only, data_only=True)

    basliklar, satirlar = [], []
    sayfa_adlari = "-"
    for read_only in (True, False):
        try:
            wb = ac(read_only)
        except Exception as e:
            yaz(f"    Excel acilamadi ({type(e).__name__})", log)
            return basliklar, satirlar
        try:
            sayfa_adlari = ", ".join(ws.title for ws in wb.worksheets) or "-"
            basliklar, satirlar = _sayfalari_oku(wb, sadece_ilk, satir_en_az)
        except Exception as e:
            yaz(f"    Excel okunurken hata ({type(e).__name__})", log)
        finally:
            try:
                wb.close()
            except Exception:
                pass
        if satirlar:
            break
    if not satirlar:
        yaz(f"    Excel bos geldi (sayfalar: {sayfa_adlari}, baslik: {len(basliklar)} sutun)", log)
    return basliklar, satirlar


def sutun_indeksi(basliklar, *anahtarlar):
    """Basligi anahtari iceren ilk sutunun sirasi; yoksa None."""
    for i, baslik in enumerate(basliklar):
        duz = sadelestir(baslik)
        if any(a in duz for a in anahtarlar):
            return i
    return None


def sutun_tam_indeksi(basliklar, ad):
    """Basligi birebir eslesen sutun (e-Arsiv Alis ile Satis karismasin diye)."""
    aranan = sadelestir(ad)
    for i, baslik in enumerate(basliklar):
        if sadelestir(baslik) == aranan:
            return i
    return None


def sutunlu_satirlar(basliklar, satirlar, *anahtarlar):
    """Belirtilen sutunu dolu olan satirlar; sutun yoksa None doner."""
    i = sutun_indeksi(basliklar, *anahtarlar)
    if i is None:
        return None
    return [s for s in satirlar
            if i < len(s) and sadelestir(s[i]) not in BOS_DEGERLER]


# --- tutarlar -------------------------------------------------------------

def tutar_cozumle(metin):
    """'1.234,56' / '1234,56' / '1234.56' (openpyxl'in dogrudan verdigi sayi) -> float.

    Bos ya da sayi olmayan metin 0.0 sayilir.
    """
    s = str(metin or "").strip()
    if not s:
        return 0.0
    s = re.sub(r"[^0-9,.\-]", "", s)  # TL isareti, bosluk vb. temizlenir
    if not s or s in ("-", ".", ","):
        return 0.0
    if "," in s:  # Turkce bicim: binlik nokta, ondalik virgul
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _tutar_sutunlari(basliklar, anahtarlar, haric=()):
    """Basligi anahtarlardan birini iceren (ve haric olanlari icermeyen) tum sutun indeksleri.

    Bazi ekranlarda KDV oran basina ayri sutunlarda gosteriliyor
    (orn. "KDV Matrahı (%20)", "KDV Matrahı (%10)"); oran onemli olmadigi
    icin ayni turden birden fazla sutun varsa hepsi toplanir.
    """
    idx = []
    for i, baslik in enumerate(basliklar):
        duz = sadelestir(baslik)
        if any(h in duz for h in haric):
            continue
        if any(a in duz for a in anahtarlar):
            idx.append(i)
    return idx


def matrah_kdv_sutunlari(basliklar):
    """(matrah sutunlari, KDV tutari sutunlari) indeksleri."""
    matrah = _tutar_sutunlari(basliklar, ("MATRAH",))
    if not matrah:  # bazi ekranlarda sutun adi "matrah" gecmiyor
        matrah = _tutar_sutunlari(basliklar, ("MAL HIZMET TOPLAM TUTARI", "MAL HIZMET TUTARI"))
    kdv = _tutar_sutunlari(basliklar, ("KDV",), haric=("ORAN",))
    return matrah, kdv


def _satirlarin_tutari(satirlar, sutunlar):
    return sum(tutar_cozumle(s[i]) for s in satirlar for i in sutunlar if i < len(s))


TOPLAM_ANAHTARLARI = ("GENEL TOPLAM", "VERGILER DAHIL TOPLAM TUTAR", "VERGILER DAHIL TUTAR",
                      "ODENECEK TUTAR", "ODENECEK", "FATURA TUTARI", "TOPLAM TUTAR")


def satir_toplam_tutari(basliklar, satir):
    """Bir faturanin genel toplam tutari (fatura listesi ve eksik/fazla listesinde gosterilir).

    Once "Genel Toplam / Odenecek Tutar" gibi bir sutun aranir ("Mal Hizmet
    Toplam Tutari" matrahtir, sayilmaz). Luca'nin Excel'inde boyle bir sutun
    yoksa tutar matrah + KDV olarak hesaplanir.
    """
    for anahtar in TOPLAM_ANAHTARLARI:
        for i in _tutar_sutunlari(basliklar, (anahtar,), haric=("MAL HIZMET", "MATRAH", "KDV")):
            if i < len(satir) and tutar_cozumle(satir[i]):
                return tutar_cozumle(satir[i])
    matrah, kdv = matrah_kdv_sutunlari(basliklar)
    return round(_satirlarin_tutari([satir], matrah) + _satirlarin_tutari([satir], kdv), 2)


def fatura_kimlikleri_tutarli(basliklar, satirlar):
    """(unvan, no, tutar) uclusu: eksik/fazla fatura listesinde tutar da gorunsun."""
    sonuc = []
    for s in satirlar:
        k = fatura_kimligi(s)
        if not k:
            continue
        unvan, no = k
        sonuc.append((unvan, no, satir_toplam_tutari(basliklar, s)))
    return sonuc


def _iki_kaynaktan(sutundan, metinden, sutun_adi):
    """Sutun ve satir metni bulgularini birlestirir.

    Luca durumu bazen kendi sutununda degil "Onay Durumu" gibi bir sutunda
    yaziyor; yalnizca sutuna bakilinca bos sutun bulgusu satir metnindeki
    kaydi orten bir sonuc veriyordu.
    """
    metinden = metinden or []
    if sutundan is None:
        return metinden, "satir metni"
    birlesik = list(sutundan)
    bilinen = {id(s) for s in sutundan}
    birlesik += [s for s in metinden if id(s) not in bilinen]
    if len(birlesik) == len(sutundan):
        return birlesik, sutun_adi
    return birlesik, f"{sutun_adi} + satir metni" if sutundan else "satir metni"


def excelden_sonuca_isle(sonuc, yol, klasor, log):
    """Inen Excel'i asil kaynak alir: satirlar, tevkifat, iptal/itiraz, matrah/KDV.

    Ekran kazimaya gore guvenilir, cunku sutun basliklari belli. Okunan
    satirlari dondurur (okunamazsa bos liste; sonuc degistirilmez).
    """
    basliklar, satirlar = excelden_tablo(yol, log)
    if not satirlar:
        return []
    yaz(f"    Excel'den {len(satirlar)} satir okundu", log)
    sonuc["fatura_sayisi"] = len(satirlar)
    sonuc["faturalar"] = [list(k) for k in fatura_kimlikleri_tutarli(basliklar, satirlar)]
    csv_yaz(klasor / "liste.csv", [basliklar] + satirlar)

    if sonuc.get("belge_tipi") in rapor.TEVKIFAT_EKRANLARI:
        tevkifatlilar, kaynak = _iki_kaynaktan(
            sutunlu_satirlar(basliklar, satirlar, "TEVKIFAT"),
            tevkifatli_satirlar(satirlar), "Tevkifat sutunu")
    else:  # satis ekrani: tevkifat KDV2 beyanina girmiyor
        tevkifatlilar, kaynak = [], ""
    if len(tevkifatlilar) > sonuc.get("tevkifat", 0):
        sonuc["tevkifat"] = len(tevkifatlilar)
    if tevkifatlilar:
        csv_yaz(klasor / "tevkifatli.csv", [basliklar] + tevkifatlilar)
        yaz(f"    DIKKAT: {len(tevkifatlilar)} tevkifatli alis faturasi ({kaynak}, KDV2)", log)

    iptaller, _ = _iki_kaynaktan(
        sutunlu_satirlar(basliklar, satirlar, "IPTAL", "ITIRAZ"),
        iptal_itiraz_satirlari(satirlar), "Iptal sutunu")
    if len(iptaller) > sonuc.get("iptal_itiraz", 0):
        sonuc["iptal_itiraz"] = len(iptaller)
    if iptaller:
        csv_yaz(klasor / "iptal-itiraz.csv", [basliklar] + iptaller)
        yaz(f"    DIKKAT: {len(iptaller)} faturada iptal/itiraz var", log)

    # KDV2 kontrolu icin toplam matrah/KDV: iptal/itiraz olan faturalar
    # zaten hicbir zaman gerceklesmedigi icin toplama dahil edilmez.
    iptal_id = {id(s) for s in iptaller}
    sayilan_satirlar = [s for s in satirlar if id(s) not in iptal_id]
    matrah_sutunlari, kdv_sutunlari = matrah_kdv_sutunlari(basliklar)
    if sayilan_satirlar and not (matrah_sutunlari or kdv_sutunlari):
        yaz("    UYARI: Matrah/KDV sutunu bulunamadi, tutar toplanamadi"
            f" (basliklar: {', '.join(b for b in basliklar if b)})", log)
    else:
        sonuc["matrah"] = _satirlarin_tutari(sayilan_satirlar, matrah_sutunlari)
        sonuc["kdv"] = _satirlarin_tutari(sayilan_satirlar, kdv_sutunlari)
    return satirlar
