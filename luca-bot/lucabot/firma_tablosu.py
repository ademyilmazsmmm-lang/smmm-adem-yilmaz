# -*- coding: utf-8 -*-
"""firmalar.xlsx sablonu ve arayuzdeki "Firma / Ekran Secimi" tablosu.

Ekran sutunlarinda: ✓ (ya da bos) = o ekran sorgulanir, X = sorgulanmaz
(bkz. sabitler.ATLA_DEGERLERI; botun okudugu taraf firma_listesi.py).
Tabloyu yazarken dosyanin yalnizca firma ve ekran / Devreden KDV hucreleri
degisir; diger sutunlar, sayfalar ve bicimler korunur, once yedek alinir.
"""

import shutil
from datetime import date, datetime
from pathlib import Path

from .musteri_listesi import tarih_metni
from .ortak import KOK, karsilastir, sadelestir
from .sabitler import ATLA_DEGERLERI, EKRAN_SUTUNLARI

SORGULA = "✓"
ATLA = "X"
AD_SUTUNU = "Kısa Adı"
KAPANIS_SUTUNU = "Kapanış Tarihi"
ACILIS_SUTUNU = "Açılış Tarihi"
DEVREDEN_SUTUNU = "Devreden KDV"
SABLON_SUTUNLARI = [AD_SUTUNU, KAPANIS_SUTUNU, DEVREDEN_SUTUNU] + list(EKRAN_SUTUNLARI)

ACIKLAMA = [
    ("Kısa Adı", "Firmanın Luca'daki firma listesinde görünen adı (Luca adı kısaltıyorsa baş kısmı yeter)."),
    ("Kapanış Tarihi", "Firma kapandıysa GG/AA/YYYY. Dönem başlamadan kapanmış firma atlanır. Boşsa açık."),
    ("Devreden KDV", "Önceki dönemden devreden KDV (TL). 'KDV ödemesi çıkabilir' uyarısında düşülür. Boşsa 0."),
    ("Ekran sütunları", f"{SORGULA} (ya da boş) = o ekran bu firmada sorgulanır, {ATLA} = sorgulanmaz."),
    ("Listede olmayan firma", "Hiç işlenmez. Yalnızca ilk sayfa okunur; bu açıklama sayfası okunmaz."),
]


def _yol(yol):
    yol = Path(yol)
    return yol if yol.is_absolute() else KOK / yol


def sablon_olustur(yol, firmalar=(), ornek=True):
    """Bos (ya da verilen firmalarla dolu) firmalar.xlsx sablonu yazar; bossa ornek bir satir."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Firmalar"
    ws.append(SABLON_SUTUNLARI)
    for ad in list(firmalar) or (["ÖRNEK FİRMA A.Ş."] if ornek else []):
        ws.append([ad, "", ""] + [SORGULA] * len(EKRAN_SUTUNLARI))
    baslik = PatternFill("solid", fgColor="1F4E78")
    for h in ws[1]:
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = baslik
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    for i in range(4, len(SABLON_SUTUNLARI) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 12
    ws.freeze_panes = "B2"
    secim = DataValidation(type="list", formula1=f'"{SORGULA},{ATLA}"', allow_blank=True)
    secim.error = f"{SORGULA} (sorgula) ya da {ATLA} (sorgulama) yazın"
    ws.add_data_validation(secim)
    son_harf = ws.cell(row=1, column=len(SABLON_SUTUNLARI)).column_letter
    secim.add(f"D2:{son_harf}2000")
    for satir in ws.iter_rows(min_row=2, min_col=4):
        for h in satir:
            h.alignment = Alignment(horizontal="center")

    a = wb.create_sheet("Açıklama")
    a.append(["Sütun", "Ne yazılır"])
    for s in ACIKLAMA:
        a.append(list(s))
    a.column_dimensions["A"].width = 24
    a.column_dimensions["B"].width = 100
    for h in a[1]:
        h.font = Font(bold=True)
    yol = Path(yol)
    wb.save(yol)
    return yol


def _baslik_satiri(ws):
    """Ilk en az iki dolu hucreli satir (firma_listesi ile ayni kural): (satir no, basliklar)."""
    for satir in ws.iter_rows():
        degerler = [h.value for h in satir]
        if len([d for d in degerler if d not in (None, "")]) >= 2:
            return satir[0].row, ["" if d is None else str(d).strip() for d in degerler]
    return None, []


def _sutun(basliklar, ad):
    aranan = sadelestir(ad)
    for i, b in enumerate(basliklar):
        if sadelestir(b) == aranan:
            return i
    return None


def _ad_sutunu(basliklar):
    for i, b in enumerate(basliklar):
        if "KISA AD" in sadelestir(b):
            return i
    return 0


def tabloyu_oku(yol):
    """[{"ad", "ekranlar": {sorgulanacak tipler}, "devreden": metin}] (dosyadaki sirayla)."""
    from openpyxl import load_workbook
    wb = load_workbook(_yol(yol), data_only=True)
    try:
        ws = wb.worksheets[0]
        satir_no, basliklar = _baslik_satiri(ws)
        if satir_no is None:
            return []
        ad_i = _ad_sutunu(basliklar)
        ekran_i = {tip: _sutun(basliklar, baslik) for baslik, tip in EKRAN_SUTUNLARI.items()}
        dev_i = _sutun(basliklar, DEVREDEN_SUTUNU)
        firmalar = []
        for satir in ws.iter_rows(min_row=satir_no + 1, values_only=True):
            ad = str(satir[ad_i]).strip() if ad_i < len(satir) and satir[ad_i] is not None else ""
            if not ad:
                continue
            ekranlar = set()
            for tip, i in ekran_i.items():
                deger = satir[i] if i is not None and i < len(satir) else None
                if sadelestir("" if deger is None else str(deger)) not in ATLA_DEGERLERI:
                    ekranlar.add(tip)
            dev = satir[dev_i] if dev_i is not None and dev_i < len(satir) else None
            firmalar.append({"ad": ad, "ekranlar": ekranlar,
                             "devreden": "" if dev is None else str(dev).strip()})
        return firmalar
    finally:
        wb.close()


def tabloyu_yaz(yol, firmalar, silinenler=()):
    """Arayuzde yapilan secimleri dosyaya yazar; once yedek alir. Yedegin yolunu dondurur.

    Dosyada olmayan firma sona eklenir; yalnizca `silinenler`deki adlarin
    satiri silinir. Dosya Excel'de acikken PermissionError verir.
    """
    from openpyxl import load_workbook
    yol = _yol(yol)
    wb = load_workbook(yol)
    ws = wb.worksheets[0]
    satir_no, basliklar = _baslik_satiri(ws)
    if satir_no is None:
        raise ValueError("Dosyada başlık satırı bulunamadı")
    ad_i = _ad_sutunu(basliklar)

    def sutun(ad):
        i = _sutun(basliklar, ad)
        if i is None:  # eski listede olmayan sutun sona eklenir
            basliklar.append(ad)
            i = len(basliklar) - 1
            ws.cell(row=satir_no, column=i + 1, value=ad)
        return i

    ekran_i = {tip: sutun(baslik) for baslik, tip in EKRAN_SUTUNLARI.items()}
    dev_i = sutun(DEVREDEN_SUTUNU)
    secimler = {f["ad"]: f for f in firmalar}
    silinecek = set(silinenler)

    def doldur(r, f):
        for tip, i in ekran_i.items():
            ws.cell(row=r, column=i + 1, value=SORGULA if tip in f["ekranlar"] else ATLA)
        dev = (f.get("devreden") or "").strip()
        ws.cell(row=r, column=dev_i + 1, value=_sayi_ya_da_metin(dev) if dev else None)

    bulunan, silinecek_satirlar = set(), []
    son_satir = ws.max_row  # yeni firmalar en alta: araya ya da notlarin ustune yazilmaz
    for satir in ws.iter_rows(min_row=satir_no + 1):
        hucre = satir[ad_i] if ad_i < len(satir) else None
        ad = str(hucre.value).strip() if hucre is not None and hucre.value is not None else ""
        if not ad:
            continue
        if ad in silinecek:
            silinecek_satirlar.append(hucre.row)
        elif ad in secimler:
            bulunan.add(ad)
            doldur(hucre.row, secimler[ad])
    for f in firmalar:  # arayuzde eklenen firmalar
        if f["ad"] not in bulunan and f["ad"] not in silinecek:
            son_satir += 1
            ws.cell(row=son_satir, column=ad_i + 1, value=f["ad"])
            doldur(son_satir, f)
    for r in sorted(silinecek_satirlar, reverse=True):
        ws.delete_rows(r)

    yedek = yol.with_name(f"{yol.stem}.yedek-{datetime.now():%Y%m%d-%H%M%S}{yol.suffix}")
    shutil.copy2(yol, yedek)
    wb.save(yol)
    return yedek


def _sayi_ya_da_metin(metin):
    from .fatura_analiz import tutar_cozumle
    t = metin.replace(" ", "").replace("TL", "")
    if not t:
        return None
    if "," not in t and t.count(".") >= 1 and all(len(p) == 3 for p in t.split(".")[1:]):
        t = t.replace(".", "")  # "5.000" binlik nokta
    deger = tutar_cozumle(t)
    return deger if deger or t.strip("0,.") == "" else metin


# --- Luca musteri listesiyle birlestirme ---------------------------------------------

def luca_kaydi_bul(ad, kayitlar):
    """Tablodaki kisa adin Luca musteri listesindeki kaydi; bulunamaz ya da belirsizse None.

    Once kisa ad / unvanla birebir, sonra (Luca adlari kisaltabildigi icin) en az
    8 harflik baslangic eslesmesi; ikinci kisi de uyuyorsa belirsiz sayilir.
    """
    k = karsilastir(ad)
    if not k:
        return None
    for alan in ("ad", "unvan"):
        tam = [r for r in kayitlar if karsilastir(r.get(alan, "")) == k]
        if len(tam) == 1:
            return tam[0]
        if len(tam) > 1:
            return None
    adaylar = []
    for r in kayitlar:
        for alan in ("ad", "unvan"):
            d = karsilastir(r.get(alan, ""))
            if d and min(len(d), len(k)) >= 8 and (d.startswith(k) or k.startswith(d)):
                adaylar.append(r)
                break
    return adaylar[0] if len(adaylar) == 1 else None


def _kapanis(kayit, yil):
    """Firmayi gercekten kapatan tarih; donem sonu (31/12/yil) ya da sonrasi 'acik' sayilir."""
    metin = kayit.get("kapanis", "")
    if not metin:
        return ""
    gun, ay, y = (int(x) for x in metin.split("/"))
    return "" if date(y, ay, gun) >= date(yil, 12, 31) else metin


def luca_plani(yol, yil, kayitlar):
    """Tabloyla Luca musteri listesinin farki; dosya degismez.

    Dondurur: {"guncellenecek": [{ad, yeni_ad, kayit, kapanis: (eski, yeni), acilis: (eski, yeni)}],
               "yeni": [kayit], "luca_da_yok": [tablodaki ad], "ayni": sayi}
    "yeni_ad": Luca'daki kisa ad (tablodaki ad onun baslangiciysa ona cevrilir).
    Luca'da kapanisi bos olan firmanin tablodaki kapanisi silinmez.
    """
    from openpyxl import load_workbook
    wb = load_workbook(_yol(yol), data_only=True)
    try:
        ws = wb.worksheets[0]
        satir_no, basliklar = _baslik_satiri(ws)
        if satir_no is None:
            raise ValueError("Dosyada başlık satırı bulunamadı")
        ad_i = _ad_sutunu(basliklar)
        kap_i, acil_i = _sutun(basliklar, KAPANIS_SUTUNU), _sutun(basliklar, ACILIS_SUTUNU)
        satirlar = []
        for satir in ws.iter_rows(min_row=satir_no + 1, values_only=True):
            ad = str(satir[ad_i]).strip() if ad_i < len(satir) and satir[ad_i] is not None else ""
            if ad:
                def deger(i):
                    return tarih_metni(satir[i]) if i is not None and i < len(satir) else ""
                satirlar.append((ad, deger(kap_i), deger(acil_i)))
    finally:
        wb.close()

    plan = {"guncellenecek": [], "yeni": [], "luca_da_yok": [], "ayni": 0}
    eslesen = set()
    for ad, kap, acil in satirlar:
        r = luca_kaydi_bul(ad, kayitlar)
        if r is None:
            plan["luca_da_yok"].append(ad)
            continue
        ilk_kez = id(r) not in eslesen  # ayni firmaya iki satir uyarsa ikincisine dokunulmaz
        eslesen.add(id(r))
        yeni_kap = _kapanis(r, yil) or kap
        yeni_acil = r.get("acilis", "") or acil
        yeni_ad = r["ad"] if ilk_kez and r.get("ad") else ad
        if (yeni_kap, yeni_acil, yeni_ad) == (kap, acil, ad):
            plan["ayni"] += 1
        else:
            plan["guncellenecek"].append({"ad": ad, "yeni_ad": yeni_ad, "kayit": r,
                                          "kapanis": (kap, yeni_kap), "acilis": (acil, yeni_acil)})
    plan["yeni"] = [dict(r, kapanis=_kapanis(r, yil)) for r in kayitlar
                    if id(r) not in eslesen and r.get("ad")]
    return plan


def luca_plani_uygula(yol, plan, yeni_ekle=True, eksikleri_sil=False):
    """Plani dosyaya yazar: kapanis/acilis guncellenir, yeni firmalar tum ekranlar isaretli eklenir.

    Listede kalan firmalarin ekran secimleri ve Devreden KDV'sine dokunulmaz.
    eksikleri_sil: Luca listesinde olmayan firmalarin satiri silinir (liste
    Luca'yi yansitsin). Once yedek alinir; yedegin yolunu dondurur.
    """
    from openpyxl import load_workbook
    yol = _yol(yol)
    wb = load_workbook(yol)
    ws = wb.worksheets[0]
    satir_no, basliklar = _baslik_satiri(ws)
    if satir_no is None:
        raise ValueError("Dosyada başlık satırı bulunamadı")
    ad_i = _ad_sutunu(basliklar)

    def sutun(ad):
        i = _sutun(basliklar, ad)
        if i is None:
            basliklar.append(ad)
            i = len(basliklar) - 1
            ws.cell(row=satir_no, column=i + 1, value=ad)
        return i

    kap_i, acil_i = sutun(KAPANIS_SUTUNU), sutun(ACILIS_SUTUNU)
    ekran_i = {tip: sutun(baslik) for baslik, tip in EKRAN_SUTUNLARI.items()}
    degisecek = {g["ad"]: g for g in plan["guncellenecek"]}
    silinecek = set(plan["luca_da_yok"]) if eksikleri_sil else set()
    silinecek_satirlar = []
    son_satir = ws.max_row
    for satir in ws.iter_rows(min_row=satir_no + 1):
        hucre = satir[ad_i] if ad_i < len(satir) else None
        ad = str(hucre.value).strip() if hucre is not None and hucre.value is not None else ""
        g = degisecek.get(ad)
        if g:
            hucre.value = g["yeni_ad"]
            ws.cell(row=hucre.row, column=kap_i + 1, value=g["kapanis"][1] or None)
            ws.cell(row=hucre.row, column=acil_i + 1, value=g["acilis"][1] or None)
        elif ad in silinecek:
            silinecek_satirlar.append(hucre.row)
    if yeni_ekle:
        for r in plan["yeni"]:
            son_satir += 1
            ws.cell(row=son_satir, column=ad_i + 1, value=r["ad"])
            ws.cell(row=son_satir, column=kap_i + 1, value=r.get("kapanis") or None)
            ws.cell(row=son_satir, column=acil_i + 1, value=r.get("acilis") or None)
            for i in ekran_i.values():
                ws.cell(row=son_satir, column=i + 1, value=SORGULA)
    for r in sorted(silinecek_satirlar, reverse=True):
        ws.delete_rows(r)

    yedek = yol.with_name(f"{yol.stem}.yedek-{datetime.now():%Y%m%d-%H%M%S}{yol.suffix}")
    shutil.copy2(yol, yedek)
    wb.save(yol)
    return yedek
