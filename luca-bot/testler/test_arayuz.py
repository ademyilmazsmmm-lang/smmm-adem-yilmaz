# -*- coding: utf-8 -*-
"""Arayuz (luca_arayuz.py) testleri: gercek bot yerine ciktisini taklit eden bir betik calisir.

Ekran gerekir (Linux'ta: xvfb-run -a python -m unittest testler.test_arayuz);
tkinter ya da ekran yoksa testler atlanir.
"""

import json
import os
import sys
import tempfile
import textwrap
import time
import unittest
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOK))

try:
    import tkinter as tk
    _kok = tk.Tk()
    _kok.destroy()
    TK_VAR = True
except Exception:  # tkinter yok ya da ekran yok
    TK_VAR = False

SAHTE_BOT = textwrap.dedent('''
    import sys, time
    print("ARGS " + " ".join(sys.argv[1:]), flush=True)
    print("Luca'da 2 firma bulundu", flush=True)
    try:
        for i, ad in enumerate(["BIRLIK TIC", "YILDIZ OTO"], 1):
            print(f"[{i}/2] {ad}  | tahmini kalan: {3 - i} dk", flush=True)
            print("  [OK] e-Arşiv Alış Faturaları: tamam (3 fatura)", flush=True)
            print("    UYARI: ornek uyari", flush=True)
            time.sleep(float(sys.argv[-1]) if sys.argv[-1].replace(".", "").isdigit() else 0.2)
    except KeyboardInterrupt:
        print("Durduruldu. O ana kadarki sonuclar kaydedildi", flush=True)
        sys.exit(130)
    print("son satir (yeni satir yok)", end="", flush=True)
''')


# --firma-listesi-cek ile cagrilinca Luca'dan musteri listesi cekmis gibi JSON yazar
SAHTE_LISTE_BOTU = textwrap.dedent('''
    import json, os, sys
    from pathlib import Path
    print("ARGS " + " ".join(sys.argv[1:]), flush=True)
    kok = Path(json.loads(Path(os.environ["LUCA_BOT_AYAR"]).read_text(encoding="utf-8"))["indirme_klasoru"])
    firmalar = [
        {"ad": "BIRLIK TICARET", "unvan": "BIRLIK TİCARET", "vkn": "1", "acilis": "01/01/2020", "kapanis": "31/03/2026"},
        {"ad": "YENI FIRMA LTD", "unvan": "YENİ FİRMA", "vkn": "2", "acilis": "01/06/2026", "kapanis": ""},
    ]
    (kok / "luca-musteri-listesi.json").write_text(
        json.dumps({"yil": 2026, "alinma": "x", "firmalar": firmalar}), encoding="utf-8")
    print("[OK] 2026 yilinda 2 firma okundu", flush=True)
''')


# --beyanname-cek ile cagrilinca PDF'leri indirmis gibi dosya + JSON yazar
SAHTE_BEYANNAME_BOTU = textwrap.dedent('''
    import json, os, sys
    from pathlib import Path
    print("ARGS " + " ".join(sys.argv[1:]), flush=True)
    kok = Path(json.loads(Path(os.environ["LUCA_BOT_AYAR"]).read_text(encoding="utf-8"))["indirme_klasoru"])
    (kok / "beyannameler").mkdir()
    yollar = []
    for ad in ("CAGRI DENE_1_KDV1_2026-08_1.pdf", "BIRLIK TICARET_2_KDV1_2026-08_1.pdf"):
        (kok / "beyannameler" / ad).write_bytes(b"%PDF-1.4")
        yollar.append(str(kok / "beyannameler" / ad))
    (kok / "luca-beyannameler.json").write_text(json.dumps({"dosyalar": yollar + [str(kok / "yok.pdf")]}))
    print("[OK] 2 beyanname PDF'i alindi", flush=True)
''')

# kar-zarar/kar_zarar.py yerine: ilerleme satirlari basar, cikti/kar-zarar.json yazar
SAHTE_KAR_ZARAR = textwrap.dedent('''
    import json, os, sys
    from pathlib import Path
    print("ARGS " + " ".join(sys.argv[1:]), flush=True)
    ayar = json.loads(Path(os.environ["LUCA_BOT_AYAR"]).read_text(encoding="utf-8"))
    print("AYAR " + json.dumps({k: ayar.get(k) for k in ("parola", "defterbeyan_kullanici", "luca_kaynagi")}), flush=True)
    firmalar = [
        {"firma": "BIRLIK TIC", "kaynak": "Luca (Mizan)", "kar": 100000.0, "satis": 200000.0, "mal_alis": 40000.0, "gider": 60000.0},
        {"firma": "YILDIZ OTO", "kaynak": "Defter Beyan", "kar": -3000.0, "satis": 1.0, "mal_alis": 1.0, "gider": 1.0},
        {"firma": "HATALI LTD", "kaynak": "Luca (Mizan)", "kar": None, "hata": "HATA: mizan inmedi"},
    ]
    cikti = Path(__file__).resolve().parent / "cikti"
    cikti.mkdir(exist_ok=True)
    for i in range(1, 4):
        print(f"[{i}/3] {firmalar[i - 1]['firma']}  (Luca mizan)", flush=True)
        (cikti / "kar-zarar.json").write_text(json.dumps({
            "donem": sys.argv[sys.argv.index("--tarih") + 1], "alinma": "x", "firmalar": firmalar[:i]}), encoding="utf-8")
    print("[OK] 2 firmanin kar/zarar tahmini hazir", flush=True)
''')


class ArayuzZemini(unittest.TestCase):
    """Ortak kurulum (sahte ayarlar, rapor.json, sahte bot); kendi testi yok."""

    def setUp(self):
        from openpyxl import Workbook
        import luca_arayuz
        self.mod = luca_arayuz
        self.d = Path(tempfile.mkdtemp(prefix="arayuz-test-"))
        wb = Workbook()
        wb.active.append(["Kısa Adı", "Devreden KDV"])
        wb.active.append(["BIRLIK TICARET", 200])
        wb.save(self.d / "firmalar.xlsx")
        (self.d / "indir").mkdir()
        donem = "01/09/2026-30/09/2026"
        (self.d / "indir" / "rapor.json").write_text(json.dumps({
            "BIRLIK TIC": {"donem": donem, "sayilar": {"esmm-alis": 1}, "matrah": {"esmm-alis": 1500},
                           "tevkifat": {"e-fatura-alis": 2}, "tevkifat_kdv": {"e-fatura-alis": 700},
                           "kdv": {"e-fatura-satis": 1000, "e-fatura-alis": 300}},
        }), encoding="utf-8")
        self.ayar = self.d / "ayarlar.json"
        self.ayar.write_text(json.dumps({
            "uye_no": "1", "kullanici_adi": "deneme", "parola": "x", "baska_ayar": 42,
            "indirme_klasoru": str(self.d / "indir"), "firma_listesi": str(self.d / "firmalar.xlsx"),
            "baslangic_tarihi": "01/09/2026", "bitis_tarihi": "30/09/2026"}), encoding="utf-8")
        os.environ["LUCA_BOT_AYAR"] = str(self.ayar)
        self.bot = self.d / "sahte_bot.py"
        self.bot.write_text(SAHTE_BOT, encoding="utf-8")
        self.kok = tk.Tk()
        self.app = luca_arayuz.Arayuz(self.kok)
        self.app.BOT = self.bot

    def tearDown(self):
        from lucabot.ortak import DURDUR_DOSYASI
        try:
            DURDUR_DOSYASI.unlink()  # sahte bot dosyayi tuketmez
        except OSError:
            pass
        if self.app.surec:
            self.app.surec.kill()
            self.app.surec.wait()
            self.app.surec.stdout.close()
        self.kok.after_cancel(self.app.dongu_id)
        self.kok.destroy()
        os.environ.pop("LUCA_BOT_AYAR", None)

    def _bekle(self, kosul, saniye=20):
        son = time.time() + saniye
        while time.time() < son:
            self.kok.update()
            if kosul():
                return True
            time.sleep(0.05)
        return False


@unittest.skipUnless(TK_VAR, "tkinter ya da ekran yok")
class ArayuzTestleri(ArayuzZemini):
    def test_kutular_rapordan_doluyor(self):
        k = self.app.kutu
        self.assertEqual(k["tevkifat"].deger.cget("text"), "700,00 TL")
        self.assertEqual(k["smm"].deger.cget("text"), "1.500,00 TL")
        self.assertEqual(k["fark"].deger.cget("text"), "0 fatura")
        self.assertEqual(k["kdv"].deger.cget("text"), "1 firma")  # 1000 - 300 - 200 devreden
        self.assertEqual(self.app.gostergeler["kdv"][0]["odeme"], 500)

    def test_detay_penceresi_acilir(self):
        once = len(self.kok.winfo_children())
        self.app.detay("kdv")
        self.kok.update()
        self.assertEqual(len(self.kok.winfo_children()), once + 1)

    def test_komut(self):
        komut = self.app.komut()
        self.assertIn("--hepsi", komut)
        self.assertIn("--bitince-kapat", komut)
        self.assertEqual(komut[komut.index("--baslangic") + 1], "01/09/2026")
        self.app.v_ekran["e-arsiv-satis"].set(False)
        self.app.v_devam.set(False)
        self.app.v_firma.set("BIRLIK")
        komut = self.app.komut()
        self.assertNotIn("--hepsi", komut)
        self.assertEqual(komut.count("--belge-tipi"), 9)
        self.assertIn("--bastan", komut)
        self.assertEqual(komut[komut.index("--firma") + 1], "BIRLIK")
        self.app.v_bit.set("31/08/2026")
        with self.assertRaises(ValueError):
            self.app.komut()
        self.app.v_bit.set("30-09")
        with self.assertRaises(ValueError):
            self.app.komut()

    def test_calistir_canli_log_ve_bitis(self):
        self.app.calistir()
        self.assertIsNotNone(self.app.surec)
        self.assertTrue(self._bekle(lambda: self.app.surec is None))
        log = self.app.log_metni()
        self.assertIn("--bitince-kapat", log)
        self.assertIn("[OK] e-Arşiv Alış Faturaları", log)
        self.assertIn("son satir (yeni satir yok)", log)
        self.assertIn("Tamamlandı", self.app.ilerleme_etiketi.cget("text"))
        self.assertEqual(self.app.durum_etiketi.cget("text").strip(), "Tamamlandı")
        self.assertEqual(str(self.app.calistir_dugmesi.cget("state")), "normal")
        # ayarlar kaydedildi, arayuzun bilmedigi anahtarlar korundu
        ayar = json.loads(self.ayar.read_text(encoding="utf-8"))
        self.assertEqual(ayar["baska_ayar"], 42)
        self.assertEqual(len(ayar["arayuz_ekranlar"]), 10)
        # renklendirme: [OK] satiri yesil etiketli
        self.assertTrue(self.app.log.tag_ranges("ok"))
        self.assertTrue(self.app.log.tag_ranges("firma"))

    def test_ilerleme_satiri(self):
        self.app._satir("[13/93] BIRLIK TIC  | tahmini kalan: 2 sa 10 dk")
        self.assertIn("BIRLIK TIC", self.app.ilerleme_etiketi.cget("text"))
        self.assertIn("13 / 93", self.app.ilerleme_etiketi.cget("text"))
        self.assertEqual(self.app.kalan_etiketi.cget("text"), "Tahmini kalan: 2 sa 10 dk")
        self.assertAlmostEqual(float(self.app.ilerleme.cget("value")), 12 * 100 / 93, places=3)

    def test_su_an_gostergesi(self):
        self.app.surec = object()  # calisiyor gibi
        try:
            self.app._satir("    GİB'den Getir aciliyor (01/09/2026 - 07/09/2026)")
            self.app._satir("================")  # ayirici cizgiler "su an" yazisi olmaz
            metin, an = self.app.son_islem
            self.assertIn("GİB'den Getir aciliyor", metin)
            self.app.son_islem = (metin, an - self.mod.UZUN_BEKLEME * 3)
            self.app._son_islemi_goster()
            yazi = self.app.son_islem_etiketi.cget("text")
            self.assertTrue(yazi.startswith("Şu an: GİB'den Getir aciliyor"), yazi)
            self.assertEqual(self.app.son_islem_etiketi.cget("fg"), self.mod.KIRMIZI)
        finally:
            self.app.surec = None
        self.app._son_islemi_goster()
        self.assertEqual(self.app.son_islem_etiketi.cget("text"), "")

    def test_durdur_sonuclari_kaydederek_durdurur(self):
        self.bot.write_text(SAHTE_BOT.replace('float(sys.argv[-1]) if sys.argv[-1].replace(".", "").isdigit() else 0.2', "30"),
                            encoding="utf-8")
        self.app.calistir()
        self.assertTrue(self._bekle(lambda: "[1/2]" in self.app.log_metni()))
        self.app.durdur()
        self.assertTrue(self._bekle(lambda: self.app.surec is None))
        self.assertIn("Durduruldu. O ana kadarki sonuclar kaydedildi", self.app.log_metni())
        self.assertEqual(self.app.durum_etiketi.cget("text").strip(), "Durduruldu")

    def test_durdur_dosyasi_yazilir_ve_eskisi_silinir(self):
        from lucabot.ortak import DURDUR_DOSYASI
        self.addCleanup(lambda: DURDUR_DOSYASI.exists() and DURDUR_DOSYASI.unlink())
        DURDUR_DOSYASI.write_text("kalinti")  # onceki calismadan kalmis istek
        self.bot.write_text(SAHTE_BOT.replace('float(sys.argv[-1]) if sys.argv[-1].replace(".", "").isdigit() else 0.2', "30"),
                            encoding="utf-8")
        self.app.calistir()
        self.assertFalse(DURDUR_DOSYASI.exists())  # yeni calisma kalintiyla durmasin
        self.assertTrue(self._bekle(lambda: "[1/2]" in self.app.log_metni()))
        self.app.durdur()
        self.assertTrue(DURDUR_DOSYASI.exists())   # bot bunu gorup duracak
        self.assertTrue(self._bekle(lambda: self.app.surec is None))

    def test_konsolsuz_yeniden_baslatma(self):
        """Konsollu python.exe ile acilan arayuz pythonw ile yeniden baslar (arkada siyah pencere kalmaz)."""
        import types
        from unittest import mock
        self.assertFalse(self.mod.konsolsuz_yeniden_baslat())  # Windows disi: dokunmaz
        d = Path(tempfile.mkdtemp())
        (d / "python.exe").write_text("")
        (d / "pythonw.exe").write_text("")

        def dene(konsol, calisti=True, exe="python.exe"):
            sahte_ctypes = types.SimpleNamespace(windll=types.SimpleNamespace(
                kernel32=types.SimpleNamespace(GetConsoleWindow=lambda: konsol)))
            with mock.patch.dict(sys.modules, {"ctypes": sahte_ctypes}), \
                    mock.patch.object(sys, "platform", "win32"), \
                    mock.patch.object(sys, "executable", str(d / exe)), \
                    mock.patch("subprocess.Popen") as popen, mock.patch("time.sleep"):
                popen.return_value.poll.return_value = None if calisti else 1
                sonuc = self.mod.konsolsuz_yeniden_baslat()
                return sonuc, popen
        sonuc, popen = dene(konsol=123)
        self.assertTrue(sonuc)
        self.assertTrue(popen.call_args[0][0][0].endswith("pythonw.exe"))
        self.assertFalse(dene(konsol=0)[0])                    # zaten konsolsuz: yeniden baslatmaz
        self.assertFalse(dene(konsol=123, calisti=False)[0])   # yeni surec dustuyse bu surecte acilir
        self.assertFalse(dene(konsol=123, exe="pythonw.exe")[0])
        (d / "pythonw.exe").unlink()
        self.assertFalse(dene(konsol=123)[0])                  # pythonw yoksa eski yontem

    def test_firma_ekran_penceresi_kaydeder(self):
        from lucabot import firma_tablosu
        from lucabot.firma_listesi import firma_listesini_oku
        firma_tablosu.sablon_olustur(self.d / "firmalar.xlsx", ["BIRLIK TICARET", "KAYA INSAAT"])
        fp = self.app.firma_ekran_penceresi()
        self.kok.update()
        self.assertEqual(len(fp.satirlar), 2)
        fp.v_ara.set("kaya")
        self.kok.update()
        self.assertEqual([s["ad"] for s in fp.gorunen()], ["KAYA INSAAT"])
        fp._sutunu_cevir("gib-5000")           # yalnizca gorunen (KAYA) firmada kapanir
        fp.v_ara.set("")
        fp._satiri_cevir(fp.satirlar[0])       # BIRLIK: tum ekranlar kapanir
        fp.satirlar[0]["secim"]["e-arsiv-alis"].set(True)
        fp.satirlar[0]["dev"].set("1.250,50")
        fp.kaydet()
        okunan = firma_listesini_oku(self.d / "firmalar.xlsx")
        self.assertEqual(okunan["KAYA INSAAT"][1], {"gib-5000"})
        self.assertEqual(len(okunan["BIRLIK TICARET"][1]), 9)
        self.assertNotIn("e-arsiv-alis", okunan["BIRLIK TICARET"][1])
        from lucabot.firma_listesi import devreden_kdvleri
        self.assertEqual(devreden_kdvleri(self.d / "firmalar.xlsx"), {"BIRLIK TICARET": 1250.5})
        # kaydedince kutular yenilenir: 1000 - 300 - 1250,50 devreden < 0, odeme cikmaz
        self.assertEqual(self.app.gostergeler["kdv"], [])

    def test_luca_dan_firma_listesi_cekilir_ve_uygulanir(self):
        from lucabot import firma_tablosu, musteri_listesi
        from lucabot.firma_listesi import firma_listesini_oku
        from tkinter import messagebox
        from datetime import date
        firma_tablosu.sablon_olustur(self.d / "firmalar.xlsx", ["BIRLIK TICARET"])
        self.bot.write_text(SAHTE_LISTE_BOTU, encoding="utf-8")
        sorular = []
        eski, messagebox.askyesno = messagebox.askyesno, lambda *a, **k: sorular.append(a) or True
        try:
            self.app.firma_listesi_cek()
        finally:
            messagebox.askyesno = eski
        self.assertIn("2026", sorular[0][1])
        self.assertEqual(str(self.app.luca_liste_dugmesi.cget("state")), "disabled")
        self.assertTrue(self._bekle(lambda: self.app.surec is None))
        self.assertIn("--firma-listesi-cek --yil 2026", self.app.log_metni())
        self.assertEqual(str(self.app.luca_liste_dugmesi.cget("state")), "normal")
        pencere = self.app.luca_penceresi
        self.assertIsNotNone(pencere)
        self.kok.update()
        self.assertEqual(len(pencere.plan["yeni"]), 1)
        self.assertEqual(len(pencere.agac.get_children()), 2)   # yeni + guncellenecek
        pencere.uygula()
        okunan = firma_listesini_oku(self.d / "firmalar.xlsx")
        self.assertEqual(okunan["BIRLIK TICARET"][0], date(2026, 3, 31))
        self.assertIn("YENI FIRMA LTD", okunan)
        self.assertEqual(self.app.mod, "calisma")

    def test_luca_dan_devir_cekilir_ve_beyannameden_alinir(self):
        from lucabot import firma_tablosu
        from tkinter import messagebox
        firma_tablosu.sablon_olustur(self.d / "firmalar.xlsx", ["BIRLIK TICARET", "CAGRI DENE"])
        self.bot.write_text(SAHTE_BEYANNAME_BOTU, encoding="utf-8")
        fp = self.app.firma_ekran_penceresi()
        alinan = []
        fp.beyannameden_al = lambda yollar=None: alinan.append([Path(y).name for y in yollar])
        eski, messagebox.askyesno = messagebox.askyesno, lambda *a, **k: True
        try:
            self.app.beyanname_cek(fp)
        finally:
            messagebox.askyesno = eski
        self.assertTrue(self._bekle(lambda: self.app.surec is None))
        self.assertIn("--beyanname-cek", self.app.log_metni())
        # yalniz var olan dosyalar iletilir
        self.assertEqual(alinan, [["CAGRI DENE_1_KDV1_2026-08_1.pdf", "BIRLIK TICARET_2_KDV1_2026-08_1.pdf"]])
        self.assertEqual(self.app.mod, "calisma")
        fp.w.destroy()

    @staticmethod
    def _agac(pencere):
        """Pencerenin icindeki ilk Treeview (ic ice cercevelerde aranir)."""
        from tkinter import ttk
        bekleyen = [pencere]
        while bekleyen:
            w = bekleyen.pop(0)
            if isinstance(w, ttk.Treeview):
                return w
            bekleyen += w.winfo_children()

    def _rapora_yaz(self, kayitlar):
        (self.d / "indir" / "rapor.json").write_text(json.dumps(kayitlar), encoding="utf-8")
        self.app.gostergeleri_yenile()

    def test_fark_penceresi_eksik_faturalari_gosterir(self):
        donem = "01/09/2026-30/09/2026"
        self._rapora_yaz({"ALEV SEZEN": {
            "donem": donem, "sayilar": {"e-arsiv-alis": 2, "e-arsiv-interaktif": 3},
            "faturalar": {
                "e-arsiv-alis": [["TURKCELL", "AAA2026000000001", 100.0], ["SHELL", "AAA2026000000002", 50.0]],
                "e-arsiv-interaktif": [["TURKCELL", "AAA2026000000001", 100.0], ["SHELL", "AAA2026000000002", 50.0],
                                       ["VODAFONE", "AAA2026000000007", 1200.5]]}}})
        w = self.app.detay("fark")
        self.kok.update()
        agac = self._agac(w)
        satirlar = [agac.item(i, "values") for i in agac.get_children()]
        self.assertEqual(satirlar[0][0], "ALEV SEZEN")
        self.assertIn("VODAFONE", satirlar[1][0])
        self.assertIn("AAA..007", satirlar[1][0])
        self.assertIn("1.200,50 TL", satirlar[1][0])
        self.assertIn("eksik", satirlar[1][3])
        # sutunlar sigacak genislikte (sayi sutunu kesilmesin)
        self.assertGreaterEqual(int(agac.column("İnteraktif", "width")), 90)
        w.destroy()

    def test_hata_kutusu_ve_tekrar_sorgula(self):
        from tkinter import messagebox
        donem = "01/09/2026-30/09/2026"
        self._rapora_yaz({
            "ALEV SEZEN": {"donem": donem, "durumlar": {"e-arsiv-alis": "tamam"}, "inmeyen": {"e-arsiv-alis": 2}},
            "METIN BALT": {"donem": donem, "durumlar": {"e-arsiv-satis": "hata: LookupError",
                                                       "gib-5000": "tamam"}}})
        self.assertIn("2 ekran", self.app.kutu["hata"].deger.cget("text"))
        w = self.app.detay("hata")
        self.kok.update()
        eski, messagebox.askyesno = messagebox.askyesno, lambda *a, **k: True
        try:
            self.app.tekrar_sorgula(self.app.gostergeler["hata"], w)
        finally:
            messagebox.askyesno = eski
        self.assertTrue(self._bekle(lambda: self.app.surec is None))
        istek = json.loads((self.d / "indir" / "tekrar-listesi.json").read_text(encoding="utf-8"))
        self.assertEqual(istek, {"ALEV SEZEN": ["e-arsiv-alis"], "METIN BALT": ["e-arsiv-satis"]})
        log = self.app.log_metni()
        self.assertIn("--tekrar-listesi", log)
        self.assertIn("--baslangic 01/09/2026 --bitis 30/09/2026", log)
        self.assertNotIn("--belge-tipi", log)

    def test_luca_listesi_cekilemezse_hata_gosterilir(self):
        from tkinter import messagebox
        self.bot.write_text("import sys\nprint('HATA: Filtre penceresi kullanilamadi')\nsys.exit(1)\n",
                            encoding="utf-8")
        eski, messagebox.askyesno = messagebox.askyesno, lambda *a, **k: True
        try:
            self.app.firma_listesi_cek()
        finally:
            messagebox.askyesno = eski
        self.assertTrue(self._bekle(lambda: self.app.surec is None))
        self.assertIsNone(self.app.luca_penceresi)
        self.assertIn("alınamadı", self.app.ilerleme_etiketi.cget("text"))

    def test_beyannameden_devir_alinir(self):
        from lucabot import beyanname, firma_tablosu
        from lucabot.firma_listesi import devreden_kdvleri
        from testler.test_birim import ORNEK_BEYANNAME
        firma_tablosu.sablon_olustur(self.d / "firmalar.xlsx", ["BIRLIK TICARET", "CAGRI DENE"])
        fp = self.app.firma_ekran_penceresi()

        def sahte_oku(yol):
            if "bozuk" in yol:
                raise beyanname.BeyannameDegil("KDV1 beyannamesi değil")
            return beyanname.cozumle(ORNEK_BEYANNAME, yol)

        eski, beyanname.oku = beyanname.oku, sahte_oku
        try:
            once = len(self.kok.winfo_children())
            sonuc = fp.beyannameden_al(["CAGRI_DENE_1_KDV1_01082026-31082026.pdf", "bozuk.pdf"])
        finally:
            beyanname.oku = eski
        self.kok.update()
        self.assertEqual(len(self.kok.winfo_children()), once + 1)  # sonuc penceresi
        self.assertEqual(fp.satirlar[1]["dev"].get(), "34.027,69")   # eylul: agustosun sonraki devri
        self.assertEqual(fp.satirlar[0]["dev"].get(), "")
        self.assertEqual([r[2] for r in sonuc], ["Yazıldı", "Okunamadı"])
        fp.kaydet()
        self.assertEqual(devreden_kdvleri(self.d / "firmalar.xlsx")["CAGRI DENE"], 34027.69)

    def test_arayuzden_firma_eklenir_ve_silinir(self):
        from lucabot import firma_tablosu
        from lucabot.firma_listesi import devreden_kdvleri, firma_listesini_oku
        firma_tablosu.sablon_olustur(self.d / "firmalar.xlsx", ["BIRLIK TICARET", "KAYA INSAAT"])
        fp = self.app.firma_ekran_penceresi()
        fp.v_yeni.set("  YENI   FIRMA ")
        fp.firma_ekle()
        fp.v_yeni.set("yeni firma")       # ayni firma ikinci kez eklenmez
        uyari = []
        eski = self.mod.messagebox.showinfo
        self.mod.messagebox.showinfo = lambda *a, **k: uyari.append(a)
        eski_soru = self.mod.messagebox.askyesno
        self.mod.messagebox.askyesno = lambda *a, **k: True
        try:
            fp.firma_ekle()
            yeni = fp.satirlar[-1]
            yeni["dev"].set("750")
            yeni["secim"]["turmob-alis"].set(False)
            fp.firma_sil(fp.satirlar[1])  # KAYA
        finally:
            self.mod.messagebox.showinfo = eski
            self.mod.messagebox.askyesno = eski_soru
        self.assertTrue(uyari)
        self.assertEqual(fp.bilgi.cget("text"), "2 firma")
        fp.kaydet()
        okunan = firma_listesini_oku(self.d / "firmalar.xlsx")
        self.assertEqual(set(okunan), {"BIRLIK TICARET", "YENI FIRMA"})
        self.assertEqual(okunan["YENI FIRMA"][1], {"turmob-alis"})
        self.assertEqual(devreden_kdvleri(self.d / "firmalar.xlsx"), {"YENI FIRMA": 750.0})

    def test_liste_yoksa_olusturulur(self):
        self.app.v_liste.set("")
        hedef = self.mod.KOK / "firmalar.xlsx"
        if hedef.exists():
            self.skipTest("program klasorunde gercek firmalar.xlsx var")
        eski = self.mod.messagebox.askyesno
        self.mod.messagebox.askyesno = lambda *a, **k: True
        try:
            fp = self.app.firma_ekran_penceresi()
        finally:
            self.mod.messagebox.askyesno = eski
            self.addCleanup(lambda: hedef.unlink(missing_ok=True))
        # rapor.json'daki firma hazir gelir, sablonun ornek satiri gelmez
        self.assertEqual([s["ad"] for s in fp.satirlar], ["BIRLIK TIC"])
        self.assertEqual(self.app.v_liste.get(), "firmalar.xlsx")
        self.assertEqual(json.loads(self.ayar.read_text(encoding="utf-8"))["firma_listesi"], "firmalar.xlsx")

    def test_liste_excel_olarak_iner(self):
        from openpyxl import load_workbook
        yol = self.d / "liste.xlsx"
        self.mod.liste_excel_yaz(yol, "Alış Tevkifat KDV — firmalar", "01/09/2026-30/09/2026",
                                 ("Firma", "Fatura", "Tevkifat KDV"),
                                 [("A", 2, "1.234,50 TL *"), ("B", 1, "100,00 TL")],
                                 ("Toplam", 3, "1.334,50 TL"), ["not"])
        ws = load_workbook(yol).active
        self.assertEqual(ws["A4"].value, "Firma")
        self.assertEqual(ws["C5"].value, 1234.5)
        self.assertEqual(ws["C7"].value, 1334.5)

    def test_giris_bilgisi_yoksa_calismaz(self):
        self.app.ayarlar["parola"] = ""
        import luca_arayuz
        uyari = []
        eski = luca_arayuz.messagebox.showwarning
        luca_arayuz.messagebox.showwarning = lambda *a, **k: uyari.append(a)
        try:
            self.app.calistir()
        finally:
            luca_arayuz.messagebox.showwarning = eski
        self.assertIsNone(self.app.surec)
        self.assertTrue(uyari)


@unittest.skipUnless(TK_VAR, "tkinter ya da ekran yok")
class KarZararSekmesiTestleri(ArayuzZemini):
    """Kâr / Zarar sekmesi: kar-zarar/kar_zarar.py yerine sahte program calisir."""

    def setUp(self):
        import kar_zarar_sekmesi
        self.kz_klasor = Path(tempfile.mkdtemp(prefix="kz-test-"))
        (self.kz_klasor / "kar_zarar.py").write_text(SAHTE_KAR_ZARAR, encoding="utf-8")
        (self.kz_klasor / "ayarlar.ornek.json").write_text(json.dumps({"luca_kaynagi": "mizan"}), encoding="utf-8")
        self._eski = kar_zarar_sekmesi.ADAYLAR
        kar_zarar_sekmesi.ADAYLAR = (self.kz_klasor,)
        super().setUp()
        self.kz = self.app.kz

    def tearDown(self):
        import kar_zarar_sekmesi
        if self.kz.surec:
            self.kz.surec.kill()
            self.kz.surec.wait()
            self.kz.surec.stdout.close()
        kar_zarar_sekmesi.ADAYLAR = self._eski
        super().tearDown()

    def _bilgi_sustur(self):
        """showinfo modal pencere acip testi bekletir; cagrilari toplar."""
        import luca_arayuz
        import kar_zarar_sekmesi
        self.bilgi = []
        for mod in (luca_arayuz, kar_zarar_sekmesi):
            eski = mod.messagebox.showinfo
            mod.messagebox.showinfo = lambda *a, **k: self.bilgi.append(a)
            self.addCleanup(setattr, mod.messagebox, "showinfo", eski)

    def test_sekme_gecisi(self):
        self.app.sekme_sec("kz")
        self.kok.update()
        self.assertTrue(self.kz.govde.winfo_ismapped())
        self.assertFalse(self.app.fatura_govde.winfo_ismapped())
        self.app.sekme_sec("fatura")
        self.kok.update()
        self.assertTrue(self.app.fatura_govde.winfo_ismapped())

    def test_calistir_tablo_ve_faturalar_dahil(self):
        self.app.sekme_sec("kz")
        self.kz.v_bas.set("01/01/2026")
        self.kz.v_bit.set("31/08/2026")
        self.kz.v_db_kod.set("dbkod")
        self.kz.v_kaynak.set("Hesap Planı Listesi")
        # Eylul faturalari: satis 1500 - alis (500 e-fatura + 1500 e-SMM) => -500
        rapor = self.d / "indir" / "rapor.json"
        veri = json.loads(rapor.read_text(encoding="utf-8"))
        veri["BIRLIK TIC"]["matrah"].update({"e-arsiv-satis": 1500, "e-fatura-alis": 500})
        rapor.write_text(json.dumps(veri), encoding="utf-8")
        # gecmis donem kaydi: Ekim faturalari (satis 5000, alis yok) -> kutuda iki ay olur
        (self.d / "indir" / "rapor-donemler.json").write_text(json.dumps({"01/10/2026-31/10/2026": {
            "BIRLIK TIC": {"firma": "BIRLIK TIC", "donem": "01/10/2026-31/10/2026",
                           "matrah": {"e-arsiv-satis": 5000}, "durumlar": {"e-arsiv-satis": "tamam"}}}}),
            encoding="utf-8")
        self.assertFalse(self.kz.v_fatura.get())  # baslangicta faturalar hesaba katilmaz
        self.kz.calistir()
        self.assertIsNotNone(self.kz.surec)
        self._bilgi_sustur()
        self.app.calistir()  # fatura calismasi basindan engellenir
        self.assertIsNone(self.app.surec)
        self.assertTrue(self._bekle(lambda: self.kz.surec is None))
        log = self.kz.log_metni()
        self.assertIn("--tarih 01/01/2026-31/08/2026", log)
        self.assertIn('"defterbeyan_kullanici": "dbkod"', log)
        self.assertIn('"luca_kaynagi": "hesap-plani"', log)
        self.assertIn('"parola": "x"', log)  # luca-bot girisi kar-zarar ayarina aktarildi
        satirlar = [self.kz.agac.item(i)["values"] for i in self.kz.agac.get_children()]
        self.assertEqual([s[0] for s in satirlar], ["BIRLIK TIC", "YILDIZ OTO", "HATALI LTD"])
        self.assertEqual(satirlar[0][5], "—")  # dugmeye basilmadan faturalar dahil degil
        self.assertEqual(list(self.kz.fatura_kutusu.cget("values")),
                         ["Eylül 2026 · 1 firma", "Ekim 2026 · 1 firma"])
        self.assertEqual(self.kz.v_fatura_donem.get(), "Ekim 2026 · 1 firma")  # en yeni ay varsayilan
        self.kz.v_fatura_donem.set("Eylül 2026 · 1 firma")  # kullanici ayi kendisi secer
        self.kz.fatura_kutusu.event_generate("<<ComboboxSelected>>")
        self.kok.update()
        self.kz.fatura_dugmesi.invoke()
        self.assertEqual(self.kz.fatura_dugmesi.cget("text"), "Faturaları Hariç Tut")
        satirlar = [self.kz.agac.item(i)["values"] for i in self.kz.agac.get_children()]
        self.assertIn("100.000,00 TL kâr", satirlar[0][3])
        self.assertIn("-500,00 TL", satirlar[0][4])  # fatura farki (satis - alis)
        self.assertIn("99.500,00 TL kâr", satirlar[0][5])  # faturalar dahil
        self.assertIn("faturalar dahil edilince 96.500,00 TL kâr", self.kz.ozet_etiketi.cget("text"))
        self.assertIn("Dahil edilen: Eylül 2026", self.kz.fatura_etiketi.cget("text"))
        self.assertIn("zarar", satirlar[1][3])
        self.assertEqual(satirlar[1][5], "—")
        self.assertIn("mizan inmedi", satirlar[2][6])
        self.kz.v_fatura.set(True)
        self.kz.v_fatura_donem.set("Ekim 2026 · 1 firma")
        self.kz.fatura_kutusu.event_generate("<<ComboboxSelected>>")
        self.kok.update()
        satirlar = [self.kz.agac.item(i)["values"] for i in self.kz.agac.get_children()]
        self.assertIn("100.000,00 TL kâr", satirlar[0][3])
        self.assertIn("105.000,00 TL kâr", satirlar[0][5])  # Ekim: +5000
        self.kz.fatura_dugmesi.invoke()  # tekrar basinca faturalar cikarilir
        self.assertEqual(self.kz.fatura_dugmesi.cget("text"), "Taranan Faturaları Dahil Et")
        satirlar = [self.kz.agac.item(i)["values"] for i in self.kz.agac.get_children()]
        self.assertEqual(satirlar[0][5], "—")
        self.assertIn("hesaba katılmıyor", self.kz.fatura_etiketi.cget("text"))
        self.kz.agac.selection_set("0")
        self.kok.update()
        self.assertIn("BIRLIK TIC", self.kz.ozet_etiketi.cget("text"))

    def test_eski_sorgular_tabloda_kalir_toplam_gorunenden(self):
        cikti = self.kz.cikti_klasoru()
        cikti.mkdir(parents=True, exist_ok=True)
        (cikti / "kar-zarar.json").write_text(json.dumps({"donem": "01/01/2026-31/08/2026", "firmalar": [
            {"firma": "ESKI FIRMA", "donem": "01/01/2026-31/07/2026", "kaynak": "Luca (Mizan)", "kar": 500.0},
            {"firma": "BIRLIK TIC", "donem": "01/01/2026-31/08/2026", "kaynak": "Luca (Mizan)", "kar": 100.0},
            {"firma": "YILDIZ OTO", "donem": "01/01/2026-31/08/2026", "kaynak": "Defter Beyan", "kar": -30.0,
             "son_hata": "HATA: oturum"}]}), encoding="utf-8")
        self.kz.sonucu_yukle()
        t = self.kz.tablo
        self.assertEqual(len(t.satirlar), 3)                       # eski sorgunun firmasi da duruyor
        self.assertEqual(t.filtre_degiskenleri["Dönem"].get(), "01/01/2026-31/08/2026")  # son sorgu secili
        self.assertEqual([t.satirlar[i][0] for i in t.gorunen_indeksler()], ["BIRLIK TIC", "YILDIZ OTO"])
        self.assertIn("Toplam (2 firma): dönem 70,00 TL kâr", self.kz.ozet_etiketi.cget("text"))
        yildiz = next(s for s in t.satirlar if s[0] == "YILDIZ OTO")
        self.assertIn("Son sorguda hata: HATA: oturum", yildiz[6])
        t.filtre_ayarla("Dönem", "Tümü")                           # iki donem birlikte: toplanmaz
        self.assertIn("Birden fazla dönem", self.kz.ozet_etiketi.cget("text"))
        t.filtre_ayarla("Dönem", "01/01/2026-31/07/2026")
        self.assertIn("Toplam (1 firma): dönem 500,00 TL kâr", self.kz.ozet_etiketi.cget("text"))

    def test_fatura_calismasi_surerken_kar_zarar_baslamaz(self):
        self.app.calistir()
        self.assertIsNotNone(self.app.surec)
        self._bilgi_sustur()
        self.kz.calistir()
        self.assertIsNone(self.kz.surec)
        self.assertTrue(self.bilgi)

    def test_tema_varsayilan_acik_ve_degisir(self):
        import luca_arayuz
        self.assertEqual(luca_arayuz.TEMA, "acik")
        self.assertEqual(luca_arayuz.ZEMIN, luca_arayuz.TEMALAR["acik"]["ZEMIN"])
        eski_kz = self.app.kz
        self.app.v_bas.set("02/09/2026")  # ekrandaki degerler tema degisince kaybolmaz
        self.app.tema_degistir()
        self.assertEqual(luca_arayuz.TEMA, "koyu")
        self.assertEqual(json.loads(self.ayar.read_text(encoding="utf-8"))["tema"], "koyu")
        self.assertEqual(self.app.v_bas.get(), "02/09/2026")
        self.assertIsNot(self.app.kz, eski_kz)                       # arayuz yeniden kuruldu
        self.assertEqual(self.kok.cget("bg"), luca_arayuz.TEMALAR["koyu"]["ZEMIN"])
        self.app.sekme_sec("kz")
        self.kok.update()
        self.assertTrue(self.app.kz.govde.winfo_ismapped())
        self.app.tema_degistir()                                      # geri
        self.assertEqual(luca_arayuz.TEMA, "acik")
        self.assertEqual(self.kok.cget("bg"), luca_arayuz.TEMALAR["acik"]["ZEMIN"])

    def test_tema_calisirken_degismez(self):
        import luca_arayuz
        self.app.calistir()
        self.assertIsNotNone(self.app.surec)
        bilgi = []
        eski = luca_arayuz.messagebox.showinfo
        luca_arayuz.messagebox.showinfo = lambda *a, **k: bilgi.append(a)
        try:
            self.app.tema_degistir()
        finally:
            luca_arayuz.messagebox.showinfo = eski
        self.assertEqual(luca_arayuz.TEMA, "acik")
        self.assertTrue(bilgi)

    def test_baslik_metinleri(self):
        yazilar = []

        def tara(w):
            for c in w.winfo_children():
                try:
                    yazilar.append(c.cget("text"))
                except Exception:
                    pass
                tara(c)
        tara(self.kok)
        for beklenen in ("Dijital Stajyer", "SMMM OFİSİ - DİJİTAL ASİSTAN", "Lisans sahibi: S. Adem Yılmaz",
                         "Luca · e-Fatura / e-Arşiv · Kâr / Zarar"):
            self.assertIn(beklenen, yazilar)
        self.assertEqual(self.kok.title(), "Dijital Stajyer")
        self.assertNotIn("smmmyilmaz.com", yazilar)

    def test_firma_listesi_yenile_secenegi(self):
        self.assertNotIn("--listeyi-yenile", self.kz.komut())  # varsayilan: kayitli liste
        self.assertIn("Kayıtlı liste yok", self.kz.liste_bilgisi.cget("text"))
        cikti = self.kz.cikti_klasoru()
        cikti.mkdir(parents=True, exist_ok=True)
        (cikti / "musteri-listeleri.json").write_text(json.dumps({"yil": 2026, "alinma": "2026-10-05T10:00:00",
            "siniflar": {"1.Sınıf": [{"ad": "A"}, {"ad": "B"}], "İşletme Defteri": [{"ad": "B"}]}}), encoding="utf-8")
        self.kz.liste_bilgisi_yaz()
        self.assertIn("05/10/2026 tarihinde alındı, 2 firma", self.kz.liste_bilgisi.cget("text"))
        self.kz.v_yenile.set(True)
        self.assertIn("--listeyi-yenile", self.kz.komut())
        self.kz.liste_bilgisi_yaz()
        self.assertIn("yeniden okunur", self.kz.liste_bilgisi.cget("text"))

    def test_varsayilan_donem(self):
        import kar_zarar_sekmesi
        from datetime import date
        self.assertEqual(kar_zarar_sekmesi.varsayilan_donem(date(2026, 10, 5)), ("01/01/2026", "31/08/2026"))
        self.assertEqual(kar_zarar_sekmesi.varsayilan_donem(date(2027, 1, 10)), ("01/01/2026", "30/11/2026"))


@unittest.skipUnless(TK_VAR, "tkinter ya da ekran yok")
class TablolarVeSekmelerTestleri(ArayuzZemini):
    """Fatura Indirme alt sekmeleri (Surec / Firma Durumu / Hatali), siralama-filtre ve yer tutucu sekme."""

    def setUp(self):
        super().setUp()
        donem = "01/09/2026-30/09/2026"
        (self.d / "indir" / "rapor.json").write_text(json.dumps({
            "BIRLIK TIC": {"firma": "BIRLIK TIC", "donem": donem,
                           "durumlar": {"e-arsiv-alis": "tamam", "e-arsiv-satis": "tamam"},
                           "sayilar": {"e-arsiv-alis": 12, "e-arsiv-satis": 3},
                           "matrah": {"e-arsiv-satis": 15000, "e-arsiv-alis": 2500}},
            "YILDIZ OTO": {"firma": "YILDIZ OTO", "donem": donem,
                           "durumlar": {"e-arsiv-alis": "hata: zaman asimi", "e-fatura-alis": "tamam"},
                           "sayilar": {"e-arsiv-alis": 0, "e-fatura-alis": 40}, "inmeyen": {"e-fatura-alis": 2},
                           "notlar": {"e-arsiv-alis": "zaman asimi"}},
            "ZEYTIN LTD": {"firma": "ZEYTIN LTD", "donem": donem,
                           "durumlar": {"e-fatura-satis": "excel inmedi"}, "sayilar": {"e-fatura-satis": 7}},
            "ESKI AS": {"firma": "ESKI AS", "donem": "01/08/2026-31/08/2026", "durumlar": {"e-arsiv-alis": "tamam"},
                        "sayilar": {"e-arsiv-alis": 5}},
        }), encoding="utf-8")
        self.app.gostergeleri_yenile()

    def test_firma_durumu_tablosu_rapordaki_verilerle_dolu(self):
        t = self.app.durum_tablosu
        self.assertEqual(t.sutunlar[:4], ["Firma", "Dönem", "Durum", "Aksiyon"])
        # varsayilan Donem filtresi ana penceredeki ay (Eylul): Agustos'a ait ESKI AS gorunmez
        gorunen = [t.satirlar[i][0] for i in t.gorunen_indeksler()]
        self.assertEqual(sorted(gorunen), ["BIRLIK TIC", "YILDIZ OTO", "ZEYTIN LTD"])
        t.filtre_ayarla("Dönem", "Tümü")
        self.assertEqual(len(t.gorunen_indeksler()), 4)
        self.assertIn("4 / 4", t.sayac.cget("text"))
        etiket = {t.satirlar[i][0]: t.satir_etiketleri[i] for i in range(4)}
        self.assertEqual(etiket["BIRLIK TIC"], "tamam")
        self.assertEqual(etiket["YILDIZ OTO"], "hata")

    def test_arama_filtre_ve_siralama(self):
        t = self.app.durum_tablosu
        t.filtre_ayarla("Dönem", "Tümü")
        t.arama_degiskeni.set("yildiz")  # Turkce harf/buyuk-kucuk fark etmez
        self.assertEqual([t.satirlar[i][0] for i in t.gorunen_indeksler()], ["YILDIZ OTO"])
        t.arama_degiskeni.set("")
        t.filtre_ayarla("Durum", "tamam")
        self.assertTrue(all(t.satirlar[i][2] == "tamam" for i in t.gorunen_indeksler()))
        t.temizle()
        t.filtre_ayarla("Dönem", "Tümü")
        t.sirala(0)  # Firma artan
        self.assertEqual([t.satirlar[i][0] for i in t.gorunen_indeksler()],
                         ["BIRLIK TIC", "ESKI AS", "YILDIZ OTO", "ZEYTIN LTD"])
        t.sirala(0)  # azalan
        self.assertEqual([t.satirlar[i][0] for i in t.gorunen_indeksler()][0], "ZEYTIN LTD")
        t.sirala(0)  # siralama kalkar
        self.assertIsNone(t.siralama)

    def test_siralama_anahtari_sayi_tutar_tarih(self):
        from tablo_gorunumu import siralama_anahtari as a
        self.assertLess(a("9"), a("12"))                       # sayisal, metin degil
        self.assertLess(a("1.234,50 TL"), a("10.000,00 TL"))   # Turk bicimli tutar
        self.assertLess(a("-3.000,00 TL zarar"), a("100,00 TL kâr"))
        self.assertLess(a("30/08/2026"), a("01/09/2026"))
        self.assertLess(a("12"), a("abc"))                     # sayilar metinden once
        self.assertLess(a("zeytin"), a(""))                    # bos degerler sonda

    def test_hatali_sekmesi_ve_tekrar_sorgu(self):
        t = self.app.hata_tablosu
        firmalar = sorted(t.satirlar[i][0] for i in range(len(t.satirlar)))
        self.assertEqual(firmalar, ["YILDIZ OTO", "YILDIZ OTO", "ZEYTIN LTD"])  # hata + inmeyen + excel inmedi
        self.assertIn("3 ekranda sorun var", self.app.hata_bilgisi.cget("text"))
        istekler = []
        self.app.tekrar_sorgula = lambda liste, pencere=None: istekler.append(liste)
        # hicbiri secili degil: uyari, sorgu yok
        import luca_arayuz
        bilgi = []
        eski = luca_arayuz.messagebox.showinfo
        luca_arayuz.messagebox.showinfo = lambda *a, **k: bilgi.append(a)
        try:
            self.app._hatalilari_sorgula(False)
        finally:
            luca_arayuz.messagebox.showinfo = eski
        self.assertTrue(bilgi)
        self.assertEqual(istekler, [])
        t.filtre_ayarla("Durum", "Excel inmedi")                 # gorunenlerin hepsi
        self.app._hatalilari_sorgula(True)
        self.assertEqual([(x["firma"], x["ekran"]) for x in istekler[0]], [("ZEYTIN LTD", "e-fatura-satis")])
        t.temizle()
        t.agac.selection_set([str(i) for i in t.gorunen_indeksler()[:2]])   # secilenler
        self.app._hatalilari_sorgula(False)
        self.assertEqual(len(istekler[1]), 2)

    def test_alt_sekmeler_ve_yer_tutucu(self):
        for ad in ("durum", "hata", "surec"):
            self.app.alt_sekme_sec(ad)
            self.kok.update()
            self.assertTrue(self.app.alt_cerceveler[ad].winfo_ismapped())
            self.assertEqual([c for k, c in self.app.alt_cerceveler.items() if c.winfo_ismapped()],
                             [self.app.alt_cerceveler[ad]])
        self.app.sekme_sec("muavin")
        self.kok.update()
        self.assertTrue(self.app.muavin_govde.winfo_ismapped())
        yazilar = [w.cget("text") for w in self.app.muavin_govde.winfo_children()[0].winfo_children()]
        self.assertIn("Çalışma var", yazilar)
        self.assertFalse(self.app.fatura_govde.winfo_ismapped())


if __name__ == "__main__":
    unittest.main()
