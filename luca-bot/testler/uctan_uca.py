# -*- coding: utf-8 -*-
"""Botu bastan sona sahte Luca uzerinde calistirir (gercek Luca'ya baglanmaz).

    python testler/uctan_uca.py [--ekran-yok]   # Linux'ta: xvfb-run -a python ...

Gecici bir klasorde ayar dosyasi ve firmalar.xlsx olusturur, sahte Luca'yi
acar, luca_bot.py'yi gece modunda (--bitince-kapat) calistirir ve sonunda
rapor.xlsx'in sayfalarini ve beklenen sonuclari kontrol eder.
"""

import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOK))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import sahte_luca  # noqa: E402


def firma_listesi_yaz(yol):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Kısa Adı", "Kapanış Tarihi", "e-Fatura Alış", "İnteraktif V.D."])
    ws.append(["AKIN COBAN", "", "", ""])
    ws.append(["DENTAL SAGLIK", "", "", "X"])   # interaktif ekrani bu firmada acilmaz
    ws.append(["ESKI DONEM LTD", "", "", ""])
    ws.append(["FATURASIZ AS", "", "", ""])
    ws.append(["KEREM TICARET", "01/01/2026", "", ""])  # donemden once kapanmis
    wb.save(yol)


def main():
    klasor = Path(tempfile.mkdtemp(prefix="lucabot-test-"))
    sunucu, adres = sahte_luca.baslat()
    firma_listesi_yaz(klasor / "firmalar.xlsx")
    ayar = {
        "giris_adresi": adres,
        "tarayici": "chromium",
        "tarayici_yolu": os.environ.get("LUCA_TEST_CHROMIUM", ""),
        "tarayici_sandbox": os.name == "nt",  # Linux test kapsayicisinda root ile sandbox acilmiyor
        "indirme_klasoru": str(klasor / "indirilenler"),
        "firma_listesi": str(klasor / "firmalar.xlsx"),
        "profil_yerel": True,
        "sorgu_azami_dakika": 2,
        "durgunluk_dakika": 1,
        "indirme_bekleme_saniye": 15,
        "azami_fatura": 0,
        "mail_yontemi": "outlook",
        "mail_otomatik_gonder": True,
        "mail_alici": "test@example.com",
    }
    ayar_yolu = klasor / "ayarlar.json"
    ayar_yolu.write_text(json.dumps(ayar), encoding="utf-8")
    ortam = dict(os.environ, LUCA_BOT_AYAR=str(ayar_yolu), LOCALAPPDATA=str(klasor / "yerel"),
                 PYTHONIOENCODING="utf-8")
    komut = [sys.executable, str(KOK / "luca_bot.py"), "--karsilastir",
             "--baslangic", "01/08/2026", "--bitis", "10/08/2026", "--bitince-kapat"]
    print("Calistiriliyor:", " ".join(komut))
    sonuc = subprocess.run(komut, env=ortam, cwd=str(KOK))
    sunucu.shutdown()

    gun = next((klasor / "indirilenler").iterdir())
    rapor = json.loads((gun / "rapor.json").read_text(encoding="utf-8"))
    from openpyxl import load_workbook
    wb = load_workbook(gun / "rapor.xlsx")
    print("\nrapor.xlsx sayfalari:", wb.sheetnames)
    hatalar = []

    def bekle(kosul, mesaj):
        if not kosul:
            hatalar.append(mesaj)

    bekle(sonuc.returncode == 0, f"cikis kodu {sonuc.returncode}")
    bekle(wb.sheetnames == ["Özet", "Firma Durumu", "İndirilen Faturalar", "Dosyalar",
                            "Hatalar ve Uyarılar"], "sayfa adlari")
    bekle(set(rapor) == {"AKIN COBAN", "DENTAL SAGLIK", "ESKI DONEM LTD", "FATURASIZ AS"},
          f"islenen firmalar: {sorted(rapor)}")
    akin = rapor.get("AKIN COBAN", {})
    bekle(akin.get("durumlar") == {"e-arsiv-alis": "tamam", "e-arsiv-interaktif": "tamam"},
          f"AKIN COBAN durumlari: {akin.get('durumlar')}")
    bekle(akin.get("iptal", {}).get("e-arsiv-alis") == 1, "AKIN COBAN iptal/itiraz")
    bekle("e-arsiv-interaktif" not in rapor.get("DENTAL SAGLIK", {}).get("durumlar", {}),
          "DENTAL SAGLIK interaktif ekrani X ile atlanmaliydi")
    bekle(rapor.get("FATURASIZ AS", {}).get("durumlar", {}).get("e-arsiv-alis") == "fatura yok",
          "FATURASIZ AS 'fatura yok' olmali")
    mutabakat = [r[6].value for r in wb["İndirilen Faturalar"].iter_rows(min_row=2)
                 if r[0].value == "AKIN COBAN"]
    bekle("e-Arşiv'de YOK (İnteraktif'te var)" in mutabakat,
          f"mutabakat sutunu eksik fatura gostermeli: {mutabakat}")
    print("Calisma klasoru:", gun)
    if hatalar:
        print("\nBASARISIZ:\n  - " + "\n  - ".join(hatalar))
        return 1
    print("\nUCTAN UCA TEST BASARILI")
    return 0


if __name__ == "__main__":
    sys.exit(main())
