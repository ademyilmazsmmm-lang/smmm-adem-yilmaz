# -*- coding: utf-8 -*-
"""Tarayici gerektirmeyen birim testleri.

    python -m unittest discover -s testler -v
"""

import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from lucabot import eposta, rapor  # noqa: E402
from lucabot.calisma import ozetle  # noqa: E402
from lucabot.fatura_analiz import (excelden_sonuca_isle, excelden_tablo,  # noqa: E402
                                   fatura_kimligi, tevkifatli_satirlar,
                                   tutar_cozumle, zipten_tevkifatlilar)
from lucabot.firma_listesi import (bugun_tamamlananlar, firma_listesini_oku,  # noqa: E402
                                   firmalari_suz, listede_bul)
from lucabot.ortak import (ekran_tamamlanmis_mi, hedef_ay_araligi,  # noqa: E402
                           sure_yaz, tarih_araliklari, tarih_cozumle, yeni_sonuc)


def _xlsx(yol, satirlar):
    from openpyxl import Workbook
    wb = Workbook()
    for s in satirlar:
        wb.active.append(s)
    wb.save(yol)
    return yol


class TarihTestleri(unittest.TestCase):
    def test_yedi_gunluk_parcalar(self):
        araliklar = tarih_araliklari(date(2026, 8, 1), date(2026, 8, 31))
        self.assertEqual(araliklar[0], ("01/08/2026", "08/08/2026"))
        self.assertEqual(araliklar[-1][1], "31/08/2026")
        # her parca bir oncekinin bitisinden baslar
        for (_, bit), (bas, _) in zip(araliklar, araliklar[1:]):
            self.assertEqual(bit, bas)

    def test_tek_gun(self):
        self.assertEqual(tarih_araliklari(date(2026, 8, 1), date(2026, 8, 1)),
                         [("01/08/2026", "01/08/2026")])

    def test_ters_aralik_hata(self):
        with self.assertRaises(ValueError):
            tarih_araliklari(date(2026, 8, 2), date(2026, 8, 1))

    def test_hedef_ay_tamami(self):
        # kisa sorgu araliginda da listeleme ayin tamami olmali
        self.assertEqual(hedef_ay_araligi(date(2026, 8, 1), date(2026, 8, 2)),
                         (date(2026, 8, 1), date(2026, 8, 31)))
        self.assertEqual(hedef_ay_araligi(date(2028, 2, 1), date(2028, 3, 15))[1], date(2028, 2, 29))

    def test_tarih_cozumle(self):
        self.assertEqual(tarih_cozumle(" 05/08/2026 "), date(2026, 8, 5))

    def test_sure_yaz(self):
        self.assertEqual(sure_yaz(75), "1 dk 15 sn")
        self.assertEqual(sure_yaz(4000), "1 sa 6 dk")


class TutarVeFaturaTestleri(unittest.TestCase):
    def test_tutar(self):
        self.assertEqual(tutar_cozumle("1.234,56"), 1234.56)
        self.assertEqual(tutar_cozumle("1234,5 TL"), 1234.5)
        self.assertEqual(tutar_cozumle(1234.5), 1234.5)
        self.assertEqual(tutar_cozumle(""), 0.0)
        self.assertEqual(tutar_cozumle("-"), 0.0)

    def test_fatura_kimligi(self):
        satir = ["05/08/2026", "TURKCELL ILETISIM A.S.", "TCL2026000000123", "1.200,00"]
        self.assertEqual(fatura_kimligi(satir), ("TURKCELL", "TCL2026000000123"))
        self.assertIsNone(fatura_kimligi(["05/08/2026", "numarasiz satir", "12"]))

    def test_tevkifat_sadece_alis(self):
        satirlar = [["X", "TEVKİFATLI"], ["Y", "SATIS"]]
        self.assertEqual(len(tevkifatli_satirlar(satirlar, "e-arsiv-alis")), 1)
        self.assertEqual(tevkifatli_satirlar(satirlar, "e-arsiv-satis"), [])

    def test_zipten_tevkifat(self):
        import zipfile
        with tempfile.TemporaryDirectory() as d:
            yol = Path(d) / "b.zip"
            with zipfile.ZipFile(yol, "w") as z:
                z.writestr("ABC2026000000001.xml", "<x><WithholdingTaxTotal/></x>")
                z.writestr("ABC2026000000002.xml", "<x/>")
            self.assertEqual(zipten_tevkifatlilar(yol), {"ABC2026000000001"})
            self.assertEqual(zipten_tevkifatlilar(Path(d) / "yok.zip"), set())

    def test_excelden_sonuca_isle(self):
        with tempfile.TemporaryDirectory() as d:
            klasor = Path(d)
            yol = _xlsx(klasor / "liste.xlsx", [
                ["Fatura No", "Unvan", "Tarih", "Tip", "Matrah", "KDV Tutarı", "KDV Oranı",
                 "Genel Toplam", "Durum"],
                ["AAA2026000000001", "TURKCELL", "05/08/2026", "SATIS", 1000, 200, 20, 1200, "ONAY"],
                ["AAA2026000000002", "VODAFONE", "06/08/2026", "SATIS", 500, 100, 20, 600, "İPTAL"],
                ["AAA2026000000003", "TRUGO", "07/08/2026", "TEVKIFAT", 2000, 400, 20, 2400, "ONAY"],
            ])
            sonuc = yeni_sonuc("F", "e-arsiv-alis")
            satirlar = excelden_sonuca_isle(sonuc, yol, klasor, None)
            self.assertEqual(len(satirlar), 3)
            self.assertEqual(sonuc["fatura_sayisi"], 3)
            self.assertEqual(sonuc["tevkifat"], 1)
            self.assertEqual(sonuc["iptal_itiraz"], 1)
            # iptal edilen fatura toplama girmez; KDV oran sutunu toplanmaz
            self.assertEqual(sonuc["matrah"], 3000)
            self.assertEqual(sonuc["kdv"], 600)
            self.assertEqual(sonuc["faturalar"][0], ["TURKCELL", "AAA2026000000001", 1200.0])
            self.assertTrue((klasor / "iptal-itiraz.csv").exists())

    def test_fatura_tutari(self):
        from lucabot.fatura_analiz import satir_toplam_tutari
        # Luca'nin Excel'inde genel toplam sutunu yoksa matrah + KDV
        basliklar = ["Fatura No", "Mal Hizmet Toplam Tutarı", "Hesaplanan KDV", "KDV Oranı"]
        self.assertEqual(satir_toplam_tutari(basliklar, ["X", "1.000,00", "200,00", "20"]), 1200.0)
        # genel toplam sutunu varsa o kullanilir
        basliklar = ["Fatura No", "Matrah", "KDV", "Ödenecek Tutar"]
        self.assertEqual(satir_toplam_tutari(basliklar, ["X", "1000", "200", "1150,5"]), 1150.5)

    def test_gib_hata_mesaji_taninir(self):
        from lucabot.gib_sorgu import gib_hatasi, hata_satiri, yetki_hatasi
        metin = ("İşlem Takip\nGİB e-Arşiv Sistemi Hata Mesajı:Doğrulama hatası Internet vergi"
                 " dairesinden kimlik doğrulanamadı.")
        self.assertTrue(gib_hatasi(metin))
        self.assertFalse(yetki_hatasi(metin))
        self.assertIn("kimlik", hata_satiri(metin))
        self.assertFalse(gib_hatasi("01/08 sorgulandı. 2 belge kaydı bulundu. İşlem sona erdi."))

    def test_bozuk_excel_cokertmez(self):
        with tempfile.TemporaryDirectory() as d:
            yol = Path(d) / "bozuk.xlsx"
            yol.write_bytes(b"bu bir excel degil")
            self.assertEqual(excelden_tablo(yol), ([], []))


class FirmaListesiTestleri(unittest.TestCase):
    def test_sadece_adi_olan_satirlar_okunur(self):
        with tempfile.TemporaryDirectory() as d:
            yol = _xlsx(Path(d) / "firmalar.xlsx", [
                ["FIRMA LISTESI"],  # tek hucrelik baslik baslik sayilmaz
                ["Kısa Adı", "Kapanış Tarihi", "e-Fatura Alış", "e-Arşiv Alış"],
                ["AKIN COBAN"],
                ["DENTAL", "", "X", ""],
                ["KAPALI", "01/03/2026", "", ""],
            ])
            liste = firma_listesini_oku(yol)
            self.assertEqual(set(liste), {"AKIN COBAN", "DENTAL", "KAPALI"})
            self.assertEqual(liste["DENTAL"][1], {"e-fatura-alis"})
            self.assertEqual(liste["KAPALI"][0], date(2026, 3, 1))

    def test_listede_bul_en_uzun_eslesme(self):
        liste = {"ADEM": 1, "ADEM MERGE": 2}
        self.assertEqual(listede_bul("ADEM MERGE", liste), "ADEM MERGE")
        self.assertEqual(listede_bul("ADEM", liste), "ADEM")

    def test_firmalari_suz(self):
        with tempfile.TemporaryDirectory() as d:
            yol = _xlsx(Path(d) / "firmalar.xlsx", [
                ["Kısa Adı", "Kapanış Tarihi", "İnteraktif V.D."],
                ["AKIN COBAN", "", "X"], ["DENTAL", "", ""], ["KAPALI", "01/03/2026", ""],
            ])
            ayarlar = {"firma_listesi": str(yol), "atlanacak_firmalar": ["DENT"]}
            secim = firmalari_suz(["AKIN COBAN", "DENTAL", "KAPALI", "LISTEDE YOK"], ayarlar,
                                  baslangic=date(2026, 8, 1))
            self.assertEqual(secim.firmalar, ["AKIN COBAN"])
            self.assertEqual(secim.atlanan_ekranlar, {"AKIN COBAN": {"e-arsiv-interaktif"}})
            bos = firmalari_suz(["AKIN COBAN"], {}, aranan=["zzz"])
            self.assertEqual(bos.firmalar, [])
            self.assertTrue(bos.bos_sebep)

    def test_bugun_tamamlananlar(self):
        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "rapor.json").write_text(json.dumps({
                "A": {"durumlar": {"e-arsiv-alis": "tamam", "gib-5000": "hata: X"}},
                "B": {"durumlar": {"e-arsiv-alis": "bekliyor"}},
            }), encoding="utf-8")
            self.assertEqual(bugun_tamamlananlar(d, ["e-arsiv-alis", "gib-5000"]),
                             {"A": {"e-arsiv-alis"}})
            self.assertTrue(ekran_tamamlanmis_mi("fatura yok"))
            self.assertFalse(ekran_tamamlanmis_mi("kaynaktan inmedi"))


class RaporTestleri(unittest.TestCase):
    def _sonuc(self, firma, tip, **alanlar):
        s = yeni_sonuc(firma, tip)
        s.update(alanlar)
        return s

    def test_ortusen_ekranlar_cifte_sayilmaz(self):
        sonuclar = [
            self._sonuc("A", "e-arsiv-alis", durum="tamam", fatura_sayisi=3, tevkifat=1),
            self._sonuc("A", "e-arsiv-interaktif", durum="tamam", fatura_sayisi=4, tevkifat=1),
            self._sonuc("A", "e-fatura-alis", durum="fatura yok"),
            self._sonuc("A", "gib-5000", durum="hata: LookupError", **{"not": "menu yok"}),
        ]
        oz = ozetle(sonuclar)
        self.assertEqual((oz.firma, oz.ekran, oz.basarili, oz.bos, oz.sorunlu), (1, 4, 2, 1, 1))
        self.assertEqual(oz.fatura, 4)     # e-arsiv-alis ile interaktif ayni faturalar
        self.assertEqual(oz.tevkifat, 1)
        self.assertIn("menu yok", oz.sorunlu_liste[0])

    def test_rapor_dosyalari_ve_mutabakat(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            klasor = Path(d)
            (klasor / "rapor.json").write_text(json.dumps({  # eski bicimde kayit
                "ESKI": {"firma": "ESKI", "donem": "", "durumlar": {"e-arsiv-alis": "tamam"},
                         "sayilar": {"e-arsiv-alis": 1}, "iptal": {}, "tevkifat": {}, "inmeyen": {},
                         "faturalar": {"e-arsiv-alis": [["X", "AAA2026000000009"]]},
                         "dosya": 2, "not": "", "son": ""}}), encoding="utf-8")
            sonuclar = [
                self._sonuc("A", "e-arsiv-alis", durum="tamam", fatura_sayisi=1,
                            faturalar=[["TURKCELL", "AAA2026000000001", 1200.0]], matrah=1000, kdv=200),
                self._sonuc("A", "e-arsiv-interaktif", durum="tamam", fatura_sayisi=2,
                            faturalar=[["TURKCELL", "AAA2026000000001", 1200.0],
                                       ["SHELL", "AAA2026000000004", 360.0]]),
                self._sonuc("B", "gib-5000", durum="hata: LookupError", **{"not": "menu yok"}),
            ]
            rapor.guncelle(klasor, sonuclar, ["C"], "e-arsiv-alis")
            wb = load_workbook(klasor / "rapor.xlsx")
            self.assertEqual(wb.sheetnames, ["Özet", "Firma Durumu", "İndirilen Faturalar",
                                             "Dosyalar", "Hatalar ve Uyarılar"])
            mutabakat = {(r[0].value, r[1].value, r[4].value): r[6].value
                         for r in wb["İndirilen Faturalar"].iter_rows(min_row=2)}
            self.assertEqual(mutabakat[("A", "İnteraktif V.D.", "AAA2026000000004")],
                             "e-Arşiv'de YOK (İnteraktif'te var)")
            self.assertEqual(mutabakat[("A", "e-Arşiv Alış", "AAA2026000000001")], "iki ekranda da var")
            hatalar = [(r[0].value, r[2].value) for r in wb["Hatalar ve Uyarılar"].iter_rows(min_row=2)]
            self.assertIn(("B", "hata: LookupError"), hatalar)
            self.assertIn(("C", "bekliyor"), hatalar)
            # Firma Durumu'ndaki tutarlar sayi olarak yazilir (Excel'de toplanabilsin)
            basliklar = [c.value for c in wb["Firma Durumu"][1]]
            j = basliklar.index("Alış Matrah")
            degerler = {r[0].value: r[j].value for r in wb["Firma Durumu"].iter_rows(min_row=2)}
            self.assertEqual(degerler["A"], 1000)
            self.assertTrue((klasor / "rapor.csv").exists())
            self.assertFalse((klasor / "rapor.json.tmp").exists())

    def test_eposta_metni(self):
        sonuclar = [self._sonuc("A", "e-arsiv-alis", durum="tamam", fatura_sayisi=3, tevkifat=2),
                    self._sonuc("B", "esmm-alis", durum="tamam", fatura_sayisi=1)]
        metin, uyari = eposta.ozet_metni(sonuclar, "01/08/2026-31/08/2026")
        self.assertTrue(uyari)
        self.assertIn("TEVKIFATLI ALIS", metin)
        self.assertIn("e-SMM ALIS", metin)
        self.assertNotIn("*", metin)


if __name__ == "__main__":
    unittest.main()
